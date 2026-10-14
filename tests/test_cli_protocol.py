# Copyright 2025-2026, Euro NCAP IVZW
# Created by IVEX NV (https://ivex.ai)
#
# Licensed under the Apache License 2.0.
# See http://www.apache.org/licenses/LICENSE-2.0 for details.

"""cli.build_cli: one command line per registered rating protocol over the
same five domains.

- generate-template under a non-default protocol stamps the protocol marker;
  the Euro NCAP command line writes no marker, so its output is unchanged.
- preprocess / compute-score refuse a file generated for another protocol,
  and a file whose marker names a protocol no installed package provides.
- A command object invoked directly (no CLI group) asserts nothing.
- Footer and help name the programme.

The test protocol registers a copy of the Euro NCAP Safe Driving master,
under another name and outside the package, as its own: the five templates
it generates differ from Euro NCAP's only by the marker, and generate-template
has to find the registered file where it was registered."""

import os
import shutil
import tempfile
import unittest
from importlib.resources import files

import openpyxl
from click.testing import CliRunner

from euroncap_rating_2026 import common
from euroncap_rating_2026.cli import build_cli, cli
from euroncap_rating_2026.safe_driving import integrity as sd_integrity
from euroncap_rating_2026.safe_driving.preprocess import preprocess as sd_preprocess

TEST_PROTOCOL = "test-26-28"
TEST_SD_TEMPLATE = "test_sd_template.xlsx"
DOMAIN_TEMPLATES = {
    "crash_protection": "cp_template.xlsx",
    "crash_avoidance": "ca_template.xlsx",
    "safe_driving": "sd_template.xlsx",
    "post_crash": "pc_template.xlsx",
    "overall": "overall_template.xlsx",
}
# What the test protocol's command line writes per domain.
TEST_DOMAIN_TEMPLATES = {**DOMAIN_TEMPLATES, "safe_driving": TEST_SD_TEMPLATE}

_master_dir = None


def _register_test_protocol():
    global _master_dir
    common.register_protocol(
        common.ProtocolSpec(TEST_PROTOCOL, "Test NCAP", "Test NCAP Ltd")
    )
    _master_dir = tempfile.TemporaryDirectory()
    master = os.path.join(_master_dir.name, TEST_SD_TEMPLATE)
    shutil.copyfile(str(files("data").joinpath("sd_template.xlsx")), master)
    sd_integrity.register_pristine_template(TEST_PROTOCOL, master)


def _unregister_test_protocol():
    global _master_dir
    sd_integrity.PRISTINE_TEMPLATE_PATHS.pop(TEST_PROTOCOL, None)
    common.PROTOCOLS.pop(TEST_PROTOCOL, None)
    sd_integrity._pristine_dfs.cache_clear()
    sd_integrity._pristine_wb.cache_clear()
    if _master_dir is not None:
        _master_dir.cleanup()
        _master_dir = None


def _generate(group, domain, tmpdir):
    runner = CliRunner()
    cwd = os.getcwd()
    os.chdir(tmpdir)
    try:
        result = runner.invoke(group, [domain, "generate-template"])
    finally:
        os.chdir(cwd)
    assert result.exit_code == 0, result.output
    return result


class TestGenerateTemplateStampsProtocol(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        _register_test_protocol()
        cls.test_cli = build_cli(TEST_PROTOCOL)

    @classmethod
    def tearDownClass(cls):
        _unregister_test_protocol()

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmpdir = self._tmp.name

    def tearDown(self):
        self._tmp.cleanup()

    def test_registered_protocol_templates_carry_the_marker_in_every_domain(self):
        for domain, name in TEST_DOMAIN_TEMPLATES.items():
            with self.subTest(domain=domain):
                _generate(self.test_cli, domain, self.tmpdir)
                path = os.path.join(self.tmpdir, name)
                self.assertTrue(os.path.exists(path), os.listdir(self.tmpdir))
                self.assertEqual(common.get_protocol_from_file(path), TEST_PROTOCOL)
                wb = openpyxl.load_workbook(path)
                self.assertEqual(wb["Version"]["A1"].value, "Template version")

    def test_euro_ncap_templates_carry_no_marker(self):
        for domain, name in DOMAIN_TEMPLATES.items():
            with self.subTest(domain=domain):
                _generate(cli, domain, self.tmpdir)
                path = os.path.join(self.tmpdir, name)
                self.assertEqual(common.get_protocol_from_file(path), "encap-26-28")
                ws = openpyxl.load_workbook(path)["Version"]
                labels = [ws.cell(r, 1).value for r in range(1, ws.max_row + 1)]
                self.assertNotIn(common.PROTOCOL_LABEL, labels)

    def test_templates_differ_from_euro_ncap_only_by_the_marker(self):
        for domain, name in DOMAIN_TEMPLATES.items():
            with self.subTest(domain=domain):
                _generate(self.test_cli, domain, self.tmpdir)
                test_wb = openpyxl.load_workbook(
                    os.path.join(self.tmpdir, TEST_DOMAIN_TEMPLATES[domain])
                )
                encap_wb = openpyxl.load_workbook(str(files("data").joinpath(name)))
                self.assertEqual(test_wb.sheetnames, encap_wb.sheetnames)
                for sheet in encap_wb.sheetnames:
                    if sheet == "Version":
                        continue
                    self.assertEqual(
                        [[c.value for c in row] for row in encap_wb[sheet].iter_rows()],
                        [[c.value for c in row] for row in test_wb[sheet].iter_rows()],
                        f"{domain}/{sheet}",
                    )

    def test_footer_and_help_name_the_programme(self):
        encap = _generate(cli, "post_crash", self.tmpdir)
        self.assertIn("Copyright 2025-2026, Euro NCAP IVZW", encap.output)
        other = _generate(self.test_cli, "post_crash", self.tmpdir)
        self.assertIn("Copyright 2025-2026, Test NCAP Ltd", other.output)
        self.assertNotIn("Euro NCAP IVZW", other.output)
        runner = CliRunner()
        self.assertIn("Euro NCAP Rating", runner.invoke(cli, ["--help"]).output)
        self.assertIn(
            "Test NCAP Rating", runner.invoke(self.test_cli, ["--help"]).output
        )

    def test_build_cli_requires_a_registered_protocol(self):
        with self.assertRaises(ValueError):
            build_cli("other-26-28")


class TestProtocolMismatchIsRefused(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        _register_test_protocol()
        cls.test_cli = build_cli(TEST_PROTOCOL)
        cls._tmp = tempfile.TemporaryDirectory()
        root = cls._tmp.name
        cls.encap_dir = os.path.join(root, "encap")
        cls.other_dir = os.path.join(root, "other")
        os.makedirs(cls.encap_dir)
        os.makedirs(cls.other_dir)
        _generate(cli, "safe_driving", cls.encap_dir)
        _generate(cls.test_cli, "safe_driving", cls.other_dir)
        cls.encap_template = os.path.join(cls.encap_dir, "sd_template.xlsx")
        cls.other_template = os.path.join(cls.other_dir, TEST_SD_TEMPLATE)

    @classmethod
    def tearDownClass(cls):
        cls._tmp.cleanup()
        _unregister_test_protocol()

    def test_euro_ncap_command_refuses_another_protocols_file(self):
        result = CliRunner().invoke(
            cli, ["safe_driving", "preprocess", "-i", self.other_template]
        )
        self.assertEqual(result.exit_code, 1)
        self.assertIn("Protocol mismatch", result.output)
        self.assertIn("Test NCAP", result.output)

    def test_other_command_refuses_a_euro_ncap_file(self):
        result = CliRunner().invoke(
            self.test_cli, ["safe_driving", "preprocess", "-i", self.encap_template]
        )
        self.assertEqual(result.exit_code, 1)
        self.assertIn("Protocol mismatch", result.output)
        self.assertIn("Euro NCAP", result.output)

    def test_marker_of_an_uninstalled_protocol_is_refused_with_a_hint(self):
        path = os.path.join(self.other_dir, "uninstalled.xlsx")
        wb = openpyxl.load_workbook(self.encap_template)
        wb["Version"][f"A{common.PROTOCOL_VERSION_ROW}"] = common.PROTOCOL_LABEL
        wb["Version"][f"B{common.PROTOCOL_VERSION_ROW}"] = "other-26-28"
        wb.save(path)
        result = CliRunner().invoke(cli, ["safe_driving", "preprocess", "-i", path])
        self.assertEqual(result.exit_code, 1)
        self.assertIn("install the package", result.output)

    def test_direct_command_invocation_checks_nothing(self):
        """A command object run outside a CLI group (tests, embedding) has no
        protocol to assert: the file's own marker drives the scoring. The
        test protocol registers no Speed Assistance profile, so the run must
        get past the protocol check and fail on exactly that, for the
        marker's protocol."""
        self.assertIsNone(common.get_cli_protocol())
        common.check_protocol(self.other_template)  # no exit
        with tempfile.TemporaryDirectory() as out_dir:
            result = CliRunner().invoke(
                sd_preprocess, ["-i", self.other_template, "-o", out_dir]
            )
        self.assertNotIn("Protocol mismatch", result.output)
        self.assertIsInstance(result.exception, ValueError)
        self.assertIn("No Speed Assistance profile", str(result.exception))
        self.assertIn(TEST_PROTOCOL, str(result.exception))


class TestRegisterPristineTemplate(unittest.TestCase):
    def test_relative_path_is_stored_absolute(self):
        common.register_protocol(
            common.ProtocolSpec(TEST_PROTOCOL, "Test NCAP", "Test NCAP Ltd")
        )
        cwd = os.getcwd()
        try:
            with tempfile.TemporaryDirectory() as tmp:
                shutil.copyfile(
                    str(files("data").joinpath("sd_template.xlsx")),
                    os.path.join(tmp, TEST_SD_TEMPLATE),
                )
                os.chdir(tmp)
                stored = sd_integrity.register_pristine_template(
                    TEST_PROTOCOL, TEST_SD_TEMPLATE
                )
                os.chdir(cwd)
                self.assertTrue(os.path.isabs(stored))
                self.assertEqual(
                    sd_integrity.pristine_template_name(TEST_PROTOCOL),
                    TEST_SD_TEMPLATE,
                )
                self.assertTrue(
                    os.path.isfile(sd_integrity.pristine_template_path(TEST_PROTOCOL))
                )
        finally:
            os.chdir(cwd)
            _unregister_test_protocol()


if __name__ == "__main__":
    unittest.main()
