# Copyright 2025-2026, Euro NCAP IVZW
# Created by IVEX NV (https://ivex.ai)
#
# Licensed under the Apache License 2.0.
# See http://www.apache.org/licenses/LICENSE-2.0 for details.

import importlib
import unittest
from importlib.resources import files

import openpyxl

from euroncap_rating_2026 import common
from euroncap_rating_2026.version import VERSION

# template file -> module holding its COMPUTED_SHEET_COLUMNS
TEMPLATES = {
    "ca_template.xlsx": "euroncap_rating_2026.crash_avoidance.integrity",
    "sd_template.xlsx": "euroncap_rating_2026.safe_driving.integrity",
    "cp_template.xlsx": "euroncap_rating_2026.crash_protection.integrity",
    "pc_template.xlsx": "euroncap_rating_2026.post_crash.integrity",
    "overall_template.xlsx": "euroncap_rating_2026.overall.sheets",
}


class TestDataTemplatesAreGenerateReady(unittest.TestCase):
    """The packaged data/*.xlsx masters must already be exactly what
    generate-template produces, so the repo/wheel file is directly usable
    and generate-template is a pure copy. After a version bump, re-stamp the
    masters with common.add_version_sheet_to_wb (CliCommand.GENERATE_TEMPLATE)."""

    def test_version_sheet_is_stamped_with_current_version(self):
        for name in TEMPLATES:
            with self.subTest(template=name):
                wb = openpyxl.load_workbook(str(files("data").joinpath(name)))
                self.assertIn("Version", wb.sheetnames)
                ws = wb["Version"]
                self.assertEqual(ws["A1"].value, "Template version")
                self.assertEqual(ws["B1"].value, VERSION)

    def test_generate_template_mutations_are_noops(self):
        for name, module_name in TEMPLATES.items():
            with self.subTest(template=name):
                columns = importlib.import_module(module_name).COMPUTED_SHEET_COLUMNS
                wb = openpyxl.load_workbook(str(files("data").joinpath(name)))
                before = {
                    ws.title: [[cell.value for cell in row] for row in ws.iter_rows()]
                    for ws in wb.worksheets
                }
                common.blank_computed_columns(wb, columns)
                common.add_version_sheet_to_wb(wb, common.CliCommand.GENERATE_TEMPLATE)
                after = {
                    ws.title: [[cell.value for cell in row] for row in ws.iter_rows()]
                    for ws in wb.worksheets
                }
                self.assertEqual(before, after)


if __name__ == "__main__":
    unittest.main()
