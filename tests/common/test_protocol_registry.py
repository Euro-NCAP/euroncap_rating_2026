# Copyright 2025-2026, Euro NCAP IVZW
# Created by IVEX NV (https://ivex.ai)
#
# Licensed under the Apache License 2.0.
# See http://www.apache.org/licenses/LICENSE-2.0 for details.

"""The protocol registry and the Version-sheet protocol marker
(common.Protocol / PROTOCOLS / register_protocol, PROTOCOL_LABEL).

A rating programme built on this package registers its protocol at import
time; a workbook records the protocol it was generated for in its Version
sheet, and a workbook without that row is a Euro NCAP workbook."""

import contextlib
import io
import os
import tempfile
import unittest

import openpyxl
import pandas as pd

from euroncap_rating_2026 import common

TEST_PROTOCOL = "test-26-28"
TEST_SPEC = common.ProtocolSpec(TEST_PROTOCOL, "Test NCAP", "Test NCAP Ltd")


class RegistryTestCase(unittest.TestCase):
    def setUp(self):
        common.register_protocol(TEST_SPEC)

    def tearDown(self):
        common.PROTOCOLS.pop(TEST_PROTOCOL, None)


class TestNormalizeProtocol(RegistryTestCase):
    def test_default_and_aliases_resolve_to_euro_ncap(self):
        for value in (None, "", " ", "nan", common.Protocol.ENCAP_26_28, "ENCAP-26-28"):
            self.assertEqual(common.normalize_protocol(value), "encap-26-28")
        self.assertEqual(common.normalize_protocol("encap-23-25"), "encap-26-28")

    def test_registered_protocol_resolves(self):
        self.assertEqual(common.normalize_protocol(" Test-26-28 "), TEST_PROTOCOL)
        self.assertIn(TEST_PROTOCOL, common.known_protocols())

    def test_unknown_protocol_raises_with_install_hint(self):
        with self.assertRaises(ValueError) as ctx:
            common.normalize_protocol("other-26-28")
        self.assertIn("install the package", str(ctx.exception))
        self.assertIn("encap-26-28", str(ctx.exception))

    def test_unregistered_protocol_is_unknown_again(self):
        common.PROTOCOLS.pop(TEST_PROTOCOL)
        with self.assertRaises(ValueError):
            common.normalize_protocol(TEST_PROTOCOL)


class TestRegisterProtocol(RegistryTestCase):
    def test_identical_registration_is_a_noop(self):
        common.register_protocol(TEST_SPEC)
        self.assertEqual(common.PROTOCOLS[TEST_PROTOCOL], TEST_SPEC)

    def test_conflicting_registration_raises(self):
        with self.assertRaises(ValueError):
            common.register_protocol(
                common.ProtocolSpec(TEST_PROTOCOL, "Another", "Another Ltd")
            )
        self.assertEqual(common.PROTOCOLS[TEST_PROTOCOL], TEST_SPEC)

    def test_reserved_ids_cannot_be_registered(self):
        # normalize_protocol resolves these before consulting the registry.
        for reserved in ("", "  ", "nan", "encap-23-25"):
            with self.subTest(reserved=reserved):
                with self.assertRaises(ValueError):
                    common.register_protocol(
                        common.ProtocolSpec(reserved, "Ghost", "Ghost Ltd")
                    )
                self.assertNotIn(reserved.strip().lower(), common.PROTOCOLS)

    def test_euro_ncap_is_built_in(self):
        spec = common.PROTOCOLS["encap-26-28"]
        self.assertEqual(spec.programme_name, "Euro NCAP")
        self.assertEqual(spec.copyright_holder, "Euro NCAP IVZW")
        self.assertEqual(common.programme_name("encap-26-28"), "Euro NCAP")
        self.assertEqual(common.programme_name(TEST_PROTOCOL), "Test NCAP")
        self.assertEqual(common.programme_name("nobody-26-28"), "nobody-26-28")


class TestVersionSheetMarker(RegistryTestCase):
    def _workbook(self, marker=None):
        wb = openpyxl.Workbook()
        common.add_version_sheet_to_wb(wb, common.CliCommand.GENERATE_TEMPLATE)
        if marker is not None:
            common.add_protocol_row_to_version_sheet(wb, marker)
        return wb

    def test_no_marker_means_euro_ncap(self):
        self.assertEqual(
            common.get_protocol_from_workbook(self._workbook()), "encap-26-28"
        )
        self.assertEqual(common.get_protocol_from_dfs({}), "encap-26-28")
        self.assertEqual(common.get_protocol_from_dfs({"Version": None}), "encap-26-28")

    def test_marker_round_trips_through_workbook_and_dataframe(self):
        wb = self._workbook(TEST_PROTOCOL)
        self.assertEqual(common.get_protocol_from_workbook(wb), TEST_PROTOCOL)
        ws = wb["Version"]
        self.assertEqual(ws.cell(common.PROTOCOL_VERSION_ROW, 1).value, "Protocol")
        # read_excel_file_to_dfs shape: the first sheet row becomes the header.
        df = pd.DataFrame(
            [[ws.cell(r, 1).value, ws.cell(r, 2).value] for r in range(2, 5)],
            columns=[ws["A1"].value, ws["B1"].value],
        )
        self.assertEqual(common.get_protocol_from_dfs({"Version": df}), TEST_PROTOCOL)

    def test_explicit_protocol_wins_over_marker(self):
        wb = self._workbook(TEST_PROTOCOL)
        self.assertEqual(common.resolve_protocol("encap-26-28", wb=wb), "encap-26-28")
        self.assertEqual(common.resolve_protocol(None, wb=wb), TEST_PROTOCOL)

    def test_check_protocol_mismatch_and_unknown_marker_exit(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "f.xlsx")
            self._workbook(TEST_PROTOCOL).save(path)
            out = io.StringIO()
            with contextlib.redirect_stdout(out), self.assertRaises(SystemExit):
                common.check_protocol(path, cli_protocol="encap-26-28")
            self.assertIn("Protocol mismatch", out.getvalue())
            self.assertIn("Test NCAP", out.getvalue())
            common.check_protocol(path, cli_protocol=TEST_PROTOCOL)  # no exit
            common.check_protocol(path)  # no CLI protocol asserted: no check

            # A marker no installed package provides.
            wb = self._workbook()
            wb["Version"][f"A{common.PROTOCOL_VERSION_ROW}"] = common.PROTOCOL_LABEL
            wb["Version"][f"B{common.PROTOCOL_VERSION_ROW}"] = "other-26-28"
            wb.save(path)
            out = io.StringIO()
            with contextlib.redirect_stdout(out), self.assertRaises(SystemExit):
                common.check_protocol(path, cli_protocol="encap-26-28")
            self.assertIn("install the package", out.getvalue())


if __name__ == "__main__":
    unittest.main()
