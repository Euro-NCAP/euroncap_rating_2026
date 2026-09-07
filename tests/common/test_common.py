# Copyright 2025-2026, Euro NCAP IVZW
# Created by IVEX NV (https://ivex.ai)
#
# Licensed under the Apache License 2.0.
# See http://www.apache.org/licenses/LICENSE-2.0 for details.

import unittest
import pandas as pd
import numpy as np
import os
import tempfile
import warnings

from unittest.mock import patch, MagicMock
from click.testing import CliRunner

# Import modules to test
from euroncap_rating_2026 import common
from euroncap_rating_2026 import config
from euroncap_rating_2026 import version
from euroncap_rating_2026.cli import cli

# Suppress ResourceWarnings for unclosed files in tests
warnings.filterwarnings("ignore", category=ResourceWarning)


# =============================================================================
# Tests for version.py
# =============================================================================
class TestVersion(unittest.TestCase):
    def test_version_is_string(self):
        """VERSION should be a string."""
        self.assertIsInstance(version.VERSION, str)

    def test_version_not_empty(self):
        """VERSION should not be empty."""
        self.assertGreater(len(version.VERSION), 0)

    def test_version_format(self):
        """VERSION should follow semantic versioning pattern."""
        parts = version.VERSION.split(".")
        self.assertGreaterEqual(
            len(parts), 2, "Version should have at least major.minor"
        )


# =============================================================================
# Tests for config.py
# =============================================================================
class TestConfig(unittest.TestCase):
    def test_settings_default_log_level(self):
        """Settings should have INFO as default log level."""
        settings = config.Settings()
        self.assertEqual(settings.log_level, "INFO")

    def test_settings_env_override(self):
        """Settings should respect environment variable override."""
        with patch.dict(os.environ, {"euroncap_rating_2026_log_level": "DEBUG"}):
            settings = config.Settings()
            self.assertEqual(settings.log_level, "DEBUG")

    def test_logging_config_creates_file_handler(self):
        """logging_config should create a log file."""
        import logging

        with tempfile.TemporaryDirectory() as tmpdir:
            original_cwd = os.getcwd()
            os.chdir(tmpdir)
            try:
                config.logging_config()
                log_file = os.path.join(tmpdir, "euroncap_rating_2026.log")
                self.assertTrue(os.path.exists(log_file))
            finally:
                os.chdir(original_cwd)
                # Close the FileHandler logging_config attached: on Windows
                # the open handle makes the TemporaryDirectory cleanup fail
                # with PermissionError.
                root_logger = logging.getLogger()
                for handler in root_logger.handlers:
                    handler.close()
                root_logger.handlers.clear()


# =============================================================================
# Tests for common.py
# =============================================================================
class TestWithFooter(unittest.TestCase):
    def test_with_footer_decorator(self):
        """with_footer decorator should print footer after function execution."""

        @common.with_footer
        def sample_func():
            print("test output")
            return "result"

        import io
        import sys

        captured_output = io.StringIO()
        sys.stdout = captured_output
        result = sample_func()
        sys.stdout = sys.__stdout__
        output = captured_output.getvalue()

        self.assertEqual(result, "result")
        self.assertIn("Generated with version", output)
        self.assertIn("Copyright 2025-2026, Euro NCAP IVZW", output)


class TestReadExcelFileToDfs(unittest.TestCase):
    def test_read_excel_file_to_dfs(self):
        """Should read Excel file and return dictionary of DataFrames."""
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            filepath = f.name

        try:
            # Create a test Excel file
            df1 = pd.DataFrame({"A": [1, 2], "B": [3, 4]})
            df2 = pd.DataFrame({"X": [5, 6], "Y": [7, 8]})
            with pd.ExcelWriter(filepath) as writer:
                df1.to_excel(writer, sheet_name="Sheet1", index=False)
                df2.to_excel(writer, sheet_name="Sheet2", index=False)

            result = common.read_excel_file_to_dfs(filepath)

            self.assertIn("Sheet1", result)
            self.assertIn("Sheet2", result)
            self.assertEqual(list(result["Sheet1"].columns), ["A", "B"])
            self.assertEqual(list(result["Sheet2"].columns), ["X", "Y"])
        finally:
            os.unlink(filepath)

    def test_preserve_na_sheets_keeps_literal_na_strings(self):
        """pd.read_excel's default na_values silently collapse a literal
        "N/A" cell into NaN -- indistinguishable from blank. Sheet columns
        named in preserve_na_sheets must get the string restored
        (safe_driving's "Input parameters" dropdowns offer "N/A" as a real
        choice); sheets not named must keep the historical NaN behavior."""
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            filepath = f.name

        try:
            df = pd.DataFrame({"Input parameter": ["Type of system"], "Value": ["N/A"]})
            with pd.ExcelWriter(filepath) as writer:
                df.to_excel(writer, sheet_name="Input parameters", index=False)
                df.to_excel(writer, sheet_name="Other", index=False)

            result = common.read_excel_file_to_dfs(
                filepath, preserve_na_sheets={"Input parameters": ("Value",)}
            )

            self.assertEqual(result["Input parameters"]["Value"].iloc[0], "N/A")
            self.assertTrue(pd.isna(result["Other"]["Value"].iloc[0]))
        finally:
            os.unlink(filepath)

    def test_preserve_na_sheets_is_column_scoped(self):
        """Only the named columns of a preserved sheet get "N/A" restored: a
        stray "N/A" typed anywhere else (e.g. a PASS/FAIL Value cell of a
        verif sheet whose Scenario column is preserved) must keep the
        historical NaN/blank behavior."""
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            filepath = f.name

        try:
            df = pd.DataFrame(
                {
                    "Category": ["Some category"],
                    "Scenario": ["N/A"],
                    "Value": ["N/A"],
                }
            )
            with pd.ExcelWriter(filepath) as writer:
                df.to_excel(writer, sheet_name="Verif", index=False)

            result = common.read_excel_file_to_dfs(
                filepath, preserve_na_sheets={"Verif": ("Scenario",)}
            )

            self.assertEqual(result["Verif"]["Scenario"].iloc[0], "N/A")
            self.assertTrue(pd.isna(result["Verif"]["Value"].iloc[0]))
        finally:
            os.unlink(filepath)


class TestPredictionDfToBgDf(unittest.TestCase):
    def test_prediction_df_to_bg_df_normalizes_prediction_values(self):
        """Should convert supported string values to PredictionColor enums."""
        df = pd.DataFrame(
            {
                "prediction": [
                    "blue",
                    "PredictionColor.RED",
                    " yellow ",
                    None,
                    3.14,
                    "unknown",
                ]
            }
        )

        result = common.prediction_df_to_bg_df(df)

        expected = pd.Series(
            [
                common.PredictionColor.BLUE,
                common.PredictionColor.RED,
                common.PredictionColor.YELLOW,
                None,
                3.14,
                None,
            ],
            name="prediction",
        )
        pd.testing.assert_series_equal(result["prediction"], expected)


class TestMergePredictionColorSources(unittest.TestCase):
    """merge_prediction_color_sources reads a prediction grid's colour from
    either shape a sheet can have: the OEM's typed colour words (text frame)
    or the cell fills preprocess re-encodes them into ("(bg)" frame)."""

    RED = common.PredictionColor.RED
    GREEN = common.PredictionColor.GREEN

    def _colors(self, df):
        return df.to_numpy(dtype=object).tolist()

    def test_text_words_only_when_bg_absent(self):
        text_df = pd.DataFrame([["CPLA", "Red", np.nan], ["60 km/h", "green", "X"]])

        result = common.merge_prediction_color_sources(text_df, None)

        self.assertEqual(
            self._colors(result),
            [[None, self.RED, None], [None, self.GREEN, None]],
        )
        self.assertEqual(list(result.columns), list(text_df.columns))
        self.assertTrue((result.dtypes == object).all())

    def test_empty_bg_frame_behaves_like_absent(self):
        text_df = pd.DataFrame([["CPLA", "Red"]])

        result = common.merge_prediction_color_sources(text_df, pd.DataFrame())

        self.assertEqual(self._colors(result), [[None, self.RED]])

    def test_fill_colours_fill_in_blank_text(self):
        # Preprocessed shape: text blanked, only an "X" on selected points.
        text_df = pd.DataFrame([["CPLA", np.nan, None], ["60 km/h", "X", np.nan]])
        bg_df = pd.DataFrame(
            [[None, self.RED, self.RED], [None, self.GREEN, self.RED]],
            columns=[None, None, None],  # header-derived labels may repeat
        )

        result = common.merge_prediction_color_sources(text_df, bg_df)

        self.assertEqual(
            self._colors(result),
            [[None, self.RED, self.RED], [None, self.GREEN, self.RED]],
        )

    def test_text_colour_wins_over_fill(self):
        text_df = pd.DataFrame([["Green", "Red"]])
        bg_df = pd.DataFrame([[self.RED, self.GREEN]])

        result = common.merge_prediction_color_sources(text_df, bg_df)

        self.assertEqual(self._colors(result), [[self.GREEN, self.RED]])

    def test_non_colour_text_never_masks_a_fill(self):
        # "X", "N/A", numbers, NaN and pd.NA are all "no colour word here".
        text_df = pd.DataFrame([["X", "N/A", 1.25, np.nan, pd.NA]], dtype=object)
        bg_df = pd.DataFrame([[self.RED] * 5])

        result = common.merge_prediction_color_sources(text_df, bg_df)

        self.assertEqual(self._colors(result), [[self.RED] * 5])

    def test_larger_bg_frame_is_cropped_to_text_shape(self):
        text_df = pd.DataFrame([["CPLA", np.nan], ["60 km/h", np.nan]])
        bg_df = pd.DataFrame(
            [
                [None, self.RED, self.GREEN],
                [None, self.RED, self.GREEN],
                [self.GREEN, self.GREEN, self.GREEN],
            ]
        )

        result = common.merge_prediction_color_sources(text_df, bg_df)

        self.assertEqual(result.shape, text_df.shape)
        self.assertEqual(self._colors(result), [[None, self.RED], [None, self.RED]])

    def test_smaller_bg_frame_is_padded_to_text_shape(self):
        text_df = pd.DataFrame([["CPLA", np.nan, "Green"], ["60 km/h", np.nan, np.nan]])
        bg_df = pd.DataFrame([[None, self.RED]])

        result = common.merge_prediction_color_sources(text_df, bg_df)

        self.assertEqual(result.shape, text_df.shape)
        self.assertEqual(
            self._colors(result),
            [[None, self.RED, self.GREEN], [None, None, None]],
        )

    def test_string_dtype_text_frame_is_recognised(self):
        # Parquet-sourced frames arrive as StringDtype with pd.NA blanks.
        text_df = pd.DataFrame({"a": pd.array(["Red", None], dtype="string")})

        result = common.merge_prediction_color_sources(text_df, None)

        self.assertEqual(self._colors(result), [[self.RED], [None]])

    def test_result_keeps_text_index(self):
        text_df = pd.DataFrame([["Red"], ["Green"]], index=[5, 9])

        result = common.merge_prediction_color_sources(text_df, None)

        self.assertEqual(list(result.index), [5, 9])


class TestGetDmPredictionRequiredMask(unittest.TestCase):
    def _build_ws(self):
        import openpyxl
        from openpyxl.worksheet.datavalidation import DataValidation

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "DE - DM pred."
        headers = [
            "Task",
            "Movement type",
            "Gaze location",
            "Warning",
            "Forward support",
            "Lane support",
        ]
        ws.append(headers)
        ws.append(
            ["Non-driving task", "Owl", "Driver side window", "Green", "Green", "Green"]
        )
        ws.append(
            [
                "Non-driving task",
                "Owl",
                "In-vehicle infotainment",
                None,
                "Green",
                "Green",
            ]
        )
        ws.append([None, None, None, None, None, None])
        ws.append(["Driving task", "Owl", "Rear view mirror", None, None, None])

        dv = DataValidation(type="list", formula1='"N/A,Green,Red"')
        ws.add_data_validation(dv)
        dv.add("D2:F3")
        return wb, ws

    def test_marks_required_cells_true(self):
        _, ws = self._build_ws()
        result = common.get_dm_prediction_required_mask(ws)
        self.assertTrue(result.iloc[0, 3])
        self.assertTrue(result.iloc[1, 3])

    def test_marks_non_validated_cells_false(self):
        _, ws = self._build_ws()
        result = common.get_dm_prediction_required_mask(ws)
        self.assertFalse(result.iloc[3, 3])

    def test_preserves_blank_separator_rows(self):
        _, ws = self._build_ws()
        result = common.get_dm_prediction_required_mask(ws)
        self.assertTrue(result.iloc[2].isnull().all())

    def test_ignores_unrelated_list_validations(self):
        import openpyxl
        from openpyxl.worksheet.datavalidation import DataValidation

        _, ws = self._build_ws()
        dv2 = DataValidation(type="list", formula1='"Functional,Non Functional"')
        ws.add_data_validation(dv2)
        dv2.add("A2:A2")
        result = common.get_dm_prediction_required_mask(ws)
        self.assertTrue(result.iloc[0, 3])
        self.assertFalse(result.iloc[3, 3])


class TestGetDefaultDmPredictionRequiredMask(unittest.TestCase):
    def test_matches_mask_built_from_packaged_template(self):
        import openpyxl
        from importlib.resources import files

        template_path = str(files("data").joinpath("sd_template.xlsx"))
        wb = openpyxl.load_workbook(template_path, data_only=True)
        try:
            expected = common.get_dm_prediction_required_mask(wb["DE - DM pred."])
        finally:
            wb.close()

        result = common.get_default_dm_prediction_required_mask()

        pd.testing.assert_frame_equal(result, expected)


class TestFrequencyProportionalSample(unittest.TestCase):
    def test_sample_returns_correct_count(self):
        """Should return exactly N samples when N < len(data)."""
        data = [
            (0, 0, "red"),
            (0, 1, "red"),
            (0, 2, "red"),
            (1, 0, "blue"),
            (1, 1, "blue"),
            (2, 0, "green"),
        ]
        result = common.frequency_proportional_sample(data, 3, seed=42)
        self.assertEqual(len(result), 3)

    def test_sample_returns_all_when_n_greater_than_len(self):
        """Should return all data when N >= len(data)."""
        data = [(0, 0, "red"), (0, 1, "blue")]
        result = common.frequency_proportional_sample(data, 10, seed=42)
        self.assertEqual(len(result), 2)

    def test_sample_preserves_proportions(self):
        """Should approximately preserve color proportions."""
        data = [(i, 0, "red") for i in range(60)] + [(i, 0, "blue") for i in range(40)]
        result = common.frequency_proportional_sample(data, 10, seed=42)

        red_count = sum(1 for _, _, c in result if c == "red")
        blue_count = sum(1 for _, _, c in result if c == "blue")

        # Should be approximately 6 red and 4 blue
        self.assertGreaterEqual(red_count, 5)
        self.assertLessEqual(red_count, 7)
        self.assertGreaterEqual(blue_count, 3)
        self.assertLessEqual(blue_count, 5)

    def test_sample_with_seed_reproducible(self):
        """Same seed should produce same results."""
        data = [(i, j, f"color{i%3}") for i in range(10) for j in range(3)]
        result1 = common.frequency_proportional_sample(data, 5, seed=123)
        result2 = common.frequency_proportional_sample(data, 5, seed=123)
        self.assertEqual(result1, result2)


class TestGetInitials(unittest.TestCase):
    def test_get_initials_normal(self):
        """Should return initials of words."""
        self.assertEqual(common.get_initials("Frontal Impact"), "FI")

    def test_get_initials_single_word(self):
        """Should return single initial for single word."""
        self.assertEqual(common.get_initials("Test"), "T")

    def test_get_initials_nan(self):
        """Should return empty string for NaN."""
        self.assertEqual(common.get_initials(pd.NA), "")
        self.assertEqual(common.get_initials(np.nan), "")

    def test_get_initials_lowercase(self):
        """Should uppercase initials."""
        self.assertEqual(common.get_initials("frontal impact"), "FI")


class TestGetFirstWord(unittest.TestCase):
    def test_get_first_word_normal(self):
        """Should return first word."""
        self.assertEqual(common.get_first_word("Frontal Impact"), "Frontal")

    def test_get_first_word_single(self):
        """Should return the word if only one word."""
        self.assertEqual(common.get_first_word("Test"), "Test")

    def test_get_first_word_nan(self):
        """Should return empty string for NaN."""
        self.assertEqual(common.get_first_word(pd.NA), "")
        self.assertEqual(common.get_first_word(np.nan), "")


class TestGetParamDfFromInputParameters(unittest.TestCase):
    def test_get_param_df_from_input_parameters(self):
        """Should extract parameters with stage info."""
        input_df = pd.DataFrame(
            {
                "Stage": ["Crash Protection", np.nan, np.nan],
                "Stage element": ["Frontal", np.nan, np.nan],
                "Stage Subelement": ["Offset", np.nan, "FW"],
                "Input parameter": ["param1", "param2", "param3"],
                "Value": [1.0, 2.0, 3.0],
            }
        )

        result = common.get_param_df_from_input_parameters(input_df)

        self.assertEqual(len(result), 3)
        self.assertIn("param_code", result.columns)
        self.assertIn("Input parameter", result.columns)
        self.assertIn("Value", result.columns)

    def test_empty_values_returns_empty_df_with_correct_columns(self):
        """Should return empty DataFrame with correct columns when all Values are NaN."""
        input_df = pd.DataFrame(
            {
                "Stage": ["Crash Protection", np.nan, np.nan],
                "Stage element": ["Frontal", np.nan, np.nan],
                "Stage subelement": [np.nan, np.nan, np.nan],
                "Input parameter": [
                    "Countermeasure?",
                    "Red line >125 mm",
                    "Torso angle (rear)",
                ],
                "Value": [np.nan, np.nan, np.nan],
            }
        )
        result = common.get_param_df_from_input_parameters(input_df)
        self.assertEqual(len(result), 0)
        self.assertIn("param_code", result.columns)
        self.assertIn("Input parameter", result.columns)
        self.assertIn("Value", result.columns)

    def test_empty_sheet_returns_empty_df_with_correct_columns(self):
        """Should return empty DataFrame with correct columns when input sheet has no rows."""
        input_df = pd.DataFrame(
            columns=[
                "Stage",
                "Stage element",
                "Stage subelement",
                "Input parameter",
                "Value",
            ]
        )
        result = common.get_param_df_from_input_parameters(input_df)
        self.assertEqual(len(result), 0)
        self.assertIn("param_code", result.columns)
        self.assertIn("Input parameter", result.columns)
        self.assertIn("Value", result.columns)


class TestGetParamDf(unittest.TestCase):
    def test_get_param_df_missing_sheet_raises(self):
        """Should raise ValueError if 'Input parameters' sheet is missing."""
        dfs = {"Other Sheet": pd.DataFrame()}

        with self.assertRaisesRegex(ValueError, "Input parameters"):
            common.get_param_df(dfs)

    def test_get_param_df_all_nan_values_does_not_crash(self):
        """Regression: all-NaN Value column must not raise KeyError.

        Mirrors the real call stack: data_loader.load_data → common.get_param_df.
        When every row in 'Input parameters' has a NaN Value (i.e. the sheet was
        submitted blank), the function must return an empty DataFrame with the
        expected columns rather than raising KeyError: 'Stage'.
        """
        input_df = pd.DataFrame(
            {
                "Stage": ["Crash Protection", np.nan, np.nan, np.nan, np.nan, np.nan],
                "Stage element": ["Frontal", np.nan, np.nan, np.nan, np.nan, np.nan],
                "Stage subelement": [np.nan, np.nan, np.nan, np.nan, np.nan, np.nan],
                "Input parameter": [
                    "Countermeasure?",
                    "Red line >125 mm outboard of the orange line",
                    "Torso angle (rear)",
                    "Number of verification tests (min 10)",
                    "Number of verification tests (min 5)",
                    "Number of verification tests (min 5)",
                ],
                "Value": [np.nan, np.nan, np.nan, np.nan, np.nan, np.nan],
            }
        )
        dfs = {"Input parameters": input_df}
        result = common.get_param_df(dfs)
        self.assertEqual(len(result), 0)
        self.assertIn("param_code", result.columns)
        self.assertIn("Input parameter", result.columns)
        self.assertIn("Value", result.columns)


class TestHardCopySheet(unittest.TestCase):
    def test_hard_copy_sheet(self):
        """Should copy a sheet between workbook objects."""
        import openpyxl

        input_wb = openpyxl.Workbook()
        ws = input_wb.active
        ws.title = "TestSheet"
        ws["A1"] = "test value"

        output_wb = openpyxl.Workbook()

        common.hard_copy_sheet(input_wb, "TestSheet", output_wb)

        self.assertIn("TestSheet", output_wb.sheetnames)
        self.assertEqual(output_wb["TestSheet"]["A1"].value, "test value")

    def test_load_or_create_workbook_creates_destination_from_empty_file(self):
        """Should create a workbook when the destination file is empty."""
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f1:
            output_file = f1.name

        try:
            with open(output_file, "wb"):
                pass

            wb_result = common.load_or_create_workbook(output_file)
            self.assertEqual(wb_result.sheetnames, [])
        finally:
            os.unlink(output_file)

    def test_hard_copy_sheet_missing_sheet_skips_with_warning(self):
        """Should skip a missing sheet with a warning instead of raising,
        so a workbook uploaded against an older template revision (e.g.
        without "Documentation"/"Data") never aborts write_report."""
        import openpyxl

        input_wb = openpyxl.Workbook()
        output_wb = openpyxl.Workbook()

        with self.assertLogs("euroncap_rating_2026.common", level="WARNING") as cm:
            result = common.hard_copy_sheet(input_wb, "NonExistent", output_wb)

        self.assertFalse(result)
        self.assertNotIn("NonExistent", output_wb.sheetnames)
        self.assertTrue(any("NonExistent" in message for message in cm.output))

    def test_hard_copy_sheet_returns_true_on_copy(self):
        import openpyxl

        input_wb = openpyxl.Workbook()
        input_wb.active.title = "Present"
        output_wb = openpyxl.Workbook()

        self.assertTrue(common.hard_copy_sheet(input_wb, "Present", output_wb))


class TestDeleteRowsKeepingMerges(unittest.TestCase):
    """
    common.delete_rows_keeping_merges: openpyxl's Worksheet.delete_rows moves
    cell values/styles up but leaves merged ranges pinned to their original
    coordinates, detaching any merge below the deletion. Used by
    report_writer.write_report to physically remove CPLA/CBLA rows
    collapsed by test_info.collapse_cpla_cbla_aeb_duplicate_rows without
    corrupting header merges further down the same sheet.
    """

    def _build_ws(self):
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        for row in range(1, 11):
            ws.cell(row=row, column=1, value=f"row{row}")
        return wb, ws

    def test_merge_below_deleted_row_shifts_up(self):
        wb, ws = self._build_ws()
        ws.merge_cells(start_row=8, start_column=3, end_row=8, end_column=5)
        common.delete_rows_keeping_merges(ws, [3])
        # Row 4's old value is now at row 3 after the shift.
        self.assertEqual(ws.cell(row=3, column=1).value, "row4")
        merged = [str(r) for r in ws.merged_cells.ranges]
        self.assertIn("C7:E7", merged)
        self.assertNotIn("C8:E8", merged)

    def test_non_contiguous_deletion_shifts_correctly(self):
        wb, ws = self._build_ws()
        ws.merge_cells(start_row=10, start_column=1, end_row=10, end_column=2)
        common.delete_rows_keeping_merges(ws, [2, 5])
        # Two rows removed above row 10 -> merge shifts to row 8.
        merged = [str(r) for r in ws.merged_cells.ranges]
        self.assertIn("A8:B8", merged)

    def test_merge_entirely_inside_deleted_rows_is_dropped(self):
        wb, ws = self._build_ws()
        ws.merge_cells(start_row=4, start_column=1, end_row=4, end_column=2)
        common.delete_rows_keeping_merges(ws, [4])
        self.assertEqual(list(ws.merged_cells.ranges), [])

    def test_merge_above_deleted_rows_is_unaffected(self):
        wb, ws = self._build_ws()
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2)
        common.delete_rows_keeping_merges(ws, [8, 9])
        merged = [str(r) for r in ws.merged_cells.ranges]
        self.assertIn("A1:B1", merged)

    def test_empty_rows_is_a_no_op(self):
        wb, ws = self._build_ws()
        ws.merge_cells(start_row=5, start_column=1, end_row=5, end_column=2)
        common.delete_rows_keeping_merges(ws, [])
        merged = [str(r) for r in ws.merged_cells.ranges]
        self.assertIn("A5:B5", merged)
        self.assertEqual(ws.cell(row=5, column=1).value, "row5")


class TestHideDocumentationDataSheets(unittest.TestCase):
    """hard_copy_sheet (used to carry "Documentation"/"Data" through
    preprocess/compute-score, see each domain's report_writer.py) doesn't
    preserve sheet_state -- a hard-copied sheet always comes back visible.
    common.hide_documentation_data_sheets is the explicit fix-up called once
    before every domain's final save."""

    def test_hides_both_sheets_when_present(self):
        import openpyxl

        wb = openpyxl.Workbook()
        wb.create_sheet("Documentation")
        wb.create_sheet("Data")
        self.assertEqual(wb["Documentation"].sheet_state, "visible")
        self.assertEqual(wb["Data"].sheet_state, "visible")

        common.hide_documentation_data_sheets(wb)

        self.assertEqual(wb["Documentation"].sheet_state, "hidden")
        self.assertEqual(wb["Data"].sheet_state, "hidden")

    def test_no_error_when_sheets_absent(self):
        import openpyxl

        wb = openpyxl.Workbook()
        common.hide_documentation_data_sheets(wb)  # must not raise


class TestBlankComputedColumns(unittest.TestCase):
    """common.blank_computed_columns sets every column named in its
    sheet->columns mapping to "-", used by each domain's generate_template()
    (the pristine template handed to the OEM)."""

    def _make_wb(self):
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Test Scores"
        ws.append(["Stage", "Stage element", "Score", "Max score"])
        ws.append(["Crash Avoidance", "Frontal Collisions", 0, 40])
        ws.append(["Crash Avoidance", "Lane Departure", 0, 10])

        ws2 = wb.create_sheet("CP - Dummy Scores")
        ws2.append(["Stage element", "Dummy", "Capping?", "Score", "Max score"])
        ws2.append(["Frontal Impact", "THOR-50", None, None, 5])

        return wb

    def test_sets_requested_columns_across_multiple_sheets(self):
        wb = self._make_wb()

        common.blank_computed_columns(
            wb,
            {
                "Test Scores": ["Score"],
                "CP - Dummy Scores": ["Capping?", "Score"],
            },
        )

        ws = wb["Test Scores"]
        self.assertEqual(
            [c.value for c in ws[2]], ["Crash Avoidance", "Frontal Collisions", "-", 40]
        )
        self.assertEqual(
            [c.value for c in ws[3]], ["Crash Avoidance", "Lane Departure", "-", 10]
        )

        ws2 = wb["CP - Dummy Scores"]
        self.assertEqual(
            [c.value for c in ws2[2]], ["Frontal Impact", "THOR-50", "-", "-", 5]
        )

    def test_preserves_header_and_unrequested_columns(self):
        wb = self._make_wb()

        common.blank_computed_columns(wb, {"Test Scores": ["Score"]})

        ws = wb["Test Scores"]
        self.assertEqual(
            [c.value for c in ws[1]], ["Stage", "Stage element", "Score", "Max score"]
        )
        # "Max score" and the identity columns must survive untouched.
        self.assertEqual(ws.cell(row=2, column=1).value, "Crash Avoidance")
        self.assertEqual(ws.cell(row=2, column=2).value, "Frontal Collisions")
        self.assertEqual(ws.cell(row=2, column=4).value, 40)
        # A sheet not named in the mapping is untouched entirely.
        ws2 = wb["CP - Dummy Scores"]
        self.assertEqual(ws2.cell(row=2, column=4).value, None)

    def test_missing_sheet_is_skipped_without_error(self):
        wb = self._make_wb()

        common.blank_computed_columns(
            wb, {"Category Scores": ["Score"]}
        )  # must not raise

        # The named sheet doesn't exist in wb; the existing sheet is untouched.
        self.assertEqual(wb["Test Scores"].cell(row=2, column=3).value, 0)

    def test_missing_column_is_skipped_without_error(self):
        wb = self._make_wb()

        common.blank_computed_columns(
            wb, {"Test Scores": ["Nonexistent column", "Score"]}
        )

        ws = wb["Test Scores"]
        self.assertEqual(ws.cell(row=2, column=3).value, "-")

    def test_does_not_write_beyond_populated_bounds(self):
        wb = self._make_wb()  # only 2 data rows (rows 2-3)

        common.blank_computed_columns(wb, {"Test Scores": ["Score"]})

        ws = wb["Test Scores"]
        self.assertIsNone(ws.cell(row=10, column=3).value)


class TestResetComputedColumnsToDash(unittest.TestCase):
    """common.reset_computed_columns_to_dash is the DataFrame-native sibling
    of blank_computed_columns, shared by every domain's preprocess (both the
    CLI path and the DataFrame path)."""

    def _make_dfs(self):
        return {
            "Test Scores": pd.DataFrame(
                {
                    "Stage": ["Crash Avoidance", "Safe Driving"],
                    "Score": [0, 0],
                    "Max score": [40, 30],
                }
            ),
            "CP - Dummy Scores": pd.DataFrame(
                {"Dummy": ["THOR-50"], "Capping?": [""], "Score": [np.nan]}
            ),
            "Input parameters": pd.DataFrame(
                {"Input parameter": ["VIN"], "Value": ["x"]}
            ),
        }

    def test_sets_requested_columns_across_multiple_sheets(self):
        dfs = self._make_dfs()

        result = common.reset_computed_columns_to_dash(
            dfs, {"Test Scores": ["Score"], "CP - Dummy Scores": ["Capping?", "Score"]}
        )

        self.assertTrue((result["Test Scores"]["Score"] == "-").all())
        self.assertTrue((result["CP - Dummy Scores"]["Score"] == "-").all())
        self.assertTrue((result["CP - Dummy Scores"]["Capping?"] == "-").all())
        # Non-requested columns and sheets survive untouched.
        self.assertEqual(list(result["Test Scores"]["Max score"]), [40, 30])
        self.assertEqual(list(result["Input parameters"]["Value"]), ["x"])

    def test_missing_sheet_and_column_are_skipped(self):
        dfs = self._make_dfs()

        result = common.reset_computed_columns_to_dash(
            dfs, {"Category Scores": ["Score"], "Test Scores": ["Nonexistent", "Score"]}
        )  # must not raise

        self.assertNotIn("Category Scores", result)
        self.assertTrue((result["Test Scores"]["Score"] == "-").all())

    def test_input_dfs_not_mutated(self):
        dfs = self._make_dfs()

        common.reset_computed_columns_to_dash(dfs, {"Test Scores": ["Score"]})

        self.assertEqual(list(dfs["Test Scores"]["Score"]), [0, 0])


class TestNumericizeComputedColumns(unittest.TestCase):
    """common.numericize_computed_columns undoes the "-" placeholder for
    callers that need the pristine template's computed columns numeric
    (each domain's _pristine_dfs, overall's calculate_score)."""

    def _make_dfs(self):
        return {
            "Test Scores": pd.DataFrame(
                {
                    "Stage": ["Crash Avoidance", "Safe Driving"],
                    "Score": ["-", 2.5],
                    "Max score": [40, 30],
                }
            ),
            "Rating": pd.DataFrame({"Star rating": ["-"]}),
        }

    def test_dash_becomes_zero_float64(self):
        dfs = self._make_dfs()

        result = common.numericize_computed_columns(
            dfs, {"Test Scores": ["Score"], "Rating": ["Star rating"]}
        )

        self.assertEqual(result["Test Scores"]["Score"].dtype, "float64")
        self.assertEqual(list(result["Test Scores"]["Score"]), [0.0, 2.5])
        self.assertEqual(list(result["Rating"]["Star rating"]), [0.0])

    def test_missing_sheet_and_column_are_skipped(self):
        dfs = self._make_dfs()

        result = common.numericize_computed_columns(
            dfs, {"Category Scores": ["Score"], "Test Scores": ["Nonexistent", "Score"]}
        )  # must not raise

        self.assertNotIn("Category Scores", result)
        self.assertEqual(list(result["Test Scores"]["Score"]), [0.0, 2.5])

    def test_input_dfs_not_mutated(self):
        dfs = self._make_dfs()

        common.numericize_computed_columns(dfs, {"Test Scores": ["Score"]})

        self.assertEqual(list(dfs["Test Scores"]["Score"]), ["-", 2.5])


class TestGetVersionFromSheet(unittest.TestCase):
    def test_get_version_from_sheet_success(self):
        """Should read version from Version sheet."""
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            filepath = f.name

        try:
            import openpyxl

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Version"
            ws["A1"] = "Generated with version"
            ws["B1"] = "1.2.3"
            wb.save(filepath)

            result, message = common.get_version_from_sheet(
                filepath, common.CliCommand.GENERATE_TEMPLATE
            )
            self.assertEqual(result, "1.2.3")
            self.assertIn("successfully", message)
        finally:
            os.unlink(filepath)

    def test_get_version_from_sheet_missing_sheet(self):
        """Should return None if Version sheet doesn't exist."""
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            filepath = f.name

        try:
            import openpyxl

            wb = openpyxl.Workbook()
            wb.save(filepath)

            result, message = common.get_version_from_sheet(filepath)
            self.assertIsNone(result)
            self.assertIn("not found", message)
        finally:
            os.unlink(filepath)

    def test_get_version_from_sheet_preprocess_row(self):
        """Should read version from row 2 for PREPROCESS command."""
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            filepath = f.name

        try:
            import openpyxl

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Version"
            ws["A1"] = "Generated with version"
            ws["B1"] = "1.0.0"
            ws["A2"] = "Preprocessed with version"
            ws["B2"] = "2.0.0"
            wb.save(filepath)

            result, message = common.get_version_from_sheet(
                filepath, common.CliCommand.PREPROCESS
            )
            self.assertEqual(result, "2.0.0")
        finally:
            os.unlink(filepath)

    def test_get_version_from_sheet_compute_score_row(self):
        """Should read version from row 3 for COMPUTE_SCORE command."""
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            filepath = f.name

        try:
            import openpyxl

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Version"
            ws["A3"] = "Score computed with version"
            ws["B3"] = "3.0.0"
            wb.save(filepath)

            result, message = common.get_version_from_sheet(
                filepath, common.CliCommand.COMPUTE_SCORE
            )
            self.assertEqual(result, "3.0.0")
        finally:
            os.unlink(filepath)

    def test_get_version_from_sheet_labels_in_any_order(self):
        """Should locate the label regardless of which row it is on."""
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            filepath = f.name

        try:
            import openpyxl

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Version"
            ws["A1"] = "Score computed with version"
            ws["B1"] = "3.0.0"
            ws["A2"] = "Generated with version"
            ws["B2"] = "1.0.0"
            ws["A3"] = "Preprocessed with version"
            ws["B3"] = "2.0.0"
            wb.save(filepath)

            result, _ = common.get_version_from_sheet(
                filepath, common.CliCommand.GENERATE_TEMPLATE
            )
            self.assertEqual(result, "1.0.0")
            result, _ = common.get_version_from_sheet(
                filepath, common.CliCommand.PREPROCESS
            )
            self.assertEqual(result, "2.0.0")
            result, _ = common.get_version_from_sheet(
                filepath, common.CliCommand.COMPUTE_SCORE
            )
            self.assertEqual(result, "3.0.0")
        finally:
            os.unlink(filepath)

    def test_get_version_from_sheet_alternate_label_wording(self):
        """Should match a label that contains extra wording, e.g. 'Generated with SC version'."""
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            filepath = f.name

        try:
            import openpyxl

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Version"
            ws["A1"] = "Version"
            ws["B1"] = "1.5"
            ws["A2"] = "Date"
            ws["B2"] = "07/26"
            ws["A3"] = "Generated with SC version"
            ws["B3"] = "5.4.3"
            wb.save(filepath)

            result, message = common.get_version_from_sheet(
                filepath, common.CliCommand.GENERATE_TEMPLATE
            )
            self.assertEqual(result, "5.4.3")
            self.assertIn("successfully", message)
        finally:
            os.unlink(filepath)

    def test_get_version_from_sheet_label_not_found(self):
        """Should return an error message (not raise) when no row matches the expected label."""
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            filepath = f.name

        try:
            import openpyxl

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Version"
            ws["A1"] = "Some unrelated label"
            ws["B1"] = "1.0.0"
            wb.save(filepath)

            result, message = common.get_version_from_sheet(
                filepath, common.CliCommand.GENERATE_TEMPLATE
            )
            self.assertIsNone(result)
            self.assertIn("not found", message)
        finally:
            os.unlink(filepath)


class TestGetVersionFromWorkbook(unittest.TestCase):
    """Workbook-native sibling of get_version_from_sheet -- same cases,
    against an already-open Workbook instead of a file path."""

    def test_get_version_from_workbook_success(self):
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Version"
        ws["A1"] = "Generated with version"
        ws["B1"] = "1.2.3"

        result, message = common.get_version_from_workbook(
            wb, common.CliCommand.GENERATE_TEMPLATE
        )
        self.assertEqual(result, "1.2.3")
        self.assertIn("successfully", message)

    def test_get_version_from_workbook_missing_sheet(self):
        import openpyxl

        wb = openpyxl.Workbook()

        result, message = common.get_version_from_workbook(wb)
        self.assertIsNone(result)
        self.assertIn("not found", message)

    def test_get_version_from_workbook_preprocess_row(self):
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Version"
        ws["A1"] = "Generated with version"
        ws["B1"] = "1.0.0"
        ws["A2"] = "Preprocessed with version"
        ws["B2"] = "2.0.0"

        result, message = common.get_version_from_workbook(
            wb, common.CliCommand.PREPROCESS
        )
        self.assertEqual(result, "2.0.0")

    def test_get_version_from_workbook_label_not_found(self):
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Version"
        ws["A1"] = "Some unrelated label"
        ws["B1"] = "1.0.0"

        result, message = common.get_version_from_workbook(
            wb, common.CliCommand.GENERATE_TEMPLATE
        )
        self.assertIsNone(result)
        self.assertIn("not found", message)

    def test_get_version_from_sheet_delegates_to_workbook(self):
        """get_version_from_sheet should be a thin load_workbook + delegate wrapper."""
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            filepath = f.name

        try:
            import openpyxl

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Version"
            ws["A1"] = "Generated with version"
            ws["B1"] = "1.2.3"
            wb.save(filepath)

            result, message = common.get_version_from_sheet(
                filepath, common.CliCommand.GENERATE_TEMPLATE
            )
            self.assertEqual(result, "1.2.3")
            self.assertIn("successfully", message)
        finally:
            os.unlink(filepath)


class TestCheckVersion(unittest.TestCase):
    def test_check_version_mismatch_major_exits(self):
        """Should exit if major version doesn't match."""
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            filepath = f.name

        try:
            import openpyxl

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Version"
            ws["B1"] = "0.0.0"  # Different major version
            wb.save(filepath)

            with self.assertRaises(SystemExit):
                with patch("sys.stdout", new=MagicMock()):
                    common.check_version(filepath, common.CliCommand.GENERATE_TEMPLATE)
        finally:
            os.unlink(filepath)

    def test_check_version_mismatch_minor_exits(self):
        """Should exit if minor version doesn't match."""
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            filepath = f.name

        try:
            import openpyxl

            # Get current version parts
            current_parts = version.VERSION.split(".")
            major = current_parts[0]
            # Use a different minor version
            different_minor = (
                str(int(current_parts[1]) + 1) if len(current_parts) > 1 else "99"
            )

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Version"
            ws["B1"] = f"{major}.{different_minor}.0"
            wb.save(filepath)

            with self.assertRaises(SystemExit):
                with patch("sys.stdout", new=MagicMock()):
                    common.check_version(filepath, common.CliCommand.GENERATE_TEMPLATE)
        finally:
            os.unlink(filepath)

    def test_check_version_patch_mismatch_warns_but_continues(self):
        """Should warn but not exit if only patch version differs."""
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            filepath = f.name

        try:
            import openpyxl

            # Get current version parts
            current_parts = version.VERSION.split(".")
            major = current_parts[0]
            minor = current_parts[1] if len(current_parts) > 1 else "0"
            # Use a different patch version
            different_patch = (
                str(int(current_parts[2]) + 1) if len(current_parts) > 2 else "99"
            )

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Version"
            ws["A1"] = "Generated with version"
            ws["B1"] = f"{major}.{minor}.{different_patch}"
            wb.save(filepath)

            # Should not raise SystemExit, just warn
            with patch("sys.stdout", new=MagicMock()):
                common.check_version(filepath, common.CliCommand.GENERATE_TEMPLATE)
            # If we get here without exception, test passes
        finally:
            os.unlink(filepath)

    def test_check_version_exact_match_passes(self):
        """Should pass if version matches exactly."""
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            filepath = f.name

        try:
            import openpyxl

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Version"
            ws["A1"] = "Generated with version"
            ws["B1"] = version.VERSION
            wb.save(filepath)

            # Should not raise any exception
            common.check_version(filepath, common.CliCommand.GENERATE_TEMPLATE)
        finally:
            os.unlink(filepath)


class TestFindVersionMismatch(unittest.TestCase):
    """Non-exiting sibling of check_version: same major/minor-vs-patch
    policy, surfaced as a finding instead of sys.exit, plus a special case
    for the known-incompatible 5.4.0 template."""

    def _wb_with_version(self, version_string, label="Generated with version"):
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Version"
        ws["A1"] = label
        ws["B1"] = version_string
        return wb

    def test_exact_match_returns_no_finding(self):
        wb = self._wb_with_version(version.VERSION)
        self.assertEqual(common.find_version_mismatch(wb), [])

    def test_no_version_sheet_returns_no_finding(self):
        import openpyxl

        wb = openpyxl.Workbook()
        self.assertEqual(common.find_version_mismatch(wb), [])

    def test_label_not_found_returns_no_finding(self):
        wb = self._wb_with_version("1.0.0", label="Some unrelated label")
        self.assertEqual(common.find_version_mismatch(wb), [])

    def test_new_template_version_label_is_recognized(self):
        with patch.object(common, "VERSION", "5.4.6"):
            wb = self._wb_with_version("6.0.0", label="Template version")
            findings = common.find_version_mismatch(wb)

        self.assertEqual(len(findings), 1)
        self.assertEqual(findings[0].check, common.CHECK_VERSION_MISMATCH)

    def test_legacy_generated_with_label_still_recognized(self):
        """Files stamped by versions <= 5.4.5 carry "Generated with version"
        instead of "Template version" and must keep being read."""
        with patch.object(common, "VERSION", "5.4.6"):
            wb = self._wb_with_version("6.0.0", label="Generated with version")
            findings = common.find_version_mismatch(wb)

        self.assertEqual(len(findings), 1)

    def test_get_version_reads_both_template_label_generations(self):
        for label in ("Template version", "Generated with version"):
            with self.subTest(label=label):
                wb = self._wb_with_version("1.2.3", label=label)
                found, message = common.get_version_from_workbook(
                    wb, common.CliCommand.GENERATE_TEMPLATE
                )
                self.assertEqual(found, "1.2.3", message)

    def test_major_minor_mismatch_returns_finding(self):
        with patch.object(common, "VERSION", "5.4.6"):
            wb = self._wb_with_version("6.0.0")
            findings = common.find_version_mismatch(wb)

        self.assertEqual(len(findings), 1)
        finding = findings[0]
        self.assertEqual(finding.check, common.CHECK_VERSION_MISMATCH)
        self.assertEqual(finding.sheet, "Version")
        self.assertEqual(finding.cell, "B1")
        self.assertIn("6.0.0", finding.message)
        self.assertIn("5.4.6", finding.message)

    def test_ordinary_patch_mismatch_returns_no_finding(self):
        """A generic patch bump is not treated as breaking -- same
        warn-and-proceed policy as check_version."""
        with patch.object(common, "VERSION", "5.4.6"):
            wb = self._wb_with_version("5.4.5")
            findings = common.find_version_mismatch(wb)

        self.assertEqual(findings, [])

    def test_known_incompatible_5_4_0_flagged_at_5_4_4(self):
        """5.4.0 is a patch-only difference from 5.4.4+ but structurally
        incompatible -- special-cased regardless."""
        with patch.object(common, "VERSION", "5.4.4"):
            wb = self._wb_with_version("5.4.0")
            findings = common.find_version_mismatch(wb)

        self.assertEqual(len(findings), 1)
        self.assertEqual(findings[0].check, common.CHECK_VERSION_MISMATCH)
        self.assertIn("5.4.0", findings[0].message)
        self.assertIn("5.4.4", findings[0].message)

    def test_known_incompatible_5_4_0_flagged_above_5_4_4(self):
        with patch.object(common, "VERSION", "5.4.6"):
            wb = self._wb_with_version("5.4.0")
            findings = common.find_version_mismatch(wb)

        self.assertEqual(len(findings), 1)

    def test_5_4_0_not_flagged_before_5_4_4(self):
        """The special case only applies once the running library reached
        the version that actually broke the template (5.4.4)."""
        with patch.object(common, "VERSION", "5.4.3"):
            wb = self._wb_with_version("5.4.0")
            findings = common.find_version_mismatch(wb)

        self.assertEqual(findings, [])


class TestAddVersionSheet(unittest.TestCase):
    def test_add_version_sheet(self):
        """Should add Version sheet with correct content."""
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            filepath = f.name

        try:
            import openpyxl

            wb = openpyxl.Workbook()
            wb.save(filepath)

            common.add_version_sheet(filepath, common.CliCommand.GENERATE_TEMPLATE)

            wb = openpyxl.load_workbook(filepath)
            self.assertIn("Version", wb.sheetnames)
            self.assertEqual(wb["Version"]["A1"].value, "Template version")
            self.assertEqual(wb["Version"]["B1"].value, version.VERSION)
        finally:
            os.unlink(filepath)

    def test_add_version_sheet_preprocess(self):
        """Should add version info in row 2 for PREPROCESS command."""
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            filepath = f.name

        try:
            import openpyxl

            wb = openpyxl.Workbook()
            wb.save(filepath)

            common.add_version_sheet(filepath, common.CliCommand.PREPROCESS)

            wb = openpyxl.load_workbook(filepath)
            self.assertIn("Version", wb.sheetnames)
            self.assertEqual(wb["Version"]["A2"].value, "Preprocessed with version")
            self.assertEqual(wb["Version"]["B2"].value, version.VERSION)
        finally:
            os.unlink(filepath)

    def test_add_version_sheet_compute_score(self):
        """Should add version info in row 3 for COMPUTE_SCORE command."""
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            filepath = f.name

        try:
            import openpyxl

            wb = openpyxl.Workbook()
            wb.save(filepath)

            common.add_version_sheet(filepath, common.CliCommand.COMPUTE_SCORE)

            wb = openpyxl.load_workbook(filepath)
            self.assertIn("Version", wb.sheetnames)
            self.assertEqual(wb["Version"]["A3"].value, "Score computed with version")
            self.assertEqual(wb["Version"]["B3"].value, version.VERSION)
        finally:
            os.unlink(filepath)

    def test_add_version_sheet_appends_to_existing(self):
        """Should append to existing Version sheet without overwriting."""
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            filepath = f.name

        try:
            import openpyxl

            wb = openpyxl.Workbook()
            wb.save(filepath)

            # Add first version entry
            common.add_version_sheet(filepath, common.CliCommand.GENERATE_TEMPLATE)
            # Add second version entry
            common.add_version_sheet(filepath, common.CliCommand.PREPROCESS)

            wb = openpyxl.load_workbook(filepath)
            self.assertEqual(wb["Version"]["A1"].value, "Template version")
            self.assertEqual(wb["Version"]["B1"].value, version.VERSION)
            self.assertEqual(wb["Version"]["A2"].value, "Preprocessed with version")
            self.assertEqual(wb["Version"]["B2"].value, version.VERSION)
        finally:
            os.unlink(filepath)


class TestVersionInfo(unittest.TestCase):
    def test_version_info_from_string(self):
        """Should parse version string correctly."""
        v = common.VersionInfo.from_string("1.2.3")
        self.assertEqual(v.major, 1)
        self.assertEqual(v.minor, 2)
        self.assertEqual(v.patch, 3)

    def test_version_info_str(self):
        """Should convert back to string correctly."""
        v = common.VersionInfo(1, 2, 3)
        self.assertEqual(str(v), "1.2.3")

    def test_version_info_invalid_format_raises(self):
        """Should raise ValueError for invalid format."""
        with self.assertRaises(ValueError):
            common.VersionInfo.from_string("1.2")
        with self.assertRaises(ValueError):
            common.VersionInfo.from_string("1.2.3.4")
        with self.assertRaises(ValueError):
            common.VersionInfo.from_string("invalid")


class TestParseVersion(unittest.TestCase):
    def test_parse_version(self):
        """Should parse version string into VersionInfo."""
        v = common.parse_version("10.20.30")
        self.assertEqual(v.major, 10)
        self.assertEqual(v.minor, 20)
        self.assertEqual(v.patch, 30)


class TestCliCommand(unittest.TestCase):
    def test_cli_command_values(self):
        """CliCommand enum should have expected values."""
        self.assertEqual(common.CliCommand.UNKNOWN.value, "unknown")
        self.assertEqual(common.CliCommand.GENERATE_TEMPLATE.value, "generate-template")
        self.assertEqual(common.CliCommand.PREPROCESS.value, "preprocess")
        self.assertEqual(common.CliCommand.COMPUTE_SCORE.value, "compute-score")


class TestGetVersionLabelForCommand(unittest.TestCase):
    def test_get_version_label_for_generate_template(self):
        """Should return correct label for GENERATE_TEMPLATE."""
        label = common.get_version_label_for_command(
            common.CliCommand.GENERATE_TEMPLATE
        )
        self.assertEqual(label, "Template version")

    def test_get_version_label_for_preprocess(self):
        """Should return correct label for PREPROCESS."""
        label = common.get_version_label_for_command(common.CliCommand.PREPROCESS)
        self.assertEqual(label, "Preprocessed with version")

    def test_get_version_label_for_compute_score(self):
        """Should return correct label for COMPUTE_SCORE."""
        label = common.get_version_label_for_command(common.CliCommand.COMPUTE_SCORE)
        self.assertEqual(label, "Score computed with version")


class TestGetVersionRowForCommand(unittest.TestCase):
    def test_get_version_row_for_generate_template(self):
        """Should return row 1 for GENERATE_TEMPLATE."""
        row = common.get_version_row_for_command(common.CliCommand.GENERATE_TEMPLATE)
        self.assertEqual(row, 1)

    def test_get_version_row_for_preprocess(self):
        """Should return row 2 for PREPROCESS."""
        row = common.get_version_row_for_command(common.CliCommand.PREPROCESS)
        self.assertEqual(row, 2)

    def test_get_version_row_for_compute_score(self):
        """Should return row 3 for COMPUTE_SCORE."""
        row = common.get_version_row_for_command(common.CliCommand.COMPUTE_SCORE)
        self.assertEqual(row, 3)


class TestOppositeFfill(unittest.TestCase):
    def test_opposite_ffill(self):
        df = pd.DataFrame({"A": [1, 1, 2, 2, 3], "Value": ["x"] * 5})
        result = common.opposite_ffill(df.copy())
        self.assertTrue(pd.isna(result.loc[1, "A"]))
        self.assertEqual(result.loc[2, "A"], 2)


def _write_row(ws, row_idx, values_by_header, header):
    for col_idx, column in enumerate(header, start=1):
        if column in values_by_header:
            ws.cell(row=row_idx, column=col_idx, value=values_by_header[column])


class TestFindMissingRequiredInputs(unittest.TestCase):
    """find_missing_required_inputs is schema-driven: required cells are
    declared in code (GreyCellSheetSchema.prediction_required_columns), not
    read from the submitted file's own fill colors. See
    crash_protection/integrity.py's SHEET_SCHEMAS for a real example."""

    def _input_params_sheet(self, wb, rows):
        ws = wb.active
        ws.title = "Input parameters"
        header = ["Stage element", "Stage subelement", "Input parameter", "Value"]
        for col_idx, column in enumerate(header, start=1):
            ws.cell(row=1, column=col_idx, value=column)
        for i, row in enumerate(rows, start=2):
            _write_row(ws, i, row, header)
        return ws

    def _schema(self, applies_to_stage_elements=None):
        return {
            "Input parameters": common.GreyCellSheetSchema(
                identity_columns=[
                    "Stage element",
                    "Stage subelement",
                    "Input parameter",
                ],
                grey_columns=["Value"],
                prediction_required_columns=["Value"],
                applies_to_stage_elements=applies_to_stage_elements,
            )
        }

    def test_reports_empty_required_cell(self):
        import openpyxl

        wb = openpyxl.Workbook()
        self._input_params_sheet(
            wb,
            [
                {
                    "Stage element": "Frontal Impact",
                    "Stage subelement": "Offset",
                    "Input parameter": "Mass",
                }
            ],
        )

        missing = common.find_missing_required_inputs(wb, self._schema())

        self.assertEqual(len(missing), 1)
        entry = missing[0]
        self.assertEqual(entry.sheet, "Input parameters")
        self.assertEqual(entry.column, "Value")
        self.assertEqual(entry.cell, "D2")
        self.assertEqual(
            entry.row_identity,
            {
                "Stage element": "Frontal Impact",
                "Stage subelement": "Offset",
                "Input parameter": "Mass",
            },
        )

    def test_missing_input_description_is_semantic_not_a_cell_reference(self):
        import openpyxl

        wb = openpyxl.Workbook()
        self._input_params_sheet(
            wb,
            [
                {
                    "Stage element": "Frontal Impact",
                    "Stage subelement": "Offset",
                    "Input parameter": "Mass",
                }
            ],
        )

        missing = common.find_missing_required_inputs(wb, self._schema())

        self.assertEqual(
            str(missing[0]),
            "Input parameters — Frontal Impact / Offset / Mass / Value is missing",
        )

    def test_ignores_filled_required_cell(self):
        import openpyxl

        wb = openpyxl.Workbook()
        self._input_params_sheet(
            wb,
            [
                {
                    "Stage element": "Frontal Impact",
                    "Stage subelement": "Offset",
                    "Input parameter": "Mass",
                    "Value": "1500",
                }
            ],
        )

        missing = common.find_missing_required_inputs(wb, self._schema())

        self.assertEqual(missing, [])

    def test_skips_row_with_blank_identity(self):
        """A row whose own identity (here: Input parameter name) is blank is
        not a real field to fill in, even if its Value cell is grey/empty --
        e.g. a row with no Input parameter name at all."""
        import openpyxl

        wb = openpyxl.Workbook()
        self._input_params_sheet(
            wb,
            [
                {
                    "Stage element": "Frontal Impact",
                    "Stage subelement": "Offset",
                    "Input parameter": None,
                }
            ],
        )

        missing = common.find_missing_required_inputs(wb, self._schema())

        self.assertEqual(missing, [])

    def test_scopes_by_stage_element_and_subelement(self):
        import openpyxl

        wb = openpyxl.Workbook()
        self._input_params_sheet(
            wb,
            [
                {
                    "Stage element": "Frontal Impact",
                    "Stage subelement": "Offset",
                    "Input parameter": "Mass",
                },
                {
                    "Stage element": "Side Impact",
                    "Stage subelement": "MDB",
                    "Input parameter": "Speed",
                },
            ],
        )

        missing = common.find_missing_required_inputs(
            wb,
            self._schema(),
            stage_element="Frontal Impact",
            stage_subelement="Offset",
        )

        self.assertEqual(len(missing), 1)
        self.assertEqual(missing[0].row_identity["Input parameter"], "Mass")

    def test_forward_fills_identity_across_rows(self):
        import openpyxl

        wb = openpyxl.Workbook()
        self._input_params_sheet(
            wb,
            [
                {
                    "Stage element": "Frontal Impact",
                    "Stage subelement": "Offset",
                    "Input parameter": "Mass",
                },
                {
                    "Input parameter": "Speed"
                },  # Stage element/subelement blank -> same as row above
            ],
        )

        missing = common.find_missing_required_inputs(wb, self._schema())

        self.assertEqual(len(missing), 2)
        self.assertEqual(missing[1].row_identity["Stage element"], "Frontal Impact")
        self.assertEqual(missing[1].row_identity["Stage subelement"], "Offset")

    def test_forward_fill_resets_when_an_ancestor_column_changes(self):
        """A blank "Scenario" cell must not leak a stale value forward once
        an ancestor column ("Category") has since changed -- regression for
        a real bug found against safe_driving's real Input parameters
        sheet, where "Speed control function" incorrectly inherited the
        unrelated "Child enters unlocked vehicle" scenario from several
        rows above once Category changed underneath it."""
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Input parameters"
        header = ["Category", "Scenario", "Input parameter", "Value"]
        for col_idx, column in enumerate(header, start=1):
            ws.cell(row=1, column=col_idx, value=column)
        rows = [
            {
                "Category": "Child presence detection",
                "Scenario": "Child left behind",
                "Input parameter": "Type of system",
            },
            {
                "Input parameter": "Seat coverage"
            },  # inherits Scenario "Child left behind"
            {
                "Category": "Speed control function",
                "Input parameter": "Type of system",
            },  # no Scenario of its own
        ]
        for i, row in enumerate(rows, start=2):
            _write_row(ws, i, row, header)

        schema = {
            "Input parameters": common.GreyCellSheetSchema(
                identity_columns=["Category", "Scenario", "Input parameter"],
                grey_columns=["Value"],
            )
        }
        missing = common.find_missing_required_inputs(wb, schema)

        self.assertEqual(len(missing), 3)
        self.assertEqual(missing[1].row_identity["Scenario"], "Child left behind")
        last_entry = missing[2]
        self.assertEqual(last_entry.row_identity["Category"], "Speed control function")
        self.assertIsNone(last_entry.row_identity["Scenario"])
        self.assertNotIn("Child left behind", last_entry.description)

    def test_applies_to_stage_elements_scopes_whole_sheet(self):
        """Sheets with no per-row Stage element/subelement column of their
        own (e.g. CP's "CP - Frontal Offset") are scoped as a whole via
        applies_to_stage_elements, not by row identity."""
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "CP - Frontal Offset"
        header = ["Criteria", "Value"]
        for col_idx, column in enumerate(header, start=1):
            ws.cell(row=1, column=col_idx, value=column)
        ws.cell(row=2, column=1, value="HIC15")

        schema = {
            "CP - Frontal Offset": common.GreyCellSheetSchema(
                identity_columns=["Criteria"],
                grey_columns=["Value"],
                prediction_required_columns=["Value"],
                applies_to_stage_elements=[("Frontal Impact", "Offset")],
            )
        }

        in_scope = common.find_missing_required_inputs(
            wb, schema, stage_element="Frontal Impact", stage_subelement="Offset"
        )
        out_of_scope = common.find_missing_required_inputs(
            wb, schema, stage_element="Side Impact", stage_subelement="MDB"
        )

        self.assertEqual(len(in_scope), 1)
        self.assertEqual(out_of_scope, [])

    def test_ignores_non_anchor_merged_cells(self):
        import openpyxl

        wb = openpyxl.Workbook()
        ws = self._input_params_sheet(
            wb,
            [
                {
                    "Stage element": "Frontal Impact",
                    "Stage subelement": "Offset",
                    "Input parameter": "Mass",
                }
            ],
        )
        # Merge the identity column (anchor, has a real value) with the
        # required "Value" column, so the Value cell (D2) becomes a
        # non-anchor MergedCell -- it must be skipped, not reported empty.
        ws.merge_cells("C2:D2")

        missing = common.find_missing_required_inputs(wb, self._schema())

        self.assertEqual(missing, [])


class TestValidateScope(unittest.TestCase):
    PAIRS = [("Frontal Impact", "Offset"), ("Side Impact", "MDB")]

    def test_no_scope_is_valid(self):
        common.validate_scope(None, None, self.PAIRS, "demo")

    def test_declared_pair_is_valid(self):
        common.validate_scope("Frontal Impact", "Offset", self.PAIRS, "demo")

    def test_element_alone_is_validated(self):
        common.validate_scope("Side Impact", None, self.PAIRS, "demo")
        with self.assertRaises(ValueError):
            common.validate_scope("frontal_impact", None, self.PAIRS, "demo")

    def test_unknown_subelement_raises_and_lists_declared_pairs(self):
        with self.assertRaises(ValueError) as ctx:
            common.validate_scope(None, "offset", self.PAIRS, "demo")
        self.assertIn("Offset", str(ctx.exception))

    def test_subelement_under_wrong_element_raises(self):
        with self.assertRaises(ValueError):
            common.validate_scope("Side Impact", "Offset", self.PAIRS, "demo")


class TestRequiredWhenColumnsPresent(unittest.TestCase):
    """GreyCellSheetSchema.required_when_columns_present gates a row's
    required columns on reference columns (e.g. HPL/LPL) being present in
    the PRISTINE template's counterpart row -- never in the user's file."""

    HEADER = ["Criteria", "HPL", "LPL", "OEM Prediction"]

    def _criteria_sheet(self, wb, rows):
        import openpyxl

        ws = wb.active
        ws.title = "Criteria"
        for col_idx, column in enumerate(self.HEADER, start=1):
            ws.cell(row=1, column=col_idx, value=column)
        for row_idx, row in enumerate(rows, start=2):
            for col_idx, column in enumerate(self.HEADER, start=1):
                ws.cell(row=row_idx, column=col_idx, value=row.get(column))
        return ws

    def _wbs(self):
        import openpyxl

        rows = [
            {"Criteria": "HIC15", "HPL": 500, "LPL": 700},
            {"Criteria": "Modifier"},  # no limits -> not OEM-predictable
        ]
        pristine_wb = openpyxl.Workbook()
        self._criteria_sheet(pristine_wb, rows)
        user_wb = openpyxl.Workbook()
        self._criteria_sheet(user_wb, rows)
        return pristine_wb, user_wb

    def _schema(self):
        return {
            "Criteria": common.GreyCellSheetSchema(
                identity_columns=["Criteria"],
                grey_columns=["OEM Prediction"],
                prediction_required_columns=["OEM Prediction"],
                required_when_columns_present=["HPL", "LPL"],
            )
        }

    def test_rows_without_gate_columns_are_not_required(self):
        pristine_wb, user_wb = self._wbs()

        missing = common.find_missing_required_inputs(
            user_wb, self._schema(), pristine_wb=pristine_wb
        )

        self.assertEqual([m.row_identity["Criteria"] for m in missing], ["HIC15"])

    def test_gate_reads_pristine_not_user_file(self):
        """Blanking the reference columns in the submitted file must not
        make its row optional."""
        pristine_wb, user_wb = self._wbs()
        ws = user_wb["Criteria"]
        ws["B2"] = None  # HPL of the HIC15 row
        ws["C2"] = None  # LPL of the HIC15 row

        missing = common.find_missing_required_inputs(
            user_wb, self._schema(), pristine_wb=pristine_wb
        )

        self.assertEqual([m.row_identity["Criteria"] for m in missing], ["HIC15"])

    def test_gate_fails_closed_without_pristine_workbook(self):
        _, user_wb = self._wbs()

        missing = common.find_missing_required_inputs(user_wb, self._schema())

        self.assertEqual(
            [m.row_identity["Criteria"] for m in missing], ["HIC15", "Modifier"]
        )


class TestFindMissingRequiredInputsInFile(unittest.TestCase):
    def test_round_trips_through_a_real_file(self):
        import openpyxl

        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            filepath = f.name

        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Input parameters"
            header = ["Stage element", "Stage subelement", "Input parameter", "Value"]
            for col_idx, column in enumerate(header, start=1):
                ws.cell(row=1, column=col_idx, value=column)
            _write_row(
                ws,
                2,
                {
                    "Stage element": "Frontal Impact",
                    "Stage subelement": "Offset",
                    "Input parameter": "Mass",
                },
                header,
            )
            wb.save(filepath)

            schemas = {
                "Input parameters": common.GreyCellSheetSchema(
                    identity_columns=[
                        "Stage element",
                        "Stage subelement",
                        "Input parameter",
                    ],
                    grey_columns=["Value"],
                    prediction_required_columns=["Value"],
                )
            }
            missing = common.find_missing_required_inputs_in_file(filepath, schemas)

            self.assertEqual(len(missing), 1)
            self.assertEqual(missing[0].sheet, "Input parameters")
            self.assertEqual(missing[0].cell, "D2")
        finally:
            os.unlink(filepath)


class TestFindMissingGridInputs(unittest.TestCase):
    """find_missing_grid_inputs backs the dynamic/matrix sheets (e.g.
    crash_avoidance's "...pred."/"...robust. pred.") that have no
    forward-fillable identity columns for GreyCellSheetSchema -- required
    cells come from the pristine template's own Data Validation ranges."""

    def _grid_wb(self):
        import openpyxl
        from openpyxl.worksheet.datavalidation import DataValidation

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Grid"
        ws["A1"] = "Test name"
        ws["B1"] = "10 km/h"
        ws["C1"] = "20 km/h"
        ws["A2"] = "Scenario 1"
        ws["A3"] = "Scenario 2"

        dv = DataValidation(type="list", formula1='"N/A,Green,Red"')
        dv.add("B2:C3")
        ws.add_data_validation(dv)
        return wb, ws

    def test_get_required_dropdown_coordinates_reads_list_validations(self):
        _, ws = self._grid_wb()

        coords = common.get_required_dropdown_coordinates(ws)

        self.assertEqual(coords, {"B2", "C2", "B3", "C3"})

    def test_reports_empty_required_cell_with_inferred_identity(self):
        pristine_wb, pristine_ws = self._grid_wb()
        user_wb, user_ws = self._grid_wb()
        user_ws["B2"] = "Green"
        # C2 left blank -- should be reported.

        missing = common.find_missing_grid_inputs(user_ws, pristine_ws, "Grid")

        self.assertEqual(len(missing), 3)
        entry = next(m for m in missing if m.cell == "C2")
        self.assertEqual(entry.sheet, "Grid")
        self.assertEqual(entry.column, "20 km/h")
        self.assertEqual(entry.row_identity, {"A": "Scenario 1"})

    def test_filled_text_cell_is_not_reported(self):
        pristine_wb, pristine_ws = self._grid_wb()
        user_wb, user_ws = self._grid_wb()
        for coord in ("B2", "C2", "B3", "C3"):
            user_ws[coord] = "Green"

        self.assertEqual(
            common.find_missing_grid_inputs(user_ws, pristine_ws, "Grid"), []
        )

    def test_colored_but_textless_cell_is_not_reported(self):
        """After preprocess, report_writer converts the OEM's typed
        dropdown value into a background fill and clears the text -- a
        colored cell must count as answered even with no text left."""
        from openpyxl.styles import PatternFill

        pristine_wb, pristine_ws = self._grid_wb()
        user_wb, user_ws = self._grid_wb()
        for coord in ("B2", "C2", "B3", "C3"):
            user_ws[coord].fill = PatternFill(
                start_color="FF009933", end_color="FF009933", fill_type="solid"
            )

        self.assertEqual(
            common.find_missing_grid_inputs(user_ws, pristine_ws, "Grid"), []
        )

    def test_alpha_prefixed_prediction_color_fill_is_not_reported(self):
        """format_prediction_cells writes 6-digit start_colors, which
        openpyxl stores with a 00 alpha prefix -- both alpha spellings of a
        real prediction color must count as answered."""
        from openpyxl.styles import PatternFill

        pristine_wb, pristine_ws = self._grid_wb()
        user_wb, user_ws = self._grid_wb()
        for coord in ("B2", "C2", "B3", "C3"):
            user_ws[coord].fill = PatternFill(
                start_color="00009933", end_color="00009933", fill_type="solid"
            )

        self.assertEqual(
            common.find_missing_grid_inputs(user_ws, pristine_ws, "Grid"), []
        )

    def test_grey_placeholder_fill_is_reported(self):
        """Several sheets paint their *unanswered* dropdown cells literal
        grey (PREDICTION_COLOR_MAP[GREY]) rather than a theme fill, and no
        dropdown offers "Grey" as an option -- a grey fill must never count
        as an OEM answer."""
        from openpyxl.styles import PatternFill

        pristine_wb, pristine_ws = self._grid_wb()
        user_wb, user_ws = self._grid_wb()
        for coord in ("B2", "C2", "B3", "C3"):
            user_ws[coord].fill = PatternFill(
                start_color="FF808080", end_color="FF808080", fill_type="solid"
            )

        missing = common.find_missing_grid_inputs(user_ws, pristine_ws, "Grid")

        self.assertEqual(len(missing), 4)

    def test_input_cell_paint_is_reported(self):
        """The D9D9D9 "you must fill this in" input paint is not a
        prediction color -- a cell carrying only that fill is unanswered."""
        from openpyxl.styles import PatternFill

        pristine_wb, pristine_ws = self._grid_wb()
        user_wb, user_ws = self._grid_wb()
        for coord in ("B2", "C2", "B3", "C3"):
            user_ws[coord].fill = PatternFill(
                start_color="FFD9D9D9", end_color="FFD9D9D9", fill_type="solid"
            )

        missing = common.find_missing_grid_inputs(user_ws, pristine_ws, "Grid")

        self.assertEqual(len(missing), 4)

    def test_non_prediction_color_fill_is_reported(self):
        """An arbitrary fill that maps to no PREDICTION_COLOR_MAP entry is
        not an answer -- only real prediction colors count."""
        from openpyxl.styles import PatternFill

        pristine_wb, pristine_ws = self._grid_wb()
        user_wb, user_ws = self._grid_wb()
        for coord in ("B2", "C2", "B3", "C3"):
            user_ws[coord].fill = PatternFill(
                start_color="FF123456", end_color="FF123456", fill_type="solid"
            )

        missing = common.find_missing_grid_inputs(user_ws, pristine_ws, "Grid")

        self.assertEqual(len(missing), 4)

    def test_default_theme_fill_with_no_text_is_reported(self):
        missing_pristine, missing_pristine_ws = self._grid_wb()
        missing_user_wb, missing_user_ws = self._grid_wb()

        missing = common.find_missing_grid_inputs(
            missing_user_ws, missing_pristine_ws, "Grid"
        )

        self.assertEqual(len(missing), 4)

    def test_required_columns_derived_from_pristine_not_from_user_file(self):
        """Removing the dropdown/fill from the user's own file must not
        make a cell optional -- required-ness only ever comes from the
        pristine template."""
        pristine_wb, pristine_ws = self._grid_wb()
        user_wb, user_ws = self._grid_wb()
        user_ws.data_validations.dataValidation.clear()

        missing = common.find_missing_grid_inputs(user_ws, pristine_ws, "Grid")

        self.assertEqual(len(missing), 4)

    def test_dv_range_outside_populated_bounds_is_not_required(self):
        """A Data Validation sqref can carry a stray coordinate pointing
        into empty space beyond the sheet's content (a copy/paste
        artifact). Such a cell must not become an unsatisfiable
        requirement -- the required set is bounded to the pristine
        sheet's populated area."""
        from openpyxl.worksheet.datavalidation import DataValidation

        pristine_wb, pristine_ws = self._grid_wb()
        stray = DataValidation(type="list", formula1='"N/A,Green,Red"')
        stray.add("F15")  # populated area is A1:C3
        pristine_ws.add_data_validation(stray)
        user_wb, user_ws = self._grid_wb()

        missing = common.find_missing_grid_inputs(user_ws, pristine_ws, "Grid")

        self.assertEqual({m.cell for m in missing}, {"B2", "C2", "B3", "C3"})

    def test_no_required_coordinates_returns_empty(self):
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        self.assertEqual(common.find_missing_grid_inputs(ws, ws, "Grid"), [])

    def test_labels_come_from_pristine_not_from_neighbouring_answers(self):
        """On a text-shape user file, the cell above a gap is often another
        required cell holding an *answer* ("Green") -- the inferred column
        label must be the real header from the pristine sheet, never a
        neighbour's answer, and an edited label in the user's file must
        not misdescribe the finding."""
        pristine_wb, pristine_ws = self._grid_wb()
        user_wb, user_ws = self._grid_wb()
        user_ws["B2"] = "Green"  # answered cell directly above the gap
        user_ws["A3"] = "Tampered label"  # edited row label
        # B3 left blank -- should be reported with pristine-derived labels.

        missing = common.find_missing_grid_inputs(user_ws, pristine_ws, "Grid")

        entry = next(m for m in missing if m.cell == "B3")
        self.assertEqual(entry.column, "10 km/h")
        self.assertEqual(entry.row_identity, {"A": "Scenario 2"})

    def test_rows_filter_restricts_required_cells(self):
        pristine_wb, pristine_ws = self._grid_wb()
        user_wb, user_ws = self._grid_wb()

        missing = common.find_missing_grid_inputs(
            user_ws, pristine_ws, "Grid", rows={2}
        )

        self.assertEqual({m.cell for m in missing}, {"B2", "C2"})


class TestRowsInSections(unittest.TestCase):
    """rows_in_sections partitions a sheet into its labelled blocks (e.g.
    "CP - VRU Prediction"'s "Headforms"/"Legform") by the section headers'
    own text in column A -- never by declared row numbers."""

    def _sectioned_ws(self):
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "Headforms"
        ws["A3"] = "WAD on Centreline"
        ws["B4"] = "row data"
        ws["A6"] = "Legform"
        ws["A7"] = "Upper legform"
        ws["B8"] = "row data"
        return ws

    def test_partitions_rows_between_declared_headers(self):
        ws = self._sectioned_ws()

        sections = common.rows_in_sections(ws, ["Headforms", "Legform"])

        self.assertEqual(sections["Headforms"], {1, 2, 3, 4, 5})
        self.assertEqual(sections["Legform"], {6, 7, 8})

    def test_inner_labels_are_not_section_boundaries(self):
        """Column-A values that aren't declared headers ("WAD on
        Centreline", "Upper legform") must not end a section."""
        ws = self._sectioned_ws()

        sections = common.rows_in_sections(ws, ["Headforms", "Legform"])

        self.assertIn(4, sections["Headforms"])
        self.assertIn(8, sections["Legform"])

    def test_missing_declared_header_raises(self):
        """A declared section that can't be found must fail loudly -- a
        silently empty section would make all its cells optional."""
        ws = self._sectioned_ws()

        with self.assertRaises(common.TemplateIntegrityError):
            common.rows_in_sections(ws, ["Headforms", "Nonexistent block"])


# =============================================================================
# Tests for cli.py
# =============================================================================
class TestCli(unittest.TestCase):
    def test_cli_help(self):
        """CLI should show help message."""
        runner = CliRunner()
        result = runner.invoke(cli, ["--help"])
        self.assertEqual(result.exit_code, 0)
        self.assertIn("Euro NCAP Rating Calculator 2026", result.output)

    def test_cli_has_crash_protection_command(self):
        """CLI should have crash_protection subcommand."""
        runner = CliRunner()
        result = runner.invoke(cli, ["--help"])
        self.assertIn("crash_protection", result.output)

    def test_cli_has_crash_avoidance_command(self):
        """CLI should have crash_avoidance subcommand."""
        runner = CliRunner()
        result = runner.invoke(cli, ["--help"])
        self.assertIn("crash_avoidance", result.output)

    def test_cli_short_help_option(self):
        """CLI should support -h for help."""
        runner = CliRunner()
        result = runner.invoke(cli, ["-h"])
        self.assertEqual(result.exit_code, 0)
        self.assertIn("Euro NCAP Rating Calculator 2026", result.output)


class GreyCellIntegrityTestCase(unittest.TestCase):
    """Shared helpers for building minimal synthetic workbooks."""

    def _make_wb(self, rows, header=("Key", "Ref", "Grey")):
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Data"
        ws.append(list(header))
        for row in rows:
            ws.append(list(row))
        return wb


class TestRebuildTrustedSheet(GreyCellIntegrityTestCase):
    def _schema(self):
        return common.GreyCellSheetSchema(
            identity_columns=["Key"], grey_columns=["Grey"]
        )

    def test_overlays_grey_and_protects_reference_columns(self):
        pristine_wb = self._make_wb([("A", 100, None), ("B", 200, None)])
        user_wb = self._make_wb([("A", 999, "userval1"), ("B", 888, "userval2")])

        common.rebuild_trusted_sheet(
            pristine_wb["Data"], user_wb["Data"], self._schema(), "Data"
        )

        ws = pristine_wb["Data"]
        # Reference column must stay pristine, ignoring the user's tamper.
        self.assertEqual(ws["B2"].value, 100)
        self.assertEqual(ws["B3"].value, 200)
        # Grey column must be overlaid from the user's file.
        self.assertEqual(ws["C2"].value, "userval1")
        self.assertEqual(ws["C3"].value, "userval2")

    def test_raises_on_reordered_rows(self):
        pristine_wb = self._make_wb([("A", 100, None), ("B", 200, None)])
        user_wb = self._make_wb([("B", 200, None), ("A", 100, None)])

        with self.assertRaises(common.TemplateIntegrityError):
            common.rebuild_trusted_sheet(
                pristine_wb["Data"], user_wb["Data"], self._schema(), "Data"
            )

    def test_raises_on_extra_row(self):
        pristine_wb = self._make_wb([("A", 100, None), ("B", 200, None)])
        user_wb = self._make_wb([("A", 100, None), ("B", 200, None), ("C", 300, None)])

        with self.assertRaises(common.TemplateIntegrityError):
            common.rebuild_trusted_sheet(
                pristine_wb["Data"], user_wb["Data"], self._schema(), "Data"
            )

    def test_raises_on_missing_row(self):
        pristine_wb = self._make_wb([("A", 100, None), ("B", 200, None)])
        user_wb = self._make_wb([("A", 100, None)])

        with self.assertRaises(common.TemplateIntegrityError):
            common.rebuild_trusted_sheet(
                pristine_wb["Data"], user_wb["Data"], self._schema(), "Data"
            )

    def test_raises_on_missing_column(self):
        pristine_wb = self._make_wb([("A", 100, None)])
        user_wb = self._make_wb(
            [("A", 100)], header=("Key", "Ref")
        )  # "Grey" column missing entirely

        with self.assertRaises(common.TemplateIntegrityError):
            common.rebuild_trusted_sheet(
                pristine_wb["Data"], user_wb["Data"], self._schema(), "Data"
            )

    def test_ignores_stray_value_beyond_pristine_rows(self):
        """A stray value sitting in a row beyond the pristine table's real
        extent (e.g. leftover data-validation fill past the last real row)
        must not be mistaken for an inserted row, as long as none of the
        identity columns are populated there."""
        pristine_wb = self._make_wb([("A", 100, None), ("B", 200, None)])
        user_wb = self._make_wb([("A", 100, "userval1"), ("B", 200, "userval2")])
        # Row 4 has no identity value (Key blank) but a stray Grey value.
        user_wb["Data"].append([None, None, "stray"])

        common.rebuild_trusted_sheet(
            pristine_wb["Data"], user_wb["Data"], self._schema(), "Data"
        )
        self.assertEqual(pristine_wb["Data"]["C2"].value, "userval1")

    def test_raises_on_extra_row_beyond_pristine_with_identity_value(self):
        pristine_wb = self._make_wb([("A", 100, None), ("B", 200, None)])
        user_wb = self._make_wb([("A", 100, None), ("B", 200, None)])
        user_wb["Data"].append(["C", 300, None])

        with self.assertRaises(common.TemplateIntegrityError):
            common.rebuild_trusted_sheet(
                pristine_wb["Data"], user_wb["Data"], self._schema(), "Data"
            )

    def test_column_reordering_does_not_smuggle_a_protected_value(self):
        """Swapping column positions in the user's file must not let a value
        land in a differently-named protected column: columns are looked up
        by header name independently in each sheet."""
        pristine_wb = self._make_wb([("A", 100, None)])
        user_wb = self._make_wb([(100, "A", None)], header=("Ref", "Key", "Grey"))

        common.rebuild_trusted_sheet(
            pristine_wb["Data"], user_wb["Data"], self._schema(), "Data"
        )
        # "Ref" must still be sourced from the pristine sheet, unaffected by
        # the user's column reordering.
        self.assertEqual(pristine_wb["Data"]["B2"].value, 100)

    def _legacy_entry(self, anchor_key="A", value="L2"):
        return common.LegacyUnlabeledCell(
            column="Label", anchor={"Key": anchor_key}, value=value
        )

    def _two_col_schema(self, legacy_unlabeled_cells=None):
        return common.GreyCellSheetSchema(
            identity_columns=["Key", "Label"],
            grey_columns=["Grey"],
            legacy_unlabeled_cells=legacy_unlabeled_cells or [],
        )

    def test_blank_identity_cell_without_explicit_whitelist_still_raises(self):
        """The exemption must never be a blanket "blank identity cell is OK
        if some other column is populated" rule: without an explicit
        LegacyUnlabeledCell entry, a blank user cell where the pristine
        template has a label is a mismatch like any other, even though the
        row's other identity column ("Key") is populated."""
        pristine_wb = self._make_wb(
            [("A", "L1", 100, None), ("A", "L2", 200, None)],
            header=("Key", "Label", "Ref", "Grey"),
        )
        user_wb = self._make_wb(
            [("A", "L1", 100, "userval1"), ("A", None, 200, "userval2")],
            header=("Key", "Label", "Ref", "Grey"),
        )

        with self.assertRaises(common.TemplateIntegrityError):
            common.rebuild_trusted_sheet(
                pristine_wb["Data"], user_wb["Data"], self._two_col_schema(), "Data"
            )

    def test_blank_identity_cell_matches_when_explicitly_whitelisted(self):
        """A template revision that names a previously-blank identity cell
        must not invalidate every prediction file produced against the
        older template -- but only once that exact (row, column) is
        declared via LegacyUnlabeledCell; the user's blank cell is then
        exempted from the comparison instead of forward-filling the user's
        own prior label into it."""
        pristine_wb = self._make_wb(
            [("A", "L1", 100, None), ("A", "L2", 200, None)],
            header=("Key", "Label", "Ref", "Grey"),
        )
        user_wb = self._make_wb(
            [("A", "L1", 100, "userval1"), ("A", None, 200, "userval2")],
            header=("Key", "Label", "Ref", "Grey"),
        )

        common.rebuild_trusted_sheet(
            pristine_wb["Data"],
            user_wb["Data"],
            self._two_col_schema([self._legacy_entry()]),
            "Data",
        )
        self.assertEqual(pristine_wb["Data"]["D2"].value, "userval1")
        self.assertEqual(pristine_wb["Data"]["D3"].value, "userval2")

    def test_whitelist_entry_does_not_match_a_different_anchor(self):
        """The anchor must match the pristine row's OTHER identity values
        exactly -- a whitelist entry for Key="B" must not accidentally
        exempt a blank Label on a Key="A" row."""
        pristine_wb = self._make_wb(
            [("A", "L1", 100, None), ("A", "L2", 200, None)],
            header=("Key", "Label", "Ref", "Grey"),
        )
        user_wb = self._make_wb(
            [("A", "L1", 100, "userval1"), ("A", None, 200, "userval2")],
            header=("Key", "Label", "Ref", "Grey"),
        )

        with self.assertRaises(common.TemplateIntegrityError):
            common.rebuild_trusted_sheet(
                pristine_wb["Data"],
                user_wb["Data"],
                self._two_col_schema([self._legacy_entry(anchor_key="B")]),
                "Data",
            )

    def test_whitelist_entry_does_not_match_a_different_value(self):
        """The declared *value* must match the pristine template's current
        value for that cell -- a stale whitelist entry (e.g. after the
        template changed again) must not silently apply to a different
        label."""
        pristine_wb = self._make_wb(
            [("A", "L1", 100, None), ("A", "L2", 200, None)],
            header=("Key", "Label", "Ref", "Grey"),
        )
        user_wb = self._make_wb(
            [("A", "L1", 100, "userval1"), ("A", None, 200, "userval2")],
            header=("Key", "Label", "Ref", "Grey"),
        )

        with self.assertRaises(common.TemplateIntegrityError):
            common.rebuild_trusted_sheet(
                pristine_wb["Data"],
                user_wb["Data"],
                self._two_col_schema([self._legacy_entry(value="L2-renamed")]),
                "Data",
            )

    def test_exempted_cell_does_not_seed_forward_fill(self):
        """The exemption must not seed the user's own forward-fill state
        with the pristine value: a later row that's genuinely blank in both
        files (a real continuation) must still inherit the user's own last
        real label, not the template's, so an unrelated shift below the
        exempted row is still caught rather than silently validated."""
        pristine_wb = self._make_wb(
            [
                ("A", "L1", 100, None),
                ("A", "L2", 200, None),
                ("A", None, 300, None),
            ],
            header=("Key", "Label", "Ref", "Grey"),
        )
        user_wb = self._make_wb(
            [
                ("A", "L1", 100, "u1"),
                ("A", None, 200, "u2"),
                ("A", None, 300, "u3"),
            ],
            header=("Key", "Label", "Ref", "Grey"),
        )

        with self.assertRaises(common.TemplateIntegrityError):
            common.rebuild_trusted_sheet(
                pristine_wb["Data"],
                user_wb["Data"],
                self._two_col_schema([self._legacy_entry()]),
                "Data",
            )

    def test_still_raises_when_row_after_exemption_is_actually_shifted(self):
        """The exemption is scoped to the one whitelisted cell; rows below
        it are still compared exactly as before, so a real mismatch after
        the exempted row is still rejected."""
        pristine_wb = self._make_wb(
            [
                ("A", "L1", 100, None),
                ("A", "L2", 200, None),
                ("A", "L3", 300, None),
            ],
            header=("Key", "Label", "Ref", "Grey"),
        )
        user_wb = self._make_wb(
            [
                ("A", "L1", 100, "u1"),
                ("A", None, 200, "u2"),
                ("A", "WRONG", 300, "u3"),
            ],
            header=("Key", "Label", "Ref", "Grey"),
        )

        with self.assertRaises(common.TemplateIntegrityError):
            common.rebuild_trusted_sheet(
                pristine_wb["Data"],
                user_wb["Data"],
                self._two_col_schema([self._legacy_entry()]),
                "Data",
            )

    def test_row_missing_entirely_from_user_file_is_not_exempted(self):
        """Even with a whitelist entry for Label on the Key="B" row, a row
        that's missing from the user's file entirely still raises: "Key"
        itself has no whitelist entry, so it's compared normally, and the
        user's blank Key at that row forward-fills the wrong ("A") value."""
        pristine_wb = self._make_wb(
            [("A", "L1", 100, None), ("B", "L2", 200, None)],
            header=("Key", "Label", "Ref", "Grey"),
        )
        user_wb = self._make_wb(
            [("A", "L1", 100, "u1")], header=("Key", "Label", "Ref", "Grey")
        )

        with self.assertRaises(common.TemplateIntegrityError):
            common.rebuild_trusted_sheet(
                pristine_wb["Data"],
                user_wb["Data"],
                self._two_col_schema([self._legacy_entry(anchor_key="B")]),
                "Data",
            )


class TestRebuildTrustedWorkbook(GreyCellIntegrityTestCase):
    def test_passthrough_sheet_copied_verbatim_and_schema_sheet_protected(self):
        import openpyxl

        pristine_wb = self._make_wb([("A", 100, None)])
        extra_ws = pristine_wb.create_sheet("Extra")
        extra_ws.append(["pristine extra content"])

        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            pristine_path = f.name
        pristine_wb.save(pristine_path)

        try:
            user_wb = self._make_wb([("A", 999, "userval")])
            user_extra_ws = user_wb.create_sheet("Extra")
            user_extra_ws.append(["real user content"])

            schema = common.GreyCellSheetSchema(
                identity_columns=["Key"], grey_columns=["Grey"]
            )
            out_wb = common.rebuild_trusted_workbook(
                pristine_path, user_wb, {"Data": schema}, ["Extra"]
            )

            self.assertEqual(out_wb["Data"]["B2"].value, 100)  # protected
            self.assertEqual(out_wb["Data"]["C2"].value, "userval")  # grey
            self.assertEqual(
                out_wb["Extra"]["A1"].value, "real user content"
            )  # passthrough
        finally:
            os.unlink(pristine_path)

    def test_raises_when_user_file_missing_required_sheet(self):
        import openpyxl

        pristine_wb = self._make_wb([("A", 100, None)])
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            pristine_path = f.name
        pristine_wb.save(pristine_path)

        try:
            user_wb = openpyxl.Workbook()  # no "Data" sheet at all
            schema = common.GreyCellSheetSchema(
                identity_columns=["Key"], grey_columns=["Grey"]
            )
            with self.assertRaises(common.TemplateIntegrityError):
                common.rebuild_trusted_workbook(
                    pristine_path, user_wb, {"Data": schema}, []
                )
        finally:
            os.unlink(pristine_path)


class TestProtectedSignature(GreyCellIntegrityTestCase):
    def _schema_map(self):
        return {
            "Data": common.GreyCellSheetSchema(
                identity_columns=["Key"], grey_columns=["Grey"]
            )
        }

    def test_stable_across_save_reload_round_trip(self):
        import openpyxl

        wb = self._make_wb([("A", 100, None), ("B", 200, None)])
        sig_in_memory = common.compute_protected_signature(wb, self._schema_map())

        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            path = f.name
        try:
            wb.save(path)
            reloaded = openpyxl.load_workbook(path, data_only=True)
            sig_reloaded = common.compute_protected_signature(
                reloaded, self._schema_map()
            )
            self.assertEqual(sig_in_memory, sig_reloaded)
        finally:
            os.unlink(path)

    def test_verify_passes_when_unmodified(self):
        wb = self._make_wb([("A", 100, None)])
        common.add_integrity_sheet_to_wb(wb, self._schema_map())
        # Should not raise.
        common.verify_protected_signature(wb, self._schema_map())

    def test_verify_raises_when_protected_cell_tampered(self):
        wb = self._make_wb([("A", 100, None)])
        common.add_integrity_sheet_to_wb(wb, self._schema_map())
        wb["Data"]["B2"].value = 999  # tamper the protected "Ref" column

        with self.assertRaises(common.TemplateIntegrityError):
            common.verify_protected_signature(wb, self._schema_map())

    def test_verify_ignores_grey_cell_edits(self):
        wb = self._make_wb([("A", 100, None)])
        common.add_integrity_sheet_to_wb(wb, self._schema_map())
        wb["Data"]["C2"].value = "a legitimate later edit"  # grey column

        # Should not raise: grey columns are excluded from the signature.
        common.verify_protected_signature(wb, self._schema_map())

    def test_verify_raises_when_integrity_sheet_missing(self):
        wb = self._make_wb([("A", 100, None)])

        with self.assertRaises(common.TemplateIntegrityError):
            common.verify_protected_signature(wb, self._schema_map())

    def test_integrity_sheet_is_hidden(self):
        wb = self._make_wb([("A", 100, None)])
        common.add_integrity_sheet_to_wb(wb, self._schema_map())
        self.assertEqual(wb[common.INTEGRITY_SHEET_NAME].sheet_state, "veryHidden")


class GreyCellIntegrityDfTestCase(unittest.TestCase):
    """Shared helpers for building minimal synthetic DataFrames, the
    dataframe-native counterpart of GreyCellIntegrityTestCase."""

    def _make_df(self, rows, header=("Key", "Ref", "Grey")):
        return pd.DataFrame(list(rows), columns=list(header))

    def _schema(self):
        return common.GreyCellSheetSchema(
            identity_columns=["Key"], grey_columns=["Grey"]
        )


class TestRebuildTrustedDf(GreyCellIntegrityDfTestCase):
    def test_overlays_grey_and_protects_reference_columns(self):
        pristine_df = self._make_df([("A", 100, None), ("B", 200, None)])
        user_df = self._make_df([("A", 999, "userval1"), ("B", 888, "userval2")])

        result = common.rebuild_trusted_df(pristine_df, user_df, self._schema(), "Data")

        # Reference column must stay pristine, ignoring the user's tamper.
        self.assertEqual(result["Ref"].tolist(), [100, 200])
        # Grey column must be overlaid from the user's file.
        self.assertEqual(result["Grey"].tolist(), ["userval1", "userval2"])

    def test_raises_on_reordered_rows(self):
        pristine_df = self._make_df([("A", 100, None), ("B", 200, None)])
        user_df = self._make_df([("B", 200, None), ("A", 100, None)])

        with self.assertRaises(common.TemplateIntegrityError):
            common.rebuild_trusted_df(pristine_df, user_df, self._schema(), "Data")

    def test_raises_on_extra_row(self):
        pristine_df = self._make_df([("A", 100, None), ("B", 200, None)])
        user_df = self._make_df([("A", 100, None), ("B", 200, None), ("C", 300, None)])

        with self.assertRaises(common.TemplateIntegrityError):
            common.rebuild_trusted_df(pristine_df, user_df, self._schema(), "Data")

    def test_raises_on_missing_row(self):
        pristine_df = self._make_df([("A", 100, None), ("B", 200, None)])
        user_df = self._make_df([("A", 100, None)])

        with self.assertRaises(common.TemplateIntegrityError):
            common.rebuild_trusted_df(pristine_df, user_df, self._schema(), "Data")

    def test_raises_on_missing_column(self):
        pristine_df = self._make_df([("A", 100, None)])
        user_df = self._make_df([("A", 100)], header=("Key", "Ref"))  # no "Grey"

        with self.assertRaises(common.TemplateIntegrityError):
            common.rebuild_trusted_df(pristine_df, user_df, self._schema(), "Data")

    def test_ignores_stray_value_beyond_pristine_rows(self):
        """A stray value sitting in a row beyond the pristine table's real
        extent must not be mistaken for an inserted row, as long as none of
        the identity columns are populated there."""
        pristine_df = self._make_df([("A", 100, None), ("B", 200, None)])
        user_df = self._make_df(
            [("A", 100, "userval1"), ("B", 200, "userval2"), (None, None, "stray")]
        )

        result = common.rebuild_trusted_df(pristine_df, user_df, self._schema(), "Data")
        self.assertEqual(result["Grey"].tolist(), ["userval1", "userval2"])

    def test_raises_on_extra_row_beyond_pristine_with_identity_value(self):
        pristine_df = self._make_df([("A", 100, None), ("B", 200, None)])
        user_df = self._make_df([("A", 100, None), ("B", 200, None), ("C", 300, None)])

        with self.assertRaises(common.TemplateIntegrityError):
            common.rebuild_trusted_df(pristine_df, user_df, self._schema(), "Data")

    def test_column_reordering_does_not_smuggle_a_protected_value(self):
        """Columns are looked up by header name, not position, so reordering
        the user frame's columns must not let a value land in a differently-
        named protected column."""
        pristine_df = self._make_df([("A", 100, None)])
        user_df = pd.DataFrame([(100, "A", None)], columns=["Ref", "Key", "Grey"])

        result = common.rebuild_trusted_df(pristine_df, user_df, self._schema(), "Data")
        # "Ref" must still be sourced from the pristine frame, unaffected by
        # the user's column reordering.
        self.assertEqual(result["Ref"].tolist(), [100])

    def test_blank_string_identity_continuation_matches(self):
        """A continuation row whose identity cell is "" instead of NaN (as a
        parquet-round-tripped dataframe can hold, see
        common._is_blank_identity_value) must still be recognised as
        carrying the previous row's identity forward, not rejected as a
        mismatched/new row."""
        pristine_df = self._make_df(
            [("A", 100, None), (None, 101, None), ("B", 200, None)]
        )
        user_df = self._make_df(
            [("A", 100, "userval1"), ("", 101, "userval2"), ("B", 200, "userval3")]
        )

        result = common.rebuild_trusted_df(pristine_df, user_df, self._schema(), "Data")
        self.assertEqual(result["Grey"].tolist(), ["userval1", "userval2", "userval3"])

    def _legacy_entry(self, anchor_key="A", value="L2"):
        return common.LegacyUnlabeledCell(
            column="Label", anchor={"Key": anchor_key}, value=value
        )

    def _two_col_schema(self, legacy_unlabeled_cells=None):
        return common.GreyCellSheetSchema(
            identity_columns=["Key", "Label"],
            grey_columns=["Grey"],
            legacy_unlabeled_cells=legacy_unlabeled_cells or [],
        )

    def test_blank_identity_cell_without_explicit_whitelist_still_raises(self):
        """DataFrame-path mirror: without an explicit LegacyUnlabeledCell
        entry, a blank user cell where the pristine template has a label is
        a mismatch like any other -- there is no blanket "some other column
        is populated" exemption."""
        pristine_df = self._make_df(
            [("A", "L1", 100, None), ("A", "L2", 200, None)],
            header=("Key", "Label", "Ref", "Grey"),
        )
        user_df = self._make_df(
            [("A", "L1", 100, "userval1"), ("A", None, 200, "userval2")],
            header=("Key", "Label", "Ref", "Grey"),
        )

        with self.assertRaises(common.TemplateIntegrityError):
            common.rebuild_trusted_df(
                pristine_df, user_df, self._two_col_schema(), "Data"
            )

    def test_blank_identity_cell_matches_when_explicitly_whitelisted(self):
        """DataFrame-path mirror of the openpyxl-path test of the same name:
        a template revision that names a previously-blank identity cell must
        not invalidate every prediction file produced against the older
        template, once that exact (row, column) is declared via
        LegacyUnlabeledCell."""
        pristine_df = self._make_df(
            [("A", "L1", 100, None), ("A", "L2", 200, None)],
            header=("Key", "Label", "Ref", "Grey"),
        )
        user_df = self._make_df(
            [("A", "L1", 100, "userval1"), ("A", None, 200, "userval2")],
            header=("Key", "Label", "Ref", "Grey"),
        )

        result = common.rebuild_trusted_df(
            pristine_df, user_df, self._two_col_schema([self._legacy_entry()]), "Data"
        )
        self.assertEqual(result["Grey"].tolist(), ["userval1", "userval2"])

    def test_whitelist_entry_does_not_match_a_different_anchor(self):
        """DataFrame-path mirror: the anchor must match the pristine row's
        other identity values exactly."""
        pristine_df = self._make_df(
            [("A", "L1", 100, None), ("A", "L2", 200, None)],
            header=("Key", "Label", "Ref", "Grey"),
        )
        user_df = self._make_df(
            [("A", "L1", 100, "userval1"), ("A", None, 200, "userval2")],
            header=("Key", "Label", "Ref", "Grey"),
        )

        with self.assertRaises(common.TemplateIntegrityError):
            common.rebuild_trusted_df(
                pristine_df,
                user_df,
                self._two_col_schema([self._legacy_entry(anchor_key="B")]),
                "Data",
            )

    def test_whitelist_entry_does_not_match_a_different_value(self):
        """DataFrame-path mirror: a stale whitelist entry must not silently
        apply to a different label."""
        pristine_df = self._make_df(
            [("A", "L1", 100, None), ("A", "L2", 200, None)],
            header=("Key", "Label", "Ref", "Grey"),
        )
        user_df = self._make_df(
            [("A", "L1", 100, "userval1"), ("A", None, 200, "userval2")],
            header=("Key", "Label", "Ref", "Grey"),
        )

        with self.assertRaises(common.TemplateIntegrityError):
            common.rebuild_trusted_df(
                pristine_df,
                user_df,
                self._two_col_schema([self._legacy_entry(value="L2-renamed")]),
                "Data",
            )

    def test_exempted_cell_does_not_seed_forward_fill(self):
        """DataFrame-path mirror: the exemption must not seed the user's own
        forward-fill state with the pristine value, so an unrelated shift
        below the exempted row is still caught."""
        pristine_df = self._make_df(
            [
                ("A", "L1", 100, None),
                ("A", "L2", 200, None),
                ("A", None, 300, None),
            ],
            header=("Key", "Label", "Ref", "Grey"),
        )
        user_df = self._make_df(
            [
                ("A", "L1", 100, "u1"),
                ("A", None, 200, "u2"),
                ("A", None, 300, "u3"),
            ],
            header=("Key", "Label", "Ref", "Grey"),
        )

        with self.assertRaises(common.TemplateIntegrityError):
            common.rebuild_trusted_df(
                pristine_df,
                user_df,
                self._two_col_schema([self._legacy_entry()]),
                "Data",
            )

    def test_still_raises_when_row_after_exemption_is_actually_shifted(self):
        """DataFrame-path mirror: the exemption is scoped to the one
        whitelisted cell; a real mismatch after the exempted row is still
        rejected."""
        pristine_df = self._make_df(
            [
                ("A", "L1", 100, None),
                ("A", "L2", 200, None),
                ("A", "L3", 300, None),
            ],
            header=("Key", "Label", "Ref", "Grey"),
        )
        user_df = self._make_df(
            [
                ("A", "L1", 100, "u1"),
                ("A", None, 200, "u2"),
                ("A", "WRONG", 300, "u3"),
            ],
            header=("Key", "Label", "Ref", "Grey"),
        )

        with self.assertRaises(common.TemplateIntegrityError):
            common.rebuild_trusted_df(
                pristine_df,
                user_df,
                self._two_col_schema([self._legacy_entry()]),
                "Data",
            )

    def test_row_all_blank_in_user_frame_is_not_exempted(self):
        """DataFrame-path mirror: a row that's all-blank across every
        identity column in the user's frame (e.g. a genuinely deleted row
        replaced by a shift) must still be rejected. Even with a whitelist
        entry for Label on the Key="B" row, "Key" itself has no whitelist
        entry and is compared normally."""
        pristine_df = self._make_df(
            [("A", "L1", 100, None), ("B", "L2", 200, None)],
            header=("Key", "Label", "Ref", "Grey"),
        )
        user_df = self._make_df(
            [("A", "L1", 100, "u1"), (None, None, 200, "u2")],
            header=("Key", "Label", "Ref", "Grey"),
        )

        with self.assertRaises(common.TemplateIntegrityError):
            common.rebuild_trusted_df(
                pristine_df,
                user_df,
                self._two_col_schema([self._legacy_entry(anchor_key="B")]),
                "Data",
            )


class TestRebuildTrustedDfs(GreyCellIntegrityDfTestCase):
    def test_passthrough_and_extra_sheets_pass_through_and_schema_sheet_protected(self):
        pristine_df = self._make_df([("A", 100, None)])
        user_df = self._make_df([("A", 999, "userval")])
        extra_df = pd.DataFrame({"col": ["real user content"]})

        result = common.rebuild_trusted_dfs(
            {"Data": pristine_df},
            {"Data": user_df, "Extra": extra_df},
            {"Data": self._schema()},
            ["Extra"],
        )

        self.assertEqual(result["Data"]["Ref"].tolist(), [100])  # protected
        self.assertEqual(result["Data"]["Grey"].tolist(), ["userval"])  # grey
        self.assertEqual(result["Extra"]["col"].tolist(), ["real user content"])

    def test_sheet_missing_from_user_dfs_is_skipped_not_raised(self):
        """Deliberate deviation from rebuild_trusted_workbook: dataframe-path
        callers legitimately pass a partial dict, so a schema'd sheet that
        simply isn't present in user_dfs must not raise."""
        pristine_df = self._make_df([("A", 100, None)])

        result = common.rebuild_trusted_dfs(
            {"Data": pristine_df},
            {"Other": pd.DataFrame({"x": [1]})},
            {"Data": self._schema()},
            [],
        )

        self.assertNotIn("Data", result)
        self.assertEqual(result["Other"]["x"].tolist(), [1])


class TestUpdateScoreSumInterval(unittest.TestCase):
    def _interval_df(self, categories, scores):
        return pd.DataFrame({"Category": categories, "Score": scores})

    def _target_df(self, categories):
        return pd.DataFrame(
            {"Category": categories, "Score": [np.nan] * len(categories)}
        )

    def test_all_assessed_sums_normally(self):
        interval_df = self._interval_df(
            ["Cat A", None, None, "Cat B", None],
            [1.0, 2.0, 3.0, 4.0, 5.0],
        )
        target_df = self._target_df(["Cat A", "Cat B"])

        result = common.update_score_sum_interval(
            target_df, interval_df, key_column="Category", interval_column="Category"
        )

        self.assertEqual(
            result.loc[result["Category"] == "Cat A", "Score"].iloc[0], 6.0
        )
        self.assertEqual(
            result.loc[result["Category"] == "Cat B", "Score"].iloc[0], 9.0
        )

    def test_all_unassessed_section_propagates_nan(self):
        interval_df = self._interval_df(
            ["Cat A", None, None],
            [np.nan, np.nan, np.nan],
        )
        target_df = self._target_df(["Cat A"])

        result = common.update_score_sum_interval(
            target_df, interval_df, key_column="Category", interval_column="Category"
        )

        self.assertTrue(
            pd.isna(result.loc[result["Category"] == "Cat A", "Score"].iloc[0])
        )

    def test_partially_unassessed_section_propagates_nan_not_partial_sum(self):
        interval_df = self._interval_df(
            ["Cat A", None, None],
            [1.0, np.nan, 3.0],
        )
        target_df = self._target_df(["Cat A"])

        result = common.update_score_sum_interval(
            target_df, interval_df, key_column="Category", interval_column="Category"
        )

        # Must NOT silently sum the assessed rows only (1.0 + 3.0 = 4.0) --
        # one unassessed scenario makes the whole Category unassessed.
        self.assertTrue(
            pd.isna(result.loc[result["Category"] == "Cat A", "Score"].iloc[0])
        )

    def test_mixed_categories_only_unassessed_one_becomes_nan(self):
        interval_df = self._interval_df(
            ["Cat A", None, "Cat B", None],
            [1.0, 2.0, np.nan, np.nan],
        )
        target_df = self._target_df(["Cat A", "Cat B"])

        result = common.update_score_sum_interval(
            target_df, interval_df, key_column="Category", interval_column="Category"
        )

        self.assertEqual(
            result.loc[result["Category"] == "Cat A", "Score"].iloc[0], 3.0
        )
        self.assertTrue(
            pd.isna(result.loc[result["Category"] == "Cat B", "Score"].iloc[0])
        )

    def test_dash_placeholder_reads_as_unassessed(self):
        """DataFrame callers pass sheets carrying the "-" placeholder (see
        reset_computed_columns_to_dash) straight in, without the CLI's
        Score-to-NaN reset -- a dash must behave exactly like NaN, in both
        the target and the interval sheet, not crash astype."""
        interval_df = self._interval_df(
            ["Cat A", None, "Cat B", None],
            [1.0, 2.0, "-", 4.0],
        )
        target_df = pd.DataFrame({"Category": ["Cat A", "Cat B"], "Score": ["-", "-"]})

        result = common.update_score_sum_interval(
            target_df, interval_df, key_column="Category", interval_column="Category"
        )

        self.assertEqual(
            result.loc[result["Category"] == "Cat A", "Score"].iloc[0], 3.0
        )
        self.assertTrue(
            pd.isna(result.loc[result["Category"] == "Cat B", "Score"].iloc[0])
        )


class TestUpdateScoringSheetDashPlaceholder(unittest.TestCase):
    """The "-" placeholder ("not yet computed", written by
    reset_computed_columns_to_dash since the dash harmonization) must read
    as unassessed NaN. The CLI path resets the Score column to NaN before
    these helpers run, but DataFrame callers invoke them directly on the
    preprocessed sheet -- both paths must produce identical results."""

    def test_update_scoring_sheet_accepts_dash_scores(self):
        df = pd.DataFrame({"Scenario": ["Scen A", "Scen B"], "Score": ["-", "-"]})

        result = common.update_scoring_sheet(df, {"Scen A": 2.5}, key_column="Scenario")

        self.assertEqual(
            result.loc[result["Scenario"] == "Scen A", "Score"].iloc[0], 2.5
        )
        self.assertTrue(
            pd.isna(result.loc[result["Scenario"] == "Scen B", "Score"].iloc[0])
        )

    def test_update_scoring_sheet_other_strings_still_raise(self):
        df = pd.DataFrame({"Scenario": ["Scen A"], "Score": ["garbage"]})

        with self.assertRaises(ValueError):
            common.update_scoring_sheet(df, {"Scen A": 1.0}, key_column="Scenario")

    def test_update_scoring_sheet_from_df_accepts_dash_scores(self):
        df = pd.DataFrame(
            {
                "Category": ["Cat A", None],
                "Scenario": ["Scen A", "Scen B"],
                "Score": ["-", "-"],
            }
        )
        score_df = pd.DataFrame(
            {"Category": ["Cat A"], "Scenario": ["Scen A"], "Score": [1.5]}
        )

        result = common.update_scoring_sheet_from_df(df, score_df)

        self.assertEqual(
            result.loc[result["Scenario"] == "Scen A", "Score"].iloc[0], 1.5
        )
        self.assertTrue(
            pd.isna(result.loc[result["Scenario"] == "Scen B", "Score"].iloc[0])
        )


if __name__ == "__main__":
    unittest.main()
