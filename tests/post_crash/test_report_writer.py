# Copyright 2025-2026, Euro NCAP IVZW
# Created by IVEX NV (https://ivex.ai)
#
# Licensed under the Apache License 2.0.
# See http://www.apache.org/licenses/LICENSE-2.0 for details.

import shutil
import tempfile
import unittest
from importlib.resources import files
from pathlib import Path

import openpyxl
import pandas as pd

from euroncap_rating_2026 import common
from euroncap_rating_2026.post_crash import report_writer as pc_report_writer


class TestPassFailDropdownAllowsBlank(unittest.TestCase):
    """A blank Value cell means "not yet assessed" and must still get a
    dropdown with allow_blank=True -- not just cells that already say
    PASS/FAIL/a leadtime band."""

    def _write_report(self, updated_dfs):
        pristine_path = str(files("data").joinpath("pc_template.xlsx"))
        with (
            tempfile.TemporaryDirectory() as input_dir,
            tempfile.TemporaryDirectory() as output_dir,
        ):
            input_file = str(Path(input_dir) / "pc_template.xlsx")
            shutil.copy(pristine_path, input_file)
            common.add_version_sheet(
                input_file, cli_command=common.CliCommand.GENERATE_TEMPLATE
            )

            pc_report_writer.write_report(
                common.CliCommand.PREPROCESS,
                input_file,
                updated_dfs=updated_dfs,
                selected_points_dict={},
                output_path=output_dir,
                format_prediction_cells=False,
            )

            output_file = Path(output_dir) / "pc_preprocessed_template.xlsx"
            return openpyxl.load_workbook(str(output_file))

    def _dv_for_cell(self, ws, cell):
        for dv in ws.data_validations.dataValidation:
            if cell.coordinate in dv.sqref:
                return dv
        return None

    def test_blank_value_cell_gets_pass_fail_dropdown_with_allow_blank(self):
        df = pd.DataFrame(
            [
                {"Category": "Advanced multi-collision brake", "Value": None},
                {
                    "Category": "Automatic activation of hazard warning lights",
                    "Value": "PASS",
                },
            ]
        )
        wb = self._write_report({"PCI - MCB & Hazard lights Verif": df})
        ws = wb["PCI - MCB & Hazard lights Verif"]

        blank_cell = ws.cell(row=2, column=2)
        pass_cell = ws.cell(row=3, column=2)

        for cell in (blank_cell, pass_cell):
            dv = self._dv_for_cell(ws, cell)
            self.assertIsNotNone(dv, f"expected a dropdown on {cell.coordinate}")
            self.assertTrue(dv.allow_blank)
            self.assertEqual(dv.formula1, '"PASS,FAIL"')

    def test_energy_management_leadtime_blank_gets_duration_dropdown_not_pass_fail(
        self,
    ):
        df = pd.DataFrame(
            [
                {
                    "Category": "Thermal propagation",
                    "Scenario": "Fulfilment of UN-R100.03 with predefined leadtime",
                    "Value": None,
                },
                {
                    "Category": "Energy isolation",
                    "Scenario": "Compliance with UN Regulation requirements",
                    "Value": None,
                },
            ]
        )
        wb = self._write_report({"VE - Energy Management Verif": df})
        ws = wb["VE - Energy Management Verif"]

        leadtime_cell = ws.cell(row=2, column=3)
        pass_fail_cell = ws.cell(row=3, column=3)

        leadtime_dv = self._dv_for_cell(ws, leadtime_cell)
        self.assertIsNotNone(leadtime_dv)
        self.assertTrue(leadtime_dv.allow_blank)
        self.assertEqual(leadtime_dv.formula1, '"≥90 min,>40 min,>20 min,≤20 min"')

        pass_fail_dv = self._dv_for_cell(ws, pass_fail_cell)
        self.assertIsNotNone(pass_fail_dv)
        self.assertTrue(pass_fail_dv.allow_blank)
        self.assertEqual(pass_fail_dv.formula1, '"PASS,FAIL"')

    def test_energy_management_leadtime_le20min_gets_duration_dropdown(self):
        """A cell already filled with '≤20 min' (the ≤ symbol used by the
        dropdown's own formula1 list) must still get the duration dropdown
        attached -- the comparison used to check the ASCII '<=20 MIN'
        instead of the actual '≤20 MIN' symbol, so this branch never
        matched and the dropdown was silently skipped."""
        df = pd.DataFrame(
            [
                {
                    "Category": "Thermal propagation",
                    "Scenario": "Fulfilment of UN-R100.03 with predefined leadtime",
                    "Value": "≤20 min",
                },
            ]
        )
        wb = self._write_report({"VE - Energy Management Verif": df})
        ws = wb["VE - Energy Management Verif"]

        leadtime_cell = ws.cell(row=2, column=3)
        leadtime_dv = self._dv_for_cell(ws, leadtime_cell)
        self.assertIsNotNone(
            leadtime_dv, f"expected a dropdown on {leadtime_cell.coordinate}"
        )
        self.assertEqual(leadtime_dv.formula1, '"≥90 min,>40 min,>20 min,≤20 min"')

    def test_input_parameters_blank_value_gets_no_pass_fail_dropdown(self):
        wb = self._write_report({})
        ws = wb["Input parameters"]

        headers = [c.value for c in ws[1]]
        value_col_idx = headers.index("Value") + 1
        for row in range(2, ws.max_row + 1):
            cell = ws.cell(row=row, column=value_col_idx)
            self.assertIsNone(self._dv_for_cell(ws, cell))


class TestMissingSheetsAreSkipped(unittest.TestCase):
    def test_preprocess_succeeds_without_documentation_data_sheets(self):
        """A prediction workbook uploaded before the "Documentation"/"Data"
        sheets existed in the template must still preprocess -- the missing
        sheets are skipped with a warning instead of aborting write_report."""
        pristine_path = str(files("data").joinpath("pc_template.xlsx"))
        with (
            tempfile.TemporaryDirectory() as input_dir,
            tempfile.TemporaryDirectory() as output_dir,
        ):
            input_file = str(Path(input_dir) / "pc_template.xlsx")
            shutil.copy(pristine_path, input_file)
            common.add_version_sheet(
                input_file, cli_command=common.CliCommand.GENERATE_TEMPLATE
            )
            wb = openpyxl.load_workbook(input_file)
            for sheet_name in common.DOCUMENTATION_DATA_SHEETS:
                wb.remove(wb[sheet_name])
            wb.save(input_file)

            pc_report_writer.write_report(
                common.CliCommand.PREPROCESS,
                input_file,
                updated_dfs={},
                selected_points_dict={},
                output_path=output_dir,
                format_prediction_cells=False,
            )

            output_file = Path(output_dir) / "pc_preprocessed_template.xlsx"
            self.assertTrue(output_file.exists())
            out_wb = openpyxl.load_workbook(str(output_file))
            for sheet_name in common.DOCUMENTATION_DATA_SHEETS:
                self.assertNotIn(sheet_name, out_wb.sheetnames)


if __name__ == "__main__":
    unittest.main()
