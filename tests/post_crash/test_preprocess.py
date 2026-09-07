# Copyright 2025-2026, Euro NCAP IVZW
# Created by IVEX NV (https://ivex.ai)
#
# Licensed under the Apache License 2.0.
# See http://www.apache.org/licenses/LICENSE-2.0 for details.

import shutil
import tempfile
import unittest
from importlib.resources import files
from pathlib import Path

import openpyxl
from click.testing import CliRunner

from euroncap_rating_2026 import common
from euroncap_rating_2026.post_crash.compute_score import compute_score
from euroncap_rating_2026.post_crash.preprocess import (
    VERIFICATION_SHEETS,
    add_verification_sheets_to_dfs,
    preprocess,
)


class TestDocumentationDataSurviveFullPipeline(unittest.TestCase):
    """post_crash has no documentation_data.py logic yet (structural parity
    only, see the repo's implementation plan): the two hidden sheets carry
    only their header row, with no ID rows. They must still exist, stay
    hidden, and survive both preprocess() and compute_score() untouched."""

    def test_preprocess_then_compute_score_e2e(self):
        pristine_path = str(files("data").joinpath("pc_template.xlsx"))
        runner = CliRunner()
        with (
            tempfile.TemporaryDirectory() as input_dir,
            tempfile.TemporaryDirectory() as output_dir,
        ):
            input_file = str(Path(input_dir) / "pc_template_versioned.xlsx")
            shutil.copy(pristine_path, input_file)
            common.add_version_sheet(
                input_file, cli_command=common.CliCommand.GENERATE_TEMPLATE
            )

            result = runner.invoke(
                preprocess, ["--input_file", input_file, "--output_path", output_dir]
            )
            self.assertEqual(result.exit_code, 0, result.output)

            preprocessed_file = str(Path(output_dir) / "pc_preprocessed_template.xlsx")
            preprocessed_wb = openpyxl.load_workbook(preprocessed_file)
            for sheet_name, id_header in [
                ("Documentation", "Documentation ID"),
                ("Data", "Data ID"),
            ]:
                ws = preprocessed_wb[sheet_name]
                self.assertEqual(ws.sheet_state, "hidden")
                self.assertEqual([c.value for c in ws[1]], [id_header, "Status"])
                self.assertEqual(ws.max_row, 1)

            result2 = runner.invoke(
                compute_score,
                ["--input_file", preprocessed_file, "--output_path", output_dir],
            )
            self.assertEqual(result2.exit_code, 0, result2.output)

            report_path = next(
                p for p in Path(output_dir).iterdir() if "report" in p.name
            )
            report_wb = openpyxl.load_workbook(str(report_path))
            for sheet_name, id_header in [
                ("Documentation", "Documentation ID"),
                ("Data", "Data ID"),
            ]:
                ws = report_wb[sheet_name]
                self.assertEqual(ws.sheet_state, "hidden")
                self.assertEqual([c.value for c in ws[1]], [id_header, "Status"])
                self.assertEqual(ws.max_row, 1)


class TestAddVerificationSheetsToDfs(unittest.TestCase):
    """add_verification_sheets_to_dfs is the DataFrame-native sibling of
    preprocess() for callers that only ever have DataFrames: it
    must generate the 6 Verif sheets and dash the not-yet-computed Score
    columns, without mutating its input."""

    def test_generates_verif_sheets_and_dashes_scores(self):
        pristine_path = str(files("data").joinpath("pc_template.xlsx"))
        dfs = common.read_excel_file_to_dfs(pristine_path)
        dfs["Test Scores"] = dfs["Test Scores"].copy()
        dfs["Test Scores"]["Score"] = 0

        result = add_verification_sheets_to_dfs(dfs)

        for sheet_name in VERIFICATION_SHEETS:
            self.assertIn(sheet_name, result)
            self.assertFalse(result[sheet_name].empty)
        for sheet_name in ["Test Scores", "Category Scores", "Scenario Scores"]:
            self.assertTrue(
                (result[sheet_name]["Score"] == "-").all(), f"{sheet_name} not dashed"
            )
        # The input dfs are not mutated.
        self.assertTrue((dfs["Test Scores"]["Score"] == 0).all())
        self.assertNotIn("RI - RS Verification", dfs)


class TestPreprocessCliDashesScoreSheets(unittest.TestCase):
    """The CLI path dashes the score sheets via report_writer's
    blank_computed_columns call (driven by integrity.COMPUTED_SHEET_COLUMNS,
    same declaration as the DataFrame path). Pin the file-level behavior:
    every populated Score cell in the preprocessed file is "-"."""

    def test_preprocessed_file_score_cells_are_dash(self):
        pristine_path = str(files("data").joinpath("pc_template.xlsx"))
        runner = CliRunner()
        with (
            tempfile.TemporaryDirectory() as input_dir,
            tempfile.TemporaryDirectory() as output_dir,
        ):
            input_file = str(Path(input_dir) / "pc_template.xlsx")
            shutil.copy(pristine_path, input_file)
            common.add_version_sheet(
                input_file, cli_command=common.CliCommand.GENERATE_TEMPLATE
            )

            result = runner.invoke(
                preprocess, ["--input_file", input_file, "--output_path", output_dir]
            )
            self.assertEqual(result.exit_code, 0, result.output)

            output_file = str(Path(output_dir) / "pc_preprocessed_template.xlsx")
            output_wb = openpyxl.load_workbook(output_file)
            for sheet_name in ["Test Scores", "Category Scores", "Scenario Scores"]:
                ws = output_wb[sheet_name]
                header = [c.value for c in ws[1]]
                score_col = header.index("Score") + 1
                values = {
                    ws.cell(row, score_col).value for row in range(2, ws.max_row + 1)
                }
                self.assertEqual(values, {"-"}, f"{sheet_name}: {values}")


if __name__ == "__main__":
    unittest.main()
