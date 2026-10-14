# Copyright 2025-2026, Euro NCAP IVZW
# Created by IVEX NV (https://ivex.ai)
#
# Licensed under the Apache License 2.0.
# See http://www.apache.org/licenses/LICENSE-2.0 for details.

"""Rear Impact Static-Front scored per seat: Driver and Front Passenger at
0.5 points each, each from its own head-restraint readings, plus the
acceptance of workbooks that still carry the single "Front" score row."""

import os
import tempfile
import unittest
from importlib.resources import files

import openpyxl

from euroncap_rating_2026 import common
from euroncap_rating_2026.crash_protection import integrity
from euroncap_rating_2026.crash_protection.compute_score import calculate_score
from tests.crash_protection.legacy_fixtures import (
    make_legacy_cp_workbook,
    make_pre_static_front_split_cp_workbook,
)
from tests.crash_protection.test_file_score import _make_dfs_max

PASSENGER = "Front Passenger"
# Whiplash points outside Static-Front: Rear-Mid 1.5 + Rear-High 1.5 +
# Static-Rear 1.0.
OTHER_WHIPLASH_POINTS = 4.0
# Effective height at the test position green (>= 825 mm), lowest and
# rearmost red (< 789.99 mm), both backsets red (> 45 / 70 mm).
ONE_GREEN_THREE_RED = {
    "Effective height - Test position": 1000,
    "Effective height - Lowest and rearmost": 0,
    "Backset - Test position": 5000,
    "Backset - Lowest and rearmost": 5000,
}
# A head restraint below 790 mm triggers the effective height modifier.
LOW_HEAD_RESTRAINT = {"Effective height modifier": 700}


def _pristine_path():
    return str(files("data").joinpath("cp_template.xlsx"))


def _score(seat_values):
    """calculate_score on the all-max fixture with {seat: {criteria: value}}
    set on the Static-Front rows."""
    dfs = _make_dfs_max()
    whiplash = dfs["CP - Rear Whiplash"].copy()
    whiplash["Value"] = whiplash["Value"].astype(object)
    loadcase = whiplash["Loadcase"].ffill()
    seat = whiplash["Seat position"].ffill()
    for seat_name, values in seat_values.items():
        for criteria, value in values.items():
            mask = (
                (loadcase == "Static-Front")
                & (seat == seat_name)
                & (whiplash["Criteria"] == criteria)
            )
            assert mask.sum() == 1, (seat_name, criteria)
            whiplash.loc[mask, "Value"] = value
    dfs["CP - Rear Whiplash"] = whiplash
    return calculate_score(dfs)


def _whiplash(result):
    ts = result["Test Scores"]
    row = ts[ts["Stage subelement"] == "Whiplash"].iloc[0]
    return float(row["Score"]), float(row["Max score"])


def _static_front_rows(result, sheet):
    df = result[sheet]
    rows = df[df["Loadcase"].ffill() == "Static-Front"]
    return {
        seat: (float(score), float(max_score))
        for seat, score, max_score in zip(
            rows["Seat position"], rows["Score"], rows["Max score"]
        )
    }


class TestStaticFrontPerSeat(unittest.TestCase):
    def _assert_seats(self, result, driver, passenger):
        for sheet in ("CP - Dummy Scores", "CP - Body region scores"):
            with self.subTest(sheet=sheet):
                self.assertEqual(
                    _static_front_rows(result, sheet),
                    {"Driver": (driver, 0.5), PASSENGER: (passenger, 0.5)},
                )
        score, max_score = _whiplash(result)
        self.assertEqual(max_score, 5.0)
        self.assertAlmostEqual(
            score - OTHER_WHIPLASH_POINTS, driver + passenger, places=5
        )

    def test_both_seats_green_score_the_full_point(self):
        self._assert_seats(_score({}), 0.5, 0.5)

    def test_each_seat_scores_its_own_half(self):
        result = _score({PASSENGER: ONE_GREEN_THREE_RED})
        self._assert_seats(result, 0.5, 0.0)

    def test_seat_order_does_not_change_the_score(self):
        result = _score({"Driver": ONE_GREEN_THREE_RED})
        self._assert_seats(result, 0.0, 0.5)

    def test_effective_height_modifier_applies_to_its_own_seat(self):
        self._assert_seats(_score({"Driver": LOW_HEAD_RESTRAINT}), 0.0, 0.5)
        self._assert_seats(_score({PASSENGER: LOW_HEAD_RESTRAINT}), 0.5, 0.0)


class TestSingleFrontRowWorkbooks(unittest.TestCase):
    """Workbooks with the single Static-Front "Front" score row are accepted:
    the row reads as the Driver row and its inspection carries over to the
    Front Passenger row, so both seats keep the value entered before."""

    INSPECTION = 25

    def _save(self, wb):
        fd, path = tempfile.mkstemp(suffix=".xlsx")
        os.close(fd)
        wb.save(path)
        self.addCleanup(os.remove, path)
        return path

    def _workbooks(self):
        pristine = _pristine_path()
        legacy = make_legacy_cp_workbook(pristine)
        ws = legacy["CP - Body region scores"]
        row = next(
            r
            for r in range(2, ws.max_row + 1)
            if ws.cell(row=r, column=3).value == "Static-Front"
        )
        ws.cell(row=row, column=9).value = self.INSPECTION
        return {
            "before Rear Impact v1.3": legacy,
            "Rear Impact v1.3 layout": make_pre_static_front_split_cp_workbook(
                pristine, self.INSPECTION
            ),
        }

    def _assert_split(self, body, dummies):
        self.assertEqual(
            body,
            {
                "Driver": (self.INSPECTION, 0.5),
                PASSENGER: (self.INSPECTION, 0.5),
            },
        )
        self.assertEqual(dummies, {"Driver": 0.5, PASSENGER: 0.5})

    def test_file_path_splits_the_front_row(self):
        for label, wb in self._workbooks().items():
            with self.subTest(label):
                path = integrity.rebuild_trusted_input(self._save(wb))
                self.addCleanup(os.remove, path)
                out = openpyxl.load_workbook(path)
                body, dummies = {}, {}
                ws = out["CP - Body region scores"]
                for r in range(2, ws.max_row + 1):
                    if ws.cell(row=r, column=3).value == "Static-Front":
                        for rr in (r, r + 1):
                            body[ws.cell(row=rr, column=4).value] = (
                                ws.cell(row=rr, column=9).value,
                                ws.cell(row=rr, column=11).value,
                            )
                ws = out["CP - Dummy Scores"]
                for r in range(2, ws.max_row + 1):
                    if ws.cell(row=r, column=4).value == "Static-Front":
                        for rr in (r, r + 1):
                            dummies[ws.cell(row=rr, column=5).value] = ws.cell(
                                row=rr, column=9
                            ).value
                self._assert_split(body, dummies)

    def test_dataframe_path_splits_the_front_row(self):
        for label, wb in self._workbooks().items():
            with self.subTest(label):
                dfs = integrity.rebuild_trusted_input_dfs(
                    common.read_excel_file_to_dfs(self._save(wb))
                )
                body_df = dfs["CP - Body region scores"]
                rows = body_df[body_df["Loadcase"].ffill() == "Static-Front"]
                body = {
                    seat: (inspection, max_score)
                    for seat, inspection, max_score in zip(
                        rows["Seat position"], rows["Inspection [%]"], rows["Max score"]
                    )
                }
                dummy_df = dfs["CP - Dummy Scores"]
                rows = dummy_df[dummy_df["Loadcase"].ffill() == "Static-Front"]
                dummies = dict(zip(rows["Seat position"], rows["Max score"]))
                self._assert_split(body, dummies)

    def test_front_row_keeps_full_score_when_both_seats_green(self):
        dfs = _make_dfs_max()
        for sheet in ("CP - Dummy Scores", "CP - Body region scores"):
            df = dfs[sheet]
            front = df.index[df["Loadcase"] == "Static-Front"][0]
            df = df.drop(index=front + 1).reset_index(drop=True)
            df.loc[front, "Seat position"] = "Front"
            df.loc[front, "Max score"] = 1.0
            dfs[sheet] = df
        dfs = integrity.rebuild_trusted_input_dfs(dfs)
        score, max_score = _whiplash(calculate_score(dfs))
        self.assertEqual((score, max_score), (5.0, 5.0))

    def test_current_layout_missing_the_passenger_row_is_rejected(self):
        wb = openpyxl.load_workbook(_pristine_path())
        ws = wb["CP - Dummy Scores"]
        row = next(
            r
            for r in range(2, ws.max_row + 1)
            if ws.cell(row=r, column=5).value == PASSENGER
        )
        ws.delete_rows(row)
        with self.assertRaises(common.TemplateIntegrityError):
            integrity.rebuild_trusted_input(self._save(wb))

    def test_front_row_with_a_passenger_row_is_rejected(self):
        wb = openpyxl.load_workbook(_pristine_path())
        ws = wb["CP - Dummy Scores"]
        row = next(
            r
            for r in range(2, ws.max_row + 1)
            if ws.cell(row=r, column=4).value == "Static-Front"
        )
        ws.cell(row=row, column=5).value = "Front"
        with self.assertRaises(common.TemplateIntegrityError):
            integrity.rebuild_trusted_input(self._save(wb))

    def test_front_on_another_loadcase_is_rejected(self):
        wb = make_pre_static_front_split_cp_workbook(_pristine_path())
        ws = wb["CP - Dummy Scores"]
        row = next(
            r
            for r in range(2, ws.max_row + 1)
            if ws.cell(row=r, column=4).value == "Rear-High"
        )
        ws.cell(row=row, column=5).value = "Front"
        with self.assertRaises(common.TemplateIntegrityError):
            integrity.rebuild_trusted_input(self._save(wb))

    def test_old_front_identity_reads_as_driver(self):
        for sheet in ("CP - Dummy Scores", "CP - Body region scores"):
            with self.subTest(sheet=sheet):
                identity = {"Loadcase": "Static-Front", "Seat position": "Front"}
                self.assertEqual(
                    integrity.to_current_identity(sheet, identity),
                    {"Loadcase": "Static-Front", "Seat position": "Driver"},
                )


if __name__ == "__main__":
    unittest.main()
