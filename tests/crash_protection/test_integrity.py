# Copyright 2025-2026, Euro NCAP IVZW
# Created by IVEX NV (https://ivex.ai)
#
# Licensed under the Apache License 2.0.
# See http://www.apache.org/licenses/LICENSE-2.0 for details.

import os
import tempfile
import unittest
from unittest.mock import patch

import openpyxl

from euroncap_rating_2026 import common, version
from euroncap_rating_2026.crash_protection import integrity


class TestRebuildTrustedInput(unittest.TestCase):
    def setUp(self):
        from importlib.resources import files

        self.pristine_path = str(files("data").joinpath("cp_template.xlsx"))

    def _save_wb(self, wb):
        fd, path = tempfile.mkstemp(suffix=".xlsx")
        os.close(fd)
        wb.save(path)
        self.addCleanup(os.remove, path)
        return path

    def test_hpl_lpl_tamper_is_neutralized(self):
        """Editing HPL/LPL in the user-supplied template must have zero
        effect on the sanitized output: it must reflect the official
        template's thresholds instead."""
        wb = openpyxl.load_workbook(self.pristine_path)
        ws = wb["CP - Frontal Offset"]
        original_hpl = ws["F3"].value
        original_lpl = ws["G3"].value
        ws["F3"].value = -999
        ws["G3"].value = 99999
        user_path = self._save_wb(wb)

        sanitized_path = integrity.rebuild_trusted_input(user_path)
        self.addCleanup(os.remove, sanitized_path)

        sanitized_wb = openpyxl.load_workbook(sanitized_path, data_only=True)
        sanitized_ws = sanitized_wb["CP - Frontal Offset"]
        self.assertEqual(sanitized_ws["F3"].value, original_hpl)
        self.assertEqual(sanitized_ws["G3"].value, original_lpl)
        self.assertNotEqual(sanitized_ws["F3"].value, -999)
        self.assertNotEqual(sanitized_ws["G3"].value, 99999)

    def test_max_score_tamper_is_neutralized(self):
        wb = openpyxl.load_workbook(self.pristine_path)
        ws = wb["CP - Dummy Scores"]
        original_max_score = ws["I2"].value
        ws["I2"].value = 99999
        user_path = self._save_wb(wb)

        sanitized_path = integrity.rebuild_trusted_input(user_path)
        self.addCleanup(os.remove, sanitized_path)

        sanitized_wb = openpyxl.load_workbook(sanitized_path, data_only=True)
        self.assertEqual(
            sanitized_wb["CP - Dummy Scores"]["I2"].value, original_max_score
        )

    def test_oem_prediction_and_value_are_preserved(self):
        """Legitimate grey-cell input must survive the rebuild unchanged."""
        wb = openpyxl.load_workbook(self.pristine_path)
        ws = wb["CP - Frontal Offset"]
        ws["I2"].value = "Green"  # OEM Prediction
        ws["J2"].value = 12.34  # Value
        user_path = self._save_wb(wb)

        sanitized_path = integrity.rebuild_trusted_input(user_path)
        self.addCleanup(os.remove, sanitized_path)

        sanitized_wb = openpyxl.load_workbook(sanitized_path, data_only=True)
        sanitized_ws = sanitized_wb["CP - Frontal Offset"]
        self.assertEqual(sanitized_ws["I2"].value, "Green")
        self.assertEqual(sanitized_ws["J2"].value, 12.34)

    def test_input_parameters_value_is_preserved(self):
        wb = openpyxl.load_workbook(self.pristine_path)
        ws = wb["Input parameters"]
        header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        value_col = header.index("Value") + 1
        ws.cell(row=2, column=value_col).value = "Yes"
        user_path = self._save_wb(wb)

        sanitized_path = integrity.rebuild_trusted_input(user_path)
        self.addCleanup(os.remove, sanitized_path)

        sanitized_wb = openpyxl.load_workbook(sanitized_path, data_only=True)
        self.assertEqual(
            sanitized_wb["Input parameters"].cell(row=2, column=value_col).value,
            "Yes",
        )

    def test_reordered_criteria_rows_are_rejected(self):
        wb = openpyxl.load_workbook(self.pristine_path)
        ws = wb["CP - Rear Whiplash"]
        row2 = [ws.cell(row=2, column=c).value for c in range(1, 11)]
        row3 = [ws.cell(row=3, column=c).value for c in range(1, 11)]
        for c in range(1, 11):
            ws.cell(row=2, column=c).value = row3[c - 1]
            ws.cell(row=3, column=c).value = row2[c - 1]
        user_path = self._save_wb(wb)

        with self.assertRaises(common.TemplateIntegrityError):
            integrity.rebuild_trusted_input(user_path)

    def test_deleted_row_is_rejected(self):
        wb = openpyxl.load_workbook(self.pristine_path)
        ws = wb["CP - Side Pole"]
        ws.delete_rows(3)
        user_path = self._save_wb(wb)

        with self.assertRaises(common.TemplateIntegrityError):
            integrity.rebuild_trusted_input(user_path)

    def test_vru_prediction_grid_passthrough_is_untouched(self):
        """The VRU prediction grid isn't schema'd -- it must pass through
        from the user's file as-is."""
        wb = openpyxl.load_workbook(self.pristine_path)
        ws = wb["CP - VRU Prediction"]
        ws["E5"].value = "sentinel-value"
        user_path = self._save_wb(wb)

        sanitized_path = integrity.rebuild_trusted_input(user_path)
        self.addCleanup(os.remove, sanitized_path)

        sanitized_wb = openpyxl.load_workbook(sanitized_path, data_only=True)
        self.assertEqual(
            sanitized_wb["CP - VRU Prediction"]["E5"].value, "sentinel-value"
        )


class TestPreprocessToComputeScoreSignature(unittest.TestCase):
    """Exercises the second hop's protection: preprocess signs its output's
    protected cells, and compute-score must reject a file whose protected
    cells were edited after signing -- using crash_protection's real
    SHEET_SCHEMAS against the packaged template."""

    def setUp(self):
        from importlib.resources import files

        self.pristine_path = str(files("data").joinpath("cp_template.xlsx"))

    def test_signed_file_verifies_and_tampered_file_is_rejected(self):
        wb = openpyxl.load_workbook(self.pristine_path)
        common.add_integrity_sheet_to_wb(wb, integrity.SHEET_SCHEMAS)

        fd, path = tempfile.mkstemp(suffix=".xlsx")
        os.close(fd)
        wb.save(path)
        self.addCleanup(os.remove, path)

        # Unmodified: verification must pass.
        clean_wb = openpyxl.load_workbook(path, data_only=True)
        common.verify_protected_signature(clean_wb, integrity.SHEET_SCHEMAS)

        # Tamper a protected HPL cell after signing: verification must reject.
        tampered_wb = openpyxl.load_workbook(path)
        tampered_wb["CP - Frontal Offset"]["F3"].value = -999
        tampered_wb.save(path)
        reloaded_wb = openpyxl.load_workbook(path, data_only=True)
        with self.assertRaises(common.TemplateIntegrityError):
            common.verify_protected_signature(reloaded_wb, integrity.SHEET_SCHEMAS)

    def test_editing_grey_cell_after_signing_still_verifies(self):
        """Editing a legitimate grey cell (e.g. a newly-added VRU 'Inspection
        [%]' style input) after signing must not trip the signature check."""
        wb = openpyxl.load_workbook(self.pristine_path)
        common.add_integrity_sheet_to_wb(wb, integrity.SHEET_SCHEMAS)

        fd, path = tempfile.mkstemp(suffix=".xlsx")
        os.close(fd)
        wb.save(path)
        self.addCleanup(os.remove, path)

        edited_wb = openpyxl.load_workbook(path)
        edited_wb["CP - Frontal Offset"]["J3"].value = 42.0  # "Value" grey column
        edited_wb.save(path)

        reloaded_wb = openpyxl.load_workbook(path, data_only=True)
        common.verify_protected_signature(reloaded_wb, integrity.SHEET_SCHEMAS)


class TestRebuildTrustedInputDfs(unittest.TestCase):
    """DataFrame-native sibling of TestRebuildTrustedInput: same tamper
    scenarios, but through rebuild_trusted_input_dfs (used by
    add_vru_sheets_to_dfs/calculate_score(dfs)) instead of rebuild_trusted_input
    (used by the file-based preprocess() CLI command)."""

    def setUp(self):
        from importlib.resources import files

        pristine_path = str(files("data").joinpath("cp_template.xlsx"))
        self.pristine_dfs = common.read_excel_file_to_dfs(pristine_path)

    def _dfs(self):
        return {name: df.copy() for name, df in self.pristine_dfs.items()}

    def test_hpl_lpl_tamper_is_neutralized(self):
        dfs = self._dfs()
        df = dfs["CP - Frontal Offset"]
        original_hpl = df.loc[1, "HPL"]
        original_lpl = df.loc[1, "LPL"]
        df.loc[1, "HPL"] = -999
        df.loc[1, "LPL"] = 99999

        sanitized = integrity.rebuild_trusted_input_dfs(dfs)

        sanitized_df = sanitized["CP - Frontal Offset"]
        self.assertEqual(sanitized_df.loc[1, "HPL"], original_hpl)
        self.assertEqual(sanitized_df.loc[1, "LPL"], original_lpl)
        self.assertNotEqual(sanitized_df.loc[1, "HPL"], -999)
        self.assertNotEqual(sanitized_df.loc[1, "LPL"], 99999)

    def test_max_score_tamper_is_neutralized(self):
        dfs = self._dfs()
        original_max_score = dfs["CP - Dummy Scores"].loc[0, "Max score"]
        dfs["CP - Dummy Scores"].loc[0, "Max score"] = 99999

        sanitized = integrity.rebuild_trusted_input_dfs(dfs)

        self.assertEqual(
            sanitized["CP - Dummy Scores"].loc[0, "Max score"], original_max_score
        )

    def test_oem_prediction_and_value_are_preserved(self):
        """Legitimate grey-cell input must survive the rebuild unchanged."""
        dfs = self._dfs()
        dfs["CP - Frontal Offset"] = dfs["CP - Frontal Offset"].astype(object)
        dfs["CP - Frontal Offset"].loc[0, "OEM Prediction"] = "Green"
        dfs["CP - Frontal Offset"].loc[0, "Value"] = 12.34

        sanitized = integrity.rebuild_trusted_input_dfs(dfs)

        sanitized_df = sanitized["CP - Frontal Offset"]
        self.assertEqual(sanitized_df.loc[0, "OEM Prediction"], "Green")
        self.assertEqual(sanitized_df.loc[0, "Value"], 12.34)

    def test_input_parameters_value_is_preserved(self):
        dfs = self._dfs()
        dfs["Input parameters"] = dfs["Input parameters"].astype(object)
        dfs["Input parameters"].loc[0, "Value"] = "Yes"

        sanitized = integrity.rebuild_trusted_input_dfs(dfs)

        self.assertEqual(sanitized["Input parameters"].loc[0, "Value"], "Yes")

    def test_reordered_criteria_rows_are_rejected(self):
        dfs = self._dfs()
        df = dfs["CP - Rear Whiplash"]
        df.iloc[[0, 1]] = df.iloc[[1, 0]].values

        with self.assertRaises(common.TemplateIntegrityError):
            integrity.rebuild_trusted_input_dfs(dfs)

    def test_deleted_row_is_rejected(self):
        dfs = self._dfs()
        dfs["CP - Side Pole"] = (
            dfs["CP - Side Pole"].drop(index=1).reset_index(drop=True)
        )

        with self.assertRaises(common.TemplateIntegrityError):
            integrity.rebuild_trusted_input_dfs(dfs)

    def test_vru_prediction_grid_passthrough_is_untouched(self):
        """The VRU prediction grid isn't schema'd -- it must pass through
        from the user's dfs as-is."""
        dfs = self._dfs()
        dfs["CP - VRU Prediction"] = dfs["CP - VRU Prediction"].astype(object)
        dfs["CP - VRU Prediction"].iloc[3, 4] = "sentinel-value"

        sanitized = integrity.rebuild_trusted_input_dfs(dfs)

        self.assertEqual(sanitized["CP - VRU Prediction"].iloc[3, 4], "sentinel-value")

    def test_sheet_missing_from_dfs_is_skipped_not_rejected(self):
        """add_vru_sheets_to_dfs/calculate_score(dfs) callers may legitimately
        pass a partial dfs dict; a schema'd sheet that isn't present must not
        raise (deviates from rebuild_trusted_input, which requires the
        complete monolithic file)."""
        dfs = {"Input parameters": self.pristine_dfs["Input parameters"].copy()}

        sanitized = integrity.rebuild_trusted_input_dfs(dfs)

        self.assertIn("Input parameters", sanitized)
        self.assertNotIn("CP - Frontal Offset", sanitized)


class TestDocumentationDataSheetCoverage(unittest.TestCase):
    """Regression test for the "sharp edge" documented in integrity.py's
    module docstring: any pristine-template sheet not covered by
    SHEET_SCHEMAS or PASSTHROUGH_SHEETS is silently reset to blank by
    rebuild_trusted_workbook/rebuild_trusted_dfs."""

    def setUp(self):
        from importlib.resources import files

        self.pristine_path = str(files("data").joinpath("cp_template.xlsx"))

    def test_every_pristine_sheet_is_covered(self):
        wb = openpyxl.load_workbook(self.pristine_path)
        covered = set(integrity.SHEET_SCHEMAS) | set(integrity.PASSTHROUGH_SHEETS)
        uncovered = set(wb.sheetnames) - covered
        self.assertEqual(uncovered, set())

    def test_documentation_and_data_sheets_are_hidden(self):
        wb = openpyxl.load_workbook(self.pristine_path)
        for name in ("Documentation", "Data"):
            self.assertIn(name, wb.sheetnames)
            self.assertEqual(wb[name].sheet_state, "hidden")

    def test_missing_input_check_covers_only_template_sheets(self):
        """find_missing_required_inputs must inspect prediction-template
        sheets only: every sheet it iterates (schemas and grid sheets) has
        to exist in the pristine template, so assessment-only sheets can
        never produce findings."""
        wb = openpyxl.load_workbook(self.pristine_path)
        checked = set(integrity.SHEET_SCHEMAS) | set(
            integrity.GRID_SHEET_SECTION_STAGE_ELEMENTS
        )
        self.assertEqual(checked - set(wb.sheetnames), set())


class TestFindMissingRequiredInputs(unittest.TestCase):
    """Exercises integrity.find_missing_required_inputs against the real
    packaged cp_template.xlsx: grey != OEM-required, and scoping to one
    protocol element must not flag other elements' sheets."""

    def setUp(self):
        from importlib.resources import files

        self.pristine_path = str(files("data").joinpath("cp_template.xlsx"))
        self.wb = openpyxl.load_workbook(self.pristine_path, data_only=True)

    def test_fresh_template_only_reports_oem_prediction_gaps(self):
        """A freshly generated, entirely-blank template must not report
        Inspection [%] (inspector-filled later) or a criteria sheet's Value
        (measured result, recomputed at scoring time, never OEM-fillable) as
        missing -- only the OEM Prediction cells and Input parameters rows
        the OEM actually owns."""
        missing = integrity.find_missing_required_inputs(self.wb)

        self.assertTrue(
            missing, "expected some OEM Prediction gaps on a blank template"
        )
        self.assertTrue(
            all(
                (m.sheet, m.column) == ("Input parameters", "Value")
                or m.column == "OEM Prediction"
                for m in missing
            )
        )
        # "CP - VRU Prediction" doesn't appear: every one of its sections
        # (Headforms, Upper legform, aPLI - Femur, aPLI - Knee & Tibia) is
        # exempt from the missing-input check --
        # consistency.find_prediction_inconsistencies'
        # no_prediction_provided check covers it instead.
        self.assertEqual(
            {m.sheet for m in missing},
            {
                "CP - Frontal Offset",
                "CP - Frontal FW",
                "CP - Side MDB",
                "CP - Side Pole",
                "Input parameters",
            },
        )

    def test_sheets_without_oem_prediction_column_report_nothing(self):
        """Sled & VT, Farside and Whiplash have no OEM Prediction column at
        all -- their only grey column (Value) is measured-result-only, so
        they must never appear regardless of how blank the sheet is."""
        missing = integrity.find_missing_required_inputs(self.wb)
        reported_sheets = {m.sheet for m in missing}

        for sheet in (
            "CP - Frontal Sled & VT",
            "CP - Side Farside",
            "CP - Rear Whiplash",
        ):
            self.assertNotIn(sheet, reported_sheets)

    def test_scoping_to_one_protocol_element_excludes_others(self):
        scoped = integrity.find_missing_required_inputs(
            self.wb, stage_element="Frontal Impact", stage_subelement="Offset"
        )

        self.assertTrue(scoped)
        self.assertEqual({m.sheet for m in scoped}, {"CP - Frontal Offset"})

    def test_scoping_to_element_with_no_prediction_column_is_empty(self):
        scoped = integrity.find_missing_required_inputs(
            self.wb, stage_element="Frontal Impact", stage_subelement="Sled & VT"
        )

        self.assertEqual(scoped, [])

    def test_only_sliding_scale_criteria_rows_are_required(self):
        """OEM Prediction is only required on rows whose criterion has
        both a higher and a lower performance limit -- modifier and
        single-limit/capping rows (DAMAGE, Pedal blocking, ...) have no
        predictable color and must not fail a complete submission. The
        gate reads HPL/LPL from the pristine template, so every reported
        row must carry both limits there."""
        missing = integrity.find_missing_required_inputs(self.wb)
        prediction_gaps = [m for m in missing if m.column == "OEM Prediction"]
        self.assertTrue(prediction_gaps)

        pristine = openpyxl.load_workbook(self.pristine_path)
        for entry in prediction_gaps:
            ws = pristine[entry.sheet]
            header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
            row = int("".join(ch for ch in entry.cell if ch.isdigit()))
            for limit in ("HPL", "LPL"):
                value = ws.cell(row=row, column=header.index(limit) + 1).value
                self.assertFalse(
                    common.is_empty_cell(value),
                    f"{entry.sheet}!{entry.cell} required without {limit}",
                )

        # And the converse spot-check: the Offset DAMAGE modifier row (no
        # limits in the pristine template) is never reported.
        offset_criteria = {
            m.row_identity.get("Criteria")
            for m in prediction_gaps
            if m.sheet == "CP - Frontal Offset"
        }
        self.assertNotIn("DAMAGE", offset_criteria)
        self.assertIn("HIC15", offset_criteria)

    def test_descriptions_are_semantic(self):
        missing = integrity.find_missing_required_inputs(
            self.wb, stage_element="Frontal Impact", stage_subelement="Offset"
        )

        entry = missing[0]
        self.assertIn("CP - Frontal Offset", entry.description)
        self.assertIn("OEM Prediction", entry.description)
        # Real row identity (Loadcase/Seat position/Dummy/Body region/
        # Criteria), not a bare cell coordinate like "I3".
        self.assertTrue(all(v is not None for v in entry.row_identity.values()))
        self.assertNotRegex(entry.description, r"^[A-Z]+\d+$")

    def test_filled_oem_prediction_is_not_reported(self):
        ws = self.wb["CP - Frontal Offset"]
        header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        col_idx = header.index("OEM Prediction") + 1
        for row in range(2, ws.max_row + 1):
            ws.cell(row=row, column=col_idx, value="Green")

        missing = integrity.find_missing_required_inputs(
            self.wb, stage_element="Frontal Impact", stage_subelement="Offset"
        )

        self.assertEqual(missing, [])

    def test_vru_grid_requires_no_sections(self):
        """No section of the VRU prediction grid is required-checked
        per-cell any more: Headforms joined the three legform sections as
        exempt -- the OEM fills only
        the cells matching the car's layout and its own prediction.
        consistency.find_prediction_inconsistencies' no_prediction_provided
        check is what now guarantees each section gets at least one
        prediction, not this one."""
        unscoped = [
            m
            for m in integrity.find_missing_required_inputs(self.wb)
            if m.sheet == "CP - VRU Prediction"
        ]
        head = [
            m
            for m in integrity.find_missing_required_inputs(
                self.wb, stage_element="VRU Impact", stage_subelement="Head Impact"
            )
            if m.sheet == "CP - VRU Prediction"
        ]

        self.assertEqual(unscoped, [])
        self.assertEqual(head, [])

    def test_vru_legform_scopes_require_only_input_parameters(self):
        """The Pelvis Impact / Leg Impact scopes (and the combined "Pelvis &
        Leg Impact" alias Test Scores uses) stay valid, but require only
        their own Input parameters rows -- never a legform matrix cell."""
        pelvis = integrity.find_missing_required_inputs(
            self.wb, stage_element="VRU Impact", stage_subelement="Pelvis Impact"
        )
        leg = integrity.find_missing_required_inputs(
            self.wb, stage_element="VRU Impact", stage_subelement="Leg Impact"
        )

        # Each scope still requires its own Input parameters row...
        self.assertEqual({m.sheet for m in pelvis}, {"Input parameters"})
        self.assertEqual({m.sheet for m in leg}, {"Input parameters"})
        self.assertEqual(len(pelvis), 1)
        self.assertEqual(len(leg), 1)
        # ...and the combined alias reports no grid cells either (it never
        # pulled in the split-named Input parameters rows).
        combined = integrity.find_missing_required_inputs(
            self.wb, stage_element="VRU Impact", stage_subelement="Pelvis & Leg Impact"
        )
        self.assertEqual(combined, [])

    def test_vru_grid_not_reported_for_other_protocol_elements(self):
        scoped = integrity.find_missing_required_inputs(
            self.wb, stage_element="Frontal Impact", stage_subelement="Offset"
        )

        self.assertNotIn("CP - VRU Prediction", {m.sheet for m in scoped})

    def test_filled_vru_grid_cell_is_not_reported(self):
        ws = self.wb["CP - VRU Prediction"]
        coords = sorted(common.get_required_dropdown_coordinates(ws))
        first_coord = coords[0]
        ws[first_coord] = "Green"

        missing = integrity.find_missing_required_inputs(
            self.wb, stage_element="VRU Impact", stage_subelement="Head Impact"
        )

        self.assertNotIn(
            first_coord,
            {m.cell for m in missing if m.sheet == "CP - VRU Prediction"},
        )

    def test_find_missing_required_inputs_in_file_round_trips(self):
        missing = integrity.find_missing_required_inputs_in_file(
            self.pristine_path,
            stage_element="Frontal Impact",
            stage_subelement="Offset",
        )

        self.assertTrue(missing)
        self.assertEqual({m.sheet for m in missing}, {"CP - Frontal Offset"})


class TestScopeValidation(unittest.TestCase):
    """An unknown (stage_element, stage_subelement) must raise instead of
    silently returning zero findings -- a typo'd caller-side mapping (e.g.
    snake_case identifiers) would otherwise silently disable the check.
    VALID_SCOPE_PAIRS must in turn cover every pair the packaged
    template's own sheets use, so no legitimate scope ever raises."""

    def setUp(self):
        from importlib.resources import files

        self.pristine_path = str(files("data").joinpath("cp_template.xlsx"))

    def test_unknown_scope_raises(self):
        wb = openpyxl.load_workbook(self.pristine_path, data_only=True)

        with self.assertRaises(ValueError):
            integrity.find_missing_required_inputs(
                wb, stage_element="frontal_impact", stage_subelement="offset"
            )
        with self.assertRaises(ValueError):
            integrity.find_missing_required_inputs(wb, stage_element="Bogus")

    def test_declared_pairs_cover_the_templates_own_literals(self):
        wb = openpyxl.load_workbook(self.pristine_path, data_only=True)
        template_pairs = set()
        for sheet in ("Test Scores", "Input parameters"):
            ws = wb[sheet]
            header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
            element_col = header.index("Stage element") + 1
            subelement_col = header.index("Stage subelement") + 1
            last_element = last_subelement = None
            for row in range(2, ws.max_row + 1):
                element = ws.cell(row=row, column=element_col).value
                subelement = ws.cell(row=row, column=subelement_col).value
                if not common.is_empty_cell(element):
                    last_element = element
                    last_subelement = None
                if not common.is_empty_cell(subelement):
                    last_subelement = subelement
                if last_element is not None and last_subelement is not None:
                    template_pairs.add((last_element, last_subelement))

        self.assertTrue(template_pairs)
        undeclared = template_pairs - set(integrity.VALID_SCOPE_PAIRS)
        self.assertEqual(undeclared, set())


class TestPristineTemplateRegression(unittest.TestCase):
    """The audit assertion that on the pure packaged
    template every required cell is empty, so the unscoped check must
    report exactly the template's grey-painted input set -- no more (a
    false requirement fails valid submissions) and no less (a false pass
    ships an unchecked element). The literal per-sheet counts pin the
    required set; evolving the template legitimately means updating them
    in the same change. The fill cross-check then verifies each reported
    cell against the independent ground truth -- the grey paint the OEM
    actually sees."""

    # "CP - VRU Prediction" doesn't appear at all: every one of its
    # sections (Headforms, Upper legform, aPLI - Femur, aPLI - Knee &
    # Tibia) is exempt from the missing-input check -- the OEM fills only
    # the cells matching the car's layout and its own prediction, so all
    # 462 of the sheet's grey cells
    # are the deliberate exception to "reported == grey".
    # consistency.find_prediction_inconsistencies' no_prediction_provided
    # check is the independent guarantee that each section still gets at
    # least one prediction.
    EXPECTED_PER_SHEET = {
        "CP - Frontal Offset": 47,
        "CP - Frontal FW": 47,
        "CP - Side MDB": 16,
        "CP - Side Pole": 8,
        "Input parameters": 6,
    }

    def setUp(self):
        from importlib.resources import files

        self.pristine_path = str(files("data").joinpath("cp_template.xlsx"))

    def test_unscoped_check_reports_exactly_the_expected_set(self):
        from collections import Counter

        wb = openpyxl.load_workbook(self.pristine_path, data_only=True)
        missing = integrity.find_missing_required_inputs(wb)

        self.assertEqual(
            dict(Counter(m.sheet for m in missing)), self.EXPECTED_PER_SHEET
        )

    def test_every_reported_cell_is_grey_painted_in_the_template(self):
        """Each reported coordinate must carry the grey input paint in the
        packaged template (a theme-shaded or literal-grey placeholder
        fill) -- a reported cell without it means the check requires
        something the template never asked the OEM to fill. Deliberately
        one-directional: some grey cells (Inspection [%], a criteria
        sheet's Value, non-sliding-scale OEM Predictions) are legitimately
        not required at prediction-upload time."""
        wb = openpyxl.load_workbook(self.pristine_path, data_only=True)
        missing = integrity.find_missing_required_inputs(wb)
        styled_wb = openpyxl.load_workbook(self.pristine_path)

        for entry in missing:
            fill = styled_wb[entry.sheet][entry.cell].fill
            self.assertEqual(
                fill.fill_type,
                "solid",
                f"{entry.sheet}!{entry.cell} reported but not painted",
            )
            if fill.fgColor.type == "theme":
                self.assertLess(
                    fill.fgColor.tint,
                    0,
                    f"{entry.sheet}!{entry.cell} has an unshaded theme fill",
                )
            else:
                grey_hex = str(fill.fgColor.rgb or "")[-6:].upper()
                self.assertIn(
                    grey_hex,
                    ("808080", "D9D9D9"),
                    f"{entry.sheet}!{entry.cell} reported but painted "
                    f"{fill.fgColor.rgb}",
                )


class TestFindMissingRequiredInputsVersionMismatch(unittest.TestCase):
    """A workbook generated by an incompatible library version must surface
    that drift as a finding, on top of -- not instead
    of -- the ordinary missing-input findings."""

    def setUp(self):
        from importlib.resources import files

        self.pristine_path = str(files("data").joinpath("cp_template.xlsx"))

    def _wb_with_version(self, version_string):
        wb = openpyxl.load_workbook(self.pristine_path, data_only=True)
        ws = wb["Version"] if "Version" in wb.sheetnames else wb.create_sheet("Version")
        ws["A1"] = "Generated with version"
        ws["B1"] = version_string
        return wb

    def test_incompatible_version_adds_finding_alongside_missing_inputs(self):
        with patch.object(common, "VERSION", "5.4.6"):
            wb = self._wb_with_version("5.4.0")
            findings = integrity.find_missing_required_inputs(wb)

        version_findings = [
            f for f in findings if f.check == common.CHECK_VERSION_MISMATCH
        ]
        self.assertEqual(len(version_findings), 1)
        self.assertTrue(any(f.check == common.CHECK_MISSING_INPUT for f in findings))

    def test_incompatible_version_surfaces_through_find_consistency_findings_too(self):
        with patch.object(common, "VERSION", "5.4.6"):
            wb = self._wb_with_version("5.4.0")
            findings = integrity.find_consistency_findings(wb)

        self.assertTrue(any(f.check == common.CHECK_VERSION_MISMATCH for f in findings))

    def test_version_finding_is_reported_even_when_scoped(self):
        """A version mismatch is a whole-workbook problem, not a per-element
        one -- it must not be dropped just because the caller scoped the
        check to one protocol element."""
        with patch.object(common, "VERSION", "5.4.6"):
            wb = self._wb_with_version("5.4.0")
            findings = integrity.find_missing_required_inputs(
                wb, stage_element="Frontal Impact", stage_subelement="Offset"
            )

        self.assertTrue(any(f.check == common.CHECK_VERSION_MISMATCH for f in findings))

    def test_matching_version_adds_no_finding(self):
        wb = self._wb_with_version(version.VERSION)
        findings = integrity.find_missing_required_inputs(wb)

        self.assertFalse(
            any(f.check == common.CHECK_VERSION_MISMATCH for f in findings)
        )

    def test_no_version_sheet_adds_no_finding(self):
        """Regression guard: the many fixtures across this suite that never
        add a Version sheet at all must not start reporting a spurious
        mismatch."""
        wb = openpyxl.load_workbook(self.pristine_path, data_only=True)
        wb.remove(wb["Version"])
        findings = integrity.find_missing_required_inputs(wb)

        self.assertFalse(
            any(f.check == common.CHECK_VERSION_MISMATCH for f in findings)
        )

    def _rename_stage_subelement_header(self, wb):
        """Reproduces a real-world shape: a renamed identity
        column ("Stage Subelement" vs "Stage subelement") that makes the
        schema walk raise TemplateIntegrityError."""
        ws = wb["Input parameters"]
        header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        col = header.index("Stage subelement") + 1
        ws.cell(row=1, column=col).value = "Stage Subelement"

    def test_structurally_broken_and_version_mismatched_returns_only_version_finding(
        self,
    ):
        """The real-world case: an old-version file's renamed header would
        otherwise raise TemplateIntegrityError before any finding --
        including the version mismatch itself -- reaches the caller."""
        with patch.object(common, "VERSION", "5.4.6"):
            wb = self._wb_with_version("5.4.0")
            self._rename_stage_subelement_header(wb)
            findings = integrity.find_missing_required_inputs(wb)

        self.assertEqual(len(findings), 1)
        self.assertEqual(findings[0].check, common.CHECK_VERSION_MISMATCH)

    def test_structurally_broken_without_version_mismatch_still_raises(self):
        """Regression guard: with no version mismatch to blame, a
        genuinely broken template must still raise as before -- the
        schema-integrity contract predates and is independent of the version finding."""
        wb = openpyxl.load_workbook(self.pristine_path, data_only=True)
        self._rename_stage_subelement_header(wb)

        with self.assertRaises(common.TemplateIntegrityError):
            integrity.find_missing_required_inputs(wb)


if __name__ == "__main__":
    unittest.main()
