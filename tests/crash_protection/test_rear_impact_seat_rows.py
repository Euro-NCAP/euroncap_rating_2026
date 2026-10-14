# Copyright 2025-2026, Euro NCAP IVZW
# Created by IVEX NV (https://ivex.ai)
#
# Licensed under the Apache License 2.0.
# See http://www.apache.org/licenses/LICENSE-2.0 for details.

"""Rear Impact static assessment with a second and a third seat row, the
"Non-Use position" criterion, per-row torso angles, and the acceptance of
workbooks produced against the previous template revision (rear seat row
"Rear", "Effective height modifier", single "Torso angle (rear)")."""

import logging
import os
import unittest
from importlib.resources import files

import numpy as np
import openpyxl
import pandas as pd

from euroncap_rating_2026 import common
from euroncap_rating_2026.crash_protection import consistency, data_loader, integrity
from euroncap_rating_2026.crash_protection.compute_score import calculate_score
from euroncap_rating_2026.crash_protection.dummy import Dummy, parse_non_use_position
from tests.crash_protection.legacy_fixtures import (
    make_legacy_cp_dfs,
    make_legacy_cp_workbook,
)
from tests.crash_protection.test_file_score import _make_dfs_max

SECOND_ROW = "Rear (second row)"
THIRD_ROW = "Rear (third row)"
# Static-Rear readings: both effective heights and both backsets pass at a
# 25 degree torso angle (backset limit 7.128 * 25 + 153 = 331.2 mm).
PASSING_GEOMETRY = (730, 780, None, 300, 300)
# The geometry points ladder: a 340 mm backset fails the 331.2 mm limit and
# a 710 mm effective height fails the 720 mm lowest-position requirement.
HEIGHT_AND_MID_BACKSET = (730, 780, None, 340, 300)
HEIGHT_ONLY = (730, 780, None, 340, 340)
HEIGHT_AND_WORST_CASE_BACKSET = (730, 780, None, 300, 340)
LOW_HEIGHT = (710, 780, None, 300, 300)
# The effective heights exactly at the 720 / 770 mm minimum pass; 0.01 mm
# below the lowest-position minimum does not.
MINIMUM_HEIGHTS = (720, 770, None, 300, 300)
JUST_BELOW_MINIMUM_HEIGHT = (719.99, 780, None, 300, 300)
OTHER_WHIPLASH_POINTS = 4.0  # Rear-Mid 1.5 + Rear-High 1.5 + Static-Front 0.5 + 0.5


def _pristine_path():
    return str(files("data").joinpath("cp_template.xlsx"))


def _with_values(values, non_use):
    values = list(values)
    values[2] = non_use
    return values


def _whiplash_scores(result):
    ts = result["Test Scores"]
    row = ts[ts["Stage subelement"] == "Whiplash"].iloc[0]
    return float(row["Score"]), float(row["Max score"])


def _static_rear_rows(result, sheet):
    df = result[sheet]
    return df[df["Loadcase"].ffill() == "Static-Rear"]


def _static_rear_dfs(second_angle, third_angle, second_values, third_values=None):
    """The all-max fixture with the Static-Rear block and torso angles set."""
    dfs = _make_dfs_max()
    params = dfs["Input parameters"].copy()
    params["Value"] = params["Value"].astype(object)
    params.loc[
        params["Input parameter"] == "Torso angle (rear second row)", "Value"
    ] = second_angle
    params.loc[params["Input parameter"] == "Torso angle (rear third row)", "Value"] = (
        third_angle
    )
    dfs["Input parameters"] = params
    whiplash = dfs["CP - Rear Whiplash"].copy()
    whiplash["Value"] = whiplash["Value"].astype(object)
    seat = whiplash["Seat position"].ffill()
    for seat_name, values in ((SECOND_ROW, second_values), (THIRD_ROW, third_values)):
        if values is None:
            continue
        whiplash.loc[seat == seat_name, "Value"] = list(values)
    dfs["CP - Rear Whiplash"] = whiplash
    return dfs


class TestStaticRearScoring(unittest.TestCase):
    """Static-Rear carries 1 point, split 75 % geometry / 25 % non-use
    position per seat row; the second row takes the whole point when the
    vehicle has no third row, 0.5 each when it has one."""

    def _score(self, second_angle, third_angle, second_values, third_values=None):
        result = calculate_score(
            _static_rear_dfs(second_angle, third_angle, second_values, third_values)
        )
        score, max_score = _whiplash_scores(result)
        self.assertEqual(max_score, 5.0)
        return score - OTHER_WHIPLASH_POINTS, result

    def test_no_third_row_pass_fail_blank(self):
        for non_use, expected in (("PASS", 1.0), ("FAIL", 0.75), (None, 0.75)):
            with self.subTest(non_use=non_use):
                score, _ = self._score(
                    25, "N/A", _with_values(PASSING_GEOMETRY, non_use)
                )
                self.assertAlmostEqual(score, expected, places=5)

    def test_numeric_encoding_1000_is_pass_and_1_is_not(self):
        score, _ = self._score(25, "N/A", _with_values(PASSING_GEOMETRY, 1000))
        self.assertAlmostEqual(score, 1.0, places=5)
        score, _ = self._score(25, "N/A", _with_values(PASSING_GEOMETRY, 1))
        self.assertAlmostEqual(score, 0.75, places=5)
        score, _ = self._score(25, "N/A", _with_values(PASSING_GEOMETRY, 0))
        self.assertAlmostEqual(score, 0.75, places=5)

    def test_third_row_both_one_neither(self):
        cases = (
            ("PASS", "PASS", 1.0),
            ("PASS", "FAIL", 0.875),
            (None, None, 0.75),
        )
        for second, third, expected in cases:
            with self.subTest(second=second, third=third):
                score, _ = self._score(
                    25,
                    25,
                    _with_values(PASSING_GEOMETRY, second),
                    _with_values(PASSING_GEOMETRY, third),
                )
                self.assertAlmostEqual(score, expected, places=5)

    def test_geometry_ladder_one_row(self):
        # Height 0.375, mid backset +0.25, worst-case backset +0.125 (only on
        # top of the mid backset), non-use +0.25.
        cases = (
            (PASSING_GEOMETRY, 1.0),
            (HEIGHT_AND_MID_BACKSET, 0.875),
            (HEIGHT_ONLY, 0.625),
            (HEIGHT_AND_WORST_CASE_BACKSET, 0.625),
            (LOW_HEIGHT, 0.0),
            (MINIMUM_HEIGHTS, 1.0),
            (JUST_BELOW_MINIMUM_HEIGHT, 0.0),
        )
        for values, expected in cases:
            with self.subTest(values=values):
                score, _ = self._score(25, "N/A", _with_values(values, "PASS"))
                self.assertAlmostEqual(score, expected, places=5)

    def test_geometry_ladder_two_rows(self):
        cases = (
            (PASSING_GEOMETRY, 0.5),
            (HEIGHT_AND_MID_BACKSET, 0.4375),
            (HEIGHT_ONLY, 0.3125),
            (HEIGHT_AND_WORST_CASE_BACKSET, 0.3125),
            (LOW_HEIGHT, 0.0),
        )
        for values, expected in cases:
            with self.subTest(values=values):
                _, result = self._score(
                    25,
                    25,
                    _with_values(values, "PASS"),
                    _with_values(PASSING_GEOMETRY, "PASS"),
                )
                rows = _static_rear_rows(result, "CP - Dummy Scores")
                self.assertEqual([float(v) for v in rows["Score"]], [expected, 0.5])

    def test_geometry_ladder_without_non_use_credit(self):
        score, _ = self._score(25, "N/A", _with_values(HEIGHT_ONLY, "FAIL"))
        self.assertAlmostEqual(score, 0.375, places=5)

    def test_geometry_ladder_body_region_score(self):
        cases = (
            (PASSING_GEOMETRY, 100.0),
            (HEIGHT_AND_MID_BACKSET, 83.33333),
            (HEIGHT_ONLY, 50.0),
            (LOW_HEIGHT, 0.0),
        )
        for values, expected in cases:
            with self.subTest(values=values):
                _, result = self._score(25, "N/A", _with_values(values, "PASS"))
                second = _static_rear_rows(result, "CP - Body region scores").iloc[0]
                self.assertAlmostEqual(
                    float(second["Body regionscore"]), expected, places=5
                )

    def test_height_not_met_gates_the_non_use_credit(self):
        score, _ = self._score(25, "N/A", _with_values(LOW_HEIGHT, "PASS"))
        self.assertAlmostEqual(score, 0.0, places=5)

    def test_inspection_is_deducted_from_the_geometry_share(self):
        for values, expected in ((PASSING_GEOMETRY, 0.925), (HEIGHT_ONLY, 0.55)):
            with self.subTest(values=values):
                dfs = _static_rear_dfs(25, "N/A", _with_values(values, "PASS"))
                regions = dfs["CP - Body region scores"].copy()
                regions["Inspection [%]"] = regions["Inspection [%]"].astype(object)
                static_rear = (regions["Loadcase"].ffill() == "Static-Rear") & (
                    regions["Seat position"] == SECOND_ROW
                )
                regions.loc[static_rear, "Inspection [%]"] = 10
                dfs["CP - Body region scores"] = regions
                score, _ = _whiplash_scores(calculate_score(dfs))
                self.assertAlmostEqual(
                    score - OTHER_WHIPLASH_POINTS, expected, places=5
                )

    def test_max_scores_agree_on_both_score_sheets(self):
        for third_angle, expected in (("N/A", [1.0, 0.0]), (25, [0.5, 0.5])):
            with self.subTest(third_angle=third_angle):
                _, result = self._score(
                    25,
                    third_angle,
                    _with_values(PASSING_GEOMETRY, "PASS"),
                    _with_values(PASSING_GEOMETRY, "PASS"),
                )
                for sheet in ("CP - Body region scores", "CP - Dummy Scores"):
                    rows = _static_rear_rows(result, sheet)
                    self.assertEqual(
                        [float(v) for v in rows["Max score"]], expected, sheet
                    )

    def test_third_row_values_are_ignored_when_no_third_row(self):
        score, result = self._score(
            25,
            "N/A",
            _with_values(PASSING_GEOMETRY, "PASS"),
            _with_values(PASSING_GEOMETRY, "PASS"),
        )
        self.assertAlmostEqual(score, 1.0, places=5)
        third = _static_rear_rows(result, "CP - Dummy Scores").iloc[1]
        self.assertEqual(float(third["Score"]), 0.0)
        self.assertEqual(float(third["Max score"]), 0.0)

    def test_blank_third_row_angle_means_no_third_row(self):
        score, _ = self._score(25, None, _with_values(PASSING_GEOMETRY, "PASS"))
        self.assertAlmostEqual(score, 1.0, places=5)

    def test_each_seat_row_uses_its_own_torso_angle(self):
        # 20 degrees -> 295.56 mm limit: both 300 mm backsets fail on the row
        # that carries that angle only, which keeps its height and non-use
        # credit (0.1875 + 0.125).
        for second_angle, third_angle, expected in (
            (25, 20, [0.5, 0.3125]),
            (20, 25, [0.3125, 0.5]),
        ):
            with self.subTest(second_angle=second_angle, third_angle=third_angle):
                _, result = self._score(
                    second_angle,
                    third_angle,
                    _with_values(PASSING_GEOMETRY, "PASS"),
                    _with_values(PASSING_GEOMETRY, "PASS"),
                )
                rows = _static_rear_rows(result, "CP - Dummy Scores")
                self.assertEqual([float(v) for v in rows["Score"]], expected)

    def test_legacy_torso_angle_name_drives_the_second_row(self):
        dfs = _static_rear_dfs(1, "N/A", _with_values(PASSING_GEOMETRY, "PASS"))
        params = dfs["Input parameters"]
        # A dataframe-path caller may still hand over the previous revision's
        # sheet: legacy labels throughout and no third-row parameter.
        params["Input parameter"] = params["Input parameter"].replace(
            {
                "Torso angle (rear second row)": "Torso angle (rear)",
                "Red line >125 mm outboard of the orange line ?": (
                    "Red line >125 mm outboard of the orange line"
                ),
            }
        )
        params.loc[params["Input parameter"] == "Torso angle (rear)", "Value"] = 25
        params = params[params["Input parameter"] != "Torso angle (rear third row)"]
        dfs["Input parameters"] = params.reset_index(drop=True)
        score, _ = _whiplash_scores(calculate_score(dfs))
        self.assertAlmostEqual(score - OTHER_WHIPLASH_POINTS, 1.0, places=5)


class TestTorsoAngleLimits(unittest.TestCase):
    def _dfs(self, second_angle, third_angle):
        return _static_rear_dfs(second_angle, third_angle, PASSING_GEOMETRY)

    def _backset_limits(self, df, seat_name):
        seat = df["Seat position"].ffill()
        rows = df[(seat == seat_name) & df["Criteria"].str.startswith("Backset")]
        return sorted(set(rows["HPL"].tolist() + rows["LPL"].tolist()))

    def test_second_and_third_row_limits_come_from_their_own_angle(self):
        df = data_loader.get_rear_whiplash_df_with_torso_angle(self._dfs(25, 20))
        self.assertEqual(self._backset_limits(df, SECOND_ROW), [331.2])
        self.assertEqual(self._backset_limits(df, THIRD_ROW), [295.56])

    def test_third_row_not_applicable_leaves_its_placeholder_limits(self):
        df = data_loader.get_rear_whiplash_df_with_torso_angle(self._dfs(25, "N/A"))
        self.assertEqual(self._backset_limits(df, SECOND_ROW), [331.2])
        self.assertEqual(self._backset_limits(df, THIRD_ROW), [0.0])
        self.assertFalse(data_loader.third_row_present(self._dfs(25, "N/A")))
        self.assertTrue(data_loader.third_row_present(self._dfs(25, 20)))
        self.assertFalse(data_loader.third_row_present(self._dfs(25, None)))
        self.assertFalse(data_loader.third_row_present(self._dfs(25, "maybe")))
        # A non-finite angle would make every backset pass: not a third row.
        self.assertFalse(data_loader.third_row_present(self._dfs(25, "1e309")))
        self.assertFalse(data_loader.third_row_present(self._dfs(25, float("inf"))))

    def test_missing_second_row_angle_is_strict_or_warns(self):
        dfs = self._dfs(None, "N/A")
        with self.assertRaises(ValueError):
            data_loader.get_rear_whiplash_df_with_torso_angle(dfs, strict=True)
        self.assertIsNone(
            data_loader.get_rear_whiplash_df_with_torso_angle(dfs, strict=False)
        )


class TestNormalizeSeatName(unittest.TestCase):
    def test_parenthesised_row_label_is_preserved(self):
        self.assertEqual(data_loader._normalize_seat_name(SECOND_ROW), SECOND_ROW)
        self.assertEqual(data_loader._normalize_seat_name(THIRD_ROW), THIRD_ROW)

    def test_plain_words_still_title_cased(self):
        self.assertEqual(
            data_loader._normalize_seat_name("driver AND front passenger"),
            "Driver and Front Passenger",
        )
        self.assertEqual(data_loader._normalize_seat_name("rear"), "Rear")


class TestParseNonUsePosition(unittest.TestCase):
    def test_words(self):
        for value in ("PASS", "pass", " Pass ", "OK", "yes", "TRUE"):
            self.assertIs(parse_non_use_position(value), True, value)
        for value in ("FAIL", "fail", "no", "False"):
            self.assertIs(parse_non_use_position(value), False, value)

    def test_numbers_1000_and_0_only(self):
        self.assertIs(parse_non_use_position(1000), True)
        self.assertIs(parse_non_use_position(1000.0), True)
        self.assertIs(parse_non_use_position("1000"), True)
        self.assertIs(parse_non_use_position(0), False)
        self.assertIs(parse_non_use_position(0.0), False)
        self.assertIsNone(parse_non_use_position(1))
        self.assertIsNone(parse_non_use_position(1.0))
        self.assertIsNone(parse_non_use_position(790))

    def test_not_assessed(self):
        for value in (None, np.nan, "", "   ", "N/A", "n/a", "-"):
            self.assertIsNone(parse_non_use_position(value), repr(value))

    def test_garbage_warns_and_is_not_a_pass(self):
        with self.assertLogs(
            "euroncap_rating_2026.crash_protection.dummy", level=logging.WARNING
        ):
            self.assertIsNone(parse_non_use_position("maybe"))
        with self.assertLogs(
            "euroncap_rating_2026.crash_protection.dummy", level=logging.WARNING
        ):
            self.assertIsNone(parse_non_use_position(True))

    def test_row_is_intercepted_on_the_dummy_not_a_measurement(self):
        df = pd.DataFrame(
            {
                "Dummy": ["HPM", None, None],
                "Body region": ["Neck", None, None],
                "Criteria": [
                    "Effective height - Lowest",
                    "Non-Use position",
                    "Backset - Lowest",
                ],
                "HPL": [720.0, None, 331.2],
                "LPL": [719.99, None, 331.2],
                "Capping": [None, None, None],
                "Value": [730.0, "PASS", 300.0],
            }
        )
        dummy, _ = Dummy.get_dummy_from_row("HPM", df, 0)
        self.assertIs(dummy.non_use_position, True)
        neck = dummy.get_body_region("Neck")
        self.assertEqual([m.name for m in neck._measurement], [])
        self.assertEqual(len(neck._criteria), 2)


class TestLegacyWorkbookMigration(unittest.TestCase):
    """A workbook produced against the previous template revision is
    accepted, its Static-Rear readings carry over, the value under the old
    modifier row is dropped and the third-row block comes back blank."""

    LEGACY_STATIC_REAR_VALUES = [730, 780, 1000, 300, 300]

    def setUp(self):
        self.pristine_path = _pristine_path()
        self.wb = make_legacy_cp_workbook(self.pristine_path)
        ws = self.wb["CP - Rear Whiplash"]
        self.legacy_static_rear_rows = range(39, 44)
        self.assertEqual(ws["A39"].value, "Static-Rear")
        self.assertEqual(ws["B39"].value, "Rear")
        for row, value in zip(
            self.legacy_static_rear_rows, self.LEGACY_STATIC_REAR_VALUES
        ):
            ws.cell(row=row, column=9).value = value
        params = self.wb["Input parameters"]
        self.assertEqual(params["D9"].value, "Torso angle (rear)")
        params["E9"].value = 25
        self.assertEqual(
            params["D8"].value, "Red line >125 mm outboard of the orange line"
        )
        params["E8"].value = "Yes"
        params["E7"].value = "Yes"
        body = self.wb["CP - Body region scores"]
        self.assertEqual(body["D89"].value, "Rear")
        body["I89"].value = 0

    def _save(self, wb):
        import tempfile

        fd, path = tempfile.mkstemp(suffix=".xlsx")
        os.close(fd)
        wb.save(path)
        self.addCleanup(os.remove, path)
        return path

    def _sanitized_wb(self):
        sanitized_path = integrity.rebuild_trusted_input(self._save(self.wb))
        self.addCleanup(os.remove, sanitized_path)
        return openpyxl.load_workbook(sanitized_path, data_only=True)

    def _sanitized_dfs(self):
        return integrity.rebuild_trusted_input_dfs(
            common.read_excel_file_to_dfs(self._save(self.wb))
        )

    def test_file_path_migrates_rows(self):
        wb = self._sanitized_wb()
        ws = wb["CP - Rear Whiplash"]
        self.assertEqual(ws["B39"].value, SECOND_ROW)
        self.assertEqual(ws["E41"].value, "Non-Use position")
        self.assertEqual(
            [ws.cell(row=r, column=9).value for r in range(39, 44)],
            [730, 780, None, 300, 300],
        )
        self.assertEqual(ws["B44"].value, THIRD_ROW)
        self.assertEqual(
            [ws.cell(row=r, column=9).value for r in range(44, 49)], [None] * 5
        )
        params = wb["Input parameters"]
        self.assertEqual(params["D9"].value, "Torso angle (rear second row)")
        self.assertEqual(params["E9"].value, 25)
        self.assertEqual(params["D10"].value, "Torso angle (rear third row)")
        self.assertIsNone(params["E10"].value)
        self.assertEqual(
            params["D8"].value, "Red line >125 mm outboard of the orange line ?"
        )
        self.assertEqual(params["E8"].value, "Yes")
        # One row further down than in the legacy file: the single
        # Static-Front row is split into Driver and Front Passenger.
        body = wb["CP - Body region scores"]
        self.assertEqual(body["D90"].value, SECOND_ROW)
        self.assertEqual(body["I90"].value, 0)
        self.assertEqual(body["D91"].value, THIRD_ROW)
        self.assertIsNone(body["I91"].value)

    def test_dataframe_path_migrates_rows(self):
        dfs = self._sanitized_dfs()
        whiplash = dfs["CP - Rear Whiplash"]
        seat = whiplash["Seat position"].ffill()
        second = whiplash[seat == SECOND_ROW]
        self.assertEqual(second["Criteria"].tolist()[2], "Non-Use position")
        values = second["Value"].tolist()
        self.assertEqual(
            [values[0], values[1], values[3], values[4]], [730, 780, 300, 300]
        )
        self.assertTrue(pd.isna(values[2]))
        self.assertTrue(whiplash[seat == THIRD_ROW]["Value"].isna().all())
        params = dfs["Input parameters"]
        by_name = dict(zip(params["Input parameter"], params["Value"]))
        self.assertEqual(by_name["Torso angle (rear second row)"], 25)
        self.assertTrue(pd.isna(by_name["Torso angle (rear third row)"]))
        self.assertEqual(
            by_name["Red line >125 mm outboard of the orange line ?"], "Yes"
        )

    def test_legacy_1000_is_not_a_non_use_pass(self):
        result = calculate_score(self._sanitized_dfs())
        rows = _static_rear_rows(result, "CP - Body region scores")
        second, third = rows["Score"].tolist()
        self.assertAlmostEqual(float(second), 0.75, places=5)
        # The third row the file never had is not scored at all.
        self.assertTrue(pd.isna(third))
        self.assertEqual([float(v) for v in rows["Max score"]], [1.0, 0.0])
        score, max_score = _whiplash_scores(result)
        self.assertAlmostEqual(score, 0.75, places=5)
        self.assertEqual(max_score, 5.0)

    def test_no_findings_asked_about_the_missing_third_row(self):
        findings = integrity.find_consistency_findings(
            self.wb, "Rear Impact", "Whiplash"
        )
        self.assertEqual(findings, [])

    def test_current_format_keeps_its_non_use_value(self):
        wb = openpyxl.load_workbook(self.pristine_path)
        ws = wb["CP - Rear Whiplash"]
        self.assertEqual(ws["E41"].value, "Non-Use position")
        ws["I41"].value = "PASS"
        ws["I46"].value = "FAIL"
        sanitized_path = integrity.rebuild_trusted_input(self._save(wb))
        self.addCleanup(os.remove, sanitized_path)
        sanitized = openpyxl.load_workbook(sanitized_path, data_only=True)
        self.assertEqual(sanitized["CP - Rear Whiplash"]["I41"].value, "PASS")
        self.assertEqual(sanitized["CP - Rear Whiplash"]["I46"].value, "FAIL")

    def test_legacy_file_claiming_a_third_row_is_rejected(self):
        ws = self.wb["CP - Rear Whiplash"]
        pristine = openpyxl.load_workbook(self.pristine_path)["CP - Rear Whiplash"]
        for row in range(44, 49):
            for col in range(1, 10):
                ws.cell(row=row, column=col).value = pristine.cell(
                    row=row, column=col
                ).value
        with self.assertRaises(common.TemplateIntegrityError):
            integrity.rebuild_trusted_input(self._save(self.wb))

    def test_legacy_file_with_a_deleted_static_rear_row_is_rejected(self):
        self.wb["CP - Rear Whiplash"].delete_rows(42)
        with self.assertRaises(common.TemplateIntegrityError):
            integrity.rebuild_trusted_input(self._save(self.wb))

    def test_legacy_file_with_an_inserted_row_is_rejected(self):
        ws = self.wb["CP - Rear Whiplash"]
        ws.insert_rows(42)
        ws["E42"].value = "Backset - Extra"
        with self.assertRaises(common.TemplateIntegrityError):
            integrity.rebuild_trusted_input(self._save(self.wb))

    def test_legacy_file_with_reordered_static_rear_rows_is_rejected(self):
        ws = self.wb["CP - Rear Whiplash"]
        ws["E42"].value, ws["E43"].value = ws["E43"].value, ws["E42"].value
        with self.assertRaises(common.TemplateIntegrityError):
            integrity.rebuild_trusted_input(self._save(self.wb))

    def test_seat_rename_applied_to_static_front_is_rejected(self):
        wb = openpyxl.load_workbook(self.pristine_path)
        ws = wb["CP - Rear Whiplash"]
        self.assertEqual(ws["A29"].value, "Static-Front")
        ws["B29"].value = "Rear"
        with self.assertRaises(common.TemplateIntegrityError):
            integrity.rebuild_trusted_input(self._save(wb))

    def test_criteria_rename_applied_to_static_front_is_rejected(self):
        wb = openpyxl.load_workbook(self.pristine_path)
        ws = wb["CP - Rear Whiplash"]
        self.assertEqual(ws["E31"].value, "Effective height modifier")
        ws["E31"].value = "Non-Use position"
        with self.assertRaises(common.TemplateIntegrityError):
            integrity.rebuild_trusted_input(self._save(wb))

    def test_mixed_revision_spelling_on_one_sheet_is_rejected(self):
        # The legacy seat label with the current criterion label: a legacy
        # file relabelling its modifier row to claim the non-use credit.
        ws = self.wb["CP - Rear Whiplash"]
        ws["E41"].value = "Non-Use position"
        with self.assertRaises(common.TemplateIntegrityError):
            integrity.rebuild_trusted_input(self._save(self.wb))

    def test_dataframe_path_legacy_dfs_helper_round_trips(self):
        dfs = integrity.rebuild_trusted_input_dfs(
            make_legacy_cp_dfs(self.pristine_path)
        )
        self.assertEqual(len(dfs["CP - Rear Whiplash"]), 47)
        self.assertEqual(len(dfs["Input parameters"]), 12)


class TestIdentityRewriting(unittest.TestCase):
    def test_to_current_identity_reaches_the_fixed_point(self):
        legacy = {
            "Loadcase": "Static-Rear",
            "Seat position": "Rear",
            "Dummy": "HPM",
            "Body region": "Neck",
            "Criteria": "Effective height modifier",
        }
        self.assertEqual(
            integrity.to_current_identity("CP - Rear Whiplash", legacy),
            {**legacy, "Seat position": SECOND_ROW, "Criteria": "Non-Use position"},
        )

    def test_static_front_modifier_is_not_translated(self):
        front = {
            "Loadcase": "Static-Front",
            "Seat position": "Driver",
            "Dummy": "HPM",
            "Body region": "Neck",
            "Criteria": "Effective height modifier",
        }
        self.assertEqual(
            integrity.to_current_identity("CP - Rear Whiplash", front), front
        )

    def test_current_identity_unknown_sheet_and_parameters(self):
        current = {"Loadcase": "Static-Rear", "Seat position": SECOND_ROW}
        self.assertEqual(
            integrity.to_current_identity("CP - Dummy Scores", current), current
        )
        self.assertEqual(
            integrity.to_current_identity("No such sheet", current), current
        )
        legacy_param = {
            "Stage": "Crash Protection",
            "Stage element": "Rear Impact",
            "Stage subelement": "Whiplash",
            "Input parameter": "Torso angle (rear)",
        }
        self.assertEqual(
            integrity.to_current_identity("Input parameters", legacy_param)[
                "Input parameter"
            ],
            "Torso angle (rear second row)",
        )

    def test_legacy_identity_spellings(self):
        self.assertEqual(
            integrity.legacy_identity_spellings("CP - Rear Whiplash"),
            [
                ("Seat position", "Rear", SECOND_ROW),
                ("Criteria", "Effective height modifier", "Non-Use position"),
            ],
        )
        self.assertEqual(
            [
                old
                for _, old, _ in integrity.legacy_identity_spellings("Input parameters")
            ],
            ["Torso angle (rear)", "Red line >125 mm outboard of the orange line"],
        )
        self.assertEqual(integrity.legacy_identity_spellings("Test Scores"), [])
        self.assertEqual(integrity.legacy_identity_spellings("No such sheet"), [])


class TestRearTorsoAngleFindings(unittest.TestCase):
    def setUp(self):
        self.wb = openpyxl.load_workbook(_pristine_path(), data_only=True)
        self.ws = self.wb["Input parameters"]
        self.assertEqual(self.ws["D9"].value, "Torso angle (rear second row)")
        self.assertEqual(self.ws["D10"].value, "Torso angle (rear third row)")

    def _findings(self):
        return consistency.find_rear_torso_angle_findings(self.wb)

    def test_valid_values_are_clean(self):
        for second, third in ((25, "N/A"), (25, 20), ("25", "n/a"), (25.5, "20")):
            with self.subTest(second=second, third=third):
                self.ws["E9"].value = second
                self.ws["E10"].value = third
                self.assertEqual(self._findings(), [])

    def test_blank_cells_are_left_to_the_missing_input_check(self):
        self.assertEqual(self._findings(), [])
        findings = integrity.find_consistency_findings(
            self.wb, "Rear Impact", "Whiplash"
        )
        self.assertEqual(
            {(f.cell, f.check) for f in findings},
            {("E9", common.CHECK_MISSING_INPUT), ("E10", common.CHECK_MISSING_INPUT)},
        )

    def test_invalid_values_are_reported(self):
        self.ws["E9"].value = "abc"
        self.ws["E10"].value = "maybe"
        findings = self._findings()
        self.assertEqual(
            [(f.cell, f.check, f.column) for f in findings],
            [
                ("E9", common.CHECK_REAR_TORSO_ANGLE, "Value"),
                ("E10", common.CHECK_REAR_TORSO_ANGLE, "Value"),
            ],
        )
        self.assertEqual(
            findings[1].row_identity["Input parameter"], "Torso angle (rear third row)"
        )
        self.assertIn("N/A", findings[1].message)

    def test_non_finite_angles_are_reported(self):
        for second, third in (("1e309", 20), (25, "inf"), (float("inf"), "N/A")):
            with self.subTest(second=second, third=third):
                self.ws["E9"].value = second
                self.ws["E10"].value = third
                self.assertEqual(len(self._findings()), 1)

    def test_current_file_missing_the_third_row_parameter_is_rejected(self):
        wb = openpyxl.load_workbook(_pristine_path())
        wb["Input parameters"].delete_rows(10)
        import tempfile

        fd, path = tempfile.mkstemp(suffix=".xlsx")
        os.close(fd)
        self.addCleanup(os.remove, path)
        wb.save(path)
        with self.assertRaises(common.TemplateIntegrityError):
            integrity.rebuild_trusted_input(path)

    def test_second_row_not_applicable_is_reported(self):
        self.ws["E9"].value = "N/A"
        self.ws["E10"].value = "N/A"
        self.assertEqual([f.cell for f in self._findings()], ["E9"])

    def test_scope_and_integration(self):
        self.ws["E9"].value = "abc"
        self.assertEqual(
            consistency.find_rear_torso_angle_findings(self.wb, "Side Impact", "Pole"),
            [],
        )
        findings = integrity.find_consistency_findings(
            self.wb, "Rear Impact", "Whiplash"
        )
        self.assertIn(common.CHECK_REAR_TORSO_ANGLE, {f.check for f in findings})

    def test_legacy_workbook_reports_nothing_about_the_third_row(self):
        wb = make_legacy_cp_workbook(_pristine_path())
        wb["Input parameters"]["E9"].value = 25
        self.assertEqual(consistency.find_rear_torso_angle_findings(wb), [])
        wb["Input parameters"]["E9"].value = "abc"
        self.assertEqual(
            [f.cell for f in consistency.find_rear_torso_angle_findings(wb)], ["E9"]
        )
