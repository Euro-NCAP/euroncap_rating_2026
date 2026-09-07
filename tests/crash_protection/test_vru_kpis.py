# Copyright 2025-2026, Euro NCAP IVZW
# Created by IVEX NV (https://ivex.ai)
#
# Licensed under the Apache License 2.0.
# See http://www.apache.org/licenses/LICENSE-2.0 for details.

"""Tests for the VRU headform KPI accessors:

- compute_score.get_vru_headform_kpis(dfs)
- compute_score.get_vru_headform_kpis_in_file(path)

Per WAD body region: colour counts of the prediction matrix and the predicted
score; plus the whole-test correction factor derived through the same
VruTestData path as scoring. No CLI wiring, no report changes.
"""

import io
import os
import tempfile
import unittest
from importlib.resources import files

import openpyxl
import pandas as pd
from openpyxl.styles import PatternFill

from euroncap_rating_2026 import common
from euroncap_rating_2026.crash_protection import vru_processing
from euroncap_rating_2026.crash_protection.compute_score import (
    get_vru_headform_kpis,
    get_vru_headform_kpis_in_file,
)
from euroncap_rating_2026.crash_protection.preprocess import add_vru_sheets_to_dfs

# df-absolute offsets of the headform matrix (slice start + 2 header rows,
# columns from 4), mirroring process_headform_test's iloc slicing.
HEAD_ROW0 = vru_processing.HEADFORMS_START_ROW_INDEX + 2
COL0 = 4
HEAD_WIDTH = vru_processing.VRU_MATRIX_COL_END_INDEX - COL0  # 21

# Symmetric pattern spanning all three WAD bands (Cyclist matrix rows 0-6,
# Adult 7-12, Child 13-18). A-pillar cells come in same-shade mirror pairs
# (matrix col j mirrors to 20 - j).
HEAD_KPI_PATTERN = {
    (0, 0): "blue",
    (2, 0): "green",
    (2, 1): "yellow",
    (2, 2): "orange",
    (8, 0): "green",
    (8, 1): "brown",
    (9, 2): "green-20",
    (9, 18): "green-20",
    (15, 0): "green",
    (15, 1): "red",
    (16, 0): "d green",
    (16, 1): "d red",
    (15, 4): "green-30",
    (15, 16): "green-30",
}

# Hand-computed expectations for HEAD_KPI_PATTERN.
EXPECTED_COUNTS = {
    "Cyclist": {"green": 1, "yellow": 1, "orange": 1, "blue": 1},
    "Adult": {"green": 1, "brown": 1, "green-20": 2},
    "Child": {"green": 1, "red": 1, "d green": 1, "d red": 1, "green-30": 2},
}
EXPECTED_PREDICTED = {"Cyclist": 2.25, "Adult": 1.25, "Child": 2.0}
# Cyclist: green+yellow+orange+blue = 4; Adult: green+brown+2x green-20 = 4;
# Child: green+red+d green+d red = 4 valid + 2 green-30 * 2 = 8.
EXPECTED_MAX_POINTS = {"Cyclist": 4.0, "Adult": 4.0, "Child": 8.0}

# A "Correct" tested value for each prediction (VRU thresholds 650/1700,
# bands at 650/1000/1350/1700); A-pillar limits are 1000/1000 and 1700/1700.
CORRECT_VALUE_BY_PREDICTION = {
    "green": 300.0,
    "yellow": 800.0,
    "orange": 1200.0,
    "brown": 1500.0,
    "red": 1800.0,
    "blue": 300.0,
    "green-20": 900.0,
    "green-30": 1500.0,
}


def _template_dfs():
    template_path = str(files("data").joinpath("cp_template.xlsx"))
    return common.read_excel_file_to_dfs(template_path)


def _parquet_roundtrip(dfs):
    def roundtrip(df):
        buf = io.BytesIO()
        df.to_parquet(buf)
        return pd.read_parquet(io.BytesIO(buf.getvalue()))

    out = {}
    for name, df in dfs.items():
        try:
            out[name] = roundtrip(df)
        except Exception:
            stringified = df.apply(
                lambda col: col.map(lambda v: str(v) if pd.notna(v) else None)
            )
            out[name] = roundtrip(stringified)
    return out


def _write_head_pattern(dfs, pattern):
    pred = dfs["CP - VRU Prediction"].astype(object)
    for i in range(19):
        for j in range(HEAD_WIDTH):
            pred.iat[HEAD_ROW0 + i, COL0 + j] = "grey"
    for (i, j), value in pattern.items():
        pred.iat[HEAD_ROW0 + i, COL0 + j] = value
    dfs["CP - VRU Prediction"] = pred


def _set_head_verification_tests(dfs, count):
    params = dfs["Input parameters"].copy()
    current_subelement = None
    for idx, row in params.iterrows():
        if not pd.isna(row.get("Stage subelement")):
            current_subelement = row["Stage subelement"]
        input_param = row.get("Input parameter")
        if pd.isna(input_param) or "Number of verification tests" not in str(
            input_param
        ):
            continue
        if current_subelement == "Head Impact":
            params.at[idx, "Value"] = count
    dfs["Input parameters"] = params


def _preprocessed_dfs(head_tests=4):
    dfs = _template_dfs()
    _write_head_pattern(dfs, HEAD_KPI_PATTERN)
    _set_head_verification_tests(dfs, head_tests)
    dfs = _parquet_roundtrip(dfs)
    _, _, result = add_vru_sheets_to_dfs(dfs)
    return result


class TestGetVruHeadformKpis(unittest.TestCase):
    def _assert_pattern_kpis(self, kpis):
        self.assertEqual(
            [region.body_region for region in kpis.regions],
            ["Child", "Adult", "Cyclist"],
        )
        for region in kpis.regions:
            expected = dict.fromkeys(region.color_counts, 0)
            expected.update(EXPECTED_COUNTS[region.body_region])
            self.assertEqual(region.color_counts, expected, region.body_region)
            self.assertAlmostEqual(
                region.predicted_score,
                EXPECTED_PREDICTED[region.body_region],
                places=6,
                msg=region.body_region,
            )
            self.assertAlmostEqual(
                region.max_points,
                EXPECTED_MAX_POINTS[region.body_region],
                places=6,
                msg=region.body_region,
            )

    def test_counts_and_predicted_score_from_prediction_points_sheet(self):
        kpis = get_vru_headform_kpis(_preprocessed_dfs())
        self._assert_pattern_kpis(kpis)

    def test_counts_survive_parquet_roundtrip_of_result(self):
        """The preprocessed dfs may round-trip through parquet storage:
        colours then arrive as plain strings, not VruPredictionColor
        members."""
        result = _parquet_roundtrip(_preprocessed_dfs())
        kpis = get_vru_headform_kpis(result)
        self._assert_pattern_kpis(kpis)

    def test_correction_factor_is_one_when_no_verification_value_filled(self):
        kpis = get_vru_headform_kpis(_preprocessed_dfs())
        self.assertEqual(kpis.correction_factor, 1.0)
        self.assertEqual(kpis.tested_score_verification, 0.0)
        # No tested value -> no blue or A-pillar credit either.
        for region in kpis.regions:
            self.assertEqual(region.blue_points, 0.0, region.body_region)
            self.assertEqual(region.a_pillar_bonus, 0.0, region.body_region)
            self.assertEqual(region.bonus, 0.0, region.body_region)

    def test_correction_factor_matches_scoring_path_when_values_filled(self):
        result = _preprocessed_dfs()
        self._fill_head_values(result, CORRECT_VALUE_BY_PREDICTION)

        kpis = get_vru_headform_kpis(result)

        # Every verification value matches its prediction, so the correction
        # factor is exactly 1.0 and tested == predicted.
        self.assertAlmostEqual(kpis.correction_factor, 1.0, places=6)
        self.assertGreater(kpis.predicted_score_verification, 0.0)
        self.assertAlmostEqual(
            kpis.tested_score_verification,
            kpis.predicted_score_verification,
            places=6,
        )

    def _fill_head_values(self, result, value_by_prediction):
        head_df = result["CP - VRU Head Impact"].copy()
        for idx, row in head_df.iterrows():
            prediction = str(row.get("OEM Prediction")).lower()
            if prediction in value_by_prediction:
                head_df.at[idx, "Value"] = value_by_prediction[prediction]
        result["CP - VRU Head Impact"] = head_df

    def test_bonus_credits_tested_blue_and_a_pillar_cells(self):
        result = _preprocessed_dfs()
        self._fill_head_values(result, CORRECT_VALUE_BY_PREDICTION)

        kpis = get_vru_headform_kpis(result)

        by_region = {r.body_region: r for r in kpis.regions}
        # Cyclist: one blue cell, HIC 300 -> green -> 100/100 = 1.0 credit.
        self.assertAlmostEqual(by_region["Cyclist"].blue_points, 1.0, places=6)
        self.assertAlmostEqual(by_region["Cyclist"].a_pillar_bonus, 0.0, places=6)
        self.assertAlmostEqual(by_region["Cyclist"].bonus, 1.0, places=6)
        # Adult: two green-20 cells, HIC 900 < 1000 -> pass -> 2 * 1.0.
        self.assertAlmostEqual(by_region["Adult"].blue_points, 0.0, places=6)
        self.assertAlmostEqual(by_region["Adult"].a_pillar_bonus, 2.0, places=6)
        self.assertAlmostEqual(by_region["Adult"].bonus, 2.0, places=6)
        # Child: two green-30 cells, HIC 1500 < 1700 -> pass -> 2 * 2.0.
        self.assertAlmostEqual(by_region["Child"].blue_points, 0.0, places=6)
        self.assertAlmostEqual(by_region["Child"].a_pillar_bonus, 4.0, places=6)
        self.assertAlmostEqual(by_region["Child"].bonus, 4.0, places=6)

    def test_bonus_is_zero_when_blue_and_a_pillar_tests_fail(self):
        result = _preprocessed_dfs()
        failing = dict(
            CORRECT_VALUE_BY_PREDICTION,
            # blue: HIC 1800 -> red -> 0; green-20: 1100 >= 1000 -> fail;
            # green-30: 1750 >= 1700 -> fail.
            blue=1800.0,
            **{"green-20": 1100.0, "green-30": 1750.0},
        )
        self._fill_head_values(result, failing)

        kpis = get_vru_headform_kpis(result)

        for region in kpis.regions:
            self.assertEqual(region.blue_points, 0.0, region.body_region)
            self.assertEqual(region.a_pillar_bonus, 0.0, region.body_region)
            self.assertEqual(region.bonus, 0.0, region.body_region)

    def test_blank_prediction_in_head_impact_sheet_does_not_crash(self):
        """A blank OEM Prediction cell (None after normalisation) must be
        skipped by the bonus computation, not crash compute_a_pillar."""
        result = _preprocessed_dfs()
        self._fill_head_values(result, CORRECT_VALUE_BY_PREDICTION)
        head_df = result["CP - VRU Head Impact"].copy()
        head_df.at[head_df.index[1], "OEM Prediction"] = None
        result["CP - VRU Head Impact"] = head_df

        kpis = get_vru_headform_kpis(result)

        self.assertEqual(
            [region.body_region for region in kpis.regions],
            ["Child", "Adult", "Cyclist"],
        )

    def test_seeded_zero_values_yield_neutral_factor_and_zero_bonus(self):
        """Value 0.0 is the preprocess-seeded placeholder: like the all-blank
        workbook it must produce the neutral fallback (factor 1.0, zero
        bonus), even though scoring itself would treat 0.0 as a tested HIC."""
        result = _preprocessed_dfs()
        head_df = result["CP - VRU Head Impact"].copy()
        head_df["Value"] = 0.0
        result["CP - VRU Head Impact"] = head_df

        kpis = get_vru_headform_kpis(result)

        self.assertEqual(kpis.correction_factor, 1.0)
        self.assertEqual(kpis.tested_score_verification, 0.0)
        for region in kpis.regions:
            self.assertEqual(region.blue_points, 0.0, region.body_region)
            self.assertEqual(region.a_pillar_bonus, 0.0, region.body_region)
            self.assertEqual(region.bonus, 0.0, region.body_region)

    def test_counts_fall_back_to_prediction_sheet_text(self):
        dfs = _template_dfs()
        _write_head_pattern(dfs, HEAD_KPI_PATTERN)
        self.assertNotIn("CP - VRU Prediction Points", dfs)
        kpis = get_vru_headform_kpis(dfs)
        self._assert_pattern_kpis(kpis)
        self.assertEqual(kpis.correction_factor, 1.0)

    def test_empty_grid_yields_zero_counts_and_factor_one(self):
        dfs = _template_dfs()
        _write_head_pattern(dfs, {})
        kpis = get_vru_headform_kpis(dfs)
        for region in kpis.regions:
            self.assertEqual(sum(region.color_counts.values()), 0)
            self.assertEqual(region.predicted_score, 0.0)
            self.assertEqual(region.max_points, 0.0)
        self.assertEqual(kpis.correction_factor, 1.0)


class TestGetVruHeadformKpisInFile(unittest.TestCase):
    def _template_workbook(self):
        template_path = str(files("data").joinpath("cp_template.xlsx"))
        return openpyxl.load_workbook(template_path)

    def _save_and_read(self, wb):
        with tempfile.TemporaryDirectory() as tmpdir:
            path = os.path.join(tmpdir, "cp_kpis.xlsx")
            wb.save(path)
            return get_vru_headform_kpis_in_file(path)

    def test_text_colours_without_fills_use_text_fallback(self):
        wb = self._template_workbook()
        ws = wb["CP - VRU Prediction"]
        # Headform matrix (i, j) -> ws row i + 4, ws col 5 + j.
        for (i, j), value in HEAD_KPI_PATTERN.items():
            ws.cell(row=i + 4, column=5 + j, value=value)

        kpis = self._save_and_read(wb)

        child = next(r for r in kpis.regions if r.body_region == "Child")
        self.assertEqual(child.color_counts["green-30"], 2)
        self.assertEqual(child.color_counts["d green"], 1)
        self.assertAlmostEqual(child.predicted_score, 2.0, places=6)
        self.assertEqual(kpis.correction_factor, 1.0)

    def test_fill_colours_preferred_over_marked_text(self):
        """After preprocess the prediction texts are replaced by X/T marks;
        the fill colours are then the only reliable source (the same one the
        compute-score CLI reads)."""

        def hex_of(color):
            rgb = vru_processing.VRU_PREDICTION_COLOR_MAP[color]
            return "FF" + "".join(f"{round(c * 255):02X}" for c in rgb)

        wb = self._template_workbook()
        ws = wb["CP - VRU Prediction"]
        filled = {
            (2, 0): vru_processing.VruPredictionColor.GREEN,
            (9, 2): vru_processing.VruPredictionColor.GREEN_20,
            (9, 18): vru_processing.VruPredictionColor.GREEN_20,
        }
        for (i, j), color in filled.items():
            cell = ws.cell(row=i + 4, column=5 + j, value="X")
            code = hex_of(color)
            cell.fill = PatternFill(start_color=code, end_color=code, fill_type="solid")

        kpis = self._save_and_read(wb)

        cyclist = next(r for r in kpis.regions if r.body_region == "Cyclist")
        adult = next(r for r in kpis.regions if r.body_region == "Adult")
        self.assertEqual(cyclist.color_counts["green"], 1)
        self.assertEqual(adult.color_counts["green-20"], 2)
        self.assertAlmostEqual(cyclist.predicted_score, 1.0, places=6)


if __name__ == "__main__":
    unittest.main()
