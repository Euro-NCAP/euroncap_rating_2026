# Copyright 2025-2026, Euro NCAP IVZW
# Created by IVEX NV (https://ivex.ai)
#
# Licensed under the Apache License 2.0.
# See http://www.apache.org/licenses/LICENSE-2.0 for details.

"""Synthesize a crash_protection workbook in the layout of the template
revision that preceded the Rear Impact v1.3 sheet, from the packaged
template: rear seat row spelled "Rear", a Static-Rear "Effective height
modifier" row instead of "Non-Use position", no "Rear (third row)" block,
a single "Torso angle (rear)" input parameter, the Farside red-line
parameter without its trailing " ?" and a single Static-Front "Front"
score row. make_pre_static_front_split_cp_workbook is the later revision
that had the Rear Impact v1.3 layout but still the single "Front" row.
Rows are located by label so the fixture keeps tracking the packaged
template."""

import os
import tempfile

import openpyxl

from euroncap_rating_2026 import common

LEGACY_TEMPLATE_VERSION = "5.4.5"


def _header(ws):
    return {cell.value: cell.column for cell in ws[1]}


def _find_row(ws, column, value, start=2):
    for row in range(start, ws.max_row + 1):
        if ws.cell(row=row, column=column).value == value:
            return row
    raise AssertionError(f"{ws.title}: no row with {value!r} in column {column}")


def _delete_rows(ws, first_row, count):
    # openpyxl.delete_rows does not move merged ranges; the packaged
    # template's only merges below these rows are empty, so drop them.
    for ref in [str(r) for r in ws.merged_cells.ranges]:
        rng = openpyxl.worksheet.cell_range.CellRange(ref)
        if rng.min_row >= first_row:
            ws.unmerge_cells(ref)
    ws.delete_rows(first_row, count)


def _fold_static_front(wb, inspection=None):
    """Static-Front score rows back to one "Front" row (max 1)."""
    for sheet_name in ("CP - Dummy Scores", "CP - Body region scores"):
        ws = wb[sheet_name]
        col = _header(ws)
        driver_row = _find_row(ws, col["Loadcase"], "Static-Front")
        assert ws.cell(row=driver_row, column=col["Seat position"]).value == "Driver"
        assert (
            ws.cell(row=driver_row + 1, column=col["Seat position"]).value
            == "Front Passenger"
        )
        _delete_rows(ws, driver_row + 1, 1)
        ws.cell(row=driver_row, column=col["Seat position"]).value = "Front"
        ws.cell(row=driver_row, column=col["Max score"]).value = 1
        if "Inspection [%]" in col:
            ws.cell(row=driver_row, column=col["Inspection [%]"]).value = inspection


def make_pre_static_front_split_cp_workbook(pristine_path: str, inspection=None):
    """The packaged cp template with only the Static-Front score rows folded
    back into one "Front" row, *inspection* on its Body region scores row."""
    wb = openpyxl.load_workbook(pristine_path)
    _fold_static_front(wb, inspection)
    return wb


def make_legacy_cp_workbook(pristine_path: str):
    """The packaged cp template rewritten into the previous revision's layout,
    stamped with a Version sheet of that revision."""
    wb = openpyxl.load_workbook(pristine_path)

    ws = wb["Input parameters"]
    col = _header(ws)
    row = _find_row(ws, col["Input parameter"], "Torso angle (rear third row)")
    _delete_rows(ws, row, 1)
    row = _find_row(ws, col["Input parameter"], "Torso angle (rear second row)")
    ws.cell(row=row, column=col["Input parameter"]).value = "Torso angle (rear)"
    row = _find_row(
        ws, col["Input parameter"], "Red line >125 mm outboard of the orange line ?"
    )
    ws.cell(row=row, column=col["Input parameter"]).value = (
        "Red line >125 mm outboard of the orange line"
    )

    ws = wb["CP - Rear Whiplash"]
    col = _header(ws)
    third_row = _find_row(ws, col["Seat position"], "Rear (third row)")
    _delete_rows(ws, third_row, 5)
    second_row = _find_row(ws, col["Seat position"], "Rear (second row)")
    ws.cell(row=second_row, column=col["Seat position"]).value = "Rear"
    non_use_row = _find_row(ws, col["Criteria"], "Non-Use position", second_row)
    ws.cell(row=non_use_row, column=col["Criteria"]).value = "Effective height modifier"

    for sheet_name in ("CP - Dummy Scores", "CP - Body region scores"):
        ws = wb[sheet_name]
        col = _header(ws)
        third_row = _find_row(ws, col["Seat position"], "Rear (third row)")
        _delete_rows(ws, third_row, 1)
        second_row = _find_row(ws, col["Seat position"], "Rear (second row)")
        ws.cell(row=second_row, column=col["Seat position"]).value = "Rear"
        ws.cell(row=second_row, column=col["Max score"]).value = 1

    _fold_static_front(wb)

    if "Version" in wb.sheetnames:
        wb["Version"]["B1"].value = LEGACY_TEMPLATE_VERSION
    return wb


def save_legacy_cp_workbook(pristine_path: str) -> str:
    """make_legacy_cp_workbook saved to a temporary file; the caller removes it."""
    wb = make_legacy_cp_workbook(pristine_path)
    fd, path = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    wb.save(path)
    return path


def make_legacy_cp_dfs(pristine_path: str) -> dict:
    """The legacy workbook as the DataFrame dict a dataframe-path caller
    holds (read the same way the tool reads any workbook)."""
    path = save_legacy_cp_workbook(pristine_path)
    try:
        return common.read_excel_file_to_dfs(path)
    finally:
        os.remove(path)
