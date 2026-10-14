# Copyright 2025-2026, Euro NCAP IVZW
# Created by IVEX NV (https://ivex.ai)
#
# Licensed under the Apache License 2.0.
# See http://www.apache.org/licenses/LICENSE-2.0 for details.

"""The crash_protection "Input parameters" sheet is matched to the official
template row by row identity (Stage / Stage element / Stage subelement /
Input parameter), not by position, so Euro NCAP's Rear Impact Lab
Assessment Sheet layouts are accepted:

- the "placeholder" layout: our rows plus a notes column;
- the "compact" layout as Euro NCAP distributes it: no Frontal/Side
  placeholder rows, and its first row (Farside) states the Stage but leaves
  the Stage element blank.

The result is always in our template's layout."""

import dataclasses
import os
import tempfile
import unittest
from importlib.resources import files

import numpy as np
import openpyxl
import pandas as pd

from euroncap_rating_2026 import common
from euroncap_rating_2026.crash_protection import integrity
from euroncap_rating_2026.crash_protection.compute_score import calculate_score
from tests.crash_protection.legacy_fixtures import (
    make_legacy_cp_dfs,
    make_legacy_cp_workbook,
)
from tests.crash_protection.test_rear_impact_seat_rows import (
    PASSING_GEOMETRY,
    _static_rear_dfs,
    _whiplash_scores,
    _with_values,
)

SHEET = "Input parameters"
SCHEMA = integrity.SHEET_SCHEMAS[SHEET]
POSITIONAL_SCHEMA = dataclasses.replace(SCHEMA, match_rows_by_identity=False)
SECOND_ROW_ANGLE = "Torso angle (rear second row)"
THIRD_ROW_ANGLE = "Torso angle (rear third row)"
COUNTERMEASURE = "Countermeasure?"
RED_LINE = "Red line >125 mm outboard of the orange line ?"
LEGACY_RED_LINE = "Red line >125 mm outboard of the orange line"
LEGACY_TORSO_ANGLE = integrity.LEGACY_TORSO_ANGLE
NOTES = "N/A in case not present"


def _pristine_path():
    return str(files("data").joinpath("cp_template.xlsx"))


def _pristine_frame():
    return common.read_excel_file_to_dfs(_pristine_path())[SHEET]


def _set_value(frame, parameter, value):
    frame = frame.copy()
    frame["Value"] = frame["Value"].astype(object)
    frame.loc[frame["Input parameter"] == parameter, "Value"] = value
    return frame


def _filled(frame):
    frame = _set_value(frame, SECOND_ROW_ANGLE, 24)
    frame = _set_value(frame, THIRD_ROW_ANGLE, "N/A")
    frame = _set_value(frame, COUNTERMEASURE, "Yes")
    return _set_value(frame, RED_LINE, "No")


def _placeholder_layout(frame):
    """Our rows plus the Lab Assessment Sheet's notes column."""
    frame = frame.copy()
    frame["Unnamed: 5"] = np.nan
    frame["Unnamed: 6"] = np.nan
    frame.loc[frame["Input parameter"] == THIRD_ROW_ANGLE, "Unnamed: 6"] = NOTES
    return frame


def _compact_layout(frame):
    """The placeholder layout without the Frontal/Side placeholder rows, its
    first row stating the Stage but not the Stage element."""
    frame = _placeholder_layout(frame)
    frame = frame[frame["Input parameter"].notna()].reset_index(drop=True)
    frame.loc[0, "Stage"] = "Crash Protection"
    frame.loc[0, "Stage element"] = np.nan
    return frame


def _rebuild_df(user_frame, schema=SCHEMA):
    return common.rebuild_trusted_df(_pristine_frame(), user_frame, schema, SHEET)


def _workbook_with_sheet_rows(rows):
    """The packaged template with its "Input parameters" sheet replaced by
    *rows* (header first)."""
    wb = openpyxl.load_workbook(_pristine_path())
    ws = wb[SHEET]
    ws.delete_rows(1, ws.max_row)
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row, start=1):
            ws.cell(row=r, column=c).value = value
    return wb


def _frame_rows(frame):
    header = [None if str(h).startswith("Unnamed") else h for h in frame.columns]
    return [header] + [
        [None if pd.isna(v) else v for v in values]
        for values in frame.itertuples(index=False)
    ]


def _save(wb):
    fd, path = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    wb.save(path)
    return path


def _values_by_parameter(frame):
    return {
        parameter: value
        for parameter, value in zip(frame["Input parameter"], frame["Value"])
        if isinstance(parameter, str)
    }


class TestDataFramePath(unittest.TestCase):
    def test_pristine_layout_is_unchanged(self):
        user = _filled(_pristine_frame())
        pd.testing.assert_frame_equal(
            _rebuild_df(user), _rebuild_df(user, POSITIONAL_SCHEMA)
        )

    def test_legacy_layout_is_unchanged(self):
        user = make_legacy_cp_dfs(_pristine_path())[SHEET]
        user = _set_value(user, "Torso angle (rear)", 24)
        pd.testing.assert_frame_equal(
            _rebuild_df(user), _rebuild_df(user, POSITIONAL_SCHEMA)
        )

    def test_both_lab_assessment_layouts_rebuild_into_our_layout(self):
        expected = _rebuild_df(_filled(_pristine_frame()))
        for name, layout in (
            ("placeholder", _placeholder_layout),
            ("compact", _compact_layout),
        ):
            with self.subTest(layout=name):
                result = _rebuild_df(layout(_filled(_pristine_frame())))
                pd.testing.assert_frame_equal(result, expected)
                self.assertEqual(list(result.columns), list(_pristine_frame().columns))
                values = _values_by_parameter(result)
                self.assertEqual(values[SECOND_ROW_ANGLE], 24)
                self.assertEqual(values[THIRD_ROW_ANGLE], "N/A")
                self.assertEqual(values[COUNTERMEASURE], "Yes")

    def test_unknown_row_is_rejected_naming_it(self):
        user = _compact_layout(_pristine_frame())
        user.loc[3, "Input parameter"] = "Seat height"
        with self.assertRaisesRegex(
            common.TemplateIntegrityError,
            r"Row 5: 'Crash Protection / Rear Impact / Whiplash / Seat height' matches no "
            "row",
        ):
            _rebuild_df(user)

    def test_duplicated_row_is_rejected(self):
        user = _compact_layout(_pristine_frame())
        user = pd.concat([user, user.iloc[[2]]], ignore_index=True)
        with self.assertRaisesRegex(
            common.TemplateIntegrityError, r"Rows 4 and 9 are both"
        ):
            _rebuild_df(user)

    def test_row_matching_several_template_rows_is_rejected(self):
        # A Whiplash row whose Input parameter is never stated could be
        # either torso angle.
        user = _compact_layout(_pristine_frame()).iloc[:1].copy()
        user.loc[0, ["Stage element", "Stage subelement", "Input parameter"]] = [
            "Rear Impact",
            "Whiplash",
            np.nan,
        ]
        with self.assertRaisesRegex(
            common.TemplateIntegrityError, "matches more than one row"
        ):
            _rebuild_df(user)

    def test_mixed_revision_spelling_is_rejected(self):
        user = _set_value(_pristine_frame(), SECOND_ROW_ANGLE, 24)
        user.loc[user["Input parameter"] == SECOND_ROW_ANGLE, "Input parameter"] = (
            "Torso angle (rear)"
        )
        with self.assertRaisesRegex(
            common.TemplateIntegrityError, "different template revision"
        ):
            _rebuild_df(user)

    def test_current_layout_without_the_third_row_is_rejected(self):
        user = _compact_layout(_pristine_frame())
        user = user[user["Input parameter"] != THIRD_ROW_ANGLE]
        with self.assertRaisesRegex(
            common.TemplateIntegrityError, "missing from your file"
        ):
            _rebuild_df(user.reset_index(drop=True))

    def test_absent_template_row_stays_blank(self):
        user = _filled(_compact_layout(_pristine_frame()))
        user = user[user["Input parameter"] != COUNTERMEASURE].reset_index(drop=True)
        user.loc[0, "Stage"] = "Crash Protection"
        values = _values_by_parameter(_rebuild_df(user))
        self.assertTrue(pd.isna(values[COUNTERMEASURE]))
        self.assertEqual(values[SECOND_ROW_ANGLE], 24)

    def _with_legacy_farside_spelling(self, frame):
        frame = frame.copy()
        frame.loc[frame["Input parameter"] == RED_LINE, "Input parameter"] = (
            LEGACY_RED_LINE
        )
        return frame

    def test_compact_layout_in_the_previous_revision_spelling(self):
        # The Farside rename is anchored on Side Impact, which the compact
        # layout never states: the anchor is read from the template row.
        user = self._with_legacy_farside_spelling(_compact_layout(_pristine_frame()))
        user = user[user["Input parameter"] != THIRD_ROW_ANGLE]
        user.loc[user["Input parameter"] == SECOND_ROW_ANGLE, "Input parameter"] = (
            LEGACY_TORSO_ANGLE
        )
        user = _set_value(user, LEGACY_RED_LINE, "No")
        user = _set_value(user, LEGACY_TORSO_ANGLE, 24)
        values = _values_by_parameter(_rebuild_df(user.reset_index(drop=True)))
        self.assertEqual(values[RED_LINE], "No")
        self.assertEqual(values[SECOND_ROW_ANGLE], 24)
        self.assertTrue(pd.isna(values[THIRD_ROW_ANGLE]))

    def test_compact_layout_mixing_revisions_is_rejected(self):
        user = self._with_legacy_farside_spelling(_compact_layout(_pristine_frame()))
        with self.assertRaisesRegex(
            common.TemplateIntegrityError, "different template revision"
        ):
            _rebuild_df(user)


class TestScoring(unittest.TestCase):
    def test_compact_layout_scores_like_our_template(self):
        dfs = _static_rear_dfs(24, "N/A", _with_values(PASSING_GEOMETRY, "PASS"))
        reference = calculate_score({k: v.copy() for k, v in dfs.items()})
        dfs[SHEET] = _compact_layout(dfs[SHEET])
        result = calculate_score(dfs)
        self.assertEqual(_whiplash_scores(result), _whiplash_scores(reference))
        pd.testing.assert_frame_equal(
            result["CP - Dummy Scores"], reference["CP - Dummy Scores"]
        )


class TestWorkbookPath(unittest.TestCase):
    def _sanitized(self, wb):
        path = _save(wb)
        self.addCleanup(os.remove, path)
        sanitized_path = integrity.rebuild_trusted_input(path)
        self.addCleanup(os.remove, sanitized_path)
        return openpyxl.load_workbook(sanitized_path, data_only=True)[SHEET]

    def _cells(self, ws):
        return [
            [cell.value for cell in row]
            for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=5)
        ]

    def test_pristine_layout_is_unchanged(self):
        wb = openpyxl.load_workbook(_pristine_path())
        wb[SHEET]["E9"].value = 24
        result = self._cells(self._sanitized(wb))
        positional = openpyxl.load_workbook(_pristine_path())[SHEET]
        common.rebuild_trusted_sheet(positional, wb[SHEET], POSITIONAL_SCHEMA, SHEET)
        self.assertEqual(result, self._cells(positional))

    def test_legacy_workbook_is_unchanged(self):
        wb = make_legacy_cp_workbook(_pristine_path())
        wb[SHEET]["E9"].value = 24
        result = self._cells(self._sanitized(wb))
        positional = openpyxl.load_workbook(_pristine_path())[SHEET]
        common.rebuild_trusted_sheet(positional, wb[SHEET], POSITIONAL_SCHEMA, SHEET)
        self.assertEqual(result, self._cells(positional))

    def test_both_lab_assessment_layouts_land_on_our_rows(self):
        for name, layout in (
            ("placeholder", _placeholder_layout),
            ("compact", _compact_layout),
        ):
            with self.subTest(layout=name):
                rows = _frame_rows(layout(_filled(_pristine_frame())))
                ws = self._sanitized(_workbook_with_sheet_rows(rows))
                self.assertEqual(ws.max_column, 5)
                self.assertEqual(ws["D7"].value, COUNTERMEASURE)
                self.assertEqual(ws["E7"].value, "Yes")
                self.assertEqual(ws["E9"].value, 24)
                self.assertEqual(ws["E10"].value, "N/A")

    def test_unknown_row_is_rejected(self):
        frame = _compact_layout(_pristine_frame())
        frame.loc[3, "Input parameter"] = "Seat height"
        wb = _workbook_with_sheet_rows(_frame_rows(frame))
        path = _save(wb)
        self.addCleanup(os.remove, path)
        with self.assertRaisesRegex(common.TemplateIntegrityError, "Seat height"):
            integrity.rebuild_trusted_input(path)


class TestMissingRequiredRows(unittest.TestCase):
    """A required row the user's sheet lacks altogether is reported, with no
    cell to point at."""

    def _findings(self, frame, *scope, check=integrity.find_consistency_findings):
        wb = _workbook_with_sheet_rows(_frame_rows(frame))
        path = _save(wb)
        self.addCleanup(os.remove, path)
        wb = openpyxl.load_workbook(path, data_only=True)
        return [f for f in check(wb, *scope) if f.sheet == SHEET]

    def test_lab_assessment_sheet_reports_nothing_when_filled(self):
        frame = _filled(_compact_layout(_pristine_frame()))
        self.assertEqual(self._findings(frame, "Rear Impact", "Whiplash"), [])

    def test_absent_second_row_angle_is_reported(self):
        frame = _filled(_compact_layout(_pristine_frame()))
        frame = frame[frame["Input parameter"] != SECOND_ROW_ANGLE]
        for check in (
            integrity.find_consistency_findings,
            integrity.find_missing_required_inputs,
        ):
            with self.subTest(check=check.__name__):
                findings = self._findings(
                    frame.reset_index(drop=True),
                    "Rear Impact",
                    "Whiplash",
                    check=check,
                )
                self.assertEqual(len(findings), 1)
                finding = findings[0]
                self.assertEqual(finding.cell, "")
                self.assertEqual(finding.check, common.CHECK_MISSING_INPUT)
                self.assertEqual(
                    finding.row_identity["Input parameter"], SECOND_ROW_ANGLE
                )
                self.assertIn("missing from your file", finding.description)

    def test_absent_row_out_of_scope_is_not_reported(self):
        frame = _filled(_compact_layout(_pristine_frame()))
        frame = frame[frame["Input parameter"] != SECOND_ROW_ANGLE]
        self.assertEqual(
            self._findings(frame.reset_index(drop=True), "Side Impact", "Farside"),
            [],
        )

    def test_compact_layout_scopes_the_farside_rows(self):
        # The Farside rows never state their Stage element; they still
        # belong to Side Impact / Farside.
        frame = _compact_layout(_pristine_frame())
        findings = self._findings(frame, "Side Impact", "Farside")
        self.assertEqual([f.cell for f in findings], ["E2", "E3"])
        self.assertEqual(
            {f.row_identity["Stage element"] for f in findings}, {"Side Impact"}
        )

    def test_legacy_file_is_not_asked_about_the_third_row(self):
        wb = make_legacy_cp_workbook(_pristine_path())
        wb[SHEET]["E9"].value = 24
        findings = [
            f
            for f in integrity.find_consistency_findings(wb, "Rear Impact", "Whiplash")
            if f.sheet == SHEET
        ]
        self.assertEqual(findings, [])


if __name__ == "__main__":
    unittest.main()
