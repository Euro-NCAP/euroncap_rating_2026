# Copyright 2025-2026, Euro NCAP IVZW
# Created by IVEX NV (https://ivex.ai)
#
# Licensed under the Apache License 2.0.
# See http://www.apache.org/licenses/LICENSE-2.0 for details.

"""crash_protection's preprocess() requires a filled-in VRU test-point
selection (its VRU pipeline errors out on a fully blank template -- a
pre-existing, unrelated limitation), so a full CLI preprocess()->
compute_score() e2e run isn't viable here the way it is for
crash_avoidance/post_crash. Instead, exercise report_writer.write_report
directly -- the actual code that hard-copies and hides the "Documentation"/
"Data" sheets -- against the pristine template."""

import shutil
import tempfile
import unittest
from importlib.resources import files
from pathlib import Path

import openpyxl

from euroncap_rating_2026 import common
from euroncap_rating_2026.crash_protection import report_writer


class TestDocumentationDataHardCopyAndHide(unittest.TestCase):
    def test_preprocess_output_has_hidden_documentation_data_sheets(self):
        pristine_path = str(files("data").joinpath("cp_template.xlsx"))
        with (
            tempfile.TemporaryDirectory() as input_dir,
            tempfile.TemporaryDirectory() as output_dir,
        ):
            input_file = str(Path(input_dir) / "cp_template.xlsx")
            shutil.copy(pristine_path, input_file)
            common.add_version_sheet(
                input_file, cli_command=common.CliCommand.GENERATE_TEMPLATE
            )

            report_writer.write_report(
                common.CliCommand.PREPROCESS,
                input_file,
                updated_dfs={},
                output_path=output_dir,
                format_prediction_cells=False,
            )

            output_file = Path(output_dir) / "cp_preprocessed_template.xlsx"
            wb = openpyxl.load_workbook(str(output_file))
            for sheet_name, id_header in [
                ("Documentation", "Documentation ID"),
                ("Data", "Data ID"),
            ]:
                ws = wb[sheet_name]
                self.assertEqual(ws.sheet_state, "hidden")
                self.assertEqual([c.value for c in ws[1]], [id_header, "Status"])


class TestVruHeadformNaTreatedAsBlank(unittest.TestCase):
    """A typed "N/A" in the VRU headform matrix (an authored dropdown
    option) means "not filled": the report must render it exactly like a
    blank cell -- grey fill, text cleared -- instead of leaking the literal
    "N/A" string into the output."""

    def test_na_cell_rendered_like_blank(self):
        pristine_path = str(files("data").joinpath("cp_template.xlsx"))
        with (
            tempfile.TemporaryDirectory() as input_dir,
            tempfile.TemporaryDirectory() as output_dir,
        ):
            input_file = str(Path(input_dir) / "cp_template.xlsx")
            shutil.copy(pristine_path, input_file)
            wb = openpyxl.load_workbook(input_file)
            vru_ws = wb["CP - VRU Prediction"]
            vru_ws["E4"] = "N/A"
            vru_ws["F4"] = "Green"
            # G4 stays blank: the reference for how E4 must render.
            wb.save(input_file)
            common.add_version_sheet(
                input_file, cli_command=common.CliCommand.GENERATE_TEMPLATE
            )

            report_writer.write_report(
                common.CliCommand.PREPROCESS,
                input_file,
                updated_dfs={},
                output_path=output_dir,
                format_prediction_cells=True,
            )

            output_file = Path(output_dir) / "cp_preprocessed_template.xlsx"
            out_ws = openpyxl.load_workbook(str(output_file))["CP - VRU Prediction"]
            na_cell = out_ws["E4"]
            blank_cell = out_ws["G4"]
            green_cell = out_ws["F4"]

            self.assertEqual(na_cell.value, blank_cell.value)
            self.assertEqual(
                na_cell.fill.start_color.rgb, blank_cell.fill.start_color.rgb
            )
            self.assertNotEqual(
                green_cell.fill.start_color.rgb, blank_cell.fill.start_color.rgb
            )


class TestMissingSheetsAreSkipped(unittest.TestCase):
    def test_preprocess_succeeds_without_documentation_data_sheets(self):
        """A prediction workbook uploaded before the "Documentation"/"Data"
        sheets existed in the template must still preprocess -- the missing
        sheets are skipped with a warning instead of aborting write_report."""
        pristine_path = str(files("data").joinpath("cp_template.xlsx"))
        with (
            tempfile.TemporaryDirectory() as input_dir,
            tempfile.TemporaryDirectory() as output_dir,
        ):
            input_file = str(Path(input_dir) / "cp_template.xlsx")
            shutil.copy(pristine_path, input_file)
            common.add_version_sheet(
                input_file, cli_command=common.CliCommand.GENERATE_TEMPLATE
            )
            wb = openpyxl.load_workbook(input_file)
            for sheet_name in common.DOCUMENTATION_DATA_SHEETS:
                wb.remove(wb[sheet_name])
            wb.save(input_file)

            report_writer.write_report(
                common.CliCommand.PREPROCESS,
                input_file,
                updated_dfs={},
                output_path=output_dir,
                format_prediction_cells=False,
            )

            output_file = Path(output_dir) / "cp_preprocessed_template.xlsx"
            self.assertTrue(output_file.exists())
            out_wb = openpyxl.load_workbook(str(output_file))
            for sheet_name in common.DOCUMENTATION_DATA_SHEETS:
                self.assertNotIn(sheet_name, out_wb.sheetnames)


if __name__ == "__main__":
    unittest.main()
