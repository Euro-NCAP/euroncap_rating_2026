# Copyright 2025-2026, Euro NCAP IVZW
# Created by IVEX NV (https://ivex.ai)
#
# Licensed under the Apache License 2.0.
# See http://www.apache.org/licenses/LICENSE-2.0 for details.

import os
import tempfile
import unittest

import openpyxl
from click.testing import CliRunner

from euroncap_rating_2026.overall import sheets
from euroncap_rating_2026.overall.generate_template import generate_template


class TestGenerateTemplate(unittest.TestCase):
    """overall's generate-template follows the four domains' pattern: every
    not-yet-computed cell (sheets.COMPUTED_SHEET_COLUMNS) reads "-" instead
    of a misleading placeholder 0, and the Version sheet is stamped."""

    def _generate(self):
        runner = CliRunner()
        tmpdir = tempfile.mkdtemp()
        cwd = os.getcwd()
        os.chdir(tmpdir)
        try:
            result = runner.invoke(generate_template, [])
            self.assertEqual(result.exit_code, 0, result.output)
            return openpyxl.load_workbook(os.path.join(tmpdir, "overall_template.xlsx"))
        finally:
            os.chdir(cwd)

    def test_computed_cells_are_dash_and_version_sheet_present(self):
        wb = self._generate()

        self.assertIn("Version", wb.sheetnames)
        for sheet_name, columns in sheets.COMPUTED_SHEET_COLUMNS.items():
            ws = wb[sheet_name]
            header = [c.value for c in ws[1]]
            for column in columns:
                col_idx = header.index(column) + 1
                values = {
                    ws.cell(row, col_idx).value for row in range(2, ws.max_row + 1)
                }
                self.assertEqual(values, {"-"}, f"{sheet_name}!{column}: {values}")


if __name__ == "__main__":
    unittest.main()
