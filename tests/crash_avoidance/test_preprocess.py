# Copyright 2025-2026, Euro NCAP IVZW
# Created by IVEX NV (https://ivex.ai)
#
# Licensed under the Apache License 2.0.
# See http://www.apache.org/licenses/LICENSE-2.0 for details.

import importlib
import shutil
import tempfile
import unittest
from importlib.resources import files
from pathlib import Path
from unittest.mock import patch

from click.testing import CliRunner

import openpyxl

from euroncap_rating_2026 import common
from euroncap_rating_2026.crash_avoidance import data_model, test_info
from euroncap_rating_2026.crash_avoidance.compute_score import compute_score
from euroncap_rating_2026.crash_avoidance.preprocess import (
    add_verification_sheets_to_dfs,
    preprocess,
)

# crash_avoidance/__init__.py does `from .preprocess import preprocess`, which
# rebinds the `crash_avoidance.preprocess` attribute from the submodule to the
# click Command object. importlib.import_module goes through sys.modules
# instead of that (shadowed) attribute, so it reliably returns the submodule
# regardless of import order -- unlike `import ... as` or `from ... import`,
# which both resolve via the same shadowed attribute.
preprocess_module = importlib.import_module(
    "euroncap_rating_2026.crash_avoidance.preprocess"
)


class TestAddVerificationSheetsToDfsRebuildsTrustedInput(unittest.TestCase):
    """Regression test for the same class of gap crash_protection's
    add_vru_sheets_to_dfs was fixed for: add_verification_sheets_to_dfs(dfs)
    -- the dataframe-in/dataframe-out entry point for test-run selection --
    must also reject a tampered "Input parameters" sheet."""

    def test_reordered_input_parameters_rows_are_rejected(self):
        pristine_path = str(files("data").joinpath("ca_template.xlsx"))
        dfs = common.read_excel_file_to_dfs(pristine_path)
        input_parameters = dfs["Input parameters"].copy()
        input_parameters.iloc[[0, 1]] = input_parameters.iloc[[1, 0]].values
        dfs["Input parameters"] = input_parameters

        with self.assertRaises(common.TemplateIntegrityError):
            add_verification_sheets_to_dfs(dfs)


class TestPreprocessCliHandlesValueErrorGracefully(unittest.TestCase):
    """A ValueError raised by preprocess_stage_subelement (e.g. the LSC
    consistency check for CBDA) must produce a clean "Error: ..." message
    and exit(1), not an unhandled traceback -- mirroring how
    common.TemplateIntegrityError is already handled."""

    def test_value_error_prints_message_and_exits_1_without_traceback(self):
        pristine_path = str(files("data").joinpath("ca_template.xlsx"))
        runner = CliRunner()
        with (
            tempfile.TemporaryDirectory() as input_dir,
            tempfile.TemporaryDirectory() as output_dir,
        ):
            # preprocess() checks for a "Version" sheet (added by
            # generate-template) before ever reaching preprocess_stage_subelement.
            input_file = str(Path(input_dir) / "ca_template_versioned.xlsx")
            shutil.copy(pristine_path, input_file)
            common.add_version_sheet(
                input_file, cli_command=common.CliCommand.GENERATE_TEMPLATE
            )

            with patch.object(
                preprocess_module,
                "preprocess_stage_subelement",
                side_effect=ValueError(
                    "LSC consistency check failed for scenario CBDA with vehicle "
                    "response VehicleResponse.INFORMATION. Please update color "
                    "specified according to LSC protocol"
                ),
            ):
                result = runner.invoke(
                    preprocess,
                    ["--input_file", input_file, "--output_path", output_dir],
                )

            self.assertEqual(result.exit_code, 1)
            self.assertIn("Error: LSC consistency check failed", result.output)
            self.assertNotIn("Traceback (most recent call last)", result.output)
            # No report was written -- the process must have bailed out before
            # write_report(), not just printed the error and continued.
            self.assertEqual(list(Path(output_dir).iterdir()), [])


class TestLscConsistencyRaisesCleanlyOnRealMismatch(unittest.TestCase):
    """The test above mocks preprocess_stage_subelement to prove the CLI's
    exception handling is wired up; it never runs the real LSC consistency
    check (test_info.check_lsc_consistency). This exercises the actual CBDA
    dooring mismatch end to end -- Vehicle response=Warning with a Green
    dooring point, which DOORING_ALLOWED_COLORS forbids -- on both entry
    points: the CLI (`crash_avoidance preprocess`) and the
    dataframe-in/dataframe-out sibling add_verification_sheets_to_dfs
    (called directly by non-CLI callers; a generic except-Exception wrapper
    around it turns this into an HTTP 422 error)."""

    def _tampered_workbook(self):
        wb = openpyxl.load_workbook(str(files("data").joinpath("ca_template.xlsx")))
        ws = wb["Input parameters"]
        header = {
            ws.cell(row=1, column=c).value: c for c in range(1, ws.max_column + 1)
        }
        for row in range(2, ws.max_row + 1):
            if (
                ws.cell(row=row, column=header["Input parameter"]).value
                == "Vehicle response"
            ):
                ws.cell(row=row, column=header["Value"]).value = "Warning"
        wb["LSC - Ped & Cyc pred."]["B24"] = "Green"
        return wb

    def test_cli_prints_clean_error_and_exits_1(self):
        runner = CliRunner()
        with (
            tempfile.TemporaryDirectory() as input_dir,
            tempfile.TemporaryDirectory() as output_dir,
        ):
            input_file = str(Path(input_dir) / "ca_tampered.xlsx")
            self._tampered_workbook().save(input_file)
            common.add_version_sheet(
                input_file, cli_command=common.CliCommand.GENERATE_TEMPLATE
            )

            result = runner.invoke(
                preprocess, ["--input_file", input_file, "--output_path", output_dir]
            )

            self.assertEqual(result.exit_code, 1)
            self.assertIn("Error: LSC consistency check failed", result.output)
            self.assertNotIn("Traceback (most recent call last)", result.output)
            self.assertEqual(list(Path(output_dir).iterdir()), [])

    def test_dataframe_path_raises_descriptive_value_error(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            path = str(Path(tmpdir) / "ca_tampered.xlsx")
            self._tampered_workbook().save(path)
            dfs = common.read_excel_file_to_dfs(path)

            with self.assertRaises(ValueError) as ctx:
                add_verification_sheets_to_dfs(dfs)

            message = str(ctx.exception)
            self.assertIn("LSC consistency check failed", message)
            self.assertIn("Please update color", message)


class TestElkReLdwExpectedValue(unittest.TestCase):
    """When the "Extended range performance" input
    parameter is LDW, the ELK RE extended-range verification rows must
    expect "DTLE @ T_LDW" (the DTLE at the warning instant, LDC v1.2
    §5.2.2.2) instead of dtle_t_end. Standard-range rows -- LDW is an
    extended-range-only concept -- and ELK/blank performance keep
    dtle_t_end."""

    def _elk_re_rows(self, erp_value):
        wb = openpyxl.load_workbook(str(files("data").joinpath("ca_template.xlsx")))

        ws = wb["Input parameters"]
        header = {
            ws.cell(row=1, column=c).value: c for c in range(1, ws.max_column + 1)
        }
        current_scenario = None
        for row in range(2, ws.max_row + 1):
            scenario = ws.cell(row=row, column=header["Scenario"]).value
            if scenario is not None and str(scenario).strip() != "":
                current_scenario = str(scenario).strip()
            if (
                current_scenario == "ELK RE"
                and ws.cell(row=row, column=header["Input parameter"]).value
                == "Extended range performance"
            ):
                ws.cell(row=row, column=header["Value"]).value = erp_value
                break
        else:
            self.fail("ELK RE / Extended range performance row not found")

        # Fill the whole ELK RE grid so selection has candidates: Green in
        # the standard range, Orange (an LDW pass) in the extended range.
        # Matrix (i, j) -> worksheet row 3+i, column B(=2)+j.
        pred = wb["LDC - Single Veh pred."]
        extended = set(data_model.EXTENDED_RANGE_CELLS["ELK RE"])
        for i in range(6):
            for j in range(6):
                pred.cell(row=3 + i, column=2 + j).value = (
                    "Orange" if (i, j) in extended else "Green"
                )

        with tempfile.TemporaryDirectory() as tmpdir:
            path = str(Path(tmpdir) / "ca_filled.xlsx")
            wb.save(path)
            dfs = common.read_excel_file_to_dfs(path)

        result = test_info.preprocess_stage_subelement(
            dfs, "Lane Departure Collisions", "Single vehicle"
        )
        df = result.test_points_df
        return df[df["Scenario"] == "ELK RE"]

    def test_ldw_switches_extended_expected_value(self):
        rows = self._elk_re_rows("LDW")
        extended = rows[rows["Range"] == "Extended"]
        standard = rows[rows["Range"] == "Standard"]
        self.assertGreater(len(extended), 0)
        self.assertGreater(len(standard), 0)
        self.assertEqual(set(extended["Expected value"]), {"DTLE @ T_LDW"})
        self.assertEqual(set(standard["Expected value"]), {"dtle_t_end"})

    def test_ldw_is_case_insensitive(self):
        rows = self._elk_re_rows("ldw")
        extended = rows[rows["Range"] == "Extended"]
        self.assertGreater(len(extended), 0)
        self.assertEqual(set(extended["Expected value"]), {"DTLE @ T_LDW"})

    def test_elk_keeps_dtle_t_end_everywhere(self):
        rows = self._elk_re_rows("ELK")
        self.assertGreater(len(rows), 0)
        self.assertEqual(set(rows["Expected value"]), {"dtle_t_end"})

    def test_blank_performance_keeps_dtle_t_end(self):
        rows = self._elk_re_rows(None)
        self.assertGreater(len(rows), 0)
        self.assertEqual(set(rows["Expected value"]), {"dtle_t_end"})


class TestDocumentationDataSurviveFullPipeline(unittest.TestCase):
    """preprocess() must compute the Documentation/Data flags once, and
    compute_score() must carry them through unchanged (not recompute, not
    drop, not blank) -- run on the pristine (blank) template, so every flag
    is expected to come out "FALSE"."""

    def test_preprocess_then_compute_score_e2e(self):
        pristine_path = str(files("data").joinpath("ca_template.xlsx"))
        runner = CliRunner()
        with (
            tempfile.TemporaryDirectory() as input_dir,
            tempfile.TemporaryDirectory() as output_dir,
        ):
            input_file = str(Path(input_dir) / "ca_template_versioned.xlsx")
            shutil.copy(pristine_path, input_file)
            common.add_version_sheet(
                input_file, cli_command=common.CliCommand.GENERATE_TEMPLATE
            )

            result = runner.invoke(
                preprocess, ["--input_file", input_file, "--output_path", output_dir]
            )
            self.assertEqual(result.exit_code, 0, result.output)

            preprocessed_file = str(Path(output_dir) / "ca_preprocessed_template.xlsx")
            preprocessed_wb = openpyxl.load_workbook(preprocessed_file)
            expected = {}
            for sheet_name, id_header in [
                ("Documentation", "Documentation ID"),
                ("Data", "Data ID"),
            ]:
                ws = preprocessed_wb[sheet_name]
                self.assertEqual(ws.sheet_state, "hidden")
                self.assertEqual([c.value for c in ws[1]], [id_header, "Status"])
                rows = {
                    ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value
                    for r in range(2, ws.max_row + 1)
                }
                self.assertTrue(rows, f"{sheet_name} has no rows")
                self.assertTrue(all(v == "FALSE" for v in rows.values()), rows)
                expected[sheet_name] = rows

            result2 = runner.invoke(
                compute_score,
                ["--input_file", preprocessed_file, "--output_path", output_dir],
            )
            self.assertEqual(result2.exit_code, 0, result2.output)

            report_path = next(
                p for p in Path(output_dir).iterdir() if "report" in p.name
            )
            report_wb = openpyxl.load_workbook(str(report_path))
            for sheet_name, id_header in [
                ("Documentation", "Documentation ID"),
                ("Data", "Data ID"),
            ]:
                ws = report_wb[sheet_name]
                self.assertEqual(ws.sheet_state, "hidden")
                self.assertEqual([c.value for c in ws[1]], [id_header, "Status"])
                rows = {
                    ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value
                    for r in range(2, ws.max_row + 1)
                }
                # Carried through byte-for-byte: same literal "FALSE" strings
                # preprocess wrote, not recomputed and not coerced to bool.
                self.assertEqual(rows, expected[sheet_name])
                for value in rows.values():
                    self.assertIsInstance(value, str)


if __name__ == "__main__":
    unittest.main()
