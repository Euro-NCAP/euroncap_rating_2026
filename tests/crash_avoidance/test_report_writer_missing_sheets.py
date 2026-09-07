# Copyright 2025-2026, Euro NCAP IVZW
# Created by IVEX NV (https://ivex.ai)
#
# Licensed under the Apache License 2.0.
# See http://www.apache.org/licenses/LICENSE-2.0 for details.

"""A prediction workbook uploaded before the "Documentation"/"Data" sheets
existed in the template must still preprocess -- write_report skips the
missing sheets with a warning instead of aborting."""

import shutil
import tempfile
import unittest
from importlib.resources import files
from pathlib import Path

import openpyxl

from euroncap_rating_2026 import common
from euroncap_rating_2026.crash_avoidance import report_writer


class TestMissingSheetsAreSkipped(unittest.TestCase):
    def _write_report(self, sheets_to_remove):
        pristine_path = str(files("data").joinpath("ca_template.xlsx"))
        with (
            tempfile.TemporaryDirectory() as input_dir,
            tempfile.TemporaryDirectory() as output_dir,
        ):
            input_file = str(Path(input_dir) / "ca_template.xlsx")
            shutil.copy(pristine_path, input_file)
            common.add_version_sheet(
                input_file, cli_command=common.CliCommand.GENERATE_TEMPLATE
            )
            wb = openpyxl.load_workbook(input_file)
            for sheet_name in sheets_to_remove:
                wb.remove(wb[sheet_name])
            wb.save(input_file)

            report_writer.write_report(
                common.CliCommand.PREPROCESS,
                input_file,
                updated_dfs={},
                output_path=output_dir,
            )

            output_file = Path(output_dir) / "ca_preprocessed_template.xlsx"
            self.assertTrue(output_file.exists())
            return openpyxl.load_workbook(str(output_file))

    def test_preprocess_succeeds_without_documentation_data_sheets(self):
        out_wb = self._write_report(common.DOCUMENTATION_DATA_SHEETS)
        for sheet_name in common.DOCUMENTATION_DATA_SHEETS:
            self.assertNotIn(sheet_name, out_wb.sheetnames)

    def test_preprocess_succeeds_without_cpla_cbla_prediction_sheet(self):
        """The CPLA/CBLA standard-range box formatting runs on the copied
        "FC - Ped & Cyc pred." sheet -- when that sheet is missing from the
        input it must be skipped too, not KeyError."""
        out_wb = self._write_report(["FC - Ped & Cyc pred."])
        self.assertNotIn("FC - Ped & Cyc pred.", out_wb.sheetnames)


if __name__ == "__main__":
    unittest.main()
