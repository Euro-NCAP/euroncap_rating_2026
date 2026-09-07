# Copyright 2025-2026, Euro NCAP IVZW
# Created by IVEX NV (https://ivex.ai)
#
# Licensed under the Apache License 2.0.
# See http://www.apache.org/licenses/LICENSE-2.0 for details.

"""Report-only prediction-consistency collectors for crash_avoidance
(consistency.find_prediction_inconsistencies) and their integration into
integrity.find_consistency_findings."""

import unittest

import openpyxl

from euroncap_rating_2026 import common
from euroncap_rating_2026.crash_avoidance import (
    consistency,
    data_model,
    integrity,
    matrix_processing,
)


def _matrix_ws_cell(df, test_name, i, j):
    """Worksheet coordinate of matrix cell (i, j), via the same
    get_matrix_indices resolution production code uses."""
    indices = matrix_processing.get_matrix_indices(df, test_name)
    from openpyxl.utils import get_column_letter

    return (
        f"{get_column_letter(indices['start_col'] + j + 1)}"
        f"{indices['start_row'] + i + 2}"
    )


class TestCrashAvoidanceConsistency(unittest.TestCase):
    def setUp(self):
        from importlib.resources import files

        self.pristine_path = str(files("data").joinpath("ca_template.xlsx"))
        self.wb = openpyxl.load_workbook(self.pristine_path, data_only=True)

    def _set_vehicle_response(self, value):
        ws = self.wb["Input parameters"]
        header = {
            ws.cell(row=1, column=c).value: c for c in range(1, ws.max_column + 1)
        }
        for row in range(2, ws.max_row + 1):
            if ws.cell(row=row, column=header["Input parameter"]).value == (
                "Vehicle response"
            ):
                ws.cell(row=row, column=header["Value"]).value = value
                return
        self.fail("Vehicle response row not found in Input parameters")

    def _set_matrix_cell(self, sheet_name, test_name, i, j, value):
        df = common.sheet_to_dataframe(self.wb[sheet_name])
        cell = _matrix_ws_cell(df, test_name, i, j)
        self.wb[sheet_name][cell] = value
        return cell

    def _set_input_parameter(self, scenario, parameter, value):
        """Set an Input parameters Value cell, scenario-aware: parameter
        names like "Prediction - Standard" repeat per scenario, so walk the
        rows tracking the last non-blank Scenario cell. Returns the Value
        cell coordinate."""
        ws = self.wb["Input parameters"]
        header = {
            ws.cell(row=1, column=c).value: c for c in range(1, ws.max_column + 1)
        }
        current_scenario = None
        for row in range(2, ws.max_row + 1):
            scenario_value = ws.cell(row=row, column=header["Scenario"]).value
            if scenario_value is not None and str(scenario_value).strip() != "":
                current_scenario = str(scenario_value).strip()
            if (
                current_scenario == scenario
                and ws.cell(row=row, column=header["Input parameter"]).value
                == parameter
            ):
                cell = ws.cell(row=row, column=header["Value"])
                cell.value = value
                return cell.coordinate
        self.fail(f"{scenario} / {parameter} row not found in Input parameters")

    def test_pristine_template_has_no_findings(self):
        self.assertEqual(consistency.find_prediction_inconsistencies(self.wb), [])

    def test_dooring_violation_reported_per_cell(self):
        self._set_vehicle_response("Warning")
        green = self._set_matrix_cell("LSC - Ped & Cyc pred.", "CBDA", 0, 0, "Green")
        brown = self._set_matrix_cell("LSC - Ped & Cyc pred.", "CBDA", 1, 2, "Brown")
        allowed = self._set_matrix_cell("LSC - Ped & Cyc pred.", "CBDA", 2, 1, "Orange")

        findings = consistency.find_prediction_inconsistencies(self.wb)

        self.assertEqual({f.check for f in findings}, {common.CHECK_LSC_DOORING})
        self.assertEqual({f.cell for f in findings}, {green, brown})
        self.assertNotIn(allowed, {f.cell for f in findings})
        for f in findings:
            self.assertEqual(f.sheet, "LSC - Ped & Cyc pred.")
            self.assertIn("Vehicle response is 'Warning'", f.description)
            self.assertEqual(f.row_identity.get("Scenario"), "CBDA")

    def test_dooring_ignores_grey_red_and_missing_response(self):
        # Red is never selected for testing, so it never reaches the
        # blocking preprocess check either; grey is simply unfilled.
        self._set_vehicle_response("Warning")
        self._set_matrix_cell("LSC - Ped & Cyc pred.", "CBDA", 0, 0, "Red")
        self.assertEqual(consistency.find_prediction_inconsistencies(self.wb), [])

        # Without a Vehicle response there is nothing to compare against --
        # the empty parameter cell is the missing-input check's finding.
        self._set_vehicle_response(None)
        self._set_matrix_cell("LSC - Ped & Cyc pred.", "CBDA", 0, 0, "Green")
        self.assertEqual(consistency.find_prediction_inconsistencies(self.wb), [])

    def test_coherence_mismatch_reported_as_one_finding_per_group(self):
        sheet = "FC - Ped & Cyc pred."
        # CPLA day: the fixed original-AEB band and the choice band overlap
        # at VUT 50/60 km/h (matrix rows 4/5 fixed, 6/7 choice).
        fixed = self._set_matrix_cell(sheet, "CPLA day", 5, 0, "Green")
        choice = self._set_matrix_cell(sheet, "CPLA day", 7, 0, "Red")
        # Function dropdown lives in worksheet column C on the choice row.
        choice_row = int("".join(ch for ch in choice if ch.isdigit()))
        self.wb[sheet][f"C{choice_row}"] = "AEB"

        findings = consistency.find_prediction_inconsistencies(self.wb)

        coherence = [f for f in findings if f.check == common.CHECK_CPLA_CBLA_COHERENCE]
        self.assertEqual(len(coherence), 1)
        finding = coherence[0]
        self.assertEqual(finding.sheet, sheet)
        self.assertEqual(finding.cell, fixed)
        self.assertIn("must predict the same color", finding.description)
        self.assertEqual(finding.row_identity.get("Scenario"), "CPLA day")

    def test_fcw_graded_color_reported(self):
        sheet = "FC - Ped & Cyc pred."
        cell = self._set_matrix_cell(sheet, "CPLA day", 6, 0, "Yellow")
        row = int("".join(ch for ch in cell if ch.isdigit()))
        self.wb[sheet][f"C{row}"] = "FCW"
        # Same colour under AEB on another choice row is fine.
        ok_cell = self._set_matrix_cell(sheet, "CPLA day", 7, 1, "Orange")
        ok_row = int("".join(ch for ch in ok_cell if ch.isdigit()))
        self.wb[sheet][f"C{ok_row}"] = "AEB"

        findings = consistency.find_prediction_inconsistencies(self.wb)

        fcw = [f for f in findings if f.check == common.CHECK_FCW_COLOR]
        self.assertEqual([f.cell for f in fcw], [cell])
        self.assertIn("Function is FCW", fcw[0].description)
        # The AEB-declared graded colour must not be flagged by any rule.
        self.assertNotIn(ok_cell, {f.cell for f in findings})

    def test_scope_filtering(self):
        self._set_vehicle_response("Warning")
        self._set_matrix_cell("LSC - Ped & Cyc pred.", "CBDA", 0, 0, "Green")
        cell = self._set_matrix_cell("FC - Ped & Cyc pred.", "CPLA day", 6, 0, "Brown")
        row = int("".join(ch for ch in cell if ch.isdigit()))
        self.wb["FC - Ped & Cyc pred."][f"C{row}"] = "FCW"

        lsc = consistency.find_prediction_inconsistencies(
            self.wb, "Low Speed Collisions", "Pedestrian & cyclist"
        )
        fc = consistency.find_prediction_inconsistencies(
            self.wb, "Frontal Collisions", "Pedestrian & cyclist"
        )
        other = consistency.find_prediction_inconsistencies(
            self.wb, "Frontal Collisions", "Car & PTW"
        )

        self.assertEqual({f.check for f in lsc}, {common.CHECK_LSC_DOORING})
        self.assertEqual({f.check for f in fc}, {common.CHECK_FCW_COLOR})
        self.assertEqual(other, [])

    def test_na_one_sided_reported(self):
        self._set_input_parameter("CCRs", "Prediction - Standard", "N/A")
        other_cell = self._set_input_parameter("CCRs", "Prediction - Extended", "VTA")

        findings = consistency.find_prediction_inconsistencies(self.wb)

        na_findings = [
            f for f in findings if f.check == common.CHECK_SCENARIO_NA_CONSISTENCY
        ]
        self.assertEqual(len(na_findings), 1)
        finding = na_findings[0]
        self.assertEqual(finding.sheet, "Input parameters")
        self.assertEqual(finding.cell, other_cell)
        self.assertEqual(finding.row_identity.get("Scenario"), "CCRs")
        self.assertEqual(
            finding.row_identity.get("Input parameter"), "Prediction - Extended"
        )
        self.assertIn("all of its input parameters must be N/A", finding.description)

    def test_na_all_na_or_blank_not_reported(self):
        # Both N/A: consistent, no finding.
        self._set_input_parameter("CCRs", "Prediction - Standard", "N/A")
        self._set_input_parameter("CCRs", "Prediction - Extended", "N/A")
        # N/A + blank: the unfilled cell is the missing-input check's finding.
        self._set_input_parameter("CCRm", "Prediction - Standard", "N/A")

        findings = consistency.find_prediction_inconsistencies(self.wb)

        self.assertEqual(
            [f for f in findings if f.check == common.CHECK_SCENARIO_NA_CONSISTENCY],
            [],
        )

    def test_na_covers_all_three_elk_re_parameters(self):
        # ELK RE has three input parameters; an N/A Extended range
        # performance with two filled prediction methods must flag both
        # prediction cells.
        std_cell = self._set_input_parameter("ELK RE", "Prediction - Standard", "VTA")
        ext_cell = self._set_input_parameter(
            "ELK RE", "Prediction - Extended", "Self claimed"
        )
        self._set_input_parameter("ELK RE", "Extended range performance", "N/A")

        findings = consistency.find_prediction_inconsistencies(self.wb)

        na_findings = [
            f for f in findings if f.check == common.CHECK_SCENARIO_NA_CONSISTENCY
        ]
        self.assertEqual({f.cell for f in na_findings}, {std_cell, ext_cell})
        for f in na_findings:
            self.assertIn("Extended range performance is 'N/A'", f.description)

    def test_na_all_three_elk_re_parameters_na_not_reported(self):
        self._set_input_parameter("ELK RE", "Prediction - Standard", "N/A")
        self._set_input_parameter("ELK RE", "Prediction - Extended", "N/A")
        self._set_input_parameter("ELK RE", "Extended range performance", "N/A")

        findings = consistency.find_prediction_inconsistencies(self.wb)

        self.assertEqual(
            [f for f in findings if f.check == common.CHECK_SCENARIO_NA_CONSISTENCY],
            [],
        )

    def test_na_predictions_na_but_erp_filled_reported(self):
        self._set_input_parameter("ELK RE", "Prediction - Standard", "N/A")
        self._set_input_parameter("ELK RE", "Prediction - Extended", "N/A")
        erp_cell = self._set_input_parameter(
            "ELK RE", "Extended range performance", "ELK"
        )

        findings = consistency.find_prediction_inconsistencies(self.wb)

        na_findings = [
            f for f in findings if f.check == common.CHECK_SCENARIO_NA_CONSISTENCY
        ]
        self.assertEqual([f.cell for f in na_findings], [erp_cell])
        self.assertEqual(
            na_findings[0].row_identity.get("Input parameter"),
            "Extended range performance",
        )

    def test_na_scope_filtering(self):
        self._set_input_parameter("CCRs", "Prediction - Standard", "N/A")
        self._set_input_parameter("CCRs", "Prediction - Extended", "VTA")
        self._set_input_parameter("ELK RE", "Prediction - Standard", "Self claimed")
        self._set_input_parameter("ELK RE", "Prediction - Extended", "N/A")

        all_findings = consistency.find_prediction_inconsistencies(self.wb)
        fc = consistency.find_prediction_inconsistencies(
            self.wb, "Frontal Collisions", "Car & PTW"
        )
        ldc = consistency.find_prediction_inconsistencies(
            self.wb, "Lane Departure Collisions", "Single vehicle"
        )
        lsc = consistency.find_prediction_inconsistencies(
            self.wb, "Low Speed Collisions", "Pedestrian & cyclist"
        )

        def na_scenarios(findings):
            return {
                f.row_identity.get("Scenario")
                for f in findings
                if f.check == common.CHECK_SCENARIO_NA_CONSISTENCY
            }

        self.assertEqual(na_scenarios(all_findings), {"CCRs", "ELK RE"})
        self.assertEqual(na_scenarios(fc), {"CCRs"})
        self.assertEqual(na_scenarios(ldc), {"ELK RE"})
        self.assertEqual(na_scenarios(lsc), set())

    def test_ldw_green_extended_reported_per_cell(self):
        # Extended range performance = LDW means the
        # extended range is covered by LDW only, so a Green (ELK pass)
        # prediction there is contradictory. Standard cells are unaffected.
        sheet = "LDC - Single Veh pred."
        self._set_input_parameter("ELK RE", "Extended range performance", "LDW")
        green_row = self._set_matrix_cell(sheet, "ELK RE", 0, 0, "Green")
        green_col = self._set_matrix_cell(sheet, "ELK RE", 3, 5, "Green")
        allowed_orange = self._set_matrix_cell(sheet, "ELK RE", 1, 1, "Orange")
        allowed_red = self._set_matrix_cell(sheet, "ELK RE", 5, 0, "Red")
        allowed_std_green = self._set_matrix_cell(sheet, "ELK RE", 2, 0, "Green")

        findings = consistency.find_prediction_inconsistencies(self.wb)

        ldw = [f for f in findings if f.check == common.CHECK_LDW_EXTENDED_COLOR]
        self.assertEqual({f.cell for f in ldw}, {green_row, green_col})
        flagged = {f.cell for f in findings}
        for cell in (allowed_orange, allowed_red, allowed_std_green):
            self.assertNotIn(cell, flagged)
        for f in ldw:
            self.assertEqual(f.sheet, sheet)
            self.assertEqual(f.row_identity.get("Scenario"), "ELK RE")
            self.assertIn("Extended range performance is 'LDW'", f.description)

    def test_ldw_rule_inactive_for_elk_or_na_performance(self):
        sheet = "LDC - Single Veh pred."
        self._set_matrix_cell(sheet, "ELK RE", 0, 0, "Green")
        for value in ("ELK", "N/A"):
            self._set_input_parameter("ELK RE", "Extended range performance", value)
            findings = consistency.find_prediction_inconsistencies(self.wb)
            self.assertEqual(
                [f for f in findings if f.check == common.CHECK_LDW_EXTENDED_COLOR],
                [],
                f"unexpected LDW finding with performance {value!r}",
            )

    def test_ldw_scope_filtering(self):
        self._set_input_parameter("ELK RE", "Extended range performance", "LDW")
        self._set_matrix_cell("LDC - Single Veh pred.", "ELK RE", 0, 0, "Green")

        ldc = consistency.find_prediction_inconsistencies(
            self.wb, "Lane Departure Collisions", "Single vehicle"
        )
        fc = consistency.find_prediction_inconsistencies(
            self.wb, "Frontal Collisions", "Car & PTW"
        )

        self.assertEqual({f.check for f in ldc}, {common.CHECK_LDW_EXTENDED_COLOR})
        self.assertEqual(
            [f for f in fc if f.check == common.CHECK_LDW_EXTENDED_COLOR], []
        )

    def test_find_consistency_findings_merges_missing_and_rules(self):
        self._set_vehicle_response("Warning")
        self._set_matrix_cell("LSC - Ped & Cyc pred.", "CBDA", 0, 0, "Green")

        findings = integrity.find_consistency_findings(
            self.wb, "Low Speed Collisions", "Pedestrian & cyclist"
        )

        checks = {f.check for f in findings}
        self.assertIn(common.CHECK_MISSING_INPUT, checks)
        self.assertIn(common.CHECK_LSC_DOORING, checks)
        # Undeclared scopes still raise rather than silently pass.
        with self.assertRaises(ValueError):
            integrity.find_consistency_findings(self.wb, "Bogus", "Bogus")


if __name__ == "__main__":
    unittest.main()
