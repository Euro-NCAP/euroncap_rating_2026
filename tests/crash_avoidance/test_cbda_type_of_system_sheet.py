# Copyright 2025-2026, Euro NCAP IVZW
# Created by IVEX NV (https://ivex.ai)
#
# Licensed under the Apache License 2.0.
# See http://www.apache.org/licenses/LICENSE-2.0 for details.

"""End-to-end check that the CBDA "Type of system" column reaches the
generated verification sheet the way a tool-filled column should: right of
"Door selected" in the header order, blank on every CBDA row, without the grey
input background and without a dropdown (the downstream evaluation tool fills
it from the measured dooring KPI, so it must not look like lab input; this
package only emits it blank and reads it back). Complements the
unit tests in test_cbda_scoring.py, which exercise the column's effect on
scoring but never write a workbook."""

import tempfile
import unittest
from importlib.resources import files
from pathlib import Path

import openpyxl
from click.testing import CliRunner

from euroncap_rating_2026 import common
from euroncap_rating_2026.crash_avoidance import test_info
from euroncap_rating_2026.crash_avoidance.preprocess import preprocess

LSC_VERIF_SHEET = "LSC - Ped & Cyc verif."
# The per-test-point header on every crash-avoidance verification sheet except
# "LDC - Single Veh verif." (row 1 carries the general-requirements block).
VERIF_HEADER_ROW = 4
CBDA_MATRIX_ROWS = range(24, 28)  # 0.50 m .. 2.00 m rear gap
CBDA_MATRIX_COLS = range(2, 5)  # 10, 15, 20 km/h


class TestCbdaTypeOfSystemColumn(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        """Preprocess a template whose CBDA matrix is fully predicted Green
        under a Retention claim -- consistent, so all 12 cells are selected."""
        wb = openpyxl.load_workbook(str(files("data").joinpath("ca_template.xlsx")))
        params = wb["Input parameters"]
        header = {
            params.cell(row=1, column=c).value: c
            for c in range(1, params.max_column + 1)
        }
        for row in range(2, params.max_row + 1):
            if (
                params.cell(row=row, column=header["Input parameter"]).value
                == "Vehicle response"
            ):
                params.cell(row=row, column=header["Value"]).value = "Retention"

        prediction = wb["LSC - Ped & Cyc pred."]
        for row in CBDA_MATRIX_ROWS:
            for col in CBDA_MATRIX_COLS:
                prediction.cell(row=row, column=col).value = "Green"

        cls._input_dir = tempfile.TemporaryDirectory()
        cls._output_dir = tempfile.TemporaryDirectory()
        input_file = str(Path(cls._input_dir.name) / "ca_cbda.xlsx")
        wb.save(input_file)
        common.add_version_sheet(
            input_file, cli_command=common.CliCommand.GENERATE_TEMPLATE
        )

        result = CliRunner().invoke(
            preprocess,
            ["--input_file", input_file, "--output_path", cls._output_dir.name],
        )
        assert result.exit_code == 0, result.output

        output_file = Path(cls._output_dir.name) / "ca_preprocessed_template.xlsx"
        cls.ws = openpyxl.load_workbook(str(output_file))[LSC_VERIF_SHEET]
        cls.headers = [
            cls.ws.cell(row=VERIF_HEADER_ROW, column=c).value
            for c in range(1, cls.ws.max_column + 1)
        ]

    @classmethod
    def tearDownClass(cls):
        cls._input_dir.cleanup()
        cls._output_dir.cleanup()

    def _cbda_rows(self):
        scenario_col = self.headers.index("Scenario") + 1
        return [
            row
            for row in range(VERIF_HEADER_ROW + 1, self.ws.max_row + 1)
            if self.ws.cell(row=row, column=scenario_col).value == "CBDA"
        ]

    def test_column_sits_next_to_door_selected(self):
        self.assertIn(test_info.TYPE_OF_SYSTEM_COLUMN, self.headers)
        self.assertEqual(
            self.headers.index(test_info.TYPE_OF_SYSTEM_COLUMN),
            self.headers.index(test_info.DOOR_SELECTED_COLUMN) + 1,
        )

    def test_column_is_not_adjacent_to_value(self):
        # The PASS/FAIL data validation targets the column left of "Value";
        # "Type of system" must not land there.
        self.assertNotEqual(
            self.headers.index(test_info.TYPE_OF_SYSTEM_COLUMN),
            self.headers.index("Value") - 1,
        )

    def test_function_is_not_written_on_cbda_rows(self):
        self.assertNotIn("Function", self.headers)

    def test_every_cbda_row_is_emitted_blank_and_not_grey(self):
        rows = self._cbda_rows()
        self.assertEqual(len(rows), len(CBDA_MATRIX_ROWS) * len(CBDA_MATRIX_COLS))
        col = self.headers.index(test_info.TYPE_OF_SYSTEM_COLUMN) + 1
        for row in rows:
            cell = self.ws.cell(row=row, column=col)
            with self.subTest(row=row):
                self.assertTrue(common.is_empty_cell(cell.value))
                # Tool-filled, not lab input: no grey background.
                self.assertNotIn(common.INPUT_CELL_RGB, str(cell.fill.start_color.rgb))

    def test_no_response_dropdown_on_the_sheet(self):
        dropdowns = [
            dv
            for dv in self.ws.data_validations.dataValidation
            if dv.formula1 == '"Information,Warning,Retention"'
        ]
        self.assertEqual(dropdowns, [])

    def test_expected_value_is_the_generic_ttc(self):
        col = self.headers.index("Expected value") + 1
        for row in self._cbda_rows():
            with self.subTest(row=row):
                self.assertEqual(self.ws.cell(row=row, column=col).value, "TTC")


if __name__ == "__main__":
    unittest.main()
