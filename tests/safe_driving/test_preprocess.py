# Copyright 2025-2026, Euro NCAP IVZW
# Created by IVEX NV (https://ivex.ai)
#
# Licensed under the Apache License 2.0.
# See http://www.apache.org/licenses/LICENSE-2.0 for details.

import shutil
import tempfile
import unittest
from importlib.resources import files
from pathlib import Path

import openpyxl
from click.testing import CliRunner

from euroncap_rating_2026 import common
from euroncap_rating_2026.safe_driving.preprocess import (
    add_verification_sheets_to_dfs,
    preprocess,
)


class TestAddVerificationSheetsToDfsRebuildsTrustedInput(unittest.TestCase):
    """Regression test for the same class of gap crash_protection's
    add_vru_sheets_to_dfs was fixed for: add_verification_sheets_to_dfs(dfs)
    -- the dataframe-in/dataframe-out entry point for ACC
    Performance/Driver Monitoring test-run selection -- must also reject a
    tampered "Input parameters" sheet."""

    def test_reordered_input_parameters_rows_are_rejected(self):
        pristine_path = str(files("data").joinpath("sd_template.xlsx"))
        dfs = common.read_excel_file_to_dfs(pristine_path)
        input_parameters = dfs["Input parameters"].copy()
        input_parameters.iloc[[0, 1]] = input_parameters.iloc[[1, 0]].values
        dfs["Input parameters"] = input_parameters

        with self.assertRaises(common.TemplateIntegrityError):
            add_verification_sheets_to_dfs(dfs)


class TestAddVerificationSheetsToDfsDashesScoreSheets(unittest.TestCase):
    """The DataFrame path must hand back score sheets whose
    not-yet-computed Score columns read "-", not the template's
    placeholder -- including the deliberately-unschema'd "Scenario
    Scores" (see integrity.COMPUTED_SHEET_COLUMNS)."""

    def test_score_sheets_read_dash(self):
        pristine_path = str(files("data").joinpath("sd_template.xlsx"))
        dfs = common.read_excel_file_to_dfs(pristine_path)
        dfs["Test Scores"] = dfs["Test Scores"].copy()
        dfs["Test Scores"]["Score"] = 0

        _, _, result = add_verification_sheets_to_dfs(dfs)

        for sheet in ["Test Scores", "Category Scores", "Scenario Scores"]:
            self.assertTrue(
                (result[sheet]["Score"] == "-").all(), f"{sheet} not dashed"
            )
        # The input dfs are not mutated.
        self.assertTrue((dfs["Test Scores"]["Score"] == 0).all())


class TestPreprocessCliDashesScoreSheets(unittest.TestCase):
    """The CLI path dashes the score sheets via report_writer's
    blank_computed_columns call (driven by integrity.COMPUTED_SHEET_COLUMNS,
    same declaration as the DataFrame path). Pin the file-level behavior:
    every populated Score cell in the preprocessed file is "-"."""

    def test_preprocessed_file_score_cells_are_dash(self):
        pristine_path = str(files("data").joinpath("sd_template.xlsx"))
        runner = CliRunner()
        with (
            tempfile.TemporaryDirectory() as input_dir,
            tempfile.TemporaryDirectory() as output_dir,
        ):
            input_file = str(Path(input_dir) / "sd_template.xlsx")
            shutil.copy(pristine_path, input_file)
            common.add_version_sheet(
                input_file, cli_command=common.CliCommand.GENERATE_TEMPLATE
            )
            # preprocess() needs "Number of rear seats" filled in; every
            # other input can stay blank.
            wb = openpyxl.load_workbook(input_file)
            ws = wb["Input parameters"]
            header = [c.value for c in ws[1]]
            param_col = header.index("Input parameter") + 1
            value_col = header.index("Value") + 1
            for row in range(2, ws.max_row + 1):
                if ws.cell(row, param_col).value == "Number of rear seats":
                    ws.cell(row, value_col).value = 2
            wb.save(input_file)

            result = runner.invoke(
                preprocess, ["--input_file", input_file, "--output_path", output_dir]
            )
            self.assertEqual(result.exit_code, 0, result.output)

            output_file = str(Path(output_dir) / "sd_preprocessed_template.xlsx")
            output_wb = openpyxl.load_workbook(output_file)
            for sheet_name in ["Test Scores", "Category Scores", "Scenario Scores"]:
                ws = output_wb[sheet_name]
                header = [c.value for c in ws[1]]
                score_col = header.index("Score") + 1
                values = {
                    ws.cell(row, score_col).value for row in range(2, ws.max_row + 1)
                }
                self.assertEqual(values, {"-"}, f"{sheet_name}: {values}")


if __name__ == "__main__":
    unittest.main()
