# Copyright 2025-2026, Euro NCAP IVZW
# Created by IVEX NV (https://ivex.ai)
#
# Licensed under the Apache License 2.0.
# See http://www.apache.org/licenses/LICENSE-2.0 for details.

"""End-to-end behaviour of a Driver monitoring prediction group that mixes Red
and Green through the real `preprocess` CLI.

A (Task, Movement type, Vehicle response) group is allowed to hold both Red and
Green cells. The preprocessed template must give those cells back exactly as the
OEM entered them, while the group is still fully suppressed: no test point is
selected on any of its gaze locations, and it still scores nothing because the
colour source scoring reads is no longer all-Green.

The unit tests in test_driver_monitoring.py work on hand-built cell dicts. This
test instead runs the packaged template through `preprocess`, which re-encodes
the typed colour words as cell fills and blanks the text -- the only shape in
which the repaint was ever visible.
"""

import shutil
import tempfile
import unittest
from importlib.resources import files
from pathlib import Path

import openpyxl
from click.testing import CliRunner

from euroncap_rating_2026 import common
from euroncap_rating_2026.safe_driving import driver_monitoring
from euroncap_rating_2026.safe_driving.preprocess import preprocess

# "Long distraction" starts at worksheet row 15; its (Task=Non-driving task,
# Movement type=Owl) block spans rows 17-21, with the Warning strategy in
# column D -- one group of five gaze locations.
GROUP_ROWS = range(17, 22)
WARNING_COLUMN = 4
GREEN_FILL_ARGB = "00009933"
RED_FILL_ARGB = "00FF3333"
# The noise-variable Performance cells (column D, rows 2-12); all mandatory
# ones must be Functional or preprocess skips test-point selection entirely.
NOISE_VARIABLE_ROWS = range(2, 13)
# Every prediction cell of the sheet (columns D-F), so that selection actually
# runs and the "nothing selected in this group" assertion is not vacuous.
PREDICTION_ROWS = range(17, 72)
PREDICTION_COLUMNS = (4, 5, 6)


def _fill_template(path):
    """The whole prediction sheet Green, except the (Non-driving task, Owl,
    Warning) group, where D19 is Red and D20 is N/A."""
    common.add_version_sheet(path, cli_command=common.CliCommand.GENERATE_TEMPLATE)
    wb = openpyxl.load_workbook(path)

    ws = wb["Input parameters"]
    header = [c.value for c in ws[1]]
    param_col = header.index("Input parameter") + 1
    value_col = header.index("Value") + 1
    for row in range(2, ws.max_row + 1):
        if ws.cell(row, param_col).value == "Number of rear seats":
            ws.cell(row, value_col).value = 2

    ws = wb["DE - DM pred."]
    for row in NOISE_VARIABLE_ROWS:
        ws.cell(row, WARNING_COLUMN).value = "Functional"
    for row in PREDICTION_ROWS:
        for column in PREDICTION_COLUMNS:
            cell = ws.cell(row, column)
            if isinstance(cell, openpyxl.cell.cell.MergedCell):
                continue
            cell.value = "Green"
    ws["D19"] = "Red"
    ws["D20"] = "N/A"
    wb.save(path)


def _run_preprocess(input_file, output_dir):
    result = CliRunner().invoke(
        preprocess, ["--input_file", input_file, "--output_path", str(output_dir)]
    )
    assert result.exit_code == 0, result.output
    return str(Path(output_dir) / "sd_preprocessed_template.xlsx")


def _fill_argb(cell):
    return cell.fill.start_color.rgb if cell.fill.fill_type else None


class TestMixedPredictionGroupThroughPreprocess(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls._tmp = tempfile.TemporaryDirectory()
        root = Path(cls._tmp.name)
        for name in ("in", "pre", "pre2"):
            (root / name).mkdir()

        input_file = str(root / "in" / "sd_template.xlsx")
        shutil.copy(str(files("data").joinpath("sd_template.xlsx")), input_file)
        _fill_template(input_file)

        cls.preprocessed_file = _run_preprocess(input_file, root / "pre")
        # Re-running preprocess on its own output must reach the same verdict.
        # The second pass reads the colours back off the sheet it is handed, so
        # dropping the repaint must not turn the group's Green cells red on the
        # way through, nor let a test point appear in it.
        cls.twice_preprocessed_file = _run_preprocess(
            cls.preprocessed_file, root / "pre2"
        )

    @classmethod
    def tearDownClass(cls):
        cls._tmp.cleanup()

    def _prediction_sheet(self, path):
        return openpyxl.load_workbook(path)["DE - DM pred."]

    def test_oem_cells_come_back_as_entered(self):
        """Also on the twice-preprocessed file: the second pass must not repaint
        what the first one preserved.

        D20 survives as the literal text "N/A" because pandas reads it as NaN
        ("DE - DM pred." is not in data_model.PRESERVE_NA_SHEETS), so it never
        enters cell_dict and no rule writes to it."""
        for path in (self.preprocessed_file, self.twice_preprocessed_file):
            ws = self._prediction_sheet(path)
            for row in (17, 18, 21):
                self.assertEqual(
                    _fill_argb(ws.cell(row, WARNING_COLUMN)),
                    GREEN_FILL_ARGB,
                    f"{path}: D{row} was repainted",
                )
            self.assertEqual(
                _fill_argb(ws.cell(19, WARNING_COLUMN)), RED_FILL_ARGB, path
            )
            self.assertEqual(ws.cell(20, WARNING_COLUMN).value, "N/A", path)

    def test_no_test_point_is_selected_in_the_group(self):
        for path in (self.preprocessed_file, self.twice_preprocessed_file):
            ws = self._prediction_sheet(path)
            marked = [
                row for row in GROUP_ROWS if ws.cell(row, WARNING_COLUMN).value == "X"
            ]
            self.assertEqual(marked, [], f"{path}: selected rows {marked}")

            # The rest of the sheet is all Green, so selection did run -- the
            # assertion above is about this group, not about an empty sheet.
            selected_elsewhere = sum(
                1
                for row in PREDICTION_ROWS
                for column in PREDICTION_COLUMNS
                if row not in GROUP_ROWS and ws.cell(row, column).value == "X"
            )
            self.assertGreater(selected_elsewhere, 0, path)

    def test_group_is_not_all_green_for_scoring(self):
        """compute_transient_score awards a Warning scenario only when every
        prediction cell of that strategy is GREEN, and it re-derives that from
        the preprocessed sheet through get_dm_colors_combined -- not from the
        selection-time rewrite. The OEM's Green cells come back Green, but the
        one Red they share the group with keeps the condition false, so the
        group still scores nothing."""
        dfs = common.read_excel_file_to_dfs(self.preprocessed_file, load_bg_colors=True)
        noise_variables_df = common.extract_table(
            dfs["DE - DM pred."], "Noise variables", 0
        )
        dm_colors = driver_monitoring.get_dm_colors_combined(
            dfs["DE - DM pred."],
            len(noise_variables_df),
            cell_colors_raw=dfs.get("DE - DM pred. (cell_colors)"),
        )

        # dm_colors is keyed by the pandas row index; Excel row N is index N-2.
        group_colors = [
            dm_colors[row - 2]["Warning"]
            for row in GROUP_ROWS
            if (row - 2) in dm_colors and "Warning" in dm_colors[row - 2]
        ]

        self.assertIn(common.PredictionColor.GREEN, group_colors)
        self.assertFalse(
            all(color == common.PredictionColor.GREEN for color in group_colors),
            f"group would score as all-GREEN: {group_colors}",
        )


if __name__ == "__main__":
    unittest.main()
