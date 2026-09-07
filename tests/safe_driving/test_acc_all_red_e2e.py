# Copyright 2025-2026, Euro NCAP IVZW
# Created by IVEX NV (https://ivex.ai)
#
# Licensed under the Apache License 2.0.
# See http://www.apache.org/licenses/LICENSE-2.0 for details.

"""End-to-end regression for the all-Red ACC scenario fix
through the real CLI pipeline.

The unit tests in test_acc_performance.py feed compute_score hand-built
dataframes. This test instead runs the packaged template through
`preprocess` -- which re-encodes the OEM's typed colour words as cell fills
and blanks the text -- and then `compute-score` on that output, the only
input shape the CLI actually accepts. A fix that only reads the text frame
passes every unit test and does nothing here.
"""

import shutil
import tempfile
import unittest
from importlib.resources import files
from pathlib import Path

import openpyxl
from click.testing import CliRunner

from euroncap_rating_2026 import common
from euroncap_rating_2026.safe_driving.compute_score import compute_score
from euroncap_rating_2026.safe_driving.preprocess import preprocess

# The four grey input cells of the CPLA grid in "VA - ACC pred." (the only
# cells carrying the "N/A,Green,Orange,Red" dropdown there); CBLA's are the
# four rows below and are left blank as the control.
CPLA_INPUT_CELLS = [f"I{row}" for row in range(103, 107)]
RED_FILL_ARGB = "00FF3333"


class TestAllRedCplaThroughPreprocess(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls._tmp = tempfile.TemporaryDirectory()
        root = Path(cls._tmp.name)
        input_dir, pre_dir, out_dir = root / "in", root / "pre", root / "out"
        for d in (input_dir, pre_dir, out_dir):
            d.mkdir()

        input_file = str(input_dir / "sd_template.xlsx")
        shutil.copy(str(files("data").joinpath("sd_template.xlsx")), input_file)
        common.add_version_sheet(
            input_file, cli_command=common.CliCommand.GENERATE_TEMPLATE
        )
        wb = openpyxl.load_workbook(input_file)
        # preprocess() needs "Number of rear seats"; everything else may stay
        # blank.
        ws = wb["Input parameters"]
        header = [c.value for c in ws[1]]
        param_col = header.index("Input parameter") + 1
        value_col = header.index("Value") + 1
        for row in range(2, ws.max_row + 1):
            if ws.cell(row, param_col).value == "Number of rear seats":
                ws.cell(row, value_col).value = 2
        ws = wb["VA - ACC pred."]
        for address in CPLA_INPUT_CELLS:
            ws[address] = "Red"
        wb.save(input_file)

        runner = CliRunner()
        result = runner.invoke(
            preprocess, ["--input_file", input_file, "--output_path", str(pre_dir)]
        )
        assert result.exit_code == 0, result.output
        cls.preprocessed_file = str(pre_dir / "sd_preprocessed_template.xlsx")

        result = runner.invoke(
            compute_score,
            ["--input_file", cls.preprocessed_file, "--output_path", str(out_dir)],
        )
        assert result.exit_code == 0, result.output
        cls.report_file = str(next(p for p in out_dir.iterdir() if "report" in p.name))

    @classmethod
    def tearDownClass(cls):
        cls._tmp.cleanup()

    def test_preprocess_moves_the_colour_into_the_cell_fill(self):
        """Pin the input shape compute-score has to cope with: no colour
        word left in the text, the Red fill is the only trace of it."""
        ws = openpyxl.load_workbook(self.preprocessed_file)["VA - ACC pred."]
        for address in CPLA_INPUT_CELLS:
            cell = ws[address]
            self.assertIn(cell.value, (None, ""), address)
            self.assertEqual(cell.fill.fgColor.rgb, RED_FILL_ARGB, address)

    def _scenario_scores(self):
        ws = openpyxl.load_workbook(self.report_file)["Scenario Scores"]
        header = [c.value for c in ws[1]]
        scenario_col, score_col = header.index("Scenario"), header.index("Score")
        return {
            row[scenario_col]: row[score_col]
            for row in ws.iter_rows(min_row=2, values_only=True)
            if row[scenario_col] in ("CPLA", "CBLA")
        }

    def test_all_red_cpla_scores_zero_in_the_report(self):
        self.assertEqual(self._scenario_scores()["CPLA"], 0)

    def test_unassessed_cbla_stays_blank_in_the_report(self):
        self.assertIsNone(self._scenario_scores()["CBLA"])


if __name__ == "__main__":
    unittest.main()
