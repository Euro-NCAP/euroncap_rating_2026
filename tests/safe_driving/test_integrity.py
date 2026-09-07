# Copyright 2025-2026, Euro NCAP IVZW
# Created by IVEX NV (https://ivex.ai)
#
# Licensed under the Apache License 2.0.
# See http://www.apache.org/licenses/LICENSE-2.0 for details.

import os
import tempfile
import unittest
from unittest.mock import patch

import openpyxl

from euroncap_rating_2026 import common, version
from euroncap_rating_2026.safe_driving import integrity


class TestRebuildTrustedInput(unittest.TestCase):
    def setUp(self):
        from importlib.resources import files

        self.pristine_path = str(files("data").joinpath("sd_template.xlsx"))

    def _save_wb(self, wb):
        fd, path = tempfile.mkstemp(suffix=".xlsx")
        os.close(fd)
        wb.save(path)
        self.addCleanup(os.remove, path)
        return path

    def test_max_score_tamper_is_neutralized(self):
        """Editing Max score in the user-supplied template must have zero
        effect on the sanitized output: it must reflect the official
        template's value instead."""
        wb = openpyxl.load_workbook(self.pristine_path)
        ws = wb["Test Scores"]
        header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        max_score_col = header.index("Max score") + 1
        original_max_score = ws.cell(row=2, column=max_score_col).value
        ws.cell(row=2, column=max_score_col).value = 99999
        user_path = self._save_wb(wb)

        sanitized_path = integrity.rebuild_trusted_input(user_path)
        self.addCleanup(os.remove, sanitized_path)

        sanitized_wb = openpyxl.load_workbook(sanitized_path, data_only=True)
        sanitized_value = (
            sanitized_wb["Test Scores"].cell(row=2, column=max_score_col).value
        )
        self.assertEqual(sanitized_value, original_max_score)
        self.assertNotEqual(sanitized_value, 99999)

    def test_input_parameters_value_is_preserved(self):
        """Legitimate grey-cell input must survive the rebuild unchanged."""
        wb = openpyxl.load_workbook(self.pristine_path)
        ws = wb["Input parameters"]
        header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        value_col = header.index("Value") + 1
        ws.cell(row=2, column=value_col).value = 3
        user_path = self._save_wb(wb)

        sanitized_path = integrity.rebuild_trusted_input(user_path)
        self.addCleanup(os.remove, sanitized_path)

        sanitized_wb = openpyxl.load_workbook(sanitized_path, data_only=True)
        self.assertEqual(
            sanitized_wb["Input parameters"].cell(row=2, column=value_col).value,
            3,
        )

    def test_reordered_rows_are_rejected(self):
        wb = openpyxl.load_workbook(self.pristine_path)
        ws = wb["Category Scores"]
        row2 = [ws.cell(row=2, column=c).value for c in range(1, 7)]
        row3 = [ws.cell(row=3, column=c).value for c in range(1, 7)]
        for c in range(1, 7):
            ws.cell(row=2, column=c).value = row3[c - 1]
            ws.cell(row=3, column=c).value = row2[c - 1]
        user_path = self._save_wb(wb)

        with self.assertRaises(common.TemplateIntegrityError):
            integrity.rebuild_trusted_input(user_path)

    def test_deleted_row_is_rejected(self):
        wb = openpyxl.load_workbook(self.pristine_path)
        ws = wb["Test Scores"]
        ws.delete_rows(3)
        user_path = self._save_wb(wb)

        with self.assertRaises(common.TemplateIntegrityError):
            integrity.rebuild_trusted_input(user_path)

    def test_scenario_scores_passthrough_is_untouched(self):
        """Scenario Scores isn't schema'd (its row structure is dynamic --
        see integrity.py) -- it must pass through from the user's file
        as-is, tamper and all."""
        wb = openpyxl.load_workbook(self.pristine_path)
        ws = wb["Scenario Scores"]
        header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        max_score_col = header.index("Max score") + 1
        ws.cell(row=2, column=max_score_col).value = 99999
        user_path = self._save_wb(wb)

        sanitized_path = integrity.rebuild_trusted_input(user_path)
        self.addCleanup(os.remove, sanitized_path)

        sanitized_wb = openpyxl.load_workbook(sanitized_path, data_only=True)
        self.assertEqual(
            sanitized_wb["Scenario Scores"].cell(row=2, column=max_score_col).value,
            99999,
        )

    def test_blank_input_parameter_from_pre_revision_template_is_tolerated(self):
        """Regression: a later template revision named "Input
        parameters" row 9's parameter ("System updates type", the
        SLIF - System updates row), which had been blank before. Every
        prediction file produced against the older template has that cell
        blank -- it must be tolerated as "unlabeled", not rejected as a
        reordered/inserted/deleted row."""
        wb = openpyxl.load_workbook(self.pristine_path)
        ws = wb["Input parameters"]
        header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        input_parameter_col = header.index("Input parameter") + 1
        self.assertEqual(
            ws.cell(row=9, column=input_parameter_col).value, "System updates type"
        )
        ws.cell(row=9, column=input_parameter_col).value = None
        user_path = self._save_wb(wb)

        sanitized_path = integrity.rebuild_trusted_input(user_path)
        self.addCleanup(os.remove, sanitized_path)

        sanitized_wb = openpyxl.load_workbook(sanitized_path, data_only=True)
        self.assertEqual(
            sanitized_wb["Input parameters"]
            .cell(row=9, column=input_parameter_col)
            .value,
            "System updates type",
        )


class TestRebuildTrustedInputDfs(unittest.TestCase):
    """DataFrame-native sibling of TestRebuildTrustedInput: same tamper
    scenarios, but through rebuild_trusted_input_dfs (used by
    get_updated_dfs/add_verification_sheets_to_dfs) instead of
    rebuild_trusted_input (used by the file-based preprocess() CLI
    command)."""

    def setUp(self):
        from importlib.resources import files

        pristine_path = str(files("data").joinpath("sd_template.xlsx"))
        self.pristine_dfs = common.read_excel_file_to_dfs(pristine_path)

    def _dfs(self):
        return {name: df.copy() for name, df in self.pristine_dfs.items()}

    def test_max_score_tamper_is_neutralized(self):
        dfs = self._dfs()
        original_max_score = dfs["Test Scores"].loc[0, "Max score"]
        dfs["Test Scores"].loc[0, "Max score"] = 99999

        sanitized = integrity.rebuild_trusted_input_dfs(dfs)

        self.assertEqual(
            sanitized["Test Scores"].loc[0, "Max score"], original_max_score
        )
        self.assertNotEqual(sanitized["Test Scores"].loc[0, "Max score"], 99999)

    def test_input_parameters_value_is_preserved(self):
        dfs = self._dfs()
        dfs["Input parameters"] = dfs["Input parameters"].astype(object)
        dfs["Input parameters"].loc[0, "Value"] = 3

        sanitized = integrity.rebuild_trusted_input_dfs(dfs)

        self.assertEqual(sanitized["Input parameters"].loc[0, "Value"], 3)

    def test_reordered_rows_are_rejected(self):
        dfs = self._dfs()
        df = dfs["Category Scores"]
        df.iloc[[0, 1]] = df.iloc[[1, 0]].values

        with self.assertRaises(common.TemplateIntegrityError):
            integrity.rebuild_trusted_input_dfs(dfs)

    def test_deleted_row_is_rejected(self):
        dfs = self._dfs()
        dfs["Test Scores"] = dfs["Test Scores"].drop(index=1).reset_index(drop=True)

        with self.assertRaises(common.TemplateIntegrityError):
            integrity.rebuild_trusted_input_dfs(dfs)

    def test_scenario_scores_passthrough_is_untouched(self):
        dfs = self._dfs()
        dfs["Scenario Scores"] = dfs["Scenario Scores"].astype(object)
        dfs["Scenario Scores"].loc[0, "Max score"] = 99999

        sanitized = integrity.rebuild_trusted_input_dfs(dfs)

        self.assertEqual(sanitized["Scenario Scores"].loc[0, "Max score"], 99999)

    def test_blank_input_parameter_from_pre_revision_template_is_tolerated(self):
        """DataFrame-path mirror of the openpyxl-path test of the same
        name -- this is the exact path taken when callers invoke
        get_updated_dfs -> rebuild_trusted_input_dfs directly with
        dataframes rather than through the file-based CLI."""
        dfs = self._dfs()
        dfs["Input parameters"] = dfs["Input parameters"].astype(object)
        self.assertEqual(
            dfs["Input parameters"].loc[7, "Input parameter"], "System updates type"
        )
        dfs["Input parameters"].loc[7, "Input parameter"] = None

        sanitized = integrity.rebuild_trusted_input_dfs(dfs)

        self.assertEqual(
            sanitized["Input parameters"].loc[7, "Input parameter"],
            "System updates type",
        )

    def test_sheet_missing_from_dfs_is_skipped_not_rejected(self):
        """get_updated_dfs/add_verification_sheets_to_dfs callers may
        legitimately pass a partial dfs dict; a schema'd sheet that isn't
        present must not raise (deviates from rebuild_trusted_input, which
        requires the complete monolithic file)."""
        dfs = {"Input parameters": self.pristine_dfs["Input parameters"].copy()}

        sanitized = integrity.rebuild_trusted_input_dfs(dfs)

        self.assertIn("Input parameters", sanitized)
        self.assertNotIn("Test Scores", sanitized)


class TestDocumentationDataSheetCoverage(unittest.TestCase):
    """Regression test for the "sharp edge" documented in integrity.py's
    module docstring: any pristine-template sheet not covered by
    SHEET_SCHEMAS or PASSTHROUGH_SHEETS is silently reset to blank by
    rebuild_trusted_workbook/rebuild_trusted_dfs -- this bit safe_driving
    once before (see integrity.py's docstring history)."""

    def setUp(self):
        from importlib.resources import files

        self.pristine_path = str(files("data").joinpath("sd_template.xlsx"))

    def test_every_pristine_sheet_is_covered(self):
        wb = openpyxl.load_workbook(self.pristine_path)
        covered = set(integrity.SHEET_SCHEMAS) | set(integrity.PASSTHROUGH_SHEETS)
        uncovered = set(wb.sheetnames) - covered
        self.assertEqual(uncovered, set())

    def test_documentation_and_data_sheets_are_hidden(self):
        wb = openpyxl.load_workbook(self.pristine_path)
        for name in ("Documentation", "Data"):
            self.assertIn(name, wb.sheetnames)
            self.assertEqual(wb[name].sheet_state, "hidden")

    def test_missing_input_check_covers_only_template_sheets(self):
        """find_missing_required_inputs must inspect prediction-template
        sheets only: every sheet it iterates (schemas and grid sheets) has
        to exist in the pristine template, so assessment-only sheets (the
        preprocess-generated '...verif.' sheets) can never produce
        findings."""
        wb = openpyxl.load_workbook(self.pristine_path)
        checked = set(integrity.SHEET_SCHEMAS) | set(
            integrity.GRID_SHEET_STAGE_ELEMENTS
        )
        self.assertEqual(checked - set(wb.sheetnames), set())


class TestFindMissingRequiredInputs(unittest.TestCase):
    """Exercises integrity.find_missing_required_inputs against the real
    packaged sd_template.xlsx, covering the schema'd "Input parameters"
    sheet plus the "DE - DM pred."/"VA - Speed assist. pred."/
    "VA - ACC pred." grid sheets (common.find_missing_grid_inputs)."""

    def setUp(self):
        from importlib.resources import files

        self.pristine_path = str(files("data").joinpath("sd_template.xlsx"))
        self.wb = openpyxl.load_workbook(self.pristine_path, data_only=True)

    def test_fresh_template_reports_grid_and_input_parameter_gaps(self):
        missing = integrity.find_missing_required_inputs(self.wb)

        self.assertTrue(missing)
        sheets = {m.sheet for m in missing}
        self.assertEqual(
            sheets,
            {
                "VA - ACC pred.",
                "VA - Speed assist. pred.",
                "DE - DM pred.",
                "Input parameters",
            },
        )

    def test_scoping_to_acc_performance_excludes_other_grid_sheets(self):
        scoped = integrity.find_missing_required_inputs(
            self.wb,
            stage_element="Vehicle Assistance",
            stage_subelement="ACC performance",
        )

        self.assertTrue(scoped)
        self.assertEqual(
            {m.sheet for m in scoped}, {"VA - ACC pred.", "Input parameters"}
        )
        for m in scoped:
            if m.sheet == "Input parameters":
                self.assertEqual(m.row_identity["Stage element"], "Vehicle Assistance")
                self.assertEqual(m.row_identity["Stage subelement"], "ACC performance")

    def test_scoping_to_driver_monitoring_only_reports_dm_pred(self):
        scoped = integrity.find_missing_required_inputs(
            self.wb,
            stage_element="Driver Engagement",
            stage_subelement="Driver monitoring",
        )

        self.assertTrue(scoped)
        self.assertEqual({m.sheet for m in scoped}, {"DE - DM pred."})

    def test_input_parameters_identity_does_not_leak_across_categories(self):
        """Regression for the real forward-fill bug found on this exact
        sheet: "Speed control function" must not inherit the unrelated
        "Child enters unlocked vehicle" scenario from a prior category."""
        missing = integrity.find_missing_required_inputs(self.wb)

        entry = next(
            m
            for m in missing
            if m.sheet == "Input parameters"
            and m.row_identity.get("Category") == "Speed control function"
        )
        self.assertIsNone(entry.row_identity.get("Scenario"))
        self.assertNotIn("Child enters unlocked vehicle", entry.description)

    def test_filling_dm_pred_cell_with_color_removes_it_from_missing(self):
        from openpyxl.styles import PatternFill

        ws = self.wb["DE - DM pred."]
        coords = sorted(common.get_required_dropdown_coordinates(ws))
        first_coord = coords[0]
        ws[first_coord].fill = PatternFill(
            start_color="FF009933", end_color="FF009933", fill_type="solid"
        )

        missing = integrity.find_missing_required_inputs(
            self.wb,
            stage_element="Driver Engagement",
            stage_subelement="Driver monitoring",
        )
        self.assertNotIn(first_coord, {m.cell for m in missing})

    def test_descriptions_are_semantic(self):
        missing = integrity.find_missing_required_inputs(
            self.wb,
            stage_element="Driver Engagement",
            stage_subelement="Driver monitoring",
        )

        entry = missing[0]
        self.assertNotRegex(entry.description, r"^[A-Z]+\d+$")
        self.assertTrue(entry.row_identity)


class TestVerifSheetsIgnored(unittest.TestCase):
    """The '...verif.' sheets exist only in preprocessed/assessment
    workbooks (preprocess generates them; sd_template.xlsx has none), and
    their Values are assessment-stage inputs filled by the assessor, not OEM
    predictions. find_missing_required_inputs covers prediction-template
    sheets only, so a blank (not yet assessed) verif Value must NOT be
    reported -- running the checks on an assessment workbook used to drown
    real findings in 'Value is missing' noise from these sheets."""

    def _make_wb(self, sheet_name, headers, rows):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_name
        ws.append(headers)
        for row in rows:
            ws.append(row)
        return wb

    def test_blank_verif_values_are_not_reported(self):
        wb = self._make_wb(
            "OM - Seatbelt usage verif.",
            ["Category", "Scenario", "Value"],
            [
                ["General requirements", None, None],
                [None, "Seatbelt buckle only", None],
            ],
        )
        for sheet_name in (
            "OM - Occ. classification verif.",
            "OM - Occ. presence verif.",
            "DE - GVC verif.",
        ):
            ws = wb.create_sheet(sheet_name)
            ws.append(["Category", "Scenario", "Value"])
            ws.append(["Some category", "Some scenario", None])

        self.assertEqual(integrity.find_missing_required_inputs(wb), [])

    def test_blank_verif_values_are_not_reported_when_scoped(self):
        wb = self._make_wb(
            "DE - GVC verif.",
            ["Category", "Scenario", "Value"],
            [["Driving", "Driving controls", None]],
        )

        scoped = integrity.find_missing_required_inputs(
            wb,
            stage_element="Driver Engagement",
            stage_subelement="General vehicle controls",
        )

        self.assertEqual(scoped, [])


class TestScopeValidation(unittest.TestCase):
    """An unknown (stage_element, stage_subelement) must raise instead of
    silently returning zero findings -- a typo'd caller-side mapping (e.g.
    snake_case identifiers) would otherwise silently disable the check.
    VALID_SCOPE_PAIRS must in turn cover every pair the packaged
    template's own sheets use, so no legitimate scope ever raises."""

    def setUp(self):
        from importlib.resources import files

        self.pristine_path = str(files("data").joinpath("sd_template.xlsx"))

    def test_unknown_scope_raises(self):
        wb = openpyxl.load_workbook(self.pristine_path, data_only=True)

        with self.assertRaises(ValueError):
            integrity.find_missing_required_inputs(
                wb,
                stage_element="vehicle_assistance",
                stage_subelement="acc_performance",
            )
        with self.assertRaises(ValueError):
            integrity.find_missing_required_inputs(wb, stage_element="Bogus")

    def test_declared_pairs_cover_the_templates_own_literals(self):
        wb = openpyxl.load_workbook(self.pristine_path, data_only=True)
        template_pairs = set()
        for sheet in ("Test Scores", "Input parameters"):
            ws = wb[sheet]
            header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
            element_col = header.index("Stage element") + 1
            subelement_col = header.index("Stage subelement") + 1
            last_element = last_subelement = None
            for row in range(2, ws.max_row + 1):
                element = ws.cell(row=row, column=element_col).value
                subelement = ws.cell(row=row, column=subelement_col).value
                if not common.is_empty_cell(element):
                    last_element = element
                    last_subelement = None
                if not common.is_empty_cell(subelement):
                    last_subelement = subelement
                if last_element is not None and last_subelement is not None:
                    template_pairs.add((last_element, last_subelement))

        self.assertTrue(template_pairs)
        undeclared = template_pairs - set(integrity.VALID_SCOPE_PAIRS)
        self.assertEqual(undeclared, set())


class TestPristineTemplateRegression(unittest.TestCase):
    """The audit assertion: on the pure packaged
    template every required cell is empty, so the unscoped check must
    report exactly the template's grey-painted input set -- no more (a
    false requirement fails valid submissions) and no less (a false pass
    ships an unchecked element). The literal per-sheet counts pin the
    required set; evolving the template legitimately means updating them
    in the same change. The fill cross-check then verifies each reported
    cell against the independent ground truth -- the grey paint the OEM
    actually sees."""

    EXPECTED_PER_SHEET = {
        "DE - DM pred.": 133,
        "VA - Speed assist. pred.": 328,
        "VA - ACC pred.": 72,
        "Input parameters": 11,
    }

    def setUp(self):
        from importlib.resources import files

        self.pristine_path = str(files("data").joinpath("sd_template.xlsx"))

    def test_unscoped_check_reports_exactly_the_expected_set(self):
        from collections import Counter

        wb = openpyxl.load_workbook(self.pristine_path, data_only=True)
        missing = integrity.find_missing_required_inputs(wb)

        self.assertEqual(
            dict(Counter(m.sheet for m in missing)), self.EXPECTED_PER_SHEET
        )

    def test_every_reported_cell_is_grey_painted_in_the_template(self):
        """Each reported coordinate must carry the grey input paint in the
        packaged template (a theme-shaded or literal-grey placeholder
        fill) -- a reported cell without it means the check requires
        something the template never asked the OEM to fill. Deliberately
        one-directional: some grey cells (Inspection [%], a criteria
        sheet's Value, non-sliding-scale OEM Predictions) are legitimately
        not required at prediction-upload time."""
        wb = openpyxl.load_workbook(self.pristine_path, data_only=True)
        missing = integrity.find_missing_required_inputs(wb)
        styled_wb = openpyxl.load_workbook(self.pristine_path)

        for entry in missing:
            fill = styled_wb[entry.sheet][entry.cell].fill
            self.assertEqual(
                fill.fill_type,
                "solid",
                f"{entry.sheet}!{entry.cell} reported but not painted",
            )
            if fill.fgColor.type == "theme":
                self.assertLess(
                    fill.fgColor.tint,
                    0,
                    f"{entry.sheet}!{entry.cell} has an unshaded theme fill",
                )
            else:
                grey_hex = str(fill.fgColor.rgb or "")[-6:].upper()
                self.assertIn(
                    grey_hex,
                    ("808080", "D9D9D9"),
                    f"{entry.sheet}!{entry.cell} reported but painted "
                    f"{fill.fgColor.rgb}",
                )


class TestFindMissingRequiredInputsVersionMismatch(unittest.TestCase):
    """A workbook generated by an incompatible library version must surface
    that drift as a finding, on top of -- not instead
    of -- the ordinary missing-input findings."""

    def setUp(self):
        from importlib.resources import files

        self.pristine_path = str(files("data").joinpath("sd_template.xlsx"))

    def _wb_with_version(self, version_string):
        wb = openpyxl.load_workbook(self.pristine_path, data_only=True)
        ws = wb["Version"] if "Version" in wb.sheetnames else wb.create_sheet("Version")
        ws["A1"] = "Generated with version"
        ws["B1"] = version_string
        return wb

    def test_incompatible_version_adds_finding_alongside_missing_inputs(self):
        with patch.object(common, "VERSION", "5.4.6"):
            wb = self._wb_with_version("5.4.0")
            findings = integrity.find_missing_required_inputs(wb)

        version_findings = [
            f for f in findings if f.check == common.CHECK_VERSION_MISMATCH
        ]
        self.assertEqual(len(version_findings), 1)
        self.assertTrue(any(f.check == common.CHECK_MISSING_INPUT for f in findings))

    def test_incompatible_version_surfaces_through_find_consistency_findings_too(self):
        with patch.object(common, "VERSION", "5.4.6"):
            wb = self._wb_with_version("5.4.0")
            findings = integrity.find_consistency_findings(wb)

        self.assertTrue(any(f.check == common.CHECK_VERSION_MISMATCH for f in findings))

    def test_version_finding_is_reported_even_when_scoped(self):
        """A version mismatch is a whole-workbook problem, not a per-element
        one -- it must not be dropped just because the caller scoped the
        check to one protocol element."""
        with patch.object(common, "VERSION", "5.4.6"):
            wb = self._wb_with_version("5.4.0")
            findings = integrity.find_missing_required_inputs(
                wb,
                stage_element="Occupant Monitoring",
                stage_subelement="Seatbelt usage",
            )

        self.assertTrue(any(f.check == common.CHECK_VERSION_MISMATCH for f in findings))

    def test_matching_version_adds_no_finding(self):
        wb = self._wb_with_version(version.VERSION)
        findings = integrity.find_missing_required_inputs(wb)

        self.assertFalse(
            any(f.check == common.CHECK_VERSION_MISMATCH for f in findings)
        )

    def test_no_version_sheet_adds_no_finding(self):
        """Regression guard: the many fixtures across this suite that never
        add a Version sheet at all must not start reporting a spurious
        mismatch."""
        wb = openpyxl.load_workbook(self.pristine_path, data_only=True)
        wb.remove(wb["Version"])
        findings = integrity.find_missing_required_inputs(wb)

        self.assertFalse(
            any(f.check == common.CHECK_VERSION_MISMATCH for f in findings)
        )

    def _rename_stage_subelement_header(self, wb):
        """Reproduces a real-world shape: a renamed identity
        column ("Stage Subelement" vs "Stage subelement") that makes the
        schema walk raise TemplateIntegrityError."""
        ws = wb["Input parameters"]
        header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        col = header.index("Stage subelement") + 1
        ws.cell(row=1, column=col).value = "Stage Subelement"

    def test_structurally_broken_and_version_mismatched_returns_only_version_finding(
        self,
    ):
        """The real-world case: an old-version file's renamed header would
        otherwise raise TemplateIntegrityError before any finding --
        including the version mismatch itself -- reaches the caller."""
        with patch.object(common, "VERSION", "5.4.6"):
            wb = self._wb_with_version("5.4.0")
            self._rename_stage_subelement_header(wb)
            findings = integrity.find_missing_required_inputs(wb)

        self.assertEqual(len(findings), 1)
        self.assertEqual(findings[0].check, common.CHECK_VERSION_MISMATCH)

    def test_structurally_broken_without_version_mismatch_still_raises(self):
        """Regression guard: with no version mismatch to blame, a
        genuinely broken template must still raise as before -- the
        schema-integrity contract predates and is independent of the version finding."""
        wb = openpyxl.load_workbook(self.pristine_path, data_only=True)
        self._rename_stage_subelement_header(wb)

        with self.assertRaises(common.TemplateIntegrityError):
            integrity.find_missing_required_inputs(wb)


if __name__ == "__main__":
    unittest.main()
