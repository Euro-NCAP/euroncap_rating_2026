# Copyright 2025-2026, Euro NCAP IVZW
# Created by IVEX NV (https://ivex.ai)
#
# Licensed under the Apache License 2.0.
# See http://www.apache.org/licenses/LICENSE-2.0 for details.

"""
The ACC matrix rows of the packaged Safe Driving template must carry the test
speeds of Safe Driving - Vehicle Assistance v1.2.

The cut-out rows lost their protocol speeds once: the pair was shifted one
column left when the target speed was set to 0, leaving the old target speeds
(50 / 70) in the VUT column. Nothing caught it --
scoring reads the cells rather than the labels -- while consuming applications
stopped creating test runs for both scenarios, because the rows no longer
matched their test spec. The protocol's own presentation is kept, so
both columns are pinned here and consistency.find_prediction_inconsistencies
reports a sheet that carries either superseded layout.
"""

import unittest
from importlib.resources import files

import openpyxl


class TestAccPredictionTemplateRows(unittest.TestCase):
    """
    Cut-out, per pp. 17 (Car-to-Car) and 21 (Car-to-Motorcyclist) of the
    protocol: VUT 70 and 90 km/h at TTC = 3.00s, with the lead vehicle (SOV)
    at 50 and 70 km/h in the "Target speed" column. The SOV is not the
    collision target -- it cuts out to reveal a standstill GVT/EMT -- but it is
    the speed the VUT is travelling at when that happens, which is why
    acc_performance scores cut-out against the "Target speed" column.
    """

    EXPECTED_CUT_OUT_ROWS = [
        ("70 km/h", "50 km/h", "3.00s"),
        ("90 km/h", "70 km/h", "3.00s"),
    ]

    @classmethod
    def setUpClass(cls):
        template_path = str(files("data").joinpath("sd_template.xlsx"))
        cls.wb = openpyxl.load_workbook(template_path)
        cls.ws = cls.wb["VA - ACC pred."]

    @classmethod
    def tearDownClass(cls):
        cls.wb.close()

    def _matrix_rows(self, scenario, row_count):
        """
        The (VUT speed, Target speed, TTC) triples of one matrix, located by
        its label in column A: the header row carries the labels, the row
        after it the impact locations, then the test points.
        """
        header_row = None
        for row in range(1, self.ws.max_row + 1):
            if self.ws.cell(row=row, column=1).value == scenario:
                header_row = row
                break
        self.assertIsNotNone(header_row, f"{scenario} matrix not found")
        first_point_row = header_row + 2
        return [
            tuple(self.ws.cell(row=row, column=col).value for col in (1, 2, 3))
            for row in range(first_point_row, first_point_row + row_count)
        ]

    def test_ccr_cut_out_vut_speeds(self):
        self.assertEqual(
            self._matrix_rows("CCR cut-out", 2), self.EXPECTED_CUT_OUT_ROWS
        )

    def test_cmr_cut_out_vut_speeds(self):
        self.assertEqual(
            self._matrix_rows("CMR cut-out", 2), self.EXPECTED_CUT_OUT_ROWS
        )

    def test_cut_in_rows_are_unchanged(self):
        """The cut-in rows were never shifted; pin them so a future edit of
        the neighbouring matrices cannot move them either."""
        self.assertEqual(
            self._matrix_rows("CCR cut-in", 2),
            [("50 km/h", "10 km/h", "0.00s"), ("120 km/h", "70 km/h", "1.50s")],
        )
        self.assertEqual(
            self._matrix_rows("CMR cut-in", 2),
            [("50 km/h", "10 km/h", "0.50s"), ("120 km/h", "70 km/h", "1.50s")],
        )


if __name__ == "__main__":
    unittest.main()
