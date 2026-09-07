# Copyright 2025-2026, Euro NCAP IVZW
# Created by IVEX NV (https://ivex.ai)
#
# Licensed under the Apache License 2.0.
# See http://www.apache.org/licenses/LICENSE-2.0 for details.

import shutil
import tempfile
import unittest
from importlib.resources import files
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Font, PatternFill

from euroncap_rating_2026 import common
from euroncap_rating_2026.safe_driving import report_writer as sd_report_writer

HEADER_FILL = PatternFill(start_color="000000", end_color="000000", fill_type="solid")
HEADER_FONT = Font(name="Calibri", color="FFFFFF", bold=True)


class TestDmVerifHeaderRowGuard(unittest.TestCase):
    """
    Reproduce the "DE - DM verif." row-4 header styling guard from
    report_writer.write_report() inline so the assertion is deterministic
    without real Excel file I/O (same pattern as
    tests/crash_protection/test_vru_blue_legform.py::TestValueColumnSAStyling).
    """

    def _apply_guard(self, ws, sheet_name):
        if sheet_name == "DE - DM verif." and ws.max_row >= 4 and ws.max_column > 1:
            common.set_header_by_row(ws, 4, HEADER_FILL, HEADER_FONT)

    def _build_short_ws(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "DE - DM verif."
        ws.cell(row=1, column=1).value = "Category"
        ws.cell(row=2, column=1).value = (
            "Mandatory noise variable Face-mask was set to Non functional."
        )
        return ws

    def _build_long_message_ws(self):
        # The single-column "Category" message table driver_monitoring.preprocess
        # writes when several mandatory noise variables fail -- has >= 4 rows
        # (one per failed variable) but only one column.
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "DE - DM verif."
        ws.cell(row=1, column=1).value = "Category"
        messages = [
            "Mandatory noise variable Daytime - Night-time was set to Non functional.",
            "Mandatory noise variable Clear sunglasses was set to Non functional.",
            "Mandatory noise variable Short facial hair was set to Non functional.",
            "Mandatory noise variable Long facial hair was set to Non functional.",
        ]
        for i, message in enumerate(messages, start=2):
            ws.cell(row=i, column=1).value = message
        return ws

    def _build_full_ws(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "DE - DM verif."
        ws.cell(row=1, column=1).value = "General requirements"
        ws.cell(row=2, column=1).value = ""
        ws.cell(row=3, column=1).value = "Category"
        ws.cell(row=3, column=2).value = "Scenario"
        ws.cell(row=4, column=1).value = "Transient - Long distraction"
        ws.cell(row=4, column=2).value = "Non-driving task - Owl - Warning"
        return ws

    def test_guard_skips_short_message_sheet(self):
        ws = self._build_short_ws()
        self.assertEqual(ws.max_row, 2)

        self._apply_guard(ws, "DE - DM verif.")

        self.assertEqual(ws.max_row, 2)
        cell_a4 = ws.cell(row=4, column=1)
        self.assertIsNone(cell_a4.value)
        self.assertIsNone(cell_a4.fill.patternType)

    def test_guard_skips_long_single_column_message_sheet(self):
        ws = self._build_long_message_ws()
        self.assertGreaterEqual(ws.max_row, 4)
        self.assertEqual(ws.max_column, 1)

        self._apply_guard(ws, "DE - DM verif.")

        cell_a4 = ws.cell(row=4, column=1)
        self.assertIsNone(cell_a4.fill.patternType)
        self.assertFalse(cell_a4.font.bold)

    def test_guard_still_applies_for_full_sheet(self):
        ws = self._build_full_ws()
        self.assertGreaterEqual(ws.max_row, 4)

        self._apply_guard(ws, "DE - DM verif.")

        for cell in ws[4]:
            self.assertEqual(cell.fill.patternType, "solid")
            self.assertTrue(cell.font.bold)


class TestDocumentationDataHardCopyAndHide(unittest.TestCase):
    """safe_driving's preprocess() requires a filled-in "Number of rear
    seats" input parameter (it errors out on a fully blank template -- a
    pre-existing, unrelated limitation), so a full CLI preprocess()->
    compute_score() e2e run isn't viable here the way it is for
    crash_avoidance/post_crash. Instead, exercise report_writer.write_report
    directly -- the actual code that hard-copies and hides the
    "Documentation"/"Data" sheets -- against the pristine template."""

    def test_preprocess_output_has_hidden_documentation_data_sheets(self):
        pristine_path = str(files("data").joinpath("sd_template.xlsx"))
        with (
            tempfile.TemporaryDirectory() as input_dir,
            tempfile.TemporaryDirectory() as output_dir,
        ):
            input_file = str(Path(input_dir) / "sd_template.xlsx")
            shutil.copy(pristine_path, input_file)
            common.add_version_sheet(
                input_file, cli_command=common.CliCommand.GENERATE_TEMPLATE
            )

            sd_report_writer.write_report(
                common.CliCommand.PREPROCESS,
                input_file,
                updated_dfs={},
                output_path=output_dir,
                format_prediction_cells=False,
            )

            output_file = Path(output_dir) / "sd_preprocessed_template.xlsx"
            wb = openpyxl.load_workbook(str(output_file))
            for sheet_name, id_header in [
                ("Documentation", "Documentation ID"),
                ("Data", "Data ID"),
            ]:
                ws = wb[sheet_name]
                self.assertEqual(ws.sheet_state, "hidden")
                self.assertEqual([c.value for c in ws[1]], [id_header, "Status"])


class _WriteReportTestBase(unittest.TestCase):
    """Shared helpers for tests that drive report_writer.write_report on
    small fixture DataFrames and inspect the written workbook."""

    def _write_report(self, updated_dfs):
        pristine_path = str(files("data").joinpath("sd_template.xlsx"))
        with (
            tempfile.TemporaryDirectory() as input_dir,
            tempfile.TemporaryDirectory() as output_dir,
        ):
            input_file = str(Path(input_dir) / "sd_template.xlsx")
            shutil.copy(pristine_path, input_file)
            common.add_version_sheet(
                input_file, cli_command=common.CliCommand.GENERATE_TEMPLATE
            )

            sd_report_writer.write_report(
                common.CliCommand.PREPROCESS,
                input_file,
                updated_dfs=updated_dfs,
                output_path=output_dir,
                format_prediction_cells=False,
            )

            output_file = Path(output_dir) / "sd_preprocessed_template.xlsx"
            return openpyxl.load_workbook(str(output_file))

    def _dv_for_cell(self, ws, cell):
        for dv in ws.data_validations.dataValidation:
            if cell.coordinate in dv.sqref:
                return dv
        return None


class TestPassFailDropdownAllowsBlank(_WriteReportTestBase):
    """A blank Value cell means "not yet assessed" and must still get the
    PASS/FAIL dropdown with allow_blank=True -- not just cells that already
    say PASS/FAIL, and not (for VA - ACC verif.) the sheet's unrelated
    numeric ACC-matrix rows."""

    def test_blank_value_cell_gets_dropdown_with_allow_blank(self):
        df = pd.DataFrame(
            [
                {"Category": "Driving", "Scenario": "Driving controls", "Value": None},
                {"Category": "", "Scenario": "Vision", "Value": "PASS"},
                {"Category": "", "Scenario": "Lights", "Value": "FAIL"},
            ]
        )
        wb = self._write_report({"DE - GVC verif.": df})
        ws = wb["DE - GVC verif."]

        blank_cell = ws.cell(row=2, column=3)
        pass_cell = ws.cell(row=3, column=3)
        fail_cell = ws.cell(row=4, column=3)

        for cell in (blank_cell, pass_cell, fail_cell):
            dv = self._dv_for_cell(ws, cell)
            self.assertIsNotNone(dv, f"expected a dropdown on {cell.coordinate}")
            self.assertTrue(dv.allow_blank)
            self.assertEqual(dv.formula1, '"PASS,FAIL"')

    def test_acc_matrix_row_blank_value_does_not_get_pass_fail_dropdown(self):
        """VA - ACC verif. mixes PASS/FAIL rows (Road features/Auto-resume)
        with numeric matrix rows in the same Value column -- a blank matrix
        row must not be caught by the PASS/FAIL blank-cell allow-list."""
        df = pd.DataFrame(
            [
                {"Category": "Road features", "Scenario": "Curves", "Value": None},
                {"Category": "CCRs", "Scenario": "CCRs straight", "Value": None},
            ]
        )
        wb = self._write_report({"VA - ACC verif.": df})
        ws = wb["VA - ACC verif."]

        road_features_cell = ws.cell(row=2, column=3)
        matrix_cell = ws.cell(row=3, column=3)

        self.assertIsNotNone(self._dv_for_cell(ws, road_features_cell))
        self.assertIsNone(self._dv_for_cell(ws, matrix_cell))


class TestNaValueRowsWhiteAndNoDropdown(_WriteReportTestBase):
    """A row whose Scenario is "N/A" (from an N/A Input parameters dropdown
    selection) needs no user input: its Value cell must stay white (no grey
    input fill) and must not get the PASS/FAIL dropdown, while sibling rows
    keep both."""

    def _is_grey_input_cell(self, cell):
        return cell.fill.patternType == "solid" and str(
            cell.fill.start_color.rgb
        ).endswith(common.INPUT_CELL_RGB)

    def test_na_passenger_airbag_row_is_white_without_dropdown(self):
        df = pd.DataFrame(
            [
                {
                    "Category": "Passenger airbag status",
                    "Scenario": "N/A",
                    "Value": None,
                },
                {
                    "Category": "Out of position",
                    "Scenario": "Close proximity to the airbag",
                    "Value": None,
                },
            ]
        )
        wb = self._write_report({"OM - Occ. classification verif.": df})
        ws = wb["OM - Occ. classification verif."]

        na_cell = ws.cell(row=2, column=3)
        sibling_cell = ws.cell(row=3, column=3)

        self.assertFalse(self._is_grey_input_cell(na_cell))
        self.assertIsNone(self._dv_for_cell(ws, na_cell))
        self.assertTrue(self._is_grey_input_cell(sibling_cell))
        self.assertIsNotNone(self._dv_for_cell(ws, sibling_cell))

    def test_real_passenger_airbag_row_keeps_grey_and_dropdown(self):
        df = pd.DataFrame(
            [
                {
                    "Category": "Passenger airbag status",
                    "Scenario": "Automatic",
                    "Value": None,
                },
            ]
        )
        wb = self._write_report({"OM - Occ. classification verif.": df})
        ws = wb["OM - Occ. classification verif."]

        cell = ws.cell(row=2, column=3)

        self.assertTrue(self._is_grey_input_cell(cell))
        self.assertIsNotNone(self._dv_for_cell(ws, cell))

    def test_na_auto_resume_row_is_white_without_dropdown(self):
        df = pd.DataFrame(
            [
                {"Category": "Road features", "Scenario": "Curves", "Value": None},
                {"Category": "Auto-resume", "Scenario": "N/A", "Value": None},
            ]
        )
        wb = self._write_report({"VA - ACC verif.": df})
        ws = wb["VA - ACC verif."]

        road_features_cell = ws.cell(row=2, column=3)
        na_cell = ws.cell(row=3, column=3)

        self.assertTrue(self._is_grey_input_cell(road_features_cell))
        self.assertIsNotNone(self._dv_for_cell(ws, road_features_cell))
        self.assertFalse(self._is_grey_input_cell(na_cell))
        self.assertIsNone(self._dv_for_cell(ws, na_cell))


if __name__ == "__main__":
    unittest.main()
