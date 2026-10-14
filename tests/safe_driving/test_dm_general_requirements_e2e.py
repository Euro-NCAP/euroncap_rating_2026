# Copyright 2025-2026, Euro NCAP IVZW
# Created by IVEX NV (https://ivex.ai)
#
# Licensed under the Apache License 2.0.
# See http://www.apache.org/licenses/LICENSE-2.0 for details.

"""Driver monitoring General requirements, one PASS/FAIL row per requirement.

The Driver Engagement protocol lists several conditions a DSM system must meet
before it is eligible for any Driver monitoring points. The preprocessed
"DE - DM verif." sheet carries one grey row per condition, and scoring gates
on all of them: any blank leaves Driver monitoring unassessed, any FAIL scores
it 0. The non-functional-system information conditions also accept N/A, which
counts as a pass; the others do not accept it. A preprocessed file from before the split, with a single
"General requirements" row, must keep scoring as it did.
"""

import shutil
import tempfile
import unittest
from importlib.resources import files
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
from click.testing import CliRunner

from euroncap_rating_2026 import common
from euroncap_rating_2026.safe_driving import data_model, driver_monitoring
from euroncap_rating_2026.safe_driving.preprocess import preprocess

SHEET = "DE - DM verif."
# Rows 2-8 of the sheet: PASS / FAIL only.
PASS_FAIL_REQUIREMENTS = [
    "Default ON at the start of every journey",
    "Deactivation is not possible with a momentary single push on a button",
    "Detection sensitivity shall not be manually adjustable",
    "Continuous monitoring of the driver state",
    "Speed activation condition for the distraction, microsleep, sleep and "
    "unresponsive driver states",
    "Speed activation condition for the impairment driver states, when the "
    "vehicle is moving forward at a speed of at least 50 km/h",
    "Assisted Driving prerequisite (≥50% of Driving collaboration and ≥50% of "
    "Driver monitoring)",
]
# Rows 9-15: non-functional-system information, PASS / FAIL / N/A.
PASS_FAIL_NA_REQUIREMENTS = [
    "Inform with visual and/or auditory cues of non functional system within "
    "10 seconds after the driver obstructs the sensor by gripping the upper "
    "portion of the steering wheel",
    "Inform with visual and/or auditory cues of non functional system within "
    "10 seconds after the driver wears a cap",
    "Inform with visual and/or auditory cues of non functional system within "
    "10 seconds after the driver wears a hat",
    "Inform with visual and/or auditory cues of non functional system within "
    "10 seconds after the driver wears dark sunglasses",
    "Inform with visual and/or auditory cues of non functional system within "
    "10 seconds for a driver with long head hair fringe obscuring eyes",
    "Inform with visual and/or auditory cues of non functional system within "
    "10 seconds for a driver with thick eyelash makeup",
    "Visual information about non-functional system shown permanently or "
    "every [10] minutes",
]
GENERAL_REQUIREMENTS = PASS_FAIL_REQUIREMENTS + PASS_FAIL_NA_REQUIREMENTS
# Worksheet rows of the requirements (row 1 is the header), then a blank
# separator row, then the test-point table header.
REQUIREMENT_ROWS = range(2, 2 + len(GENERAL_REQUIREMENTS))
PASS_FAIL_NA_ROWS = REQUIREMENT_ROWS[len(PASS_FAIL_REQUIREMENTS) :]
TEST_POINT_HEADER_ROW = REQUIREMENT_ROWS.stop + 1


def _fill_template(path):
    """Every mandatory noise variable Functional and every prediction cell
    Green, so that test points are selected and can score."""
    common.add_version_sheet(path, cli_command=common.CliCommand.GENERATE_TEMPLATE)
    wb = openpyxl.load_workbook(path)

    ws = wb["Input parameters"]
    header = [c.value for c in ws[1]]
    param_col = header.index("Input parameter") + 1
    value_col = header.index("Value") + 1
    for row in range(2, ws.max_row + 1):
        if ws.cell(row, param_col).value == "Number of rear seats":
            ws.cell(row, value_col).value = 2

    ws = wb["DE - DM pred."]
    for row in range(2, 13):
        ws.cell(row, 4).value = "Functional"
    for row in range(17, 72):
        for column in (4, 5, 6):
            cell = ws.cell(row, column)
            if isinstance(cell, openpyxl.cell.cell.MergedCell):
                continue
            cell.value = "Green"
    wb.save(path)


class TestGeneralRequirementsThroughPreprocess(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls._tmp = tempfile.TemporaryDirectory()
        root = Path(cls._tmp.name)
        (root / "in").mkdir()
        (root / "pre").mkdir()
        input_file = str(root / "in" / "sd_template.xlsx")
        shutil.copy(str(files("data").joinpath("sd_template.xlsx")), input_file)
        _fill_template(input_file)

        result = CliRunner().invoke(
            preprocess,
            ["--input_file", input_file, "--output_path", str(root / "pre")],
        )
        assert result.exit_code == 0, result.output
        cls.preprocessed_file = str(root / "pre" / "sd_preprocessed_template.xlsx")
        cls.dfs = common.read_excel_file_to_dfs(
            cls.preprocessed_file, load_bg_colors=True
        )

    @classmethod
    def tearDownClass(cls):
        cls._tmp.cleanup()

    def test_one_grey_pass_fail_row_per_requirement(self):
        wb = openpyxl.load_workbook(self.preprocessed_file)
        ws = wb[SHEET]
        header = [c.value for c in ws[1]]
        value_col = header.index("Value") + 1

        self.assertEqual(ws.cell(2, 1).value, "General requirements")
        self.assertEqual(
            [ws.cell(row, 2).value for row in REQUIREMENT_ROWS],
            GENERAL_REQUIREMENTS,
        )
        for row in REQUIREMENT_ROWS:
            cell = ws.cell(row, value_col)
            self.assertIsNone(cell.value)
            self.assertTrue(
                str(cell.fill.start_color.rgb).endswith(common.INPUT_CELL_RGB),
                f"{cell.coordinate} is not a grey input cell",
            )
            dropdowns = [
                validation.formula1
                for validation in ws.data_validations.dataValidation
                if cell.coordinate in validation.sqref
            ]
            expected = '"PASS,FAIL,N/A"' if row in PASS_FAIL_NA_ROWS else '"PASS,FAIL"'
            self.assertEqual(dropdowns, [expected], cell.coordinate)

        self.assertEqual(ws.cell(TEST_POINT_HEADER_ROW, 1).value, "Category")
        self.assertTrue(ws.cell(TEST_POINT_HEADER_ROW, 1).font.bold)

    def _verification_df(self, general_requirement_values):
        """The preprocessed verification sheet with every selected test point
        PASS and the given General requirement Values."""
        df = self.dfs[SHEET].copy()
        header_idx = np.flatnonzero((df["Category"] == "Category").to_numpy())[0]
        for idx in range(header_idx + 1, len(df)):
            if not common.is_empty_cell(df.iloc[idx]["Category"]):
                df.loc[df.index[idx], "Value"] = "PASS"
        for idx, value in enumerate(general_requirement_values):
            df.loc[df.index[idx], "Value"] = value
        return df

    def _score(self, verification_df):
        dfs = dict(self.dfs)
        dfs[SHEET] = verification_df
        return driver_monitoring.compute_score(dfs)

    def test_all_pass_scores_like_single_general_requirements_row(self):
        new_layout = self._verification_df(["PASS"] * len(GENERAL_REQUIREMENTS))
        # A file preprocessed before the split: one "General requirements"
        # row, then the blank separator and the test-point table.
        legacy_layout = new_layout.drop(
            new_layout.index[1 : len(GENERAL_REQUIREMENTS)]
        ).reset_index(drop=True)

        new_scores = self._score(new_layout)
        legacy_scores = self._score(legacy_layout)

        self.assertGreater(new_scores["Score"].sum(), 0)
        pd.testing.assert_frame_equal(new_scores, legacy_scores)

    def test_one_fail_scores_zero(self):
        values = ["PASS"] * len(GENERAL_REQUIREMENTS)
        values[3] = "FAIL"
        scores = self._score(self._verification_df(values))
        self.assertTrue((scores["Score"] == 0.0).all())

    def test_one_blank_leaves_driver_monitoring_unassessed(self):
        values = ["PASS"] * len(GENERAL_REQUIREMENTS)
        values[-1] = None
        scores = self._score(self._verification_df(values))
        self.assertTrue(scores["Score"].isna().all())

    def test_na_counts_as_pass_where_allowed(self):
        values = ["PASS"] * len(PASS_FAIL_REQUIREMENTS) + ["N/A"] * len(
            PASS_FAIL_NA_REQUIREMENTS
        )
        na_scores = self._score(self._verification_df(values))
        pass_scores = self._score(
            self._verification_df(["PASS"] * len(GENERAL_REQUIREMENTS))
        )
        self.assertGreater(na_scores["Score"].sum(), 0)
        pd.testing.assert_frame_equal(na_scores, pass_scores)

    def test_na_not_accepted_on_pass_fail_requirement(self):
        values = ["PASS"] * len(GENERAL_REQUIREMENTS)
        values[0] = "N/A"
        scores = self._score(self._verification_df(values))
        self.assertTrue(scores["Score"].isna().all())

    def test_stray_na_in_test_point_value_stays_blank(self):
        df = self._verification_df(["PASS"] * len(GENERAL_REQUIREMENTS))
        header_idx = np.flatnonzero((df["Category"] == "Category").to_numpy())[0]
        first_point = df.index[header_idx + 1]
        na_df = df.copy()
        na_df.loc[first_point, "Value"] = "N/A"
        blank_df = df.copy()
        blank_df.loc[first_point, "Value"] = np.nan
        pd.testing.assert_frame_equal(self._score(na_df), self._score(blank_df))

    def test_na_typed_in_file_survives_reading(self):
        """N/A picked from the dropdown in the preprocessed file reaches
        scoring as N/A, not as a blank cell."""
        wb = openpyxl.load_workbook(self.preprocessed_file)
        ws = wb[SHEET]
        header = [c.value for c in ws[1]]
        value_col = header.index("Value") + 1
        for row in REQUIREMENT_ROWS:
            ws.cell(row, value_col).value = (
                "N/A" if row in PASS_FAIL_NA_ROWS else "PASS"
            )
        for row in range(TEST_POINT_HEADER_ROW + 1, ws.max_row + 1):
            if ws.cell(row, 1).value:
                ws.cell(row, value_col).value = "PASS"
        with tempfile.TemporaryDirectory() as tmp:
            path = str(Path(tmp) / "filled.xlsx")
            wb.save(path)
            dfs = common.read_excel_file_to_dfs(
                path,
                load_bg_colors=True,
                preserve_na_sheets=data_model.PRESERVE_NA_SHEETS,
            )
        scores = driver_monitoring.compute_score(dfs)
        self.assertGreater(scores["Score"].sum(), 0)


if __name__ == "__main__":
    unittest.main()
