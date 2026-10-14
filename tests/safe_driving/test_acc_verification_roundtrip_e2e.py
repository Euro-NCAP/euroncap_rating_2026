# Copyright 2025-2026, Euro NCAP IVZW
# Created by IVEX NV (https://ivex.ai)
#
# Licensed under the Apache License 2.0.
# See http://www.apache.org/licenses/LICENSE-2.0 for details.

"""End-to-end check that ACC matrix points score through the real CLI
pipeline, on a realistic one-scenario assessment.

`preprocess` writes "VA - ACC verif." in the two-block report layout whose
matrix block carries its own header row inside the sheet. Read back with a
single header row, the matrix columns lose their names, so a scorer that
looks the fields up by name would see no OEM Prediction and stamp every
point Red. The unit tests in test_acc_performance.py feed compute_score
in-memory frames that already have the right column names and cannot see
this; this test runs the packaged template through `preprocess`, fills the
measured values in the written file, and runs `compute-score` on that --
the only input shape the CLI accepts.
"""

import shutil
import tempfile
import unittest
from importlib.resources import files
from pathlib import Path

import openpyxl
from click.testing import CliRunner

from euroncap_rating_2026 import common
from euroncap_rating_2026.safe_driving import acc_performance
from euroncap_rating_2026.safe_driving.compute_score import compute_score
from euroncap_rating_2026.safe_driving.preprocess import preprocess

# CCRs straight grid in "VA - ACC pred.": rows 3/4 are VUT 60/70 km/h against
# a stationary target, columns G/H/I are impact locations 50%/25%/0%.
# Test point label -> (prediction cell, prediction, measured v_reduction,
# expected colour, expected score). An 8-row matrix gives 0.125 per point.
# Green threshold is v_VUT - v_target - 2 (58 / 68); orange is > 15.
CASES = {
    "(2, 4)": ("G3", "Green", 60.0, "Green", 0.125),  # full avoidance
    "(2, 5)": ("H3", "Green", 20.0, "Orange", 0.0625),
    "(2, 6)": ("I3", "Orange", 59.0, "Green", 0.125),
    "(3, 4)": ("G4", "Orange", 10.0, "Red", 0.0),
    "(3, 5)": ("H4", "Green", 12.0, "Red", 0.0),
    "(3, 6)": ("I4", "Orange", 30.0, "Orange", 0.0625),
}
EXPECTED_SCENARIO_SCORE = 0.375


class TestAccMatrixRoundTrip(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls._tmp = tempfile.TemporaryDirectory()
        root = Path(cls._tmp.name)
        input_dir, pre_dir, out_dir = root / "in", root / "pre", root / "out"
        for d in (input_dir, pre_dir, out_dir):
            d.mkdir()

        input_file = str(input_dir / "sd_template.xlsx")
        shutil.copy(str(files("data").joinpath("sd_template.xlsx")), input_file)
        common.add_version_sheet(
            input_file, cli_command=common.CliCommand.GENERATE_TEMPLATE
        )
        wb = openpyxl.load_workbook(input_file)
        ws = wb["Input parameters"]
        header = [c.value for c in ws[1]]
        param_col = header.index("Input parameter") + 1
        value_col = header.index("Value") + 1
        for row in range(2, ws.max_row + 1):
            if ws.cell(row, param_col).value == "Number of rear seats":
                ws.cell(row, value_col).value = 2
        ws = wb["VA - ACC pred."]
        for address, prediction, *_ in CASES.values():
            ws[address] = prediction
        wb.save(input_file)

        runner = CliRunner()
        result = runner.invoke(
            preprocess, ["--input_file", input_file, "--output_path", str(pre_dir)]
        )
        assert result.exit_code == 0, result.output
        cls.preprocessed_file = str(pre_dir / "sd_preprocessed_template.xlsx")

        # Fill the measured v_reduction of each matrix row, located by its
        # "Test point" label under the interior header row.
        wb = openpyxl.load_workbook(cls.preprocessed_file)
        ws = wb["VA - ACC verif."]
        test_point_col, value_col = cls._matrix_columns(ws)
        for row in range(cls._interior_header_row(ws) + 1, ws.max_row + 1):
            case = CASES.get(ws.cell(row, test_point_col).value)
            if case is not None:
                ws.cell(row, value_col).value = case[2]
        wb.save(cls.preprocessed_file)

        result = runner.invoke(
            compute_score,
            ["--input_file", cls.preprocessed_file, "--output_path", str(out_dir)],
        )
        assert result.exit_code == 0, result.output
        cls.report_file = str(next(p for p in out_dir.iterdir() if "report" in p.name))

    @classmethod
    def tearDownClass(cls):
        cls._tmp.cleanup()

    @staticmethod
    def _interior_header_row(ws) -> int:
        for row in range(2, ws.max_row + 1):
            if ws.cell(row, 3).value == "Test point":
                return row
        raise AssertionError("no interior header row in 'VA - ACC verif.'")

    @classmethod
    def _matrix_columns(cls, ws):
        header = [c.value for c in ws[cls._interior_header_row(ws)]]
        return header.index("Test point") + 1, header.index("Value") + 1

    def _report_colours(self):
        ws = openpyxl.load_workbook(self.report_file)["VA - ACC verif."]
        test_point_col, value_col = self._matrix_columns(ws)
        # The computed colour is written right after "Value" under a blank
        # header (see acc_performance.compute_score's "Color" -> "" rename).
        return {
            ws.cell(row, test_point_col).value: ws.cell(row, value_col + 1).value
            for row in range(self._interior_header_row(ws) + 1, ws.max_row + 1)
        }

    def test_each_matrix_row_is_coloured_from_its_prediction_and_value(self):
        colours = self._report_colours()
        for test_point, (_, prediction, value, expected, _) in CASES.items():
            with self.subTest(test_point=test_point, prediction=prediction, v=value):
                self.assertEqual(colours[test_point], expected)

    def test_scenario_score_sums_the_matrix_points(self):
        ws = openpyxl.load_workbook(self.report_file)["Scenario Scores"]
        header = [c.value for c in ws[1]]
        scenario_col, score_col = header.index("Scenario"), header.index("Score")
        scores = {
            row[scenario_col]: row[score_col]
            for row in ws.iter_rows(min_row=2, values_only=True)
        }
        self.assertAlmostEqual(scores["CCRs straight"], EXPECTED_SCENARIO_SCORE)

    def test_reread_sheet_binds_the_matrix_fields(self):
        """The frame `compute-score` gets from disk names the matrix columns
        'Unnamed: N'; the binder must recover the fields the scorer reads."""
        dfs = common.read_excel_file_to_dfs(self.preprocessed_file)
        raw_df = dfs["VA - ACC verif."]
        self.assertNotIn("OEM Prediction", raw_df.columns)
        bound = acc_performance._bind_matrix_header_columns(raw_df)
        for field in ("OEM Prediction", "VUT speed", "Target speed", "Test point"):
            self.assertIn(field, bound)
        self.assertEqual(bound[0], "Category")
        self.assertEqual(bound[-1], "Value")

    def test_report_keeps_the_two_block_layout(self):
        """The rebind is internal to scoring: the written sheet still has the
        blank top header over the matrix columns and the interior header."""
        ws = openpyxl.load_workbook(self.report_file)["VA - ACC verif."]
        self.assertEqual(
            [ws.cell(1, c).value for c in range(1, 10)],
            ["Category", "Scenario", None, None, None, None, None, None, "Value"],
        )
        self.assertEqual(self._interior_header_row(ws), 9)


if __name__ == "__main__":
    unittest.main()
