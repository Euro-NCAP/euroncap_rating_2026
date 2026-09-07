# Copyright 2025-2026, Euro NCAP IVZW
# Created by IVEX NV (https://ivex.ai)
#
# Licensed under the Apache License 2.0.
# See http://www.apache.org/licenses/LICENSE-2.0 for details.

"""Grey-cell trust boundary for the safe_driving domain.

Declares, per sheet, which columns are legitimate OEM input ("grey" cells --
Value) versus protected reference/structural data (Max score, row labels)
that must always come from the official packaged template rather than a
user-supplied file. See `euroncap_rating_2026.common` for the underlying
mechanism.

NOTE: several sheet types are not covered by a schema here, left as a
follow-up if this mechanism is generalized further:
  - "Scenario Scores" is deliberately NOT schema'd here, unlike the other
    two domains, and is listed in PASSTHROUGH_SHEETS instead so the
    file-based rebuild copies it verbatim from the user's file rather than
    silently overwriting it with the (structurally different) pristine
    template's placeholder rows. Its row structure is not static:
    seatbelt_usage.py's
    update_scenario_scores_with_rear_seats replaces the single "Rear seat
    occupancy" placeholder row with one row per rear seat (2-6 rows,
    depending on the OEM's own "Number of rear seats" input parameter, with
    a "Max score" that is itself computed as 5 / num_rear_seats, not a fixed
    reference value); speed_assistance.py's
    update_scenario_scores_with_scf_intelligent_row and acc_performance.py's
    update_scenario_scores_autoresume_value likewise fill in a dynamic
    "Scenario" label post-preprocess. This reshaping is part of normal,
    always-exercised scoring (not an edge case -- every vehicle has some
    number of rear seats), so rebuild_trusted_df's fixed-pristine-row-count
    assumption would reject essentially all legitimate real submissions.
  - "DE - DM pred.", "VA - Speed assist. pred.", "VA - ACC pred." are multi-
    subtable / matrix-stacked sheets with no forward-fillable identity in
    the GreyCellSheetSchema sense, so they can't be schema'd either -- they
    are entirely OEM input (grey), so they're listed in PASSTHROUGH_SHEETS
    instead, same treatment as "Scenario Scores" above. (A sheet that is
    genuinely all-grey and exists in the pristine template but is in
    neither SHEET_SCHEMAS nor PASSTHROUGH_SHEETS gets silently reset to the
    pristine template's blank placeholder by rebuild_trusted_workbook --
    that was the original bug this comment now documents the fix for.)
  - All 8 "...verif." sheets (e.g. "VA - ACC verif.", whose "OEM Prediction"
    column is read directly by acc_performance.compute_score) have no
    pristine-template counterpart at all -- none of them exist in
    sd_template.xlsx, they're generated entirely at preprocess time. For the
    same reason they are deliberately outside find_missing_required_inputs /
    find_consistency_findings: those checks cover prediction-template sheets
    only, so a workbook that carries assessment-stage sheets (a preprocessed
    file being filled in by the assessor) doesn't get its not-yet-assessed verif
    Values reported as missing OEM inputs.
  - "Documentation"/"Data" are also PASSTHROUGH_SHEETS -- there is no
    computed logic for this domain yet (unlike crash_avoidance's
    documentation_data.py), but they're still not schema'd: nothing to
    protect until a domain-specific check exists.
"""

import functools
from importlib.resources import files
import os
import tempfile

import openpyxl

from euroncap_rating_2026 import common
from euroncap_rating_2026.safe_driving import consistency

SHEET_SCHEMAS = {
    "Input parameters": common.GreyCellSheetSchema(
        identity_columns=[
            "Stage",
            "Stage element",
            "Stage subelement",
            "Category",
            "Scenario",
            "Input parameter",
        ],
        grey_columns=["Value"],
        # A later template revision named this row's "Input
        # parameter" ("System updates type") -- it was blank before. Every
        # prediction file produced against the prior template revision has
        # that one cell blank; this narrowly whitelists exactly that
        # (row, column), not blank "Input parameter" cells in general (see
        # common.LegacyUnlabeledCell).
        legacy_unlabeled_cells=[
            common.LegacyUnlabeledCell(
                column="Input parameter",
                anchor={
                    "Stage element": "Vehicle Assistance",
                    "Stage subelement": "Speed assistance",
                    "Category": "SLIF - System updates",
                },
                value="System updates type",
            ),
        ],
    ),
    "Test Scores": common.GreyCellSheetSchema(
        identity_columns=["Stage", "Stage element", "Stage subelement"],
        computed_columns=["Score"],
    ),
    "Category Scores": common.GreyCellSheetSchema(
        identity_columns=["Stage", "Stage element", "Stage subelement", "Category"],
        computed_columns=["Score"],
    ),
}

PASSTHROUGH_SHEETS = [
    "Version",
    "Scenario Scores",
    "DE - DM pred.",
    "VA - Speed assist. pred.",
    "VA - ACC pred.",
    "Documentation",
    "Data",
]

PRISTINE_TEMPLATE_NAME = "sd_template.xlsx"

# Sheet -> columns that are only ever produced by compute-score, never by
# the OEM or preprocess. "Test Scores"/"Category Scores" come from the
# schemas that declare the trust boundary; "Scenario Scores" is deliberately
# NOT schema'd (see the module docstring -- its row structure is dynamic,
# so it's a PASSTHROUGH_SHEET instead), but its "Score" column is the exact
# same kind of not-yet-computed placeholder. Consumed by generate_template
# and report_writer's preprocess path (blank_computed_columns), the DataFrame
# path (reset_computed_columns_to_dash) and _pristine_dfs
# (numericize_computed_columns).
COMPUTED_SHEET_COLUMNS = {
    name: list(schema.computed_columns)
    for name, schema in SHEET_SCHEMAS.items()
    if schema.computed_columns
}
COMPUTED_SHEET_COLUMNS["Scenario Scores"] = ["Score"]

# "DE - DM pred.", "VA - Speed assist. pred." and "VA - ACC pred." have no
# forward-fillable identity columns (see module docstring), so they aren't
# GreyCellSheetSchema'd -- but each already declares its OEM-input cells via
# Data Validation dropdown ranges authored directly in the packaged
# template (same mechanism as crash_avoidance's grid sheets;
# "DE - DM pred." is also what _get_dm_prediction_required_coordinates
# reads for scoring). find_missing_required_inputs below reads those
# ranges (see common.find_missing_grid_inputs) instead of re-deriving
# driver_monitoring.py's/speed_assistance.py's own positional table parsing.
GRID_SHEET_STAGE_ELEMENTS = {
    "DE - DM pred.": [("Driver Engagement", "Driver monitoring")],
    "VA - Speed assist. pred.": [("Vehicle Assistance", "Speed assistance")],
    "VA - ACC pred.": [("Vehicle Assistance", "ACC performance")],
}


def rebuild_trusted_input(user_file_path: str) -> str:
    """Build a sanitized copy of *user_file_path* where every protected
    reference/structural cell (Max score, row labels, ...) is sourced from
    the official packaged template instead of the user-supplied file, and
    only the designated grey input cells (Value) are carried over from the
    user.

    Returns the path to a new temporary .xlsx file. The caller is
    responsible for deleting it.

    Raises common.TemplateIntegrityError if the user's file is missing a
    required sheet/column, or if a sheet's rows don't match the pristine
    template's row identity structure (inserted, deleted, or reordered
    rows).
    """
    pristine_path = str(files("data").joinpath(PRISTINE_TEMPLATE_NAME))
    user_wb = openpyxl.load_workbook(user_file_path, data_only=True)
    try:
        trusted_wb = common.rebuild_trusted_workbook(
            pristine_path, user_wb, SHEET_SCHEMAS, PASSTHROUGH_SHEETS
        )
    finally:
        user_wb.close()

    fd, sanitized_path = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    trusted_wb.save(sanitized_path)
    return sanitized_path


@functools.lru_cache(maxsize=1)
def _pristine_dfs() -> dict:
    """The packaged pristine template loaded as a dict of DataFrames, cached
    for the lifetime of the process. Backing rebuild_trusted_input_dfs,
    which -- unlike rebuild_trusted_input -- can be called on every single
    get_updated_dfs(dfs)/add_verification_sheets_to_dfs invocation by a long-running
    caller, not once per CLI process; re-parsing the template file on
    every call would be a real, avoidable cost."""
    pristine_path = str(files("data").joinpath(PRISTINE_TEMPLATE_NAME))
    dfs = common.read_excel_file_to_dfs(pristine_path)
    # "Scenario Scores" is PASSTHROUGH, never overlaid onto user data by
    # rebuild_trusted_dfs, so numericizing it here is safe.
    return common.numericize_computed_columns(dfs, COMPUTED_SHEET_COLUMNS)


def rebuild_trusted_input_dfs(user_dfs: dict) -> dict:
    """DataFrame-native sibling of rebuild_trusted_input: same SHEET_SCHEMAS
    and PASSTHROUGH_SHEETS, same pristine template, but operating directly
    on a dict of DataFrames instead of an on-disk .xlsx file -- for callers
    (e.g. an application that parses an upload once and works in dataframe/
    parquet land afterwards) that never have the file/openpyxl workbook
    rebuild_trusted_input needs.

    Returns a copy of *user_dfs* with every protected reference cell (Max
    score, row labels, ...) in a schema'd sheet sourced from the official
    packaged template instead of *user_dfs*; only the designated grey input
    cells are carried over from *user_dfs*. A schema'd sheet missing from
    *user_dfs* is left untouched (see common.rebuild_trusted_dfs).

    Raises common.TemplateIntegrityError if a present schema'd sheet is
    missing a required column, or if its rows don't match the pristine
    template's row identity structure (inserted, deleted, or reordered
    rows).
    """
    return common.rebuild_trusted_dfs(
        _pristine_dfs(), user_dfs, SHEET_SCHEMAS, PASSTHROUGH_SHEETS
    )


@functools.lru_cache(maxsize=1)
def _pristine_wb():
    """The packaged pristine template as an openpyxl Workbook, cached for
    the lifetime of the process -- backs find_missing_required_inputs'
    grid-sheet checks, which read the template's Data Validation ranges
    fresh from the pristine file rather than from the user's own workbook.
    """
    pristine_path = str(files("data").joinpath(PRISTINE_TEMPLATE_NAME))
    return openpyxl.load_workbook(pristine_path)


# Every (Stage element, Stage subelement) pair this domain's sheets use --
# so an unknown scope raises instead of silently reporting zero findings
# (common.validate_scope). Declared literally rather than derived from
# data_model.STAGE_SUBELEMENTS because that list capitalizes differently
# from the sheets themselves ("Speed Assistance" vs the sheets' "Speed
# assistance") and the check matches the sheets' own strings. The sheets
# also carry two spellings for steering ("Steering assistance" in Input
# parameters, "Steering assistance performance" in Test Scores) -- both
# are accepted.
VALID_SCOPE_PAIRS = [
    ("Occupant Monitoring", "Seatbelt usage"),
    ("Occupant Monitoring", "Occupant classification"),
    ("Occupant Monitoring", "Occupant presence"),
    ("Driver Engagement", "Driver monitoring"),
    ("Driver Engagement", "General vehicle controls"),
    ("Vehicle Assistance", "Speed assistance"),
    ("Vehicle Assistance", "ACC performance"),
    ("Vehicle Assistance", "Steering assistance"),
    ("Vehicle Assistance", "Steering assistance performance"),
]


def find_missing_required_inputs(
    wb, stage_element: str = None, stage_subelement: str = None
) -> list:
    """Find empty cells the OEM is required to fill in for the safe_driving
    domain, optionally scoped to one protocol element (e.g.
    stage_element="Vehicle Assistance", stage_subelement="ACC performance").
    Covers the schema'd sheets (Input parameters) and the
    "DE - DM pred."/"VA - Speed assist. pred."/"VA - ACC pred." grid
    sheets -- prediction-template sheets only. The preprocess-generated
    "...verif." sheets are deliberately not checked (see the module
    docstring): their Values are assessment-stage inputs, not OEM
    predictions. Raises ValueError for a scope the domain doesn't declare.

    A version-mismatched workbook (common.find_version_mismatch) is almost
    always also structurally different (renamed sheets/columns), which
    would otherwise raise common.TemplateIntegrityError before any finding
    -- including the version mismatch itself -- makes it back to the
    caller. When a version mismatch is detected, a resulting
    TemplateIntegrityError is swallowed in favour of just the version
    finding: a "wrong template version" message is strictly more useful
    than the low-level column-name error it would otherwise surface.
    """
    common.validate_scope(
        stage_element, stage_subelement, VALID_SCOPE_PAIRS, "safe_driving"
    )
    version_mismatch = common.find_version_mismatch(wb)
    try:
        missing = version_mismatch + common.find_missing_required_inputs(
            wb, SHEET_SCHEMAS, stage_element, stage_subelement
        )

        pristine_wb = _pristine_wb()
        for sheet_name, applies_to in GRID_SHEET_STAGE_ELEMENTS.items():
            if sheet_name not in wb.sheetnames:
                continue
            if not common.stage_elements_in_scope(
                applies_to, stage_element, stage_subelement
            ):
                continue
            missing.extend(
                common.find_missing_grid_inputs(
                    wb[sheet_name], pristine_wb[sheet_name], sheet_name
                )
            )
        return missing
    except common.TemplateIntegrityError:
        if version_mismatch:
            return version_mismatch
        raise


def find_missing_required_inputs_in_file(
    input_file: str, stage_element: str = None, stage_subelement: str = None
) -> list:
    """Load an Excel file and return find_missing_required_inputs() for it."""
    wb = openpyxl.load_workbook(input_file, data_only=True)
    try:
        return find_missing_required_inputs(wb, stage_element, stage_subelement)
    finally:
        wb.close()


def find_consistency_findings(
    wb, stage_element: str = None, stage_subelement: str = None
) -> list:
    """One list of findings covering every consistency rule this domain
    declares: the missing required inputs
    (find_missing_required_inputs) plus every report-only prediction-rule
    violation (consistency.find_prediction_inconsistencies --
    dm_group_consistency). Each finding says which rule produced it in
    ``check`` (common.CHECK_*). Same scope convention and error behaviour
    as find_missing_required_inputs: ValueError for an undeclared scope,
    common.TemplateIntegrityError for structural problems."""
    findings = find_missing_required_inputs(wb, stage_element, stage_subelement)
    findings.extend(
        consistency.find_prediction_inconsistencies(wb, stage_element, stage_subelement)
    )
    return findings


def find_consistency_findings_in_file(
    input_file: str, stage_element: str = None, stage_subelement: str = None
) -> list:
    """Load an Excel file and return find_consistency_findings() for it."""
    wb = openpyxl.load_workbook(input_file, data_only=True)
    try:
        return find_consistency_findings(wb, stage_element, stage_subelement)
    finally:
        wb.close()
