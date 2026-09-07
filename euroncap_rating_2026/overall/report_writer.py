# Copyright 2025-2026, Euro NCAP IVZW
# Created by IVEX NV (https://ivex.ai)
#
# Licensed under the Apache License 2.0.
# See http://www.apache.org/licenses/LICENSE-2.0 for details.

import os
from datetime import datetime

import pandas as pd

from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

from euroncap_rating_2026 import common

import logging

logger = logging.getLogger(__name__)

COMPUTE_SHEETS_TO_COPY = [
    "Version",
    "Input parameters",
]


def get_output_file_path(output_path: str) -> str:
    current_datetime = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    return os.path.join(output_path, f"overall_{current_datetime}_report.xlsx")


def write_report(
    command: common.CliCommand,
    input_file: str,
    updated_dfs: dict,
    selected_points_dict,
    output_path: str,
    format_prediction_cells=True,
    sheet_order=None,
):
    if command == common.CliCommand.COMPUTE_SCORE:
        output_file = get_output_file_path(output_path)
        sheets_to_copy = COMPUTE_SHEETS_TO_COPY
    else:
        logger.error(f"Unsupported command: {command}")
        return

    print(f"Saving to {output_file}...")

    # Copy base sheets from input to output via workbook objects
    input_wb = common.load_or_create_workbook(input_file)
    wb = common.load_or_create_workbook(output_file)
    for sheet in sheets_to_copy:
        common.hard_copy_sheet(input_wb, sheet, wb)
    wb.save(output_file)

    # Write updated DataFrames via ExcelWriter to preserve all rows
    try:
        with pd.ExcelWriter(
            output_file, engine="openpyxl", mode="a", if_sheet_exists="replace"
        ) as writer:
            for sheet_name, updated_df in updated_dfs.items():
                updated_df.to_excel(writer, sheet_name=sheet_name, index=False)
    except Exception as e:
        logger.error(f"Failed to write updated sheets: {e}")

    wb = common.load_or_create_workbook(output_file)

    calculation_color = PatternFill(
        start_color="FFDD04", end_color="FFDD04", fill_type="solid"
    )

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        logger.debug(f"Formatting sheet: {sheet_name}")

        if sheet_name == "Version":
            white_font = Font(color="FFFFFF", name="Aptos Narrow")
            black_fill = PatternFill(
                start_color="000000", end_color="000000", fill_type="solid"
            )
            left_align = Alignment(horizontal="left")
            ws.column_dimensions["A"].width = 10 / 3.78
            for row in ws.iter_rows():
                for cell in row:
                    cell.font = white_font
                    cell.fill = black_fill
                    cell.alignment = left_align
            continue

        ###############################################################
        # Color Score columns
        ###############################################################
        for col in ws.iter_cols():
            if col[0].value in ["Score", "Points", "Star rating"]:
                for cell in col[1:]:
                    cell.fill = calculation_color
                    if common.is_empty_cell(cell.value):
                        cell.value = ""

        common.set_number_precision(ws)
        common.set_title_capitalization(ws)

    ###############################################################
    # Adjust column widths based on content
    ###############################################################
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for column in ws.columns:
            max_length = 0
            if not hasattr(column[0], "column_letter"):
                continue
            col_letter = column[0].column_letter
            for cell in column:
                try:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                except (TypeError, AttributeError):
                    pass
            ws.column_dimensions[col_letter].width = max_length + 2

    ###############################################################
    # Borders and header formatting
    ###############################################################
    header_fill = PatternFill(
        start_color="000000", end_color="000000", fill_type="solid"
    )
    header_font = Font(name="Calibri", color="FFFFFF", bold=True)

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        if ws.max_row < 1:
            continue
        max_row = ws.max_row
        max_col = ws.max_column

        # Thick border on last column
        for row in range(1, max_row + 1):
            cell = ws.cell(row=row, column=max_col)
            cell.border = cell.border + Border(
                right=Side(border_style="thick", color="000000")
            )
        # Thick border on last row
        for col in range(1, max_col + 1):
            cell = ws.cell(row=max_row, column=col)
            cell.border = cell.border + Border(
                bottom=Side(border_style="thick", color="000000")
            )

        # Black header row
        if sheet_name != "Version":
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font

    ###############################################################
    # Reorder sheets
    ###############################################################
    def reorder_sheets(workbook, main_order=None):
        if main_order is None:
            main_order = [
                "Input parameters",
                "Rating",
                "Stage Scores",
                "Stage element Scores",
                "Test Scores",
                "Version",
            ]
        desired_order = [s for s in main_order if s in workbook.sheetnames]
        for remaining in workbook.sheetnames:
            if remaining not in desired_order:
                desired_order.append(remaining)
        for idx, sname in enumerate(desired_order):
            if workbook.sheetnames[idx] != sname:
                sheet = workbook[sname]
                workbook._sheets.insert(
                    idx, workbook._sheets.pop(workbook._sheets.index(sheet))
                )

    reorder_sheets(wb, sheet_order)
    common.add_version_sheet_to_wb(wb, command)
    wb.save(output_file)

    print(f"Log available at {os.path.join(output_path, 'euroncap_rating_2026.log')}")
    print(" " * 40)
    print(f"Final report available at {output_file}")
    logger.info(f"Final report available at {output_file}")
