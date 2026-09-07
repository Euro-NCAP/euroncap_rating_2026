# Copyright 2025-2026, Euro NCAP IVZW
# Created by IVEX NV (https://ivex.ai)
#
# Licensed under the Apache License 2.0.
# See http://www.apache.org/licenses/LICENSE-2.0 for details.

"""Report-only prediction-consistency rules for the crash_protection domain.

Two rules on "CP - VRU Prediction":

- legform_symmetry: the VRU legform T/ST/SA symmetry
  (vru_processing.check_legform_blue_symmetry, which preprocess enforces by
  raising on the first offending cell). This module runs the same shared
  rule logic (collect_legform_symmetry_violations) on the raw uploaded
  workbook and returns every violation as a common.ConsistencyFinding, so
  a consuming application can display all of them at once.
- no_prediction_provided: the Headforms/Upper legform/aPLI sections have no
  per-cell required-input check (the OEM fills only the cells matching the
  car's layout and its own prediction),
  so nothing else guarantees the OEM entered *any* prediction there. This
  rule reports a section that is entirely untouched (every cell GREY).

The legform_symmetry rule is what still validates the cells the OEM does
fill in the sections exempted from the missing-input check;
no_prediction_provided is what guarantees they filled in at least one.
"""

import logging

from openpyxl.utils import get_column_letter

from euroncap_rating_2026 import common
from euroncap_rating_2026.crash_protection import vru_processing

logger = logging.getLogger(__name__)

VRU_PREDICTION_SHEET = "CP - VRU Prediction"

# The legform matrix rows, top to bottom, with the template section each row
# renders under and the (Stage element, Stage subelement) scope it belongs
# to. Kept per-row -- not per-block -- because the symmetry rule is
# evaluated within a row, and consuming applications check Pelvis Impact and
# Leg Impact as two separate scopes whose results it merges: a block-wide
# attribution would duplicate every finding across the two calls. The
# combined "Pelvis & Leg Impact" alias (Test Scores' vocabulary, accepted by
# integrity.VALID_SCOPE_PAIRS) matches every row. no_prediction_provided
# reuses the same (section, applies_to) pairs for rows 0 (Pelvis Impact) and
# 1/2 (Leg Impact, checked together -- either aPLI row having a prediction
# satisfies the element).
_LEGFORM_ROW_SCOPES = (
    (
        "Upper legform",
        (("VRU Impact", "Pelvis Impact"), ("VRU Impact", "Pelvis & Leg Impact")),
    ),
    (
        "aPLI - Femur",
        (("VRU Impact", "Leg Impact"), ("VRU Impact", "Pelvis & Leg Impact")),
    ),
    (
        "aPLI - Knee & Tibia",
        (("VRU Impact", "Leg Impact"), ("VRU Impact", "Pelvis & Leg Impact")),
    ),
)

_HEADFORM_SCOPE = (("VRU Impact", "Head Impact"),)

# Column-A section-header row for "Headforms", confirmed against the
# packaged template (cp_template.xlsx: "Headforms" literally sits in A1).
# Used only to anchor a no_prediction_provided finding, which has no single
# offending data cell to point at.
_HEADFORMS_HEADER_ROW = 1


def _legform_row(row_idx: int) -> int:
    """Worksheet row of legform matrix row row_idx. This is also the row
    column A's own section header (Upper legform/aPLI - Femur/aPLI - Knee &
    Tibia) sits on -- confirmed against the packaged template -- so it
    doubles as the anchor for a no_prediction_provided finding.

    get_legform_matrix slices vru_prediction_df.iloc[LEGFORMS_START_ROW_INDEX:
    ...][2:, 4:...]: matrix row 0 is DataFrame row LEGFORMS_START_ROW_INDEX+2
    and matrix column 0 is DataFrame column 4; +2/+1 convert the 0-based,
    header-stripped DataFrame indices to 1-based worksheet ones.
    """
    return vru_processing.LEGFORMS_START_ROW_INDEX + 2 + row_idx + 2


def _legform_cell(row_idx: int, j: int) -> str:
    """Worksheet coordinate of legform matrix cell (row_idx, j)."""
    ws_col = 4 + j + 1
    return f"{get_column_letter(ws_col)}{_legform_row(row_idx)}"


def _legform_symmetry_message(
    v: "vru_processing.LegformSymmetryViolation", section: str, cell: str
) -> str:
    """The report-only legform_symmetry message: sheet/section/cell plus
    the symmetric cell it disagrees with -- no matrix row/col/index jargon
    (that wording stays exclusive to the blocking ValueError,
    vru_processing._legform_violation_message)."""
    if v.rule == "sa_out_of_range":
        sentence = "SA prediction; no symmetric cell exists in this matrix."
    else:
        sym_cell = _legform_cell(v.row_idx, v.j_sym)
        if v.rule == "t_st":
            sentence = (
                f"asymmetric T/ST prediction; symmetric cell {sym_cell} has "
                f"color '{v.sym_color.value}', expected T or ST."
            )
        else:  # sa_mismatch
            sentence = (
                f"SA prediction; symmetric cell {sym_cell} has color "
                f"'{v.sym_color.value}', expected SA."
            )
    return f"{VRU_PREDICTION_SHEET} — {section} {cell}: {sentence}"


def _section_has_prediction(matrix_rows) -> bool:
    """Whether any cell across the given legform/headform matrix row(s)
    carries a real prediction -- anything other than GREY. GREY is the
    literal "nothing entered" value throughout this sheet (see
    vru_processing.VruPredictionColor); other excluded-from-sampling colors
    like D_RED/D_GREEN are still real OEM predictions, just not
    verification-test candidates, so they must count here."""
    return any(
        color != vru_processing.VruPredictionColor.GREY
        for row in matrix_rows
        for color in row
    )


def _no_prediction_finding(
    section_label: str, ws_row: int, expected: str
) -> "common.ConsistencyFinding":
    cell = f"A{ws_row}"
    return common.ConsistencyFinding(
        sheet=VRU_PREDICTION_SHEET,
        column="A",
        cell=cell,
        row_identity={"Section": section_label},
        check=common.CHECK_NO_PREDICTION_PROVIDED,
        message=(
            f"{VRU_PREDICTION_SHEET} — {section_label}: no prediction provided; "
            f"at least one cell in this section must have a {expected} prediction."
        ),
    )


def find_prediction_inconsistencies(
    wb, stage_element: str = None, stage_subelement: str = None
) -> list:
    """Every prediction-rule violation in *wb* for the crash_protection
    domain, as a list of common.ConsistencyFinding (check =
    legform_symmetry or no_prediction_provided). Report-only: collects all
    violations instead of raising on the first -- the blocking behaviour
    stays in preprocess. Scope filtering follows
    find_missing_required_inputs' (Stage element, Stage subelement)
    convention; scope validity is the caller's concern
    (integrity.find_consistency_findings validates it once).
    """
    if VRU_PREDICTION_SHEET not in wb.sheetnames:
        return []

    df = common.sheet_to_dataframe(wb[VRU_PREDICTION_SHEET])
    if df.empty:
        return []

    findings = []

    legform_matrix = vru_processing.get_legform_matrix(df)
    if legform_matrix and legform_matrix[0]:
        col_center = len(legform_matrix[0]) // 2
        rows_in_scope = [
            row_idx
            for row_idx, (_section, applies_to) in enumerate(_LEGFORM_ROW_SCOPES)
            if common.stage_elements_in_scope(
                list(applies_to), stage_element, stage_subelement
            )
        ]
        for v in vru_processing.collect_legform_symmetry_violations(legform_matrix):
            if v.row_idx >= len(_LEGFORM_ROW_SCOPES) or v.row_idx not in rows_in_scope:
                continue
            section = _LEGFORM_ROW_SCOPES[v.row_idx][0]
            cell = _legform_cell(v.row_idx, v.j)
            findings.append(
                common.ConsistencyFinding(
                    sheet=VRU_PREDICTION_SHEET,
                    column=get_column_letter(4 + v.j + 1),
                    cell=cell,
                    row_identity={
                        "Section": section,
                        "Lateral position": col_center - v.j,
                    },
                    check=common.CHECK_LEGFORM_SYMMETRY,
                    message=_legform_symmetry_message(v, section, cell),
                )
            )

        if common.stage_elements_in_scope(
            list(_LEGFORM_ROW_SCOPES[0][1]), stage_element, stage_subelement
        ) and not _section_has_prediction([legform_matrix[0]]):
            findings.append(
                _no_prediction_finding(
                    "Upper legform (Pelvis Impact)", _legform_row(0), "T, ST or SA"
                )
            )

        if (
            len(legform_matrix) >= 3
            and common.stage_elements_in_scope(
                list(_LEGFORM_ROW_SCOPES[1][1]), stage_element, stage_subelement
            )
            and not _section_has_prediction(legform_matrix[1:3])
        ):
            findings.append(
                _no_prediction_finding(
                    "aPLI - Femur / aPLI - Knee & Tibia (Leg Impact)",
                    _legform_row(1),
                    "T, ST or SA",
                )
            )

    if common.stage_elements_in_scope(
        list(_HEADFORM_SCOPE), stage_element, stage_subelement
    ):
        headform_matrix = vru_processing.get_headform_matrix(df)
        if (
            headform_matrix
            and headform_matrix[0]
            and not _section_has_prediction(headform_matrix)
        ):
            findings.append(
                _no_prediction_finding("Headforms", _HEADFORMS_HEADER_ROW, "color")
            )

    return findings
