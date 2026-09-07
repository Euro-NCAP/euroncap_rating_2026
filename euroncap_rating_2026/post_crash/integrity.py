# Copyright 2025-2026, Euro NCAP IVZW
# Created by IVEX NV (https://ivex.ai)
#
# Licensed under the Apache License 2.0.
# See http://www.apache.org/licenses/LICENSE-2.0 for details.

"""Grey-cell trust boundary for the post_crash domain.

Declares, per sheet, which columns are legitimate OEM input ("grey" cells --
Value) versus protected reference/structural data (Max score, row labels)
that must always come from the official packaged template rather than a
user-supplied file. See `euroncap_rating_2026.common` for the underlying
mechanism.

NOTE: the six "Verif" sheets (RI - RS Verification, RI - ERG Verification,
PCI - Advanced eCall Verif, PCI - MCB & Hazard lights Verif,
VE - Energy Management Verif, VE - Occupant Extrication Verif) are not
covered by a schema here: they are generated at preprocess time from
hardcoded Python constants (see data_model.py / energy_management.py /
occupant_extrication.py), not from a pristine packaged .xlsx sheet, so there
is no template counterpart to rebuild against with this mechanism. Their
Category/Scenario row labels aren't verified against tampering. Left as a
follow-up if this mechanism is generalized further. For the same reason they
are deliberately outside find_missing_required_inputs /
find_consistency_findings: those checks cover prediction-template sheets
only, so a workbook that carries assessment-stage sheets (a preprocessed
file being filled in by the assessor) doesn't get its not-yet-assessed Verif
Values reported as missing OEM inputs.

"Documentation"/"Data" are PASSTHROUGH_SHEETS too -- there is no computed
logic for this domain yet (unlike crash_avoidance's documentation_data.py),
but they're still not schema'd: nothing to protect until a domain-specific
check exists.
"""

import functools
from importlib.resources import files
import os
import tempfile

import openpyxl

from euroncap_rating_2026 import common
from euroncap_rating_2026.post_crash import data_model

SHEET_SCHEMAS = {
    "Input parameters": common.GreyCellSheetSchema(
        identity_columns=[
            "Stage",
            "Stage element",
            "Stage subelement",
            "Category",
            "Input parameter",
        ],
        grey_columns=["Value"],
    ),
    "Test Scores": common.GreyCellSheetSchema(
        identity_columns=["Stage", "Stage element", "Stage subelement"],
        computed_columns=["Score"],
    ),
    "Category Scores": common.GreyCellSheetSchema(
        identity_columns=["Stage", "Stage element", "Stage subelement", "Category"],
        computed_columns=["Score"],
    ),
    "Scenario Scores": common.GreyCellSheetSchema(
        identity_columns=[
            "Stage",
            "Stage element",
            "Stage subelement",
            "Category",
            "Scenario",
        ],
        computed_columns=["Score"],
    ),
}

# Sheet -> columns that are only ever produced by compute-score, never by
# the OEM or preprocess, derived from the same schemas that declare the
# trust boundary. Consumed by generate_template and report_writer's
# preprocess path (blank_computed_columns), the DataFrame path
# (reset_computed_columns_to_dash) and _pristine_dfs
# (numericize_computed_columns).
COMPUTED_SHEET_COLUMNS = {
    name: list(schema.computed_columns)
    for name, schema in SHEET_SCHEMAS.items()
    if schema.computed_columns
}

PASSTHROUGH_SHEETS = ["Version", "Documentation", "Data"]

PRISTINE_TEMPLATE_NAME = "pc_template.xlsx"


def rebuild_trusted_input(user_file_path: str) -> str:
    """Build a sanitized copy of *user_file_path* where every protected
    reference/structural cell (Max score, row labels, ...) is sourced from
    the official packaged template instead of the user-supplied file, and
    only the designated grey input cells (Value) are carried over from the
    user.

    Returns the path to a new temporary .xlsx file. The caller is
    responsible for deleting it.

    Raises common.TemplateIntegrityError if the user's file is missing a
    required sheet/column, or if a sheet's rows don't match the pristine
    template's row identity structure (inserted, deleted, or reordered
    rows).
    """
    pristine_path = str(files("data").joinpath(PRISTINE_TEMPLATE_NAME))
    user_wb = openpyxl.load_workbook(user_file_path, data_only=True)
    try:
        trusted_wb = common.rebuild_trusted_workbook(
            pristine_path, user_wb, SHEET_SCHEMAS, PASSTHROUGH_SHEETS
        )
    finally:
        user_wb.close()

    fd, sanitized_path = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    trusted_wb.save(sanitized_path)
    return sanitized_path


@functools.lru_cache(maxsize=1)
def _pristine_dfs() -> dict:
    """The packaged pristine template loaded as a dict of DataFrames, cached
    for the lifetime of the process. Backing rebuild_trusted_input_dfs,
    which -- unlike rebuild_trusted_input -- can be called on every single
    calculate_score(dfs) invocation by a long-running caller, not once per CLI
    process; re-parsing the template file on every call would be a real,
    avoidable cost."""
    pristine_path = str(files("data").joinpath(PRISTINE_TEMPLATE_NAME))
    dfs = common.read_excel_file_to_dfs(pristine_path)
    return common.numericize_computed_columns(dfs, COMPUTED_SHEET_COLUMNS)


def rebuild_trusted_input_dfs(user_dfs: dict) -> dict:
    """DataFrame-native sibling of rebuild_trusted_input: same SHEET_SCHEMAS
    and PASSTHROUGH_SHEETS, same pristine template, but operating directly
    on a dict of DataFrames instead of an on-disk .xlsx file -- for callers
    (e.g. an application that parses an upload once and works in dataframe/
    parquet land afterwards) that never have the file/openpyxl workbook
    rebuild_trusted_input needs.

    Returns a copy of *user_dfs* with every protected reference cell (Max
    score, row labels, ...) in a schema'd sheet sourced from the official
    packaged template instead of *user_dfs*; only the designated grey input
    cells are carried over from *user_dfs*. A schema'd sheet missing from
    *user_dfs* is left untouched (see common.rebuild_trusted_dfs).

    Raises common.TemplateIntegrityError if a present schema'd sheet is
    missing a required column, or if its rows don't match the pristine
    template's row identity structure (inserted, deleted, or reordered
    rows).
    """
    return common.rebuild_trusted_dfs(
        _pristine_dfs(), user_dfs, SHEET_SCHEMAS, PASSTHROUGH_SHEETS
    )


# Every (Stage element, Stage subelement) pair this domain's sheets use --
# same literals as data_model.STAGE_SUBELEMENTS -- so an unknown scope
# raises instead of silently reporting zero findings
# (common.validate_scope).
VALID_SCOPE_PAIRS = [
    (entry["Stage element"], entry["Stage subelement"])
    for entry in data_model.STAGE_SUBELEMENTS
]


def find_missing_required_inputs(
    wb, stage_element: str = None, stage_subelement: str = None
) -> list:
    """Find empty cells the OEM is required to fill in for the post_crash
    domain, optionally scoped to one protocol element (e.g.
    stage_element="Rescue Information", stage_subelement="Rescue sheets").
    Covers prediction-template sheets only; the preprocess-generated
    "...Verif" sheets are deliberately not checked (see the module
    docstring). Raises ValueError for a scope the domain doesn't declare.
    See common.find_missing_required_inputs.

    A version-mismatched workbook (common.find_version_mismatch) is almost
    always also structurally different (renamed sheets/columns), which
    would otherwise raise common.TemplateIntegrityError before any finding
    -- including the version mismatch itself -- makes it back to the
    caller. When a version mismatch is detected, a resulting
    TemplateIntegrityError is swallowed in favour of just the version
    finding: a "wrong template version" message is strictly more useful
    than the low-level column-name error it would otherwise surface.
    """
    common.validate_scope(
        stage_element, stage_subelement, VALID_SCOPE_PAIRS, "post_crash"
    )
    version_mismatch = common.find_version_mismatch(wb)
    try:
        return version_mismatch + common.find_missing_required_inputs(
            wb, SHEET_SCHEMAS, stage_element, stage_subelement
        )
    except common.TemplateIntegrityError:
        if version_mismatch:
            return version_mismatch
        raise


def find_missing_required_inputs_in_file(
    input_file: str, stage_element: str = None, stage_subelement: str = None
) -> list:
    """Load an Excel file and return find_missing_required_inputs() for it."""
    wb = openpyxl.load_workbook(input_file, data_only=True)
    try:
        return find_missing_required_inputs(wb, stage_element, stage_subelement)
    finally:
        wb.close()


def find_consistency_findings(
    wb, stage_element: str = None, stage_subelement: str = None
) -> list:
    """One list of findings covering every consistency rule this domain
    declares. post_crash declares no prediction rules
    beyond the required inputs, so this is find_missing_required_inputs
    under the cross-domain umbrella name every domain's integrity module
    exposes -- callers can treat the four domains uniformly."""
    return find_missing_required_inputs(wb, stage_element, stage_subelement)


def find_consistency_findings_in_file(
    input_file: str, stage_element: str = None, stage_subelement: str = None
) -> list:
    """Load an Excel file and return find_consistency_findings() for it."""
    wb = openpyxl.load_workbook(input_file, data_only=True)
    try:
        return find_consistency_findings(wb, stage_element, stage_subelement)
    finally:
        wb.close()
