# Copyright 2025-2026, Euro NCAP IVZW
# Created by IVEX NV (https://ivex.ai)
#
# Licensed under the Apache License 2.0.
# See http://www.apache.org/licenses/LICENSE-2.0 for details.

"""Grey-cell trust boundary for the crash_avoidance domain.

Declares, per sheet, which columns are legitimate OEM input ("grey" cells --
Value) versus protected reference/structural data (Max score, row labels)
that must always come from the official packaged template rather than a
user-supplied file. See `euroncap_rating_2026.common` for the underlying
mechanism.

NOTE: several sheet types are not covered by a schema here, left as a
follow-up if this mechanism is generalized further:
  - The "...pred." sheets (e.g. "FC - Car & PTW pred.") carry OEM data via
    cell background color, not cell text -- the same situation as
    crash_protection's "CP - VRU Prediction" -- so they're passed through
    verbatim rather than schema'd.
  - The "...robust. pred." sheets (e.g. "FC - Car & PTW robust. pred.") have
    OEM-input grey columns that are dynamically scenario-named (derived from
    data_model.py), not a small static list GreyCellSheetSchema can express.
  - The "...verif." sheets (e.g. "FC - Car & PTW verif.") are generated at
    preprocess time (no pristine template counterpart) and use a
    double-header / "general requirements" block layout incompatible with
    rebuild_trusted_df's single-flat-header assumption.
  - "FC - Driver State Link" is copied through verbatim by rebuild_trusted_*
    (no protected reference data), but its dropdown grid is OEM input and is
    covered by find_missing_required_inputs via GRID_SHEET_STAGE_ELEMENTS.
  - "Documentation"/"Data" hold "is supporting evidence required" flags that
    are computed once in preprocess.py (see documentation_data.py) from
    already-trust-boundary-protected sheets, then simply carried through
    compute-score unchanged -- they are not user input and not reset to a
    pristine placeholder, so they don't fit GreyCellSheetSchema either.
    Listed as PASSTHROUGH_SHEETS on the same trust tier as the "...pred."
    sheets above (not covered by the HMAC integrity signature).
"""

import functools
from importlib.resources import files
import os
import tempfile

import openpyxl

from euroncap_rating_2026 import common
from euroncap_rating_2026.crash_avoidance import consistency, data_model

SHEET_SCHEMAS = {
    "Input parameters": common.GreyCellSheetSchema(
        identity_columns=[
            "Stage",
            "Stage element",
            "Stage subelement",
            "Category",
            "Scenario",
            "Input parameter",
        ],
        grey_columns=["Value"],
    ),
    "Test Scores": common.GreyCellSheetSchema(
        identity_columns=["Stage", "Stage element", "Stage subelement"],
        computed_columns=["Score"],
    ),
    "Category Scores": common.GreyCellSheetSchema(
        identity_columns=["Stage", "Stage element", "Stage subelement", "Category"],
        computed_columns=["Score"],
    ),
    "Scenario Scores": common.GreyCellSheetSchema(
        identity_columns=[
            "Stage",
            "Stage element",
            "Stage subelement",
            "Category",
            "Scenario",
            "Layer",
        ],
        computed_columns=["Score"],
    ),
}

# Sheet -> columns that are only ever produced by compute-score, never by
# the OEM or preprocess, derived from the same schemas that declare the
# trust boundary. Consumed by generate_template (blank_computed_columns),
# preprocess (reset_computed_columns_to_dash) and _pristine_dfs
# (numericize_computed_columns).
COMPUTED_SHEET_COLUMNS = {
    name: list(schema.computed_columns)
    for name, schema in SHEET_SCHEMAS.items()
    if schema.computed_columns
}

# Sheets copied verbatim from the user's file: either genuine user content
# with no fixed reference counterpart to protect (the prediction/robustness
# grids, see module docstring), or content whose integrity is already
# enforced separately (Version, via common.check_version).
PASSTHROUGH_SHEETS = [
    "Version",
    "FC - Driver State Link",
    "FC - Car & PTW pred.",
    "FC - Ped & Cyc pred.",
    "FC - Car & PTW robust. pred.",
    "FC - Ped & Cyc robust. pred.",
    "LDC - Single Veh pred.",
    "LDC - Car & PTW pred.",
    "LDC - robust. pred.",
    "LSC - Car & PTW pred.",
    "LSC - Ped & Cyc pred.",
    "Documentation",
    "Data",
]

PRISTINE_TEMPLATE_NAME = "ca_template.xlsx"

# The "...pred."/"...robust. pred." grid sheets above have no
# forward-fillable identity columns (see module docstring), so they aren't
# GreyCellSheetSchema'd -- but they do carry genuine OEM prediction input,
# declared via Data Validation dropdown ranges directly in the packaged
# template. find_missing_required_inputs below reads those ranges (see
# common.find_missing_grid_inputs) rather than re-deriving
# matrix_processing.MATRIX_INDICES/CELL_REMOVAL_MAP or
# robustness_layer.ROBUSTNESS_LAYER_TO_VERIFICATION_CONDITION. Each sheet is
# scoped to the protocol element(s) it belongs to; "LDC - robust. pred." is
# the one sheet shared across two Lane Departure Collisions subelements.
GRID_SHEET_STAGE_ELEMENTS = {
    # FC-prefixed, but the "Driver state link" loadcase belongs to Lane
    # Departure Collisions / Single vehicle in data_model.py -- the grid
    # declares, per FC scenario, which sensitivity ranges the driver-state
    # link applies to, and the OEM must fill all of it at prediction time.
    "FC - Driver State Link": [("Lane Departure Collisions", "Single vehicle")],
    "FC - Car & PTW pred.": [("Frontal Collisions", "Car & PTW")],
    "FC - Car & PTW robust. pred.": [("Frontal Collisions", "Car & PTW")],
    "FC - Ped & Cyc pred.": [("Frontal Collisions", "Pedestrian & cyclist")],
    "FC - Ped & Cyc robust. pred.": [("Frontal Collisions", "Pedestrian & cyclist")],
    "LDC - Single Veh pred.": [("Lane Departure Collisions", "Single vehicle")],
    "LDC - Car & PTW pred.": [("Lane Departure Collisions", "Car & PTW")],
    "LDC - robust. pred.": [
        ("Lane Departure Collisions", "Single vehicle"),
        ("Lane Departure Collisions", "Car & PTW"),
    ],
    "LSC - Car & PTW pred.": [("Low Speed Collisions", "Car & PTW")],
    "LSC - Ped & Cyc pred.": [("Low Speed Collisions", "Pedestrian & cyclist")],
}


def rebuild_trusted_input(user_file_path: str) -> str:
    """Build a sanitized copy of *user_file_path* where every protected
    reference/structural cell (Max score, row labels, ...) is sourced from
    the official packaged template instead of the user-supplied file, and
    only the designated grey input cells (Value) are carried over from the
    user.

    Returns the path to a new temporary .xlsx file. The caller is
    responsible for deleting it.

    Raises common.TemplateIntegrityError if the user's file is missing a
    required sheet/column, or if a sheet's rows don't match the pristine
    template's row identity structure (inserted, deleted, or reordered
    rows).
    """
    pristine_path = str(files("data").joinpath(PRISTINE_TEMPLATE_NAME))
    user_wb = openpyxl.load_workbook(user_file_path, data_only=True)
    try:
        trusted_wb = common.rebuild_trusted_workbook(
            pristine_path, user_wb, SHEET_SCHEMAS, PASSTHROUGH_SHEETS
        )
    finally:
        user_wb.close()

    fd, sanitized_path = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    trusted_wb.save(sanitized_path)
    return sanitized_path


@functools.lru_cache(maxsize=1)
def _pristine_dfs() -> dict:
    """The packaged pristine template loaded as a dict of DataFrames, cached
    for the lifetime of the process. Backing rebuild_trusted_input_dfs,
    which -- unlike rebuild_trusted_input -- can be called on every single
    calculate_score(dfs)/add_verification_sheets_to_dfs invocation by a long-running
    caller, not once per CLI process; re-parsing the template file on
    every call would be a real, avoidable cost."""
    pristine_path = str(files("data").joinpath(PRISTINE_TEMPLATE_NAME))
    dfs = common.read_excel_file_to_dfs(pristine_path)
    return common.numericize_computed_columns(dfs, COMPUTED_SHEET_COLUMNS)


def rebuild_trusted_input_dfs(user_dfs: dict) -> dict:
    """DataFrame-native sibling of rebuild_trusted_input: same SHEET_SCHEMAS
    and PASSTHROUGH_SHEETS, same pristine template, but operating directly
    on a dict of DataFrames instead of an on-disk .xlsx file -- for callers
    (e.g. an application that parses an upload once and works in dataframe/
    parquet land afterwards) that never have the file/openpyxl workbook
    rebuild_trusted_input needs.

    Returns a copy of *user_dfs* with every protected reference cell (Max
    score, row labels, ...) in a schema'd sheet sourced from the official
    packaged template instead of *user_dfs*; only the designated grey input
    cells are carried over from *user_dfs*. A schema'd sheet missing from
    *user_dfs* is left untouched (see common.rebuild_trusted_dfs).

    Raises common.TemplateIntegrityError if a present schema'd sheet is
    missing a required column, or if its rows don't match the pristine
    template's row identity structure (inserted, deleted, or reordered
    rows).
    """
    return common.rebuild_trusted_dfs(
        _pristine_dfs(), user_dfs, SHEET_SCHEMAS, PASSTHROUGH_SHEETS
    )


@functools.lru_cache(maxsize=1)
def _pristine_wb():
    """The packaged pristine template as an openpyxl Workbook, cached for
    the lifetime of the process -- backs find_missing_required_inputs'
    grid-sheet checks, which read the template's Data Validation ranges
    fresh from the pristine file rather than from the user's own workbook.
    """
    pristine_path = str(files("data").joinpath(PRISTINE_TEMPLATE_NAME))
    return openpyxl.load_workbook(pristine_path)


# Every (Stage element, Stage subelement) pair this domain's sheets use --
# same literals as data_model.STAGE_SUBELEMENTS -- so an unknown scope
# raises instead of silently reporting zero findings
# (common.validate_scope).
VALID_SCOPE_PAIRS = [
    (entry["Stage element"], entry["Stage subelement"])
    for entry in data_model.STAGE_SUBELEMENTS
]


def find_missing_required_inputs(
    wb, stage_element: str = None, stage_subelement: str = None
) -> list:
    """Find empty cells the OEM is required to fill in for the
    crash_avoidance domain, optionally scoped to one protocol element (e.g.
    stage_element="Frontal Collisions", stage_subelement="Car & PTW").
    Covers both the schema'd sheets (Input parameters) and the
    "...pred."/"...robust. pred." grid sheets. Raises ValueError for a
    scope the domain doesn't declare.

    A version-mismatched workbook (common.find_version_mismatch) is almost
    always also structurally different (renamed sheets/columns), which
    would otherwise raise common.TemplateIntegrityError before any finding
    -- including the version mismatch itself -- makes it back to the
    caller. When a version mismatch is detected, a resulting
    TemplateIntegrityError is swallowed in favour of just the version
    finding: a "wrong template version" message is strictly more useful
    than the low-level column-name error it would otherwise surface.
    """
    common.validate_scope(
        stage_element, stage_subelement, VALID_SCOPE_PAIRS, "crash_avoidance"
    )
    version_mismatch = common.find_version_mismatch(wb)
    try:
        missing = version_mismatch + common.find_missing_required_inputs(
            wb, SHEET_SCHEMAS, stage_element, stage_subelement
        )

        pristine_wb = _pristine_wb()
        for sheet_name, applies_to in GRID_SHEET_STAGE_ELEMENTS.items():
            if sheet_name not in wb.sheetnames:
                continue
            if not common.stage_elements_in_scope(
                applies_to, stage_element, stage_subelement
            ):
                continue
            missing.extend(
                common.find_missing_grid_inputs(
                    wb[sheet_name], pristine_wb[sheet_name], sheet_name
                )
            )
        return missing
    except common.TemplateIntegrityError:
        if version_mismatch:
            return version_mismatch
        raise


def find_missing_required_inputs_in_file(
    input_file: str, stage_element: str = None, stage_subelement: str = None
) -> list:
    """Load an Excel file and return find_missing_required_inputs() for it."""
    wb = openpyxl.load_workbook(input_file, data_only=True)
    try:
        return find_missing_required_inputs(wb, stage_element, stage_subelement)
    finally:
        wb.close()


def find_consistency_findings(
    wb, stage_element: str = None, stage_subelement: str = None
) -> list:
    """One list of findings covering every consistency rule this domain
    declares: the missing required inputs
    (find_missing_required_inputs) plus every report-only prediction-rule
    violation (consistency.find_prediction_inconsistencies -- lsc_dooring,
    cpla_cbla_coherence, fcw_color, scenario_na_consistency). Each finding
    says which rule produced
    it in ``check`` (common.CHECK_*). Same scope convention and error
    behaviour as find_missing_required_inputs: ValueError for an undeclared
    scope, common.TemplateIntegrityError for structural problems."""
    findings = find_missing_required_inputs(wb, stage_element, stage_subelement)
    findings.extend(
        consistency.find_prediction_inconsistencies(wb, stage_element, stage_subelement)
    )
    return findings


def find_consistency_findings_in_file(
    input_file: str, stage_element: str = None, stage_subelement: str = None
) -> list:
    """Load an Excel file and return find_consistency_findings() for it."""
    wb = openpyxl.load_workbook(input_file, data_only=True)
    try:
        return find_consistency_findings(wb, stage_element, stage_subelement)
    finally:
        wb.close()
