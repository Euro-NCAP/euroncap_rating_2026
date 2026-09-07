# Copyright 2025-2026, Euro NCAP IVZW
# Created by IVEX NV (https://ivex.ai)
#
# Licensed under the Apache License 2.0.
# See http://www.apache.org/licenses/LICENSE-2.0 for details.

"""Report-only prediction-consistency rules for the crash_avoidance domain.

The rules themselves are the ones preprocess already enforces (or will
enforce) by raising -- the LSC dooring colour rule, the CPLA/CBLA duplicate
coherence rule -- plus the FCW colour rule, the ELK RE LDW
colour rule (Extended range performance = LDW forbids Green
in the extended-range cells) and the scenario_na_consistency rule (when any
of a scenario's input parameters is N/A, all of that scenario's input
parameters must be N/A). This module runs them on the *raw uploaded workbook*
(text-valued dropdown cells, before any preprocess rewrite or CPLA/CBLA
duplicate-row collapse) and returns every violation as a
common.ConsistencyFinding instead of failing fast, so a consuming application can
display all of them at once. The rule logic is shared with preprocess
(test_info.DOORING_ALLOWED_COLORS, collect_lsc_consistency_violations,
collect_cpla_cbla_coherence_mismatches) -- only the reporting differs.

Grey (unfilled) cells are never rule violations here: an empty required cell
is the missing-input check's finding (integrity.find_missing_required_inputs),
not a colour inconsistency.
"""

import logging

import pandas as pd
from openpyxl.utils import get_column_letter

from euroncap_rating_2026 import common
from euroncap_rating_2026.crash_avoidance import (
    data_model,
    matrix_processing,
    test_info,
)

logger = logging.getLogger(__name__)

LSC_PED_CYC_SHEET = "LSC - Ped & Cyc pred."
FC_PED_CYC_SHEET = "FC - Ped & Cyc pred."
LDC_SINGLE_VEH_SHEET = "LDC - Single Veh pred."
INPUT_PARAMETERS_SHEET = "Input parameters"

# The loadcases carrying a grey-cell Function dropdown (FCW/AEB choice band)
# -- the scope of both the coherence rule and the FCW colour rule.
FCW_CHOICE_LOADCASES = ("CPLA day", "CPLA night", "CBLA")

# FCW's criterion (FCW TTC) is pass/fail: graded colours only exist for the
# AEB criterion's Figure 5-1 bands (Frontal Collisions v1.2 §5.2).
FCW_FORBIDDEN_COLORS = (
    common.PredictionColor.YELLOW,
    common.PredictionColor.ORANGE,
    common.PredictionColor.BROWN,
)


def _matrix_cell(matrix_indices, test_point) -> str:
    """Worksheet coordinate of a matrix test point: DataFrame indices are
    0-based with the header row stripped (hence +2 for the row, +1 for the
    column -- see the same arithmetic in test_info's collapse bookkeeping)."""
    row = matrix_indices["start_row"] + test_point.row + 2
    col = matrix_indices["start_col"] + test_point.col + 1
    return f"{get_column_letter(col)}{row}"


def _row_identity(loadcase_name, test_point) -> dict:
    return {"Scenario": loadcase_name, **test_point.attributes}


def _capitalized(color) -> str:
    return str(color.value).capitalize()


def _find_dooring_findings(dfs) -> list:
    """LSC dooring (CBDA): predicted colours the "Vehicle response" input
    parameter forbids (test_info.DOORING_ALLOWED_COLORS)."""
    if LSC_PED_CYC_SHEET not in dfs or INPUT_PARAMETERS_SHEET not in dfs:
        return []
    prediction_df = dfs[LSC_PED_CYC_SHEET]

    vehicle_response = test_info.get_vehicle_response(dfs)
    if vehicle_response is None:
        # Nothing to compare against; the empty/invalid parameter cell is
        # the missing-input check's finding.
        return []

    test_points = matrix_processing.get_test_matrix(
        prediction_df, "CBDA", data_model.StageSubelementKey.LSC
    )
    if not test_points:
        return []
    matrix_indices = matrix_processing.get_matrix_indices(prediction_df, "CBDA")
    if not matrix_indices:
        return []

    # Mirror preprocess' selection (get_all_lsc_points): grey cells are
    # unfilled and red ones are never selected for testing, so neither ever
    # reaches the blocking check -- don't report what preprocess accepts.
    candidates = [
        tp
        for tp in test_points
        if tp.color not in (common.PredictionColor.GREY, common.PredictionColor.RED)
    ]
    allowed = test_info.DOORING_ALLOWED_COLORS[vehicle_response]
    allowed_desc = ", ".join(_capitalized(c) for c in allowed)

    findings = []
    for tp in test_info.collect_lsc_consistency_violations(
        candidates, vehicle_response
    ):
        cell = _matrix_cell(matrix_indices, tp)
        findings.append(
            common.ConsistencyFinding(
                sheet=LSC_PED_CYC_SHEET,
                column=get_column_letter(matrix_indices["start_col"] + tp.col + 1),
                cell=cell,
                row_identity=_row_identity("CBDA", tp),
                check=common.CHECK_LSC_DOORING,
                message=(
                    f"{LSC_PED_CYC_SHEET} — CBDA {cell}: prediction "
                    f"'{_capitalized(tp.color)}' is not allowed when Vehicle "
                    f"response is '{vehicle_response.value}' "
                    f"(allowed: {allowed_desc})"
                ),
            )
        )
    return findings


def _find_ldc_single_veh_findings(dfs) -> list:
    """ELK RE under LDW: when the "Extended range
    performance" input parameter is LDW, an extended-range cell cannot
    predict Green -- a Green ELK pass contradicts the declared LDW-only
    extended-range performance (allowed: Orange, Red, N/A)."""
    if LDC_SINGLE_VEH_SHEET not in dfs or INPUT_PARAMETERS_SHEET not in dfs:
        return []

    if test_info.get_extended_range_performance(dfs) != "ldw":
        # ELK (or absent/N/A) performance puts no constraint on the colours;
        # an empty/invalid parameter cell is the missing-input check's
        # finding.
        return []

    prediction_df = dfs[LDC_SINGLE_VEH_SHEET]
    test_points = matrix_processing.get_test_matrix(
        prediction_df, "ELK RE", data_model.StageSubelementKey.LDC
    )
    if not test_points:
        return []
    matrix_indices = matrix_processing.get_matrix_indices(prediction_df, "ELK RE")
    if not matrix_indices:
        return []

    findings = []
    for tp in test_points:
        if tp.test_range != matrix_processing.TestRange.EXTENDED:
            continue
        if tp.color != common.PredictionColor.GREEN:
            continue
        cell = _matrix_cell(matrix_indices, tp)
        findings.append(
            common.ConsistencyFinding(
                sheet=LDC_SINGLE_VEH_SHEET,
                column=get_column_letter(matrix_indices["start_col"] + tp.col + 1),
                cell=cell,
                row_identity=_row_identity("ELK RE", tp),
                check=common.CHECK_LDW_EXTENDED_COLOR,
                message=(
                    f"{LDC_SINGLE_VEH_SHEET} — ELK RE {cell}: prediction "
                    f"'Green' is not allowed in the extended range when "
                    f"Extended range performance is 'LDW' "
                    f"(allowed: Orange, Red, N/A)"
                ),
            )
        )
    return findings


def _find_fc_ped_cyc_findings(dfs) -> list:
    """CPLA/CBLA rules on the FC - Ped & Cyc prediction sheet:

    - coherence: duplicate declarations (same attributes twice, via the
      Function choice band) must predict the same colour;
    - fcw_color: a row declared Function=FCW must not predict a graded
      colour (Yellow/Orange/Brown) -- FCW TTC is pass/fail.

    Runs on the uncollapsed matrices, as uploaded.
    """
    if FC_PED_CYC_SHEET not in dfs:
        return []
    prediction_df = dfs[FC_PED_CYC_SHEET]

    findings = []
    for loadcase_name in FCW_CHOICE_LOADCASES:
        test_points = matrix_processing.get_test_matrix(
            prediction_df, loadcase_name, data_model.StageSubelementKey.FC
        )
        if not test_points:
            continue
        matrix_indices = matrix_processing.get_matrix_indices(
            prediction_df, loadcase_name
        )
        if not matrix_indices:
            continue

        for non_grey, attributes in test_info.collect_cpla_cbla_coherence_mismatches(
            test_points
        ):
            cells = [(tp, _matrix_cell(matrix_indices, tp)) for tp in non_grey]
            colors_desc = ", ".join(
                f"{cell}={_capitalized(tp.color)}" for tp, cell in cells
            )
            attributes_desc = " / ".join(
                str(value) for value in attributes.values() if value is not None
            )
            findings.append(
                common.ConsistencyFinding(
                    sheet=FC_PED_CYC_SHEET,
                    column=get_column_letter(
                        matrix_indices["start_col"] + non_grey[0].col + 1
                    ),
                    cell=cells[0][1],
                    row_identity={"Scenario": loadcase_name, **attributes},
                    check=common.CHECK_CPLA_CBLA_COHERENCE,
                    message=(
                        f"{FC_PED_CYC_SHEET} — {loadcase_name}: the duplicate "
                        f"declarations of ({attributes_desc}) must predict the "
                        f"same color, got {colors_desc}"
                    ),
                )
            )

        for tp in test_points:
            if str(tp.attributes.get("Function", "")).upper() != "FCW":
                continue
            if tp.color not in FCW_FORBIDDEN_COLORS:
                continue
            cell = _matrix_cell(matrix_indices, tp)
            findings.append(
                common.ConsistencyFinding(
                    sheet=FC_PED_CYC_SHEET,
                    column=get_column_letter(matrix_indices["start_col"] + tp.col + 1),
                    cell=cell,
                    row_identity=_row_identity(loadcase_name, tp),
                    check=common.CHECK_FCW_COLOR,
                    message=(
                        f"{FC_PED_CYC_SHEET} — {loadcase_name} {cell}: "
                        f"prediction '{_capitalized(tp.color)}' is not allowed "
                        f"when Function is FCW (FCW TTC is pass/fail: N/A, "
                        f"Green or Red only)"
                    ),
                )
            )
    return findings


def _find_scenario_na_findings(
    dfs, stage_element: str = None, stage_subelement: str = None
) -> list:
    """scenario_na_consistency: when any of a scenario's input parameters is
    N/A, ALL input parameters of that scenario must be N/A (e.g. ELK RE's
    three: Prediction - Standard, Prediction - Extended and Extended range
    performance). One finding per filled non-N/A Value cell in a violating
    scenario. A blank cell is never a violation source (the missing-input
    check owns unfilled cells). Scope filtering is row-level on the sheet's
    own forward-filled Stage element / Stage subelement columns, because the
    scenarios span several stage elements in one sheet.
    """
    if INPUT_PARAMETERS_SHEET not in dfs:
        return []
    df = dfs[INPUT_PARAMETERS_SHEET]
    required_columns = {
        "Stage element",
        "Stage subelement",
        "Category",
        "Scenario",
        "Input parameter",
        "Value",
    }
    if not required_columns.issubset(df.columns):
        return []

    ffilled = df.copy()
    hierarchy_columns = ["Stage element", "Stage subelement", "Category", "Scenario"]
    ffilled[hierarchy_columns] = (
        ffilled[hierarchy_columns]
        .map(lambda v: pd.NA if common.is_empty_cell(v) else v)
        .ffill()
    )
    value_col_letter = get_column_letter(list(df.columns).index("Value") + 1)

    # {scenario: [(parameter, value, worksheet row, stage element, subelement)]}
    scenario_rows = {}
    for idx, row in ffilled.iterrows():
        parameter = row["Input parameter"]
        if common.is_empty_cell(parameter) or common.is_empty_cell(row["Scenario"]):
            continue
        scenario_rows.setdefault(row["Scenario"], []).append(
            (
                parameter,
                row["Value"],
                idx + 2,  # header row is worksheet row 1
                row["Stage element"],
                row["Stage subelement"],
            )
        )

    findings = []
    for scenario, rows in scenario_rows.items():
        na_params = [
            parameter
            for parameter, value, *_rest in rows
            if common.is_not_applicable(value)
        ]
        if not na_params:
            continue
        na_desc = ", ".join(na_params)
        for parameter, value, ws_row, row_stage_element, row_stage_subelement in rows:
            if common.is_not_applicable(value) or common.is_empty_cell(value):
                continue
            if stage_element is not None and row_stage_element != stage_element:
                continue
            if (
                stage_subelement is not None
                and row_stage_subelement != stage_subelement
            ):
                continue
            cell = f"{value_col_letter}{ws_row}"
            findings.append(
                common.ConsistencyFinding(
                    sheet=INPUT_PARAMETERS_SHEET,
                    column="Value",
                    cell=cell,
                    row_identity={
                        "Stage element": row_stage_element,
                        "Stage subelement": row_stage_subelement,
                        "Scenario": scenario,
                        "Input parameter": parameter,
                    },
                    check=common.CHECK_SCENARIO_NA_CONSISTENCY,
                    message=(
                        f"{INPUT_PARAMETERS_SHEET} — {str(scenario).strip()}: "
                        f"{parameter} is '{str(value).strip()}' but {na_desc} "
                        f"is 'N/A'; when any input parameter of a scenario is "
                        f"N/A, all of its input parameters must be N/A"
                    ),
                )
            )
    return findings


def find_prediction_inconsistencies(
    wb, stage_element: str = None, stage_subelement: str = None
) -> list:
    """Every prediction-rule violation in *wb* for the crash_avoidance
    domain, as a list of common.ConsistencyFinding (check = lsc_dooring,
    cpla_cbla_coherence, fcw_color, ldw_extended_color or
    scenario_na_consistency). Report-only:
    collects all violations instead of raising on the first, and never fails
    on content -- the blocking behaviour stays in preprocess. Scope filtering
    follows find_missing_required_inputs' (Stage element, Stage subelement)
    convention; scope validity is the caller's concern
    (integrity.find_consistency_findings validates it once).
    """
    sheet_names = [
        name
        for name in (
            LSC_PED_CYC_SHEET,
            FC_PED_CYC_SHEET,
            LDC_SINGLE_VEH_SHEET,
            INPUT_PARAMETERS_SHEET,
        )
        if name in wb.sheetnames
    ]
    dfs = {name: common.sheet_to_dataframe(wb[name]) for name in sheet_names}

    findings = []
    if common.stage_elements_in_scope(
        [("Low Speed Collisions", "Pedestrian & cyclist")],
        stage_element,
        stage_subelement,
    ):
        findings.extend(_find_dooring_findings(dfs))
    if common.stage_elements_in_scope(
        [("Frontal Collisions", "Pedestrian & cyclist")],
        stage_element,
        stage_subelement,
    ):
        findings.extend(_find_fc_ped_cyc_findings(dfs))
    if common.stage_elements_in_scope(
        [("Lane Departure Collisions", "Single vehicle")],
        stage_element,
        stage_subelement,
    ):
        findings.extend(_find_ldc_single_veh_findings(dfs))
    findings.extend(_find_scenario_na_findings(dfs, stage_element, stage_subelement))
    return findings
